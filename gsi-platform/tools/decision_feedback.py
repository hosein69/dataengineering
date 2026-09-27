# -*- coding: utf-8 -*-
"""Local, published-snapshot decision feedback; share only aggregated JSON.

Export a review workbook from GSI, collect a human assessment and later verified
outcomes, then score paired evidence quality and observed outcome cohorts. No
claim of causal financial uplift follows from either score automatically.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from gsi.core.text import clean_order_ref, clean_part_no

HEADERS = [
    "case_id", "سفارش", "متریال", "شرح‌های کارشناسان", "قبلی_متریال_قابل_مشاهده",
    "قبلی_وضعیت_موجودی", "جدید_وضعیت_موجودی", "جدید_شکاف‌ها",
    "حساس_برای_تصمیم", "شواهد_قبلی_کافی", "شواهد_جدید_کافی",
    "نسخه_استفاده_شده", "معیار_موفقیت_از_پیش", "نتیجه_تأییدشده",
    "تاریخ_نتیجه", "مبنای_تأیید_و_یادداشت",
]
YES_NO_UNKNOWN = {"بله", "خیر", "نامعلوم", ""}
ARMS = {"قبلی", "جدید", "نامعلوم", ""}
OUTCOMES = {"موفق", "ناموفق", "در انتظار", "نامعلوم", ""}


def _text(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _excel_text(value):
    """Source strings must not become formulas in the local review workbook."""
    s = _text(value)
    return "'" + s if s.startswith(("=", "+", "-", "@")) else s


def _case_id(run_id: str, order: str, material: str) -> str:
    return hashlib.sha256(f"{run_id}\0{order}\0{material}".encode("utf-8")).hexdigest()[:20]


def build_roster(main: pd.DataFrame, positions: pd.DataFrame, run_id: str):
    """Old flat-mart material representation is a proxy, not a saved old run."""
    if not run_id or positions is None or positions.empty:
        raise ValueError("Snapshot جدید، شامل دفتر expert_material_positions، لازم است.")
    required = {"KEY_ORDER", "KEY_MATERIAL", "SUPPLY_POSITION_STATUS"}
    if not required.issubset(positions):
        raise ValueError("ستون‌های دفتر متریال ناقص‌اند؛ Refresh نسخهٔ جدید لازم است.")
    old = {}
    if isinstance(main, pd.DataFrame) and not main.empty:
        for row in main.to_dict("records"):
            order = clean_order_ref(_text(row.get("CANONICAL_ORDER")) or _text(row.get("KEY_ORDER")))
            material = clean_part_no(_text(row.get("KEY_MATERIAL")))
            expert_material = clean_part_no(_text(row.get("MOGH_MATERIAL")) or _text(row.get("MOGH_MFR_PART_NO")))
            if order and material and expert_material == material:
                old.setdefault((order, material), set()).add(_text(row.get("SUPPLY_POSITION_STATUS")))
    roster = []
    seen = set()
    for row in positions.to_dict("records"):
        order = clean_order_ref(_text(row.get("KEY_ORDER")))
        material = clean_part_no(_text(row.get("KEY_MATERIAL")))
        if not material:
            continue
        pair = (order, material)
        if pair in seen:
            raise ValueError("دفتر جدید کلید سفارش×متریال تکراری دارد؛ ارزیابی متوقف شد.")
        seen.add(pair)
        statuses = old.get(pair, set())
        roster.append({
            "case_id": _case_id(run_id, order, material), "سفارش": order,
            "متریال": material, "شرح‌های کارشناسان": _text(row.get("MOGH_MATERIAL_DESCS_ALL")),
            "قبلی_متریال_قابل_مشاهده": "بله" if statuses else "خیر",
            "قبلی_وضعیت_موجودی": " | ".join(sorted(x for x in statuses if x)),
            "جدید_وضعیت_موجودی": _text(row.get("SUPPLY_POSITION_STATUS")),
            "جدید_شکاف‌ها": _text(row.get("SUPPLY_POSITION_GAPS")),
        })
    if not roster:
        raise ValueError("هیچ کد متریال قابل ارزیابی در Snapshot منتشرشده وجود ندارد.")
    return sorted(roster, key=lambda r: r["case_id"])


def load_roster():
    from gsi.warehouse.service import last_report
    report = last_report()
    if report is None:
        raise ValueError("گزارش منتشرشده یافت نشد؛ ابتدا Refresh و Verify کنید.")
    _, main, extras, _ = report
    positions = extras.get("expert_material_positions")
    run_id = _text(extras.get("warehouse_run_id"))
    return build_roster(main, positions, run_id), run_id


def export_workbook(path: Path, roster: list[dict], run_id: str, sample: int = 120, seed: int = 1405):
    if sample < 0:
        raise ValueError("sample باید صفر (همه) یا مثبت باشد.")
    chosen = (sorted(random.Random(seed).sample(roster, min(sample, len(roster))),
                     key=lambda r: r["case_id"]) if sample else roster)
    wb = Workbook()
    ws = wb.active
    ws.title = "بازبینی"
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for item in chosen:
        ws.append([_excel_text(item.get(h, "")) for h in HEADERS])
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:P{len(chosen)+1}"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="102B41")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="right", wrap_text=True)
    for col, width in {"A": 24, "B": 18, "C": 18, "D": 42, "E": 20, "F": 22,
                       "G": 22, "H": 42, "I": 19, "J": 19, "K": 19, "L": 18,
                       "M": 38, "N": 18, "O": 17, "P": 48}.items():
        ws.column_dimensions[col].width = width
    for col, values in (("I", "بله,خیر,نامعلوم"), ("J", "بله,خیر,نامعلوم"),
                        ("K", "بله,خیر,نامعلوم"), ("L", "قبلی,جدید,نامعلوم"),
                        ("N", "موفق,ناموفق,در انتظار,نامعلوم")):
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(chosen)+1}")
    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 115
    for line in (
        "فقط ستون‌های I تا P را تکمیل کنید. سفارش/متریال برای بازبینی محلی است؛ این Excel را برای تحلیل بیرونی نفرستید.",
        "حساس برای تصمیم: آیا این قلم واقعاً برای یک تصمیم خرید/مالی مهم بود؟ نامعلوم هم مجاز است.",
        "شواهد قبلی/جدید کافی: قضاوت شما با شاهد پرونده؛ هر دو را برای همان case_id ارزیابی کنید.",
        "نتیجه تأییدشده: فقط پس از تعریف معیار موفقیت از پیش و داشتن شاهد نتیجه پر شود.",
        "نسخه استفاده شده: نسخه‌ای که واقعاً مبنای تصمیم بود؛ حدس نزنید.",
        "خروجی JSON فقط درصدهای تجمیعی دارد. تفاوت دو گروه بدون تخصیص تصادفی، اثر علّی نسخه نیست.",
        "نمای قبلی یک بازسازی از شکل جدول تخت روی همین Snapshot است؛ داده اجرای تاریخی قبلی نیست.",
    ):
        guide.append([line])
    meta = wb.create_sheet("_meta")
    meta.sheet_state = "hidden"
    sample_hash = hashlib.sha256("\n".join(x["case_id"] for x in chosen).encode("ascii")).hexdigest()
    for key, value in {"schema": 1, "run_id": run_id, "population": len(roster),
                       "old_visible": sum(x["قبلی_متریال_قابل_مشاهده"] == "بله" for x in roster),
                       "new_visible": len(roster), "sample": len(chosen), "seed": seed,
                       "sample_sha256": sample_hash,
                       "generated": datetime.now().isoformat(timespec="seconds")}.items():
        meta.append([key, value])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"population": len(roster), "sample": len(chosen), "path": str(path)}


def _bootstrap_delta(pairs: list[tuple[int, int]], seed=1405):
    if len(pairs) < 20:
        return None
    rng = random.Random(seed)
    deltas = [b - a for a, b in pairs]
    samples = sorted(sum(rng.choices(deltas, k=len(deltas))) / len(deltas)
                     for _ in range(2000))
    return [round(100 * samples[49], 2), round(100 * samples[1949], 2)]


def score_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if "بازبینی" not in wb or "_meta" not in wb:
        raise ValueError("این فایل الگوی بازخورد GSI نیست.")
    meta = {str(row[0]): row[1] for row in wb["_meta"].values if row[0] is not None}
    if meta.get("schema") != 1:
        raise ValueError("نسخه الگوی بازخورد شناخته‌شده نیست.")
    values = wb["بازبینی"].values
    head = list(next(values))
    if head != HEADERS:
        raise ValueError("سرستون‌های فایل بازخورد تغییر کرده‌اند.")
    rows = [dict(zip(head, row)) for row in values]
    row_ids = [_text(r["case_id"]) for r in rows]
    digest = hashlib.sha256("\n".join(row_ids).encode("ascii")).hexdigest()
    if (len(rows) != meta.get("sample") or len(set(row_ids)) != len(rows)
            or digest != meta.get("sample_sha256")):
        raise ValueError("تعداد یا شناسه پرونده‌های بازبینی تغییر کرده است.")
    paired = []
    outcomes = {"قبلی": [], "جدید": []}
    relevant = 0
    for i, row in enumerate(rows, start=2):
        for col in ("حساس_برای_تصمیم", "شواهد_قبلی_کافی", "شواهد_جدید_کافی"):
            if _text(row[col]) not in YES_NO_UNKNOWN:
                raise ValueError(f"سطر {i}: مقدار {col} نامعتبر است.")
        arm = _text(row["نسخه_استفاده_شده"])
        outcome = _text(row["نتیجه_تأییدشده"])
        if arm not in ARMS or outcome not in OUTCOMES:
            raise ValueError(f"سطر {i}: نسخه یا نتیجه نامعتبر است.")
        if outcome in {"موفق", "ناموفق"} and arm not in outcomes:
            raise ValueError(f"سطر {i}: برای نتیجهٔ قطعی باید نسخهٔ استفاده‌شده مشخص باشد.")
        if _text(row["حساس_برای_تصمیم"]) == "بله":
            relevant += 1
            old, new = (_text(row[c]) for c in ("شواهد_قبلی_کافی", "شواهد_جدید_کافی"))
            if old in {"بله", "خیر"} and new in {"بله", "خیر"}:
                paired.append((int(old == "بله"), int(new == "بله")))
        if arm in outcomes and outcome in {"موفق", "ناموفق"}:
            if not _text(row["معیار_موفقیت_از_پیش"]) or not _text(row["تاریخ_نتیجه"]):
                raise ValueError(f"سطر {i}: نتیجه بدون معیار از پیش و تاریخ قابل شمارش نیست.")
            outcomes[arm].append(int(outcome == "موفق"))
    n = len(paired)
    old_yes = sum(a for a, _ in paired)
    new_yes = sum(b for _, b in paired)
    observed = {arm: {"n": len(vals), "success": sum(vals),
                       "rate_pct": round(100 * sum(vals) / len(vals), 2) if vals else None}
                for arm, vals in outcomes.items()}
    old_rate, new_rate = observed["قبلی"]["rate_pct"], observed["جدید"]["rate_pct"]
    return {
        "schema": 1, "comparison": "same_snapshot_old_flat_mart_proxy_vs_new_material_ledger",
        "population": int(meta["population"]), "sample": len(rows),
        "technical_coverage": {
            "old_visible": int(meta["old_visible"]), "new_visible": int(meta["new_visible"]),
            "old_pct": round(100 * int(meta["old_visible"]) / int(meta["population"]), 2),
            "new_pct": round(100 * int(meta["new_visible"]) / int(meta["population"]), 2),
        },
        "reviewer_assessed_evidence": {
            "relevant": relevant, "paired_complete": n,
            "old_sufficient": old_yes, "new_sufficient": new_yes,
            "old_pct": round(100 * old_yes / n, 2) if n else None,
            "new_pct": round(100 * new_yes / n, 2) if n else None,
            "delta_percentage_points": round(100 * (new_yes - old_yes) / n, 2) if n else None,
            "improved": sum(new > old for old, new in paired),
            "worsened": sum(new < old for old, new in paired),
            "bootstrap_95pct_interval_pp": _bootstrap_delta(paired),
        },
        "observed_outcomes": {
            "cohorts": observed,
            "descriptive_delta_percentage_points": (round(new_rate - old_rate, 2)
                                                    if min(len(outcomes["قبلی"]), len(outcomes["جدید"])) >= 20
                                                    and old_rate is not None and new_rate is not None else None),
            "causal_uplift_estimate": None,
            "reason": "A descriptive comparison needs at least 20 verified outcomes per arm; without randomized comparable assignment it is not causal."
        },
        "sharing": "aggregate_only_no_identifiers_or_notes",
    }


def main(argv=None):
    ap = argparse.ArgumentParser(description="GSI local decision feedback and honest uplift measurement")
    sub = ap.add_subparsers(dest="command", required=True)
    exp = sub.add_parser("export", help="Build a local review workbook from published GSI")
    exp.add_argument("--output", default="decision_feedback.xlsx")
    exp.add_argument("--sample", type=int, default=120, help="0 means all material positions")
    exp.add_argument("--seed", type=int, default=1405)
    score = sub.add_parser("score", help="Write shareable aggregate JSON from completed workbook")
    score.add_argument("--input", required=True)
    score.add_argument("--output", default="decision_feedback_summary.json")
    args = ap.parse_args(argv)
    try:
        if args.command == "export":
            roster, run_id = load_roster()
            result = export_workbook(Path(args.output), roster, run_id, args.sample, args.seed)
        else:
            result = score_workbook(Path(args.input))
            out = Path(args.output)
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            result = {"output": str(out), "technical_coverage": result["technical_coverage"],
                      "reviewed_pairs": result["reviewer_assessed_evidence"]["paired_complete"],
                      "causal_uplift_estimate": None}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ValueError, FileNotFoundError) as ex:
        ap.error(str(ex))


if __name__ == "__main__":
    raise SystemExit(main())
