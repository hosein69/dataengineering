# -*- coding: utf-8 -*-
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
from openpyxl import load_workbook
from tools.decision_feedback import build_roster, export_workbook, score_workbook, load_roster, _bootstrap_delta
from gsi.warehouse.store import Warehouse


class FeedbackTests(unittest.TestCase):
    def test_interval_only_after_twenty_paired_reviews(self):
        self.assertIsNone(_bootstrap_delta([(0, 1)] * 19))
        self.assertEqual(_bootstrap_delta([(0, 1)] * 20), [100.0, 100.0])

    def test_export_reads_only_a_published_snapshot(self):
        main, positions = self.fixtures()
        with tempfile.TemporaryDirectory() as d:
            with patch.dict("os.environ", {"GSI_DWH_PATH": str(Path(d) / "warehouse.sqlite")}):
                wh = Warehouse()
                with wh.run({"reference_date": "2026-09-27"}) as rid:
                    wh.frame(main, "mart", "df")
                    wh.frame(main, "mart", "main")
                    wh.frame(positions, "mart", "extras/expert_material_positions")
                self.assertIsNone(__import__("gsi.warehouse.service", fromlist=["last_report"]).last_report())
                wh.publish(rid)
                roster, run_id = load_roster()
                self.assertEqual(run_id, rid)
                self.assertEqual(len(roster), 2)

    def fixtures(self):
        main = pd.DataFrame([{"CANONICAL_ORDER": "823107D", "KEY_MATERIAL": "M1",
                              "MOGH_MATERIAL": "M1", "SUPPLY_POSITION_STATUS": "PARTIAL"}])
        positions = pd.DataFrame([
            {"KEY_ORDER": "823107D", "KEY_MATERIAL": "M1",
             "MOGH_MATERIAL_DESCS_ALL": "شرح اول | شرح دوم", "SUPPLY_POSITION_STATUS": "PARTIAL"},
            {"KEY_ORDER": "823107D", "KEY_MATERIAL": "M2",
             "MOGH_MATERIAL_DESCS_ALL": "=SUM(1,1)", "SUPPLY_POSITION_STATUS": "COMPLETE"},
        ])
        return main, positions

    def test_coverage_and_missing_feedback_do_not_invent_success(self):
        roster = build_roster(*self.fixtures(), "run-1")
        self.assertEqual(len(roster), 2)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "feedback.xlsx"
            export_workbook(path, roster, "run-1", sample=0)
            wb = load_workbook(path)
            desc_cells = [wb["بازبینی"][f"D{row}"] for row in (2, 3)]
            self.assertTrue(all(c.data_type == "s" for c in desc_cells))
            self.assertTrue(any(str(c.value).startswith("'=SUM") for c in desc_cells))
            result = score_workbook(path)
        self.assertEqual(result["technical_coverage"]["old_pct"], 50.0)
        self.assertEqual(result["technical_coverage"]["new_pct"], 100.0)
        self.assertIsNone(result["reviewer_assessed_evidence"]["delta_percentage_points"])
        self.assertIsNone(result["observed_outcomes"]["causal_uplift_estimate"])
        public = json.dumps(result, ensure_ascii=False)
        self.assertNotIn("823107D", public)
        self.assertNotIn("M1", public)
        self.assertNotIn("شرح", public)
        self.assertNotIn("run-1", public)

    def test_paired_review_and_observed_outcomes_stay_separate(self):
        roster = build_roster(*self.fixtures(), "run-1")
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "feedback.xlsx"
            export_workbook(path, roster, "run-1", sample=0)
            wb = load_workbook(path)
            ws = wb["بازبینی"]
            ws["I2"] = ws["I3"] = "بله"
            ws["J2"] = "خیر"; ws["K2"] = "بله"
            ws["J3"] = "بله"; ws["K3"] = "خیر"
            ws["L2"] = "قبلی"; ws["L3"] = "جدید"
            ws["M2"] = ws["M3"] = "رفع تأخیر تا مهلت"
            ws["N2"] = "ناموفق"; ws["N3"] = "موفق"
            ws["O2"] = ws["O3"] = "1405/07/15"
            wb.save(path)
            result = score_workbook(path)
        paired = result["reviewer_assessed_evidence"]
        self.assertEqual((paired["improved"], paired["worsened"]), (1, 1))
        self.assertEqual(paired["delta_percentage_points"], 0.0)
        self.assertIsNone(paired["bootstrap_95pct_interval_pp"])
        self.assertIsNone(result["observed_outcomes"]["descriptive_delta_percentage_points"])
        self.assertIsNone(result["observed_outcomes"]["causal_uplift_estimate"])

    def test_missing_success_definition_is_rejected(self):
        roster = build_roster(*self.fixtures(), "run-1")
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "feedback.xlsx"
            export_workbook(path, roster, "run-1", sample=0)
            wb = load_workbook(path)
            ws = wb["بازبینی"]
            ws["L2"] = "جدید"; ws["N2"] = "موفق"; ws["O2"] = "1405/07/15"
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "معیار"):
                score_workbook(path)

    def test_case_list_cannot_be_changed_after_sampling(self):
        roster = build_roster(*self.fixtures(), "run-1")
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "feedback.xlsx"
            export_workbook(path, roster, "run-1", sample=0)
            wb = load_workbook(path)
            wb["بازبینی"]["A2"] = "changed"
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "شناسه"):
                score_workbook(path)


if __name__ == "__main__":
    unittest.main()
