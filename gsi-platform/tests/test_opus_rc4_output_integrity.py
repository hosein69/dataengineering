# -*- coding: utf-8 -*-
"""ناوردای کیفیت خروجی — RC4.

هر تست اینجا یک نقص **اندازه‌گیری‌شده روی خروجی واقعی** را قفل می‌کند، نه یک
سلیقهٔ معماری. منبع اندازه‌گیری: بستهٔ RC3 و فایل OF پیوست خودِ بسته.
"""
from __future__ import annotations

import io
import json
import re

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.cashflow.engine import build_cashflow
from gsi.cashflow.inputs import legacy_of
from gsi.cashflow.report import excel_bytes, html_report, empty_reason
from gsi.design import tokens as T
from gsi.studio_core.html_export import build_dynamic_html


# ═════════ ۱) هستهٔ HTML چیزی را بی‌اعلام حذف نکند ═════════

def _wide_frame(rows: int = 40, extra: int = 30) -> pd.DataFrame:
    data = {"KEY_MATERIAL": [f"M{i}" for i in range(rows)],
            "ORG_DEPT": ["A"] * rows,
            "مانده تعهد": range(rows)}
    for j in range(extra):
        data[f"EVIDENCE_{j}"] = [f"v{j}_{i}" for i in range(rows)]
    return pd.DataFrame(data)


def _report_meta(html: str) -> dict:
    m = re.search(r"const REPORT_META=(\{.*?\});\n", html, re.S)
    assert m, "REPORT_META در خروجی نیست"
    return json.loads(m.group(1))


def test_html_keeps_every_row_in_the_payload():
    df = _wide_frame(rows=137)
    meta = _report_meta(build_dynamic_html(df, "2026-09-15"))
    assert meta["payload_rows"] == 137
    assert meta["embedded_rows"] == 137, "هیچ ردیفی نباید در payload حذف شود"


def test_html_declares_every_column_it_left_out():
    df = _wide_frame()
    html = build_dynamic_html(df, "2026-09-15")
    meta = _report_meta(html)
    assert meta["source_columns"] == len(df.columns)
    assert meta["payload_columns"] < meta["source_columns"], "این سناریو باید ستون حذف‌شده داشته باشد"
    excluded = meta["excluded_columns"]
    assert len(excluded) == meta["source_columns"] - meta["payload_columns"]
    # فهرست باید دقیق باشد، نه فقط یک عدد
    assert set(excluded) == {c for c in df.columns if c not in set(excluded)} ^ set(df.columns)
    # و باید برای انسان هم دیده شود، نه فقط داخل JS
    assert 'id="gsi_manifest"' in html
    assert f"{len(excluded):,}" in html
    for name in excluded[:3]:
        assert name in html, f"نام ستون حذف‌شده «{name}» در فهرست محتوای خروجی نیست"


def test_html_does_not_claim_excel_carries_columns_it_drops():
    html = build_dynamic_html(_wide_frame(), "2026-09-15")
    assert "خروجی Excel شامل همه نتایج است" not in html, (
        "این ادعا از نظر ستون نادرست بود: دکمه Excel فقط ستون‌های تب فعال را می‌دهد")
    assert "ستون‌های همین تب" in html


def test_html_discloses_fx_ledger_truncation():
    # دفتر ارزی پیش از تعبیه به دامنهٔ همان گزارش محدود می‌شود، پس KEY_REG باید
    # در هر دو طرف باشد؛ وگرنه اصلاً چیزی برای بریدن نمی‌ماند.
    df = _wide_frame(rows=400)
    df["KEY_REG"] = [f"R{i}" for i in range(400)]
    ledger = pd.DataFrame({"KEY_REG": [f"R{i}" for i in range(400)],
                           "FX_MONEY_STAGE": ["ALLOCATION"] * 400,
                           "FX_NTSW_BALANCE": range(400)})
    meta = _report_meta(build_dynamic_html(df, "2026-09-15",
                                           process_extras={"fx_ledger": ledger}))
    trunc = {t["section"]: t for t in meta["truncated_sections"]}
    assert "fx_ledger" in trunc, "برش ۲۵۰ ردیفی دفتر ارزی باید اعلام شود"
    assert trunc["fx_ledger"]["source_rows"] == 400
    assert trunc["fx_ledger"]["embedded_rows"] == 250


# ═════════ ۲) OF: نه گزارش خالی، نه دوبارشماری ═════════

def _of_frame() -> pd.DataFrame:
    """دو CB با دانهٔ مخلوط، دقیقاً مثل فایل واقعی.

    CB1: سه بارنامه، یک ارزش ثبت سفارش و یک پرداخت که در هر سه ردیف تکرار شده.
    CB2: یک بارنامه.
    """
    rows = []
    for bl, blv in [("BL1", "100"), ("BL2", "200"), ("BL3", "300")]:
        rows.append({"CB No.": "CB1", "CB Date": "1402/02/25", "CB Value": "1000", "Fx": "EUR",
                     "BL": bl, "BL Value": blv, "Date of Sending Documents": "1403-03-09",
                     "Payment": "900", "Fx Payment": "EUR", "Date of Buying Currency": "1402/07/03",
                     "Payment Number": "", "OF Deadline": "1403/04/07",
                     "Punishment Value": "50", "Punishment Fx": "EUR",
                     "Punishment Date to Bank": "1403/05/03",
                     "Date of Finance Receipt": "1403-04-19", "Date of Clearance": "1403-04-19"})
    rows.append({"CB No.": "CB2", "CB Date": "1402/03/01", "CB Value": "700", "Fx": "EUR",
                 "BL": "BL9", "BL Value": "700", "Date of Sending Documents": "1403-03-20",
                 "Payment": "700", "Fx Payment": "EUR", "Date of Buying Currency": "1402/08/01",
                 "Payment Number": "PN-9", "OF Deadline": "1403/05/07",
                 "Punishment Value": "", "Punishment Fx": "", "Punishment Date to Bank": "",
                 "Date of Finance Receipt": "", "Date of Clearance": ""})
    return pd.DataFrame(rows)


def _summary(result: dict, case: str) -> pd.Series:
    s = result["summary"]
    return s[s["case_id"] == case].iloc[0]


def test_of_workbook_is_not_rejected_wholesale():
    """قبل از این اصلاح، ۱۰۰٪ ردیف‌های OF واقعی UNVERIFIED_EVIDENCE می‌شدند."""
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    assert not r["events"].empty, "گزارش نباید کاملاً خالی باشد"
    assert "UNVERIFIED_EVIDENCE" not in set(r["issues"]["code"])


def test_of_case_level_columns_are_not_fanned_out():
    """ارزش ثبت سفارش در سه ردیف تکرار شده؛ جمع آن باید ۱۰۰۰ بماند نه ۳۰۰۰."""
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    row = _summary(r, "CB1")
    assert row["registration_value"] == pytest.approx(1000)
    assert row["paid"] == pytest.approx(900), "پرداخت تکرارشده نباید سه برابر شود"
    assert row["fees"] == pytest.approx(50), "جریمهٔ سطح CB نباید سه برابر شود"


def test_of_row_level_columns_keep_their_own_grain():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    assert _summary(r, "CB1")["shipment_value"] == pytest.approx(600)  # 100+200+300
    assert _summary(r, "CB2")["shipment_value"] == pytest.approx(700)


def test_of_keeps_row_traceability_after_dedup():
    ev = legacy_of(_of_frame())
    reg = ev[(ev["kind"] == "REGISTRATION") & (ev["case_id"] == "CB1")]
    assert len(reg) == 1
    assert "ردیف‌های منبع: 2, 3, 4" in reg.iloc[0]["note"]


def test_of_never_creates_bank_movements():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    assert r["movements"].empty and r["accounts"].empty
    assert (r["issues"]["code"] == "ACCOUNT_DETAIL_GAP").any()


def test_of_names_the_missing_identifier_column():
    ev = legacy_of(_of_frame())
    pay = ev[(ev["kind"] == "PAYMENT") & (ev["case_id"] == "CB1")].iloc[0]
    assert "Payment Number" in pay["note"], "خالی‌بودن شماره پرداخت باید اعلام شود"
    paid2 = ev[(ev["kind"] == "PAYMENT") & (ev["case_id"] == "CB2")].iloc[0]
    assert paid2["document"] == "PN-9"


# ═════════ ۳) بخش خالی باید علت داشته باشد ═════════

def test_empty_html_section_states_why_not_just_that_it_is_empty():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    html = html_report(r, excel_bytes(r))
    assert r["rate_register"].empty
    assert empty_reason("rate_register") in html
    assert "navgroup" in html, "بخش‌های بدون شاهد باید زیر عنوان جدا گروه شوند"
    assert 'class="blank"' in html


def test_empty_excel_sheet_states_why_instead_of_bare_headers():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    wb = load_workbook(io.BytesIO(excel_bytes(r)))
    ws = wb["شناسنامه نرخ‌ها"]
    assert ws.max_row == 2, "شیت خالی باید یک ردیف توضیح داشته باشد"
    assert [c.value for c in ws[1]][:2] == ["وضعیت", "علت"]
    assert ws.cell(2, 2).value == empty_reason("rate_register")


def test_filled_sections_come_before_empty_ones():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    html = html_report(r, excel_bytes(r))
    tabs = re.findall(r'<button role="tab" class="([a-z]*)"', html)
    first_blank = tabs.index("blank")
    assert "blank" not in tabs[:first_blank]
    assert all(t == "blank" for t in tabs[first_blank:])


def test_cashflow_report_uses_the_package_font_stack():
    r = build_cashflow(legacy_of(_of_frame()), as_of="2026-08-31")
    html = html_report(r, excel_bytes(r))
    assert T.FONT_STACK in html
    assert "font:14px Tahoma,Arial,sans-serif" not in html


def test_manifest_is_scope_disclosure_not_a_data_quality_metric():
    """فهرست محتوا نباید واژگان «کیفیت داده» را وارد نمای عملیاتی کند.

    قرارداد مخاطب (test_v27_audience_and_voice) می‌گوید سنجه‌های کیفیت داده در
    نمای کارشناس/مدیر فقط تردید می‌سازند و جایشان شیت سلامت سیستم است. این بخش
    یک سنجهٔ کیفیت نیست؛ اعلام دامنهٔ همین فایل است و باید همیشه دیده شود.
    """
    html = build_dynamic_html(_wide_frame(), "2026-09-15")
    for banned in ("پوشش داده", "درصد صحت", "شکاف اسکیما", "سلامت سورس", "کیفیت داده"):
        assert banned not in html, f"واژه «{banned}» قرارداد مخاطب را می‌شکند"
    assert 'id="gsi_manifest"' in html
