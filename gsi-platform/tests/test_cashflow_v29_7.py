from decimal import Decimal as D
import io
import pandas as pd
import pytest
from gsi.cashflow.engine import build_cashflow,select_rate
from gsi.cashflow.report import excel_bytes,html_report
from gsi.cashflow.inputs import legacy_of


def event(eid,kind,amount='100',currency='EUR',**kw):
    r=dict(event_id=eid,case_id='R1',kind=kind,amount=amount,currency=currency,date='2026-01-01',
           document='DOC-'+eid,source='bank',status='POSTED',from_account='EXTERNAL:bank',to_account='OWN:bank')
    if kind in {'PAYMENT','FEE','FX_SELL'}:r.update(from_account='OWN:bank',to_account='EXTERNAL:supplier')
    r.update(kw);return r


def run(*rows,**kw):return build_cashflow(pd.DataFrame(rows),as_of='2026-01-31',**kw)
def codes(result):return set(result['issues']['code'])
def link(i,a,b,x='60',y='60',**kw):return dict(link_id=i,from_event=a,to_event=b,from_amount=x,to_amount=y,document='link-document',**kw)


def test_exact_cash_and_opening():
    r=run(event('o','OPENING','0'),event('f','FUNDING','0.3'),event('p','PAYMENT','0.1'))
    assert r['accounts'].iloc[0]['closing']==D('0.2')
    assert r['periods'].iloc[0]['net_movement']==D('0.2')


def test_missing_opening_is_not_zero():
    r=run(event('f','FUNDING'))
    assert r['accounts'].iloc[0]['closing'] is None
    assert 'OPENING_GAP' in codes(r)


def test_duplicate_payment_does_not_fanout():
    p=event('p','PAYMENT');r=run(p,p)
    assert r['summary'].iloc[0]['paid']==D(100)
    assert len(r['events'])==1


def test_conflicting_payment_quarantines_all_versions():
    r=run(event('p','PAYMENT'),event('p','PAYMENT','101'))
    assert r['events'].empty and len(r['observations'])==2

@pytest.mark.parametrize('change,code',[
    ({'amount':''},'MISSING_AMOUNT_CURRENCY'),({'amount':'NaN'},'MISSING_AMOUNT_CURRENCY'),
    ({'amount':'-1'},'MISSING_AMOUNT_CURRENCY'),({'document':''},'UNVERIFIED_EVIDENCE'),
    ({'date':'2026-02-01'},'FUTURE_EVENT'),({'date':''},'MISSING_DATE'),
    ({'event_id':''},'MISSING_ID'),({'case_id':''},'MISSING_CASE'),
    ({'from_account':''},'MISSING_ACCOUNTS'),({'from_account':'EXTERNAL:x','to_account':'OWN:a'},'INVALID_CASH_DIRECTION')])
def test_invalid_cash_does_not_count(change,code):
    r=run(event('p','PAYMENT',**change));assert r['events'].empty;assert code in codes(r)


def test_noncash_does_not_spend_money():
    r=run(event('a','ALLOCATION','100'),event('s','CUSTOMS','90'),event('c','COMMITMENT','100'),event('z','SETTLEMENT','30'),
          links=pd.DataFrame([link('ob','c','z','30','30',relation_type='OBLIGATION_SETTLEMENT')]))
    assert r['movements'].empty
    assert r['summary'].iloc[0]['commitment_remaining']==D(70)
    assert r['summary'].iloc[0]['allocation_remaining']==D(100)


def test_unlinked_settlement_never_reduces_obligation():
    r=run(event('c','COMMITMENT','100'),event('z','SETTLEMENT','30'))
    assert r['summary'].iloc[0]['commitment_remaining']==D(100)
    assert 'UNLINKED_SETTLEMENT' in codes(r)


def test_snapshot_is_reconciliation_not_event():
    ms=pd.DataFrame([dict(measurement_id='m1',case_id='R1',metric='COMMITMENT_BALANCE',observed_at='2026-01-31',
                          amount='70',currency='EUR',source='NTSW',document='screen')])
    r=run(event('c','COMMITMENT','100'),event('z','SETTLEMENT','30'),
          links=pd.DataFrame([link('ob','c','z','30','30',relation_type='OBLIGATION_SETTLEMENT')]),measurements=ms)
    assert len(r['events'])==2 and len(r['measurements'])==1
    assert r['reconciliation'].iloc[0]['status']=='MATCH'


def test_repeated_snapshots_do_not_multiply_balance():
    ms=pd.DataFrame([dict(measurement_id=f'm{i}',case_id='R1',metric='COMMITMENT_BALANCE',observed_at=f'2026-01-{10+i:02d}',
                          amount='100',currency='EUR',source='NTSW',document='screen') for i in range(3)])
    r=run(event('c','COMMITMENT','100'),measurements=ms)
    assert r['summary'].iloc[0]['commitment_remaining']==D(100)
    assert len(r['reconciliation'])==1


def test_reversal_is_immutable_opposite_cash_event():
    p=event('p','PAYMENT','40')
    rev=event('r','REVERSAL','40',reversal_of='p',from_account='EXTERNAL:supplier',to_account='OWN:bank')
    r=run(event('o','OPENING','100'),p,rev)
    assert r['summary'].iloc[0]['paid']==0
    assert r['accounts'].iloc[0]['closing']==D(100)


def test_bad_reversal_is_quarantined():
    r=run(event('p','PAYMENT','40'),event('r','REVERSAL','40',reversal_of='p',from_account='OWN:bank',to_account='EXTERNAL:supplier'))
    assert 'INVALID_REVERSAL' in codes(r) and 'r' not in set(r['events']['event_id'])


def test_source_event_id_replay_and_conflict():
    a=event('a','PAYMENT','10',source_event_id='bank-1',document='BANK-DOC-1')
    b=event('b','PAYMENT','10',source_event_id='bank-1',document='BANK-DOC-1')
    r=run(a,b);assert len(r['events'])==1 and 'SOURCE_REPLAY_REMOVED' in codes(r)
    r=run(a,event('b','PAYMENT','11',source_event_id='bank-1'))
    assert r['events'].empty and 'CONFLICTING_SOURCE_EVENT_ID' in codes(r)


def test_cross_currency_obligation_needs_basis_and_authorization():
    r=run(event('c','COMMITMENT','100','EUR'),event('z','SETTLEMENT','110','USD'),
          links=pd.DataFrame([link('ob','c','z','100','110',relation_type='OBLIGATION_SETTLEMENT')]))
    assert r['links'].empty and 'UNSUPPORTED_CROSS_CURRENCY_SETTLEMENT' in codes(r)
    r=run(event('c','COMMITMENT','100','EUR'),event('z','SETTLEMENT','110','USD'),
          links=pd.DataFrame([link('ob','c','z','100','110',relation_type='OBLIGATION_SETTLEMENT',authorization='BANK-OK',conversion_basis='BANK-CERT')]))
    assert len(r['links'])==1 and r['summary'].query("currency == 'EUR'").iloc[0]['commitment_remaining']==0


def test_toman_is_never_silently_treated_as_rial():
    r=run(event('p','PAYMENT','100','IRT'))
    assert r['events'].empty and 'AMBIGUOUS_CURRENCY_UNIT' in codes(r)


def test_currencies_never_added():
    r=run(event('p','PAYMENT','100','EUR'),event('q','PAYMENT','200','USD'))
    assert len(r['summary'])==2


def test_refund_not_automatic_commitment_release():
    r=run(event('c','COMMITMENT'),event('r','REFUND','40'))
    assert r['summary'].iloc[0]['commitment_remaining']==D(100)


def test_allocation_and_quota_excess():
    r=run(event('a','ALLOCATION'),event('u','ALLOCATION_USE','110'),event('q','QUOTA','50'),event('v','QUOTA_USE','60'))
    assert {'EXCESS_ALLOCATION','EXCESS_QUOTA'}<=codes(r)


def test_fx_requires_both_legs():
    r=run(event('x','FX_BUY',group_id='G'));assert len(r['events'])==1
    assert r['conversions'].empty and 'INCOMPLETE_FX_PAIR' in codes(r)
    assert r['accounts'].iloc[0]['inflow']==D(100)
    r=run(event('a','FX_SELL','110','USD',group_id='G',to_account='EXTERNAL:bank'),event('b','FX_BUY','100','EUR',group_id='G'))
    assert len(r['events'])==2


def test_links_split_payment_and_preserve_capacity():
    r=run(event('f','FUNDING'),event('p','PAYMENT'),event('b1','SHIPMENT','60'),event('b2','SHIPMENT','40'),
        links=pd.DataFrame([link('1','f','p','100','100'),link('2','p','b1'),link('3','p','b2','40','40')]))
    assert len(r['links'])==3
    assert r['summary'].iloc[0]['unmatched_payment']==0
    assert r['summary'].iloc[0]['untraced_payment']==0


def test_overallocation_never_clamped_or_hidden():
    r=run(event('p','PAYMENT'),event('b1','SHIPMENT','60'),event('b2','SHIPMENT','60'),
          links=pd.DataFrame([link('1','p','b1'),link('2','p','b2')]))
    assert r['links'].empty and 'OVERALLOCATED_LINK' in codes(r)


def test_cross_case_needs_authorization():
    r=run(event('p','PAYMENT'),event('b','SHIPMENT',case_id='R2'),links=pd.DataFrame([link('1','p','b')]))
    assert r['links'].empty and 'UNAUTHORIZED_REALLOCATION' in codes(r)


def test_rates_historical_no_future_no_mixed_purpose():
    rates=pd.DataFrame([dict(date='2026-01-02',base='EUR',quote='IRR',rate='9',purpose='accounting',approved='true',source='ref',max_age_days='0'),
                        dict(date='2026-01-01',base='EUR',quote='IRR',rate='7',purpose='customs',approved='true',source='ref',max_age_days='0')])
    r=run(event('p','PAYMENT'),rates=rates)
    assert r['valuation'].iloc[0]['value'] is None


def test_exact_documented_fx_valuation():
    rates=pd.DataFrame([dict(date='2026-01-01',base='EUR',quote='IRR',rate='7',purpose='accounting',approved='true',source='ref',max_age_days='0')])
    r=run(event('p','PAYMENT'),rates=rates);assert r['valuation'].iloc[0]['value']==D(700)


def test_rate_conflict_blocks_valuation():
    rates=pd.DataFrame([dict(date='2026-01-01',base='EUR',quote='IRR',rate=v,purpose='accounting',approved='true',source='ref',max_age_days='0') for v in ['7','8']])
    assert run(event('p','PAYMENT'),rates=rates)['valuation'].iloc[0]['value'] is None


def test_old_pdf_rule_not_applied():
    r=run(event('c','COMMITMENT',due_date='2026-01-15',rule_id='old'),rules=pd.DataFrame([dict(rule_id='old',kind='COMMITMENT',source='pdf',approved='true',effective_from='2024-01-01',effective_to='2024-12-31')]))
    assert 'تأیید نشده' in r['deadlines'].iloc[0]['basis']


def test_of_missing_payment_ids_never_becomes_cash():
    """OF بدون شماره پرداخت هرگز نباید «گردش نقدی» شود.

    نسخهٔ قبلی این تست با «هیچ رویدادی پذیرفته نشود» این را می‌سنجید. اندازه‌گیری
    روی فایل واقعی OF نشان داد آن قاعده ۱۰۰٪ ردیف‌ها را با کد UNVERIFIED_EVIDENCE
    بیرون می‌اندازد و گزارش جریان وجوه کاملاً خالی می‌شود — یعنی تنها منبع موجود
    کاربر بی‌استفاده می‌ماند. بنابراین همان ناوردای واقعی سنجیده می‌شود، نه یک
    نمایندهٔ سخت‌گیرانه‌تر از آن: ردیف OF می‌تواند «شاهد منبع» باشد ولی
    هیچ‌وقت نباید حساب بانکی، گردش حساب یا مانده حساب بسازد.
    """
    source=pd.DataFrame([{'CB No.':'R1','Payment':100,'Fx Payment':'EUR','Date of Buying Currency':'1402/07/03'}]*3)
    ev=legacy_of(source)
    assert set(ev['status'])=={'SOURCE_FACT'}
    assert not ev['from_account'].fillna('').astype(str).str.strip().any()
    assert not ev['to_account'].fillna('').astype(str).str.strip().any()
    assert all('شماره پرداخت' not in str(d) for d in ev['document'])
    r=build_cashflow(ev)
    assert r['movements'].empty and r['accounts'].empty
    paid=r['events'][r['events']['kind']=='PAYMENT']
    assert len(paid)==1                      # سه ردیف یکسان، یک شناسه پایدار
    assert not str(paid.iloc[0]['from_account'] or '').strip()
    assert not str(paid.iloc[0]['to_account'] or '').strip()
    assert 'گزارش مشتق' in paid.iloc[0]['note']
    assert (r['issues']['code']=='ACCOUNT_DETAIL_GAP').any()


def test_html_and_excel_safely_escape_and_keep_long_ids():
    from openpyxl import load_workbook
    r=run(event('999999999999999999','PAYMENT',document='=1+2',note='</script><script>BAD()</script>'))
    x=excel_bytes(r);h=html_report(r,x)
    assert '</script><script>BAD()' not in h and '&lt;/script&gt;' in h
    w=load_workbook(io.BytesIO(x));ws=w['رویدادهای پذیرفته‌شده']
    assert ws.cell(2,1).value=='999999999999999999' and ws.cell(2,1).data_type=='s'
    assert all(c.data_type!='f' for row in ws for c in row)


def test_transfer_does_not_inflate_period_cash():
    r=run(event('t','TRANSFER',from_account='OWN:a',to_account='OWN:b'))
    assert r['periods'].empty
    assert sum(r['movements']['signed_amount'])==0


def test_cash_links_need_actual_account_and_currency_path():
    r=run(event('f','FUNDING','100','EUR'),event('p','PAYMENT','100','USD'),links=pd.DataFrame([link('1','f','p')]))
    assert r['links'].empty and 'ACCOUNT_PATH_GAP' in codes(r)


def test_rate_regime_must_match():
    rates=pd.DataFrame([dict(date='2026-01-01',base='EUR',quote='IRR',rate='7',purpose='accounting',approved='true',source='ref',max_age_days='0',regime='commercial')])
    assert run(event('p','PAYMENT'),rates=rates)['valuation'].iloc[0]['value'] is None
    assert run(event('p','PAYMENT',rate_regime='commercial'),rates=rates)['valuation'].iloc[0]['value']==D(700)


def test_invalid_report_date_rejected():
    with pytest.raises(ValueError):build_cashflow(pd.DataFrame(),as_of='broken')


def test_partial_bl_and_customs_currency_remain_distinct():
    r=run(event('p','PAYMENT'),event('b','SHIPMENT','80',bl_id='BL1'),event('c','CUSTOMS','90','USD',bl_id='BL1'),links=pd.DataFrame([link('1','p','b','60','60')]))
    assert len(r['documents'])==2
    eur=r['documents'].query("currency == 'EUR'").iloc[0]
    assert eur['unfunded_shipment']==D(20)
    assert 'UNMATCHED_PAYMENT' in codes(r)


def test_dashboard_financial_panel_build_and_scope(tmp_path):
    from streamlit.testing.v1 import AppTest
    from unittest.mock import patch
    from pathlib import Path
    appfile=tmp_path/'panel.py'
    appfile.write_text("import pandas as pd\nfrom gsi.cashflow.ui import render\nrender(pd.DataFrame({'KEY_REG':['R1']}),{},'2026-01-31')\n")
    source=pd.DataFrame([event('one','PAYMENT'),event('two','PAYMENT',case_id='R2')]).to_csv(index=False).encode()
    def upload(*args,**kwargs):
        if kwargs.get('key')=='cashflow_input':
            b=io.BytesIO(source);b.name='Events.csv';return b
        return None
    with patch('streamlit.file_uploader',side_effect=upload):
        app=AppTest.from_file(str(appfile)).run()
        app.button(key='cashflow_build').click().run()
        assert not app.exception
        cache=app.session_state['cashflow_result']
        assert cache[1]['events']['case_id'].tolist()==['R1']
    assert len(app.get('download_button'))==4
