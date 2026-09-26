import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.process_intelligence.core import build_bundle, profile
from gsi.process_intelligence.patterns import ActivityTaxonomy, compile_spec, match_cases


def _events():
    return pd.DataFrame({
        "_CASE_KEY": ["001", "001", "001", "002", "002", "003", "003"],
        "ACTIVITY_FA": ["ثبت سفارش", "تخصیص ارز", "تامین ارز", "ثبت سفارش", "تخصیص ارز", "ثبت سفارش", "ابطال"],
        "EVENTTIME": [
            "2026-09-01T08:00:00+03:30", "2026-09-02T08:00:00+03:30", "2026-09-04T08:00:00+03:30",
            "2026-09-01T08:00:00+03:30", "2026-09-03T08:00:00+03:30",
            "2026-09-01T08:00:00+03:30", "2026-09-01T12:00:00+03:30",
        ],
    })


def _spec():
    return {
        "version": 1,
        "taxonomy": {
            "FX": ["تخصیص ارز", "تامین ارز"],
            "FINANCE": ["@FX", "رفع تعهد"],
        },
        "patterns": [{
            "id": "reg-to-fx",
            "label": "ورود از ثبت سفارش به زنجیره ارز",
            "constraints": [
                {"activity": "ثبت سفارش", "capture": "start"},
                {"taxonomy": "FX", "quantifier": "one_or_more", "capture": "fx"},
            ],
        }],
    }


def test_taxonomy_recursive_and_cycle_rejected():
    t = ActivityTaxonomy(_spec()["taxonomy"])
    assert t.contains("FINANCE", "تخصیص ارز")
    with pytest.raises(ValueError):
        ActivityTaxonomy({"A": ["@B"], "B": ["@A"]})


def test_pattern_greedy_capture_and_source_lineage():
    facts = profile(_events(), pattern_spec=_spec())
    p = facts["patterns"]
    assert p["enabled"] and p["status"] == "observational_only"
    assert p["summary"][0]["unique_cases"] == 2
    first = next(m for m in p["matches"] if m["case_key"] == "001")
    assert first["activities"] == ["ثبت سفارش", "تخصیص ارز", "تامین ارز"]
    assert first["captures"]["fx"]["activities"] == ["تخصیص ارز", "تامین ارز"]
    assert first["source_rows"] == [2, 3, 4]
    assert "ریسک" in p["note"]


def test_wildcard_optional_exclusion_and_bad_spec():
    rows = [
        {"_CASE_KEY": "1", "ACTIVITY_FA": "ثبت سفارش", "source_row": 2},
        {"_CASE_KEY": "1", "ACTIVITY_FA": "کنترل اسناد", "source_row": 3},
        {"_CASE_KEY": "1", "ACTIVITY_FA": "تخصیص ارز", "source_row": 4},
    ]
    spec = {"version": 1, "patterns": [{"id": "x", "constraints": [
        {"activity": "ثبت*"}, {"any": True, "quantifier": "optional", "exclude_activities": ["ابطال"]},
        {"activity": "*ارز"},
    ]}]}
    assert match_cases(rows, spec)["summary"][0]["match_count"] == 1
    with pytest.raises(ValueError):
        compile_spec({"version": 1, "patterns": [{"id": "x", "constraints": [{"activity": "A", "taxonomy": "T"}]}]})


def test_transition_metrics_keep_unique_cases_and_occurrences_separate():
    f = profile(_events())
    edge = next(e for e in f["graph"]["edges"] if e["source"] == "n2" or True)
    # Find by labels instead of relying on lexical node id assignment.
    labels = {n["id"]: n["label"] for n in f["graph"]["nodes"]}
    reg_fx = next(e for e in f["graph"]["edges"]
                  if labels[e["source"]] == "ثبت سفارش" and labels[e["target"]] == "تخصیص ارز")
    assert reg_fx["count"] == 2 and reg_fx["occurrences"] == 2
    assert reg_fx["median_hours"] == 36.0
    assert reg_fx["p90_hours"] == pytest.approx(45.6)


def test_handoff_fail_closed_and_ready():
    f = profile(_events())
    assert f["handoff"]["status"] == "not_computable"
    d = _events().copy()
    d["FROM_TEAM"] = ["خرید", "", "", "خرید", "", "", ""]
    d["TO_TEAM"] = ["ارز", "", "", "ارز", "", "", ""]
    d["SENT_AT"] = ["2026-09-01T08:00:00+03:30", "", "", "2026-09-01T08:00:00+03:30", "", "", ""]
    d["ACCEPTED_AT"] = ["2026-09-01T10:00:00+03:30", "", "", "2026-09-01T12:00:00+03:30", "", "", ""]
    d["DOCUMENT_REF"] = ["DOC-1", "", "", "DOC-2", "", "", ""]
    h = profile(d)["handoff"]
    assert h["status"] == "ready" and h["summary"]["unique_cases"] == 2
    assert h["summary"]["median_acceptance_hours"] == 3.0
    d.loc[0, "ACCEPTED_AT"] = "2026-08-31T08:00:00+03:30"
    with pytest.raises(ValueError):
        profile(d)


def test_bundle_has_three_views_analysis_and_safe_excel():
    d = _events().copy()
    d.loc[0, "ACTIVITY_FA"] = "=1+1<script>alert(1)</script>"
    files = build_bundle(d, pattern_spec=_spec())
    assert {"report.html", "report.xlsx", "evidence.json", "process-analysis.json", "manifest.json"} <= set(files)
    html = files["report.html"].decode("utf-8")
    assert "۱) وضعیت و اقدام" in html and "۲) مسیر فرآیند" in html and "۳) تحویل بین واحدها" in html
    assert "<script>" not in html and 'src="http' not in html and 'href="http' not in html
    assert "data:font/" not in html and "IRANSansWeb" in html
    analysis = json.loads(files["process-analysis.json"])
    assert analysis["patterns"]["enabled"]
    wb = load_workbook(io.BytesIO(files["report.xlsx"]))
    assert "وضعیت و اقدام" in wb.sheetnames and "خلاصه الگوها" in wb.sheetnames
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
