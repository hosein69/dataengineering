# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import html
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Tuple


class _TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: List[str] = []
    def handle_data(self, data):
        if data and data.strip():
            self.parts.append(data.strip())


def _clean(text: str) -> str:
    text = str(text or "").replace("\x00", " ")
    # Persian/Arabic orthography normalization improves offline FTS recall.
    text = text.translate(str.maketrans({"ي":"ی","ى":"ی","ك":"ک","ة":"ه","ۀ":"ه","ؤ":"و","إ":"ا","أ":"ا"}))
    text = text.replace("\u200c", " ")
    text = re.sub(r"[\u064b-\u065f\u0670]", "", text)
    text = re.sub(r"[\t\r ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: Path) -> Tuple[str, Dict]:
    """Extract text without requiring heavyweight services.

    PDF/DOCX support is opportunistic: if pypdf/python-docx are installed they
    are used, otherwise that file is quarantined instead of breaking the index.
    """
    p = Path(path)
    ext = p.suffix.lower()
    meta = {"name": p.name, "ext": ext}
    if ext in {".txt", ".md", ".log", ".sql", ".py"}:
        return _clean(p.read_text(encoding="utf-8", errors="replace")), meta
    if ext in {".htm", ".html"}:
        hp = _TextHTMLParser(); hp.feed(p.read_text(encoding="utf-8", errors="replace"))
        return _clean("\n".join(hp.parts)), meta
    if ext == ".json":
        obj = json.loads(p.read_text(encoding="utf-8", errors="replace"))
        return _clean(json.dumps(obj, ensure_ascii=False, indent=2)), meta
    if ext == ".csv":
        lines=[]
        with p.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
            for i,row in enumerate(csv.reader(f)):
                lines.append(" | ".join(map(str,row)))
                if i >= 5000: break
        return _clean("\n".join(lines)), meta
    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        out=[]
        for ws in wb.worksheets:
            out.append(f"[Sheet: {ws.title}]")
            for i,row in enumerate(ws.iter_rows(values_only=True)):
                vals=["" if v is None else str(v) for v in row]
                if any(x.strip() for x in vals): out.append(" | ".join(vals))
                if i >= 3000: break
        wb.close()
        return _clean("\n".join(out)), meta
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except Exception as ex:
            raise RuntimeError("PDF parser unavailable (optional package pypdf)") from ex
        reader=PdfReader(str(p)); pages=[]
        for i,page in enumerate(reader.pages):
            pages.append(f"[Page {i+1}]\n{page.extract_text() or ''}")
        return _clean("\n".join(pages)), meta
    if ext == ".docx":
        try:
            from docx import Document
        except Exception as ex:
            raise RuntimeError("DOCX parser unavailable (optional package python-docx)") from ex
        d=Document(str(p)); parts=[x.text for x in d.paragraphs if x.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return _clean("\n".join(parts)), meta
    raise RuntimeError(f"unsupported extension: {ext or '(none)'}")


def chunk_text(text: str, max_chars: int = 1400, overlap: int = 180) -> List[str]:
    text=_clean(text)
    if not text: return []
    paras=[p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks=[]; buf=""
    for para in paras:
        candidate=(buf+"\n\n"+para).strip() if buf else para
        if len(candidate) <= max_chars:
            buf=candidate; continue
        if buf: chunks.append(buf)
        if len(para) <= max_chars:
            buf=para
        else:
            start=0
            while start < len(para):
                chunks.append(para[start:start+max_chars])
                start += max(1, max_chars-overlap)
            buf=""
    if buf: chunks.append(buf)
    return chunks
