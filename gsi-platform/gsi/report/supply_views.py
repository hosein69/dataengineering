# -*- coding: utf-8 -*-
"""سه نمای تصمیم‌ساز زنجیره تأمین: متریال، بارنامه و اداره.

اصل طراحی: **«الان کجاست؟ از کِی؟ دست کیست؟ معطل چه کسی است؟ اگر
نمی‌دانیم، دقیقاً چه داده‌ای کم است و از چه کسی باید خواست؟»**

هیچ علت یا کارشناسی از روی نبودِ داده حدس زده نمی‌شود.

## دو بازنویسی نسبت به نسخه ۲۶٫۹

**۱) موقعیت از رویداد می‌آید، نه از موجودی.** نسخه قبل موقعیت را از
ستون‌های موجودی حدس می‌زد (`موجودی در گمرک > 0 ⇒ گمرک`). آن حدس تاریخ
ندارد، پس به سؤال «از کِی؟» جواب نمی‌دهد. حالا ماژول ``resolve.part_status``
آخرین رویدادِ تاریخ‌دار را می‌یابد و «کجا/کِی/چه کسی» از همان می‌آید.
اگر آن ستون‌ها موجود نباشند (مثلاً فراخوانی مستقیم از Studio روی یک
زیرمجموعه)، همان‌جا محاسبه می‌شوند.

**۲) سرعت.** پیاده‌سازی قبلی برای هر ردیف یک ``pd.DataFrame([row])``
می‌ساخت و روی آن ``_first`` می‌زد. اندازه‌گیری واقعی روی همین ماشین:

    ۵۰۰ ردیف   → ۱۴٫۸ ثانیه
    ۲٬۰۰۰ ردیف → ۳۶٫۵ ثانیه

و این تابع سه بار در هر بازتولید Studio و یک بار در ``save()`` گزارش
رسمی صدا زده می‌شد. روی داده واقعی یعنی گزارش روزانه عملاً تمام نمی‌شد.
نسخه فعلی کاملاً برداری است (groupby + transform) و همان ۲٬۰۰۰ ردیف را
در کسری از ثانیه می‌سازد.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from ..resolve import part_status as ps
from ..resolve.expert_scope import OWNER_GAP, OWNER_NAME, SCOPES
from .palette import LuxuryPalette as P

RED = P.STATUS_CRITICAL      # از design.tokens؛ کنتراست ۷٫۴۳ روی سفید

_BLANKS = {"nan": "", "None": "", "NaT": "", "—": "", "-": ""}


def _s(df: pd.DataFrame, c: str) -> pd.Series:
    """ستون متنی تمیز؛ اگر نبود، ستون خالی هم‌طول."""
    if c not in df.columns:
        return pd.Series("", index=df.index, dtype=object)
    return df[c].fillna("").astype(str).str.strip().replace(_BLANKS)


def _num(df: pd.DataFrame, c: str) -> pd.Series:
    """ستون عددی؛ اگر نبود، ستون NaN هم‌طول.

    ``df.get(c)`` برای ستون غایب ``None`` می‌دهد و ``pd.to_numeric(None)``
    یک اسکالر برمی‌گرداند نه Series — و بعد ``.groupby`` روی آن می‌ترکد.
    """
    if c not in df.columns:
        return pd.Series(pd.NA, index=df.index, dtype="Float64")
    return pd.to_numeric(df[c], errors="coerce")


def _first_col(df: pd.DataFrame, cols, default: str = "نامشخص") -> pd.Series:
    """اولین ستون موجود که برای آن ردیف مقدار دارد — برداری."""
    out = pd.Series("", index=df.index, dtype=object)
    for c in cols:
        if c not in df.columns:
            continue
        out = out.where(out.ne(""), _s(df, c))
    return out.replace("", default)


def _group_mode(s: pd.Series, keys: pd.Series) -> pd.Series:
    """پرتکرارترین مقدار غیرخالی در هر گروه، پخش‌شده روی ردیف‌های گروه."""
    tmp = pd.DataFrame({"k": keys, "v": s})
    non = tmp[tmp["v"].ne("")]
    if non.empty:
        return pd.Series("", index=s.index, dtype=object)
    mode = non.groupby("k")["v"].agg(lambda x: x.value_counts().index[0])
    return keys.map(mode).fillna("")


def _ensure_status(df: pd.DataFrame) -> pd.DataFrame:
    """ستون‌های وضعیت را تضمین می‌کند؛ اگر نبودند، همین‌جا می‌سازد."""
    if ps.WHERE in df.columns:
        return df
    return ps.resolve(df.copy())


def _base(df: pd.DataFrame, group_col: str, kind: str) -> pd.DataFrame:
    if group_col not in df.columns:
        return pd.DataFrame()
    work = df[_s(df, group_col).ne("")].copy()
    if work.empty:
        return pd.DataFrame()
    work = _ensure_status(work)
    raw_keys = _s(work, group_col)
    # The pipeline's KEY_MATERIAL is already the canonical display identity.
    # Keep it verbatim here; search normalization belongs to the search layer,
    # not the audit/report display layer.
    keys = raw_keys

    where = _s(work, ps.WHERE).replace("", "نامشخص")
    when = _s(work, ps.WHEN)
    age = _num(work, ps.AGE)

    desc = (_first_col(work, ["MATERIAL_DESC", "CANONICAL_GOODS_DESC"], "")
            if kind == "متریال" else pd.Series("", index=work.index))
    if kind == "متریال":
        # The visible group title is the material identity.  Description is a
        # separate attribute and is resolved *within the exact material key* so
        # a stale/fanned-out description can never impersonate another material.
        title = keys
        group_desc = _group_mode(desc, keys).where(keys.ne(""), "")
        desc_n = (pd.DataFrame({"k": keys, "d": desc})
                  .loc[lambda x: x["d"].ne("")]
                  .groupby("k")["d"].nunique())
        desc_conflict = keys.map(desc_n).fillna(0).gt(1)
        group_size = keys.map(keys.value_counts()).fillna(0).astype(int)
        order_n = keys.map(pd.DataFrame({"k": keys, "v": _s(work, "CANONICAL_ORDER")})
                           .loc[lambda x: x["v"].ne("")]
                           .groupby("k")["v"].nunique()).fillna(0).astype(int)
        bl_n = keys.map(pd.DataFrame({"k": keys, "v": _s(work, "CANONICAL_BL")})
                        .loc[lambda x: x["v"].ne("")]
                        .groupby("k")["v"].nunique()).fillna(0).astype(int)
    else:
        title = keys
        group_desc = pd.Series("", index=work.index, dtype=object)
        desc_conflict = pd.Series(False, index=work.index)
        group_size = pd.Series(1, index=work.index, dtype=int)
        order_n = pd.Series(0, index=work.index, dtype=int)
        bl_n = pd.Series(0, index=work.index, dtype=int)

    out = pd.DataFrame({
        "کلید گروه": keys,
        "عنوان گروه": title,
        "شرح متریال": group_desc if kind == "متریال" else "",
        "اختلاف شرح متریال": desc_conflict.map({True: "⚠ چند شرح برای یک متریال", False: ""}) if kind == "متریال" else "",
        "تعداد ردیف تفصیلی گروه": group_size if kind == "متریال" else 1,
        "تعداد سفارش یکتای گروه": order_n if kind == "متریال" else 0,
        "تعداد بارنامه یکتای گروه": bl_n if kind == "متریال" else 0,
        "متریال": keys if kind == "متریال" else _first_col(work, ["KEY_MATERIAL"]),
        "سفارش": _first_col(work, ["CANONICAL_ORDER"]),
        "ثبت سفارش": _first_col(work, ["KEY_REG"]),
        "بارنامه": _first_col(work, ["CANONICAL_BL"]),
        # ── کجا / کِی / چه کسی ──
        "موقعیت فعلی": where,
        "آخرین فعالیت": _s(work, ps.ACTIVITY).replace("", "—"),
        "تاریخ آخرین رویداد": when.replace("", "—"),
        "سن وضعیت (روز)": age,
        "کارشناس مسئول وضعیت": _s(work, ps.WHO).replace("", "—"),
        "حوزه مسئول وضعیت": _s(work, ps.WHO_SCOPE).replace("", "—"),
        "معطل حوزه": _s(work, ps.WAITING_SCOPE).replace("", "—"),
        "معطل کارشناس": _s(work, ps.WAITING_WHO).replace("", "—"),
        "فعالیت بعدی مورد انتظار": _s(work, ps.NEXT_ACT).replace("", "—"),
        # ── مالکیت ──
        "مالک قطعه (کارشناس خرید)": _s(work, OWNER_NAME).replace("", "شناسایی نشد"),
        "شکاف داده مالک": _s(work, OWNER_GAP),
        # ── چرایی و شواهد ──
        "چرایی / مبنای وضعیت": _reason(work),
        "شواهد تعیین وضعیت": _s(work, ps.BASIS).replace("", "شاهد تاریخ‌دار ثبت نشده است"),
        "داده مفقود برای تعیین وضعیت": _s(work, ps.MISSING),
        # ── زمینه ──
        "مرحله فعلی": _first_col(work, ["STAGE_FA"]),
        "مدیریت": _first_col(work, ["ORG_DEPT"]),
        "بحرانی": _first_col(work, ["بحرانی (کوتاه)", "کد طبقه بحرانی"]),
        "مقاومت (روز)": _num(work, "مقاومت (روز)"),
        "روزهای رسوب": _num(work, "روزهای رسوب"),
    })
    for scope in SCOPES:
        out[scope.fa] = _s(work, scope.key).replace("", "—")
    return out.sort_values(["کلید گروه"], kind="stable", ignore_index=True)


def _reason(work: pd.DataFrame) -> pd.Series:
    """علت ثبت‌شده مقدم است؛ سپس علت داده‌ای و قابل ممیزی."""
    r = _first_col(work, ["BL_CRITICAL_REASON", "ORDER_CRITICAL_REASON", "روایت",
                          "شرح هشدارها", "ALERTS", "MOGH_COMMERCIAL_NOTE",
                          "MOGH_LOGISTICS_NOTE"], "")
    waiting = _s(work, ps.WAITING_SCOPE)
    nxt = _s(work, ps.NEXT_ACT)
    derived = pd.Series("علت اختصاصی در داده‌های موجود ثبت نشده است.",
                        index=work.index, dtype=object)
    has_next = waiting.ne("") & nxt.ne("")
    derived = derived.mask(
        has_next, "منتظر «" + nxt + "» است؛ این کار بر عهده " + waiting + " است.")
    no_evidence = _s(work, ps.WHERE).isin(["", "نامشخص"])
    derived = derived.mask(
        no_evidence,
        "هیچ رویداد تاریخ‌داری ثبت نشده است؛ تا ثبت اولین تاریخ، موقعیت قابل تعیین نیست.")
    return r.where(r.ne(""), derived)


def _join_nonempty(parts: list[pd.Series], sep: str = " ؛ ") -> pd.Series:
    """Join row-wise advisory text without inventing values."""
    if not parts:
        return pd.Series(dtype=object)
    idx = parts[0].index
    out = pd.Series("", index=idx, dtype=object)
    for part in parts:
        x = part.fillna("").astype(str).str.strip()
        add = x.ne("")
        both = add & out.ne("")
        out = out.mask(both, out + sep + x)
        out = out.mask(add & out.eq(""), x)
    return out


def _expert_advisory(df: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    """Commercial Expert Data is advisory-only in the HTML material dashboard.

    It may raise a warning and explain it, but it must not become the material
    dashboard's position/owner/status authority.
    """
    idx = df.index
    flags: list[pd.Series] = []
    notes: list[pd.Series] = []

    missing = (_s(df, "ORDER_MISSING_COMMERCIAL_EXPERT")
               .str.lower().isin({"true", "1", "yes", "بله"}))
    if missing.any():
        flags.append(pd.Series("", index=idx).mask(missing, "⚠ سفارش در سورس کارشناسان یافت نشد"))
        reason = _s(df, "ORDER_MISSING_COMMERCIAL_REASON")
        notes.append(reason.where(missing, ""))

    state = _s(df, "COMMERCIAL_COVERAGE_STATE")
    unavailable = state.isin({"source_unavailable", "source_schema_gap"})
    if unavailable.any():
        flags.append(pd.Series("", index=idx).mask(unavailable, "⚠ پوشش سورس کارشناسان قابل سنجش نیست"))
        notes.append(pd.Series("", index=idx).mask(
            state.eq("source_schema_gap"), "ساختار Commercial Expert Data با قرارداد مورد انتظار منطبق نیست"
        ).mask(state.eq("source_unavailable"), "Commercial Expert Data در این اجرا در دسترس نبوده است"))

    gap = _s(df, OWNER_GAP)
    has_gap = gap.ne("")
    if has_gap.any():
        flags.append(pd.Series("", index=idx).mask(has_gap, "⚠ داده کارشناسی ناقص"))
        notes.append(pd.Series("", index=idx).mask(has_gap, "کمبود ثبت کارشناس: " + gap))

    raw_alert = _s(df, "MOGH_ALERTS")
    if raw_alert.ne("").any():
        flags.append(pd.Series("", index=idx).mask(raw_alert.ne(""), "⚠ هشدار ثبت‌شده در سورس کارشناسان"))
        notes.append(raw_alert)

    for col, label in (("MOGH_COMMERCIAL_NOTE", "کامنت بازرگانی"),
                       ("MOGH_LOGISTICS_NOTE", "کامنت لجستیک"),
                       ("MOGH_INVENTORY_NOTE", "کامنت موجودی")):
        x = _s(df, col)
        notes.append(pd.Series("", index=idx).mask(x.ne(""), label + ": " + x))

    # A mismatch is only a warning; it never overwrites the dashboard identity.
    expert_mat = _first_col(df, ["MOGH_MATERIAL", "MOGH_MFR_PART_NO"], "").astype(str).str.strip()
    ref_mat = _first_col(df, ["ORC_PART_NO", "KEY_MATERIAL"], "").astype(str).str.strip()
    mismatch = expert_mat.ne("") & ref_mat.ne("") & expert_mat.ne(ref_mat)
    if mismatch.any():
        flags.append(pd.Series("", index=idx).mask(mismatch, "⚠ مغایرت کد متریال کارشناسان با مرجع"))
        notes.append(pd.Series("", index=idx).mask(
            mismatch, "متریال اعلامی کارشناس=" + expert_mat + " | متریال مرجع=" + ref_mat))

    alert = _join_nonempty(flags) if flags else pd.Series("", index=idx, dtype=object)
    comment = _join_nonempty(notes) if notes else pd.Series("", index=idx, dtype=object)
    return alert, comment


def _ntsw_advisory(df: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    """NTSW contributes only warning/comment to the HTML material dashboard."""
    idx = df.index
    flags: list[pd.Series] = []
    notes: list[pd.Series] = []

    bal = _num(df, "NTSW_BALANCE")
    open_bal = bal.fillna(0).gt(0)
    if open_bal.any():
        flags.append(pd.Series("", index=idx).mask(open_bal, "⚠ مانده تعهد NTSW باز است"))

    rel = _s(df, "NTSW_RELEASE_STATUS")
    open_rel = rel.str.contains("نشده|باز", regex=True, na=False)
    if open_rel.any():
        flags.append(pd.Series("", index=idx).mask(open_rel, "⚠ رفع تعهد NTSW باز است"))

    alloc = _s(df, "NTSW_ALLOC_STATUS")
    alloc_bad = alloc.str.contains("رد|پذیرفته نشده|ابطال", regex=True, na=False)
    if alloc_bad.any():
        flags.append(pd.Series("", index=idx).mask(alloc_bad, "⚠ وضعیت تخصیص NTSW نیازمند بررسی است"))

    cur = _s(df, "NTSW_CURRENCY")
    deadline = _s(df, "NTSW_DEADLINE")
    process = _s(df, "NTSW_ALLOC_PROCESS")
    req_amount = _num(df, "NTSW_REQ_AMOUNT")

    parts = []
    for label, ser in (("وضعیت تخصیص", alloc), ("فرآیند تخصیص", process),
                       ("وضعیت رفع تعهد", rel), ("مهلت رفع تعهد", deadline)):
        parts.append(pd.Series("", index=idx).mask(ser.ne(""), label + ": " + ser))
    bal_text = bal.map(lambda v: "" if pd.isna(v) else f"مانده تعهد: {float(v):,.2f}")
    req_text = req_amount.map(lambda v: "" if pd.isna(v) else f"مبلغ درخواست: {float(v):,.2f}")
    cur_text = pd.Series("", index=idx).mask(cur.ne(""), "ارز: " + cur)
    parts.extend([bal_text, req_text, cur_text])

    alert = _join_nonempty(flags) if flags else pd.Series("", index=idx, dtype=object)
    comment = _join_nonempty(parts)
    return alert, comment


def build_material_html_view(df: pd.DataFrame, today=None) -> pd.DataFrame:
    """Material supply view for the standalone HTML dashboard.

    Contract: Commercial Expert Data and NTSW are *advisory-only* here.  Their
    raw fields do not determine current position, owner, waiting person, group
    identity, row count or stage.  They surface only as warning/comment fields.
    The Streamlit supply view remains unchanged.
    """
    if "KEY_MATERIAL" not in df.columns or df.empty:
        return pd.DataFrame()

    # Source-prefixed columns identify the real merged pipeline, rather than a
    # caller-supplied, already-independent operational dataframe.
    work = df.copy().reset_index(drop=True)
    exp_alert, exp_comment = _expert_advisory(work)
    ntsw_alert, ntsw_comment = _ntsw_advisory(work)
    advisory = pd.DataFrame({
        "هشدار کارشناسان": exp_alert, "کامنت کارشناسان": exp_comment,
        "هشدار NTSW": ntsw_alert, "کامنت NTSW": ntsw_comment,
    })
    raw_expert = any(str(c).startswith("MOGH_") for c in work.columns)
    # A prefixed material is evidence of the legacy expert -> Oracle bridge.
    # Oracle confirms the material, NOT its association with an order/BL.
    expert_bridge = any(c in work for c in ("MOGH_MATERIAL", "MOGH_MFR_PART_NO"))
    if expert_bridge:
        work["KEY_MATERIAL"] = _s(work, "ORC_PART_NO")
        context = _join_nonempty([
            _s(df.reset_index(drop=True), c).map(lambda x: label + x if x else "")
            for c, label in (("CANONICAL_ORDER", "سفارش اعلامی: "),
                             ("CANONICAL_BL", "بارنامه اعلامی: "))
        ])
        advisory["کامنت کارشناسان"] = _join_nonempty([
            advisory["کامنت کارشناسان"], context])
        advisory["هشدار کارشناسان"] = _join_nonempty([
            advisory["هشدار کارشناسان"],
            pd.Series("⚠ ارتباط متریال با سفارش/بارنامه فقط از کارشناسان است؛ تأیید عملیاتی نشده", index=work.index)])
        # There is no independent material/order link in this adapter contract.
        # Never attach a real BL event to a material using an advisory-only link.
        for c in ("CANONICAL_ORDER", "CANONICAL_BL", "KEY_REG"):
            work[c] = ""
        for c, *_ in ps.TIMELINE:
            work[c] = ""

    # Rebuild aliases from allowed raw candidates; blanking only NTSW dates
    # leaves MOGH_PO_SENT_DATE -> PO_SENT_DATE and mixed descriptions intact.
    if raw_expert:
        from ..stages.s20_derive import DERIVED
        for target, (candidates, default, numeric) in DERIVED.items():
            if not any(c.startswith(("MOGH_", "NTSW_")) for c in candidates):
                continue
            allowed = [c for c in candidates if not c.startswith(("MOGH_", "NTSW_"))]
            work[target] = _first_col(work, allowed, "")
        work["MATERIAL_DESC"] = _s(work, "ORC_MATERIAL_DESC")
        work["CANONICAL_GOODS_DESC"] = _s(work, "ORC_MATERIAL_DESC")

    # DROP excluded timeline columns: an empty column still becomes a missing
    # milestone and incorrectly assigns the next activity/waiting department.
    work = work.drop(columns=list(ps.OUTPUT_COLUMNS) +
                     ["NTSW_COMMIT_DATE", "NTSW_ALLOC_DATE"], errors="ignore")
    if raw_expert:
        work = work.drop(columns=["PO_SENT_DATE"], errors="ignore")
    for c in ("MOGH_COMMERCIAL_NOTE", "MOGH_LOGISTICS_NOTE", "MOGH_INVENTORY_NOTE",
              "MOGH_ALERTS", "ALERTS", "شرح هشدارها", "روایت",
              "BL_CRITICAL_REASON", "ORDER_CRITICAL_REASON"):
        work[c] = ""
    if expert_bridge:
        # These derived metrics can include expert inventory or order-level
        # NTSW risk. They cannot be presented as independent material facts.
        for c in ("بحرانی (کوتاه)", "کد طبقه بحرانی", "مقاومت (روز)",
                  "روزهای رسوب", "ORG_DEPT"):
            work[c] = ""

    if expert_bridge:
        from ..engines.criticality import CriticalityEngine
        engine = CriticalityEngine()
        warehouse = pd.DataFrame({
            "STOCK_IKCO": _num(work, "ORC_STOCK_IKCO"),
            "STOCK_SAPCO": _num(work, "ORC_STOCK_SAPCO"),
            "DAILY_NEED": _num(work, "ORC_DAILY_NEED"),
        }).astype(object).where(lambda x: x.notna(), None)
        cache = {}
        results = []
        for row in warehouse.to_dict(orient="records"):
            key = tuple(row.values())
            if key not in cache:
                cache[key] = engine.evaluate(row)
            results.append(cache[key])
        work["مقاومت (روز)"] = [r.resistance_warehouse for r in results]
        work["بحرانی (کوتاه)"] = [r.band_short for r in results]

    valid = _s(work, "KEY_MATERIAL").ne("")
    work = work.loc[valid].copy()
    advisory = advisory.loc[valid].copy()
    if work.empty:
        return pd.DataFrame()
    # Sort once with the SAME positional index for both frames, including when
    # callers supply duplicate indices or rows without a material key.
    order = work.assign(__key=_s(work, "KEY_MATERIAL")).sort_values("__key", kind="stable").index
    work = ps.resolve(work, today=today)
    out = _base(work, "KEY_MATERIAL", "متریال")
    for c in advisory:
        out[c] = advisory.loc[order, c].to_numpy()
    drop = ["کارشناس مسئول وضعیت", "معطل کارشناس", "مالک قطعه (کارشناس خرید)",
            "شکاف داده مالک"] + [s.fa for s in SCOPES]
    out = out.drop(columns=drop, errors="ignore")
    out["مرحله فعلی"] = out["موقعیت فعلی"]
    if expert_bridge:
        out["داده مفقود برای تعیین وضعیت"] = "ارتباط مستقل متریال با سفارش/بارنامه در داده موجود نیست"
        out["شواهد تعیین وضعیت"] = "هویت متریال: Oracle؛ اتصال کارشناسان صرفاً کنترلی است"
        out["چرایی / مبنای وضعیت"] = "بدون اتصال مستقل، رویداد سفارش یا بارنامه به متریال نسبت داده نمی‌شود."
        for c in ("معطل حوزه", "فعالیت بعدی مورد انتظار", "حوزه مسئول وضعیت"):
            out[c] = "نامشخص"
    # Advisory fan-out must not multiply operational rows. Aggregate distinct
    # comments on the operational identity, without summing inventory/amounts.
    counts = ["تعداد ردیف تفصیلی گروه", "تعداد سفارش یکتای گروه", "تعداد بارنامه یکتای گروه"]
    operational = [c for c in out if c not in list(advisory) + counts]
    def distinct(values):
        return " ؛ ".join(dict.fromkeys(str(v) for v in values if pd.notna(v) and str(v).strip()))
    out = out.groupby(operational, dropna=False, sort=False, as_index=False).agg(
        {c: distinct for c in advisory})
    out["تعداد ردیف تفصیلی گروه"] = out.groupby("متریال")["متریال"].transform("size")
    for target, source in ((counts[1], "سفارش"), (counts[2], "بارنامه")):
        out[target] = out.groupby("متریال")[source].transform(lambda x: x[~x.isin(["", "نامشخص"])].nunique())
    leading = ["متریال", "شرح متریال", "هشدار کارشناسان", "کامنت کارشناسان",
               "هشدار NTSW", "کامنت NTSW", "موقعیت فعلی", "مرحله فعلی"]
    return out[leading + [c for c in out if c not in leading]]


def build_material_view(df: pd.DataFrame) -> pd.DataFrame:
    return _base(df, "KEY_MATERIAL", "متریال")


def build_bl_view(df: pd.DataFrame) -> pd.DataFrame:
    return _base(df, "CANONICAL_BL", "بارنامه")


def build_dept_view(df: pd.DataFrame) -> pd.DataFrame:
    """نمای اداره — یک ردیف به ازای هر اداره/مدیریت."""
    key = "ORG_DEPT"
    if key not in df.columns:
        return pd.DataFrame()
    work = df[_s(df, key).ne("")].copy()
    if work.empty:
        return pd.DataFrame()
    work = _ensure_status(work)
    keys = _s(work, key)
    work = work.assign(_k=keys)

    where_mode = _group_mode(_s(work, ps.WHERE), keys)
    wait_mode = _group_mode(_s(work, ps.WAITING_SCOPE), keys)
    age = _num(work, ps.AGE)

    def nunique(col: str) -> pd.Series:
        if col not in work.columns:
            return pd.Series(0, index=work["_k"].unique())
        s = _s(work, col).replace("", pd.NA)
        return s.groupby(work["_k"]).nunique()

    crit = (_s(work, "کد طبقه بحرانی")
            .isin(["STOCKOUT", "CRITICAL", "BECOMING_CRITICAL"])
            .groupby(work["_k"]).sum())
    gap = _s(work, OWNER_GAP).ne("").groupby(work["_k"]).sum()
    ownerless = _s(work, OWNER_NAME).eq("").groupby(work["_k"]).sum()

    idx = pd.Index(sorted(work["_k"].unique()), name=None)
    out = pd.DataFrame({
        "کلید گروه": idx,
        "اداره / مدیریت": idx,
        "معاونت": _group_mode(_s(work, "ORG_VICE"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "مدیر": _group_mode(_s(work, "ORG_MANAGER"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "رئیس": _group_mode(_s(work, "ORG_HEAD"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "موقعیت غالب پرونده‌ها": where_mode.groupby(work["_k"]).first().reindex(idx).fillna("نامشخص"),
        "معطل غالب حوزه": wait_mode.groupby(work["_k"]).first().reindex(idx).replace("", "—").fillna("—"),
        "میانگین سن وضعیت (روز)": age.groupby(work["_k"]).mean().reindex(idx).round(1),
        "بیشترین سن وضعیت (روز)": age.groupby(work["_k"]).max().reindex(idx),
        "تعداد متریال": nunique("KEY_MATERIAL").reindex(idx).fillna(0).astype(int),
        "تعداد بارنامه": nunique("CANONICAL_BL").reindex(idx).fillna(0).astype(int),
        "تعداد سفارش": nunique("CANONICAL_ORDER").reindex(idx).fillna(0).astype(int),
        "تعداد بحرانی": crit.reindex(idx).fillna(0).astype(int),
        "ردیف با شکاف داده مالک": gap.reindex(idx).fillna(0).astype(int),
        "ردیف بدون مالک شناسایی‌شده": ownerless.reindex(idx).fillna(0).astype(int),
        "کمترین مقاومت (روز)": _num(work, "مقاومت (روز)").groupby(work["_k"]).min().reindex(idx),
        "بیشترین رسوب (روز)": _num(work, "روزهای رسوب").groupby(work["_k"]).max().reindex(idx),
    }).reset_index(drop=True)
    out["چرایی بار کاری / وضعیت"] = [
        f"غالب پرونده‌ها در «{w}» است"
        + (f" و معطل {t} " if t not in ("—", "") else " ")
        + f"— {c} پرونده بحرانی، {g} ردیف با شکاف داده مالک."
        for w, t, c, g in zip(out["موقعیت غالب پرونده‌ها"], out["معطل غالب حوزه"],
                              out["تعداد بحرانی"], out["ردیف با شکاف داده مالک"])]
    return out


# ═══════════════════════════════════════════════════════════════════════════
#  نوشتن در Excel
# ═══════════════════════════════════════════════════════════════════════════
def style_grouped(ws, df: pd.DataFrame, key_col: str = "کلید گروه") -> None:
    if df.empty:
        return
    red = Side(style="thick", color=RED)
    vals = df[key_col].astype(str).tolist() if key_col in df.columns else []
    for r in range(2, len(df) + 2):
        first = r == 2 or (r - 2 < len(vals) and vals[r - 2] != vals[r - 3])
        last = r == len(df) + 1 or (r - 1 < len(vals) and vals[r - 1] != vals[r - 2])
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = P.font_body()
            cell.alignment = P.align("right", wrap=True)
            b = cell.border
            cell.border = Border(left=b.left, right=b.right,
                                 bottom=red if last else b.bottom,
                                 top=red if first else b.top)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for i in range(1, ws.max_column + 1):
        sample = [str(ws.cell(x, i).value or "") for x in range(1, min(ws.max_row, 25) + 1)]
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(14, max(map(len, sample)) + 2), 48)


def _write_table(ws, t: pd.DataFrame) -> None:
    for j, h in enumerate(t.columns, 1):
        c = ws.cell(1, j, h)
        c.font = P.font_header(1)
        c.fill = P.fill_header()
        c.alignment = P.align("center", wrap=True)
    for i, row in enumerate(t.itertuples(index=False, name=None), 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, None if pd.isna(v) else v)
    style_grouped(ws, t, "کلید گروه" if "کلید گروه" in t.columns else "__none__")


SHEETS = ("۱۴. متریال محور", "۱۵. بارنامه محور", "۱۶. اداره محور")


def write_supply_sheets(wb, df: pd.DataFrame) -> dict:
    specs = [(SHEETS[0], build_material_view(df)),
             (SHEETS[1], build_bl_view(df)),
             (SHEETS[2], build_dept_view(df))]
    for name, t in specs:
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = False
        if t.empty:
            ws["A1"] = "داده کافی برای این نما وجود ندارد."
            continue
        _write_table(ws, t)
    return {name: t for name, t in specs}
