# -*- coding: utf-8 -*-
"""تاریخچه KPI برای پاسخ به «بهتر شدیم یا بدتر؟».

هر اجرای رسمی یک snapshot کوچک و بدون جزئیات ردیفی ذخیره می‌کند. این فایل
حاوی PII نیست و برای نمودار روند مدیریتی استفاده می‌شود. مسیر با
AIBL_HISTORY_PATH قابل تغییر است؛ پیش‌فرض داخل OUTPUT_DIR است.
"""
from __future__ import annotations

__contract__ = 1

import os
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from ..config.settings import SETTINGS
from ..design import tokens as T
from ..design.excel import X
from ..studio_core.grain import safe_agg
from .financial_summary import commitment_numeric_if_single_currency

SHEET_HISTORY = "۱۸. روند مدیریتی"


def _path() -> Path:
    return Path(os.getenv("AIBL_HISTORY_PATH", os.path.join(SETTINGS.OUTPUT_DIR, "AIBL_KPI_History.csv")))


def snapshot(df: pd.DataFrame, ref_date: date | str) -> dict:
    def uniq(c: str) -> int:
        return int(df[c].replace("", pd.NA).nunique()) if c in df.columns else 0
    band = df.get("کد طبقه بحرانی", df.get("بحرانی (کوتاه)", pd.Series("", index=df.index))).astype(str)
    crit_mask = band.isin(["STOCKOUT", "CRITICAL", "توقف خط", "بحرانی"])
    stop_mask = band.isin(["STOCKOUT", "توقف خط"])
    late = pd.to_numeric(df.get("روزهای تأخیر", pd.Series(0, index=df.index)), errors="coerce")

    # Legacy datasets without currency/key metadata retain the historical numeric
    # contract.  Once the authoritative REG/currency columns exist, a naked
    # cross-currency float is intentionally not emitted.
    has_fin_contract = (
        "مانده تعهد" in df.columns
        and any(c in df.columns for c in ("KEY_REG", "CANONICAL_REG"))
        and any(c in df.columns for c in ("NTSW_CURRENCY", "CURRENCY", "ارز"))
    )
    if has_fin_contract:
        bal_value, bal_currency, bal_display = commitment_numeric_if_single_currency(df)
        overdue_df = df.loc[late > 0]
        overdue_value, overdue_currency, overdue_display = commitment_numeric_if_single_currency(overdue_df)
    else:
        bal_value = round(safe_agg(df, "مانده تعهد", "sum"), 2) if "مانده تعهد" in df.columns else 0.0
        overdue_value = round(safe_agg(df.loc[late.fillna(0) > 0], "مانده تعهد", "sum"), 2) if "مانده تعهد" in df.columns else 0.0
        bal_currency = overdue_currency = None
        bal_display = str(bal_value)
        overdue_display = str(overdue_value)

    return {
        "تاریخ": str(ref_date)[:10],
        "ردیف": len(df),
        "سفارش یکتا": uniq("CANONICAL_ORDER"),
        "بارنامه یکتا": uniq("CANONICAL_BL"),
        "متریال یکتا": uniq("KEY_MATERIAL"),
        "متریال بحرانی": int(df.loc[crit_mask, "KEY_MATERIAL"].replace("", pd.NA).nunique()) if "KEY_MATERIAL" in df.columns else int(crit_mask.sum()),
        "متریال توقف خط": int(df.loc[stop_mask, "KEY_MATERIAL"].replace("", pd.NA).nunique()) if "KEY_MATERIAL" in df.columns else int(stop_mask.sum()),
        "مانده تعهد": bal_value,
        "ارز مانده تعهد": bal_currency or ("چندارزی/نامشخص" if has_fin_contract else ""),
        "مانده تعهد نمایشی": bal_display,
        "مانده تعهد معوق": overdue_value,
        "ارز مانده تعهد معوق": overdue_currency or ("چندارزی/نامشخص" if has_fin_contract else ""),
        "مانده تعهد معوق نمایشی": overdue_display,
        "میانگین مقاومت": round(safe_agg(df, "مقاومت (روز)", "mean"), 2) if "مقاومت (روز)" in df.columns else 0.0,
        "میانگین رسوب": round(safe_agg(df, "روزهای رسوب", "mean"), 2) if "روزهای رسوب" in df.columns else 0.0,
    }


def append_snapshot(df: pd.DataFrame, ref_date: date | str) -> pd.DataFrame:
    p = _path(); p.parent.mkdir(parents=True, exist_ok=True)
    row = pd.DataFrame([snapshot(df, ref_date)])
    try:
        old = pd.read_csv(p, encoding="utf-8-sig") if p.exists() else pd.DataFrame()
    except Exception:
        old = pd.DataFrame()
    hist = pd.concat([old, row], ignore_index=True)
    if "تاریخ" in hist.columns:
        hist["تاریخ"] = hist["تاریخ"].astype(str)
        hist = hist.drop_duplicates("تاریخ", keep="last").sort_values("تاریخ")
    hist.to_csv(p, index=False, encoding="utf-8-sig")
    return hist.tail(366).reset_index(drop=True)


def write_history_sheet(wb, hist: pd.DataFrame) -> None:
    if SHEET_HISTORY in wb.sheetnames:
        del wb[SHEET_HISTORY]
    ws = wb.create_sheet(SHEET_HISTORY)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws["A1"] = "روند مدیریتی — آیا وضعیت نسبت به روزهای قبل بهتر شده است؟"
    ws["A1"].font = Font(name="IRANSans", size=14, bold=True, color=X(T.TEXT_ON_BRAND))
    ws["A1"].fill = PatternFill("solid", fgColor=X(T.BRAND_NAVY))
    cols = list(hist.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(3, j, c); cell.font = Font(name="IRANSans", bold=True, color=X(T.TEXT_ON_BRAND))
        cell.fill = PatternFill("solid", fgColor=X(T.TEAL_INK)); cell.alignment = Alignment(horizontal="center")
    for i, row in hist.iterrows():
        for j, c in enumerate(cols, 1):
            ws.cell(i + 4, j, row[c])
    ws.freeze_panes = "A4"; ws.auto_filter.ref = f"A3:{ws.cell(3, max(1,len(cols))).coordinate}"
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(3, j).column_letter].width = 20 if c != "تاریخ" else 14
    if len(hist) >= 2 and "متریال بحرانی" in cols:
        chart = LineChart(); chart.title = "روند متریال بحرانی"; chart.height = 7; chart.width = 14
        cidx = cols.index("متریال بحرانی") + 1
        chart.add_data(Reference(ws, min_col=cidx, min_row=3, max_row=3+len(hist)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3+len(hist)))
        ws.add_chart(chart, "A8")
    if len(hist) >= 2 and "مانده تعهد" in cols and pd.to_numeric(hist["مانده تعهد"], errors="coerce").notna().any():
        chart = LineChart(); chart.title = "روند مانده تعهد — فقط تک‌ارزی"; chart.height = 7; chart.width = 14
        cidx = cols.index("مانده تعهد") + 1
        chart.add_data(Reference(ws, min_col=cidx, min_row=3, max_row=3+len(hist)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3+len(hist)))
        ws.add_chart(chart, "J8")
