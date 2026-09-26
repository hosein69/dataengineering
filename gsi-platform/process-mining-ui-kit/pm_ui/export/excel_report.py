# -*- coding: utf-8 -*-
"""RTL Excel export from the same evidence-backed process facts as HTML."""
from __future__ import annotations

__contract__ = 2

import io
from typing import Mapping, Sequence

from .. import tokens as T


def build_excel_report(*, kpis: Sequence[Mapping], nodes: Sequence[Mapping], edges: Sequence[Mapping],
                       case_status: Sequence[Mapping] = (), handoff: Mapping | None = None,
                       patterns: Mapping | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor=T.INK.lstrip("#"))
    header_font = Font(color="FFFFFF", bold=True, name="IRANSansWeb", size=11)
    body_font = Font(name="IRANSansWeb", size=10.5)
    right = Alignment(horizontal="right", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    first_sheet_used = False

    def _safe(value):
        # openpyxl can preserve text beginning with '=' as a formula unless data_type is forced.
        return "" if value is None else value

    def _sheet(name: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
        nonlocal first_sheet_used
        ws = wb.active if not first_sheet_used else wb.create_sheet()
        first_sheet_used = True
        ws.title = name[:31]
        ws.sheet_view.rightToLeft = True
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                value = _safe(value)
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = body_font
                cell.alignment = right if c == 1 else center
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 24
        ws.freeze_panes = "A2"

    _sheet("KPI", ["شاخص", "مقدار", "تغییر", "وضعیت"],
           [(k.get("label", ""), k.get("value", ""), k.get("delta", ""),
             T.status_of(k.get("status")).label) for k in kpis])

    labels = {n.get("id"): n.get("label", n.get("id")) for n in nodes}
    edge_rows = [(
        labels.get(e.get("source"), e.get("source")), labels.get(e.get("target"), e.get("target")),
        int(e.get("count", 0)), int(e.get("occurrences", 0)), e.get("median_hours"), e.get("p90_hours"),
    ) for e in sorted(edges, key=lambda e: (-int(e.get("count", 0)), -int(e.get("occurrences", 0))))]
    _sheet("مسیرهای فرآیند", ["از مرحله", "به مرحله", "پرونده یکتا", "دفعات عبور", "میانه ساعت", "صدک ۹۰ ساعت"], edge_rows)

    node_rows = [(n.get("label", n.get("id")), n.get("count", ""), T.status_of(n.get("status")).label) for n in nodes]
    _sheet("مراحل فرآیند", ["مرحله", "آخرین مرحله پرونده‌ها", "وضعیت"], node_rows)

    if case_status:
        _sheet("وضعیت و اقدام", ["پرونده", "وضعیت", "مرحله فعلی", "مالک", "کیفیت شاهد", "اقدام بعدی", "زمان مشاهده", "ردیف شاهد"],
               [(r.get("case_key"), r.get("case_state"), r.get("current_activity"), r.get("owner_team"),
                 r.get("evidence_quality"), r.get("next_action"), r.get("observed_at"), r.get("source_row")) for r in case_status])

    handoff = handoff or {}
    if handoff.get("status") == "ready":
        _sheet("تحویل بین واحدها", ["پرونده", "شناسه تحویل", "فرستنده", "گیرنده", "مالک", "ارسال", "پذیرش", "تکمیل", "سند", "علت برگشت", "ردیف شاهد"],
               [(r.get("case_key"), r.get("handoff_id"), r.get("from_team"), r.get("to_team"), r.get("owner_team"),
                 r.get("sent_at"), r.get("accepted_at"), r.get("completed_at"), r.get("document_ref"), r.get("return_reason"), r.get("source_row"))
                for r in handoff.get("rows", [])])
    else:
        _sheet("تحویل بین واحدها", ["وضعیت", "توضیح"], [("قابل محاسبه نیست", handoff.get("note", "شاهد کافی وجود ندارد"))])

    patterns = patterns or {}
    if patterns.get("enabled"):
        _sheet("خلاصه الگوها", ["شناسه", "عنوان", "تعداد تطبیق", "پرونده یکتا", "تفسیر"],
               [(r.get("id"), r.get("label"), r.get("match_count"), r.get("unique_cases"), patterns.get("note"))
                for r in patterns.get("summary", [])])
        _sheet("تطبیق الگوها", ["الگو", "پرونده", "فعالیت‌ها", "ردیف‌های شاهد"],
               [(r.get("pattern_label"), r.get("case_key"), " ← ".join(r.get("activities", [])),
                 ", ".join(str(x) for x in r.get("source_rows", []))) for r in patterns.get("matches", [])])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
