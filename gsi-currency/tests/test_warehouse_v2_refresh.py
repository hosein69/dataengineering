# -*- coding: utf-8 -*-
"""سه بارفرش پشت‌سرهم روی منابع ساختگی (انبار داده نسخه ۲).

۱. بارفرش اول همه فایل‌ها را می‌خواند و خط مبنای سوابق را می‌سازد.
۲. بارفرش دوم بدون تغییر فایل: هیچ فایلی دوباره خوانده نمی‌شود، هیچ تغییری ثبت نمی‌شود،
   هسته کسب‌وکار دوباره ساخته نمی‌شود و خروجی همان است.
۳. یک خانه در Oracle عوض می‌شود: فقط همان فایل خوانده می‌شود و همان یک تغییر با مقدار قبلی
   و جدید ثبت می‌شود؛ نسخه قبلی پاک نمی‌شود و هنوز به تاریخ اجرای قبلی خواندنی است.
"""
from __future__ import annotations

import importlib.util
import json
import os
import tempfile

import pandas as pd

_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("gsi_make_synthetic_v2_refresh",
                                               os.path.join(_TESTS_DIR, "make_synthetic.py"))
_ms = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_ms)

# conftest منابع و تنظیمات را برای ماژولی که DIRS دارد به همین پوشه می‌بندد
DIRS = _ms.build(tempfile.mkdtemp(prefix="gsi_v2_refresh_"))
ORACLE = os.path.join(DIRS["bls"], "Oracle.xlsx")


def _refresh():
    from gsi.pipeline import Pipeline
    res = Pipeline().run(build_report=False)
    return res, res.extras["warehouse_run_id"]


def _edit_oracle_description(value):
    import openpyxl
    wb = openpyxl.load_workbook(ORACLE)
    ws = wb["Total_Report 14050610"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("شرح جنس") + 1, value=value)
    wb.save(ORACLE)


def _ledger(c, rid):
    return {s: (m, json.loads(f)) for s, m, f in c.execute(
        "SELECT source,mode,files FROM wh_source_load WHERE run_id=?", (rid,))}


def _mart(wh, rid, name):
    with wh.db() as c:
        fid = c.execute("SELECT id FROM wh_frame WHERE run_id=? AND layer='mart' AND name=?", (rid, name)).fetchone()[0]
    return wh.read_frame(fid)


def test_refresh_reads_only_changed_files_and_logs_each_change_exactly():
    from gsi.warehouse.business_dwh import FACT_INPUTS
    from gsi.warehouse.snapshots import as_of_sql
    from gsi.warehouse.store import Warehouse

    r1, rid1 = _refresh()
    wh = Warehouse()
    with wh.db() as c:
        first = _ledger(c, rid1)
    with_files = {s for s, (_, files) in first.items() if files}
    assert "oracle" in with_files and len(with_files) >= 5
    assert {first[s][0] for s in with_files} == {"parsed"}
    # سورسی که فایلش نیست «خوانده شد» ثبت نمی‌شود
    assert {m for s, (m, _) in first.items() if s not in with_files} <= {"missing"}

    # ── بدون تغییر فایل ──
    r2, rid2 = _refresh()
    with wh.db() as c:
        second = _ledger(c, rid2)
        logged = c.execute("SELECT count(*) FROM src_change WHERE run_id=?", (rid2,)).fetchone()[0]
        records_before = c.execute("SELECT count(*) FROM src_record").fetchone()[0]
    assert {second[s][0] for s in with_files} == {"reused"}
    assert logged == 0
    history = r2.extras["source_history"]
    assert history["unchanged_frames"] == history["frames"]
    skipped = set(r2.extras["business_dwh_counts"]["unchanged"])
    assert {"evidence", *FACT_INPUTS} <= skipped
    for name in ("df", "main"):
        pd.testing.assert_frame_equal(_mart(wh, rid1, name), _mart(wh, rid2, name))

    # ── یک خانه در Oracle ──
    _edit_oracle_description("مجموعه صفحه علایم (اصلاح‌شده)")
    r3, rid3 = _refresh()
    with wh.db() as c:
        third = _ledger(c, rid3)
        events = c.execute("SELECT source,frame,rkey,kind FROM src_change WHERE run_id=? AND field IS NULL",
                           (rid3,)).fetchall()
        fields = c.execute("SELECT field,old,new FROM src_change WHERE run_id=? AND field IS NOT NULL",
                           (rid3,)).fetchall()
        records_after = c.execute("SELECT count(*) FROM src_record").fetchone()[0]
        seq2, seq3 = wh.run_seq(rid2), wh.run_seq(rid3)
        facts = {seq: {r[0]: json.loads(r[1]) for r in c.execute(
                    as_of_sql("dwh_fact_oracle_material", seq, rid, c)).fetchall()}
                 for seq, rid in ((seq2, rid2), (seq3, rid3))}
    assert third["oracle"][0] == "parsed"
    assert {third[s][0] for s in with_files - {"oracle"}} == {"reused"}
    assert [(s, f, k) for s, f, _, k in events] == [("oracle", "main", "changed")]
    assert len(fields) == 1
    assert json.loads(fields[0][1]) == "مجموعه صفحه علایم"
    assert json.loads(fields[0][2]) == "مجموعه صفحه علایم (اصلاح‌شده)"
    # نسخه قبلی بسته شد و ماند؛ یک نسخه تازه کنارش نشست
    assert records_after == records_before + 1
    assert r3.extras["business_dwh_counts"]["versions"]["dwh_fact_oracle_material"]["changed"] == 1
    changed = [k for k in facts[seq3] if facts[seq3][k] != facts[seq2][k]]
    assert len(changed) == 1
    assert "مجموعه صفحه علایم (اصلاح‌شده)" in facts[seq3][changed[0]].values()
    assert "مجموعه صفحه علایم" in facts[seq2][changed[0]].values()
