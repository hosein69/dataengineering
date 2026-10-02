# -*- coding: utf-8 -*-
"""R10 (مالک، ۱۴۰۵/۰۷/۰۸): گزارش حمل و ترخیص پارت‌های فایل کارشناسان را در وضعیت نزد سازنده، آماده حمل،
در راه و در گمرک نشان می‌دهد، حتی بی‌بارنامه؛ بازه تاریخ روی کل گزارش است و مبنای تاریخ انتخاب می‌شود.
فقط داده ساختگی."""
import io

import pandas as pd
from openpyxl import load_workbook

from gsi.report import shipping_clearance_report as R
from tests.test_shipping_clearance_report import REF, _frame


def _ship(o, m, state, bl="", sched="", po=""):
    return {"KEY_ORDER": o, "KEY_MATERIAL": m, "MATERIAL_DESC": "شرح " + m, "PART_NO": "", "PART_STATE": state,
            "QTY_IN_PART": 5.0, "CLEARED_QTY": None, "UOM": "PC", "BL": bl, "BL_SUSPECT": "", "MODE": "SEA",
            "TRANSPORT_NO": "", "CARRIER": "", "SCHEDULED_SHIP": sched, "PO_SENT_DATE": po, "STAGE_FA": "",
            "LOGISTICS_NOTE": "", "CRIT_CODE": "STOCKOUT" if m == "M7" else "SAFE", "EXPERT": "کارشناس",
            "DATA_GAP_FILES": "Oracle.xlsx: متریال نیست" if m == "M8" else ""}


def _extras():
    ship = pd.DataFrame([
        _ship("O4", "M7", "IN_CUSTOMS", sched="1405/06/02"),               # بی‌بارنامه
        _ship("O1", "M1", "IN_TRANSIT", bl="BLA", sched="1405/05/10"),      # در فایل‌های حمل هم هست
        _ship("O5", "M8", "READY", bl="BLX"),                               # فقط در فایل کارشناسان، بی‌تاریخ
        _ship("O1", "M2", "AT_SUPPLIER", po="1405/04/01"),
        _ship("O3", "M5", "CLEARED", bl="BLB", sched="1405/04/20"),         # ترخیص کامل: این جدول نیست
        _ship("O9", "M9", "IN_TRANSIT", sched="1405/06/01"),                # بیرون از دامنه Studio
    ])
    return {"expert_shipments": ship}


def test_expert_rows_appear_even_without_bl():
    m = R.build_model(_frame(), _extras(), REF)
    e = m.expert
    assert len(e) == 4 and not m.empty
    assert e["وضعیت پارت (کارشناس)"].tolist() == ["در گمرک", "در راه", "آماده حمل", "نزد سازنده"]
    bl = dict(zip(e["متریال"], e["بارنامه در گزارش حمل"]))
    assert bl == {"M7": R.EXP_BL_NONE, "M1": R.EXP_BL_TRACKED, "M8": R.EXP_BL_EXPERT_ONLY, "M2": R.EXP_BL_NONE}
    assert m.counts["exp_IN_CUSTOMS"] == 1 and m.counts["exp_no_bl"] == 2 and m.counts["exp_expert_only"] == 1
    page = R.build_html(m, lang="fa", embed_fonts=False, embed_excel=False)
    assert 'data-key="expert"' in page and "M7" in page and 'id="sc-bs"' in page
    assert 'id="sc-u" type="checkbox">' in page                             # بی‌تاریخ پیش‌فرض بیرون
    en = R.build_html(m, lang="en", embed_fonts=False, embed_excel=False)
    assert "Expert shipments" in en and "In customs" in en
    wb = load_workbook(io.BytesIO(R.build_excel(m, "fa")))
    assert wb["پارت‌های کارشناسان"].max_row >= 5


def test_expert_rows_only_when_dossiers_are_empty():
    df = _frame()
    df = df[df["CANONICAL_BL"].eq("")]
    m = R.build_model(df, _extras(), REF)
    assert m.dossiers.empty and len(m.expert) == 2 and not m.empty
    assert "M7" in R.build_html(m, lang="fa", embed_fonts=False)


def test_date_range_rebuilds_whole_report():
    df, ex = _frame(), _extras()
    m = R.build_model(df, ex, REF)
    f = R.filtered_model(m, df, ex, REF, start="1405/05/15")
    assert f.dossiers["بارنامه"].tolist() == ["BLA"] and f.total == 1       # BLB پیش از بازه، BLC بی‌تاریخ
    assert f.counts[R.ST_FULL] == 0 and f.counts[R.ST_TRANSIT] == 0
    assert f.expert["متریال"].tolist() == ["M7"]
    g = R.filtered_model(m, df, ex, REF, start="1405/05/15", include_undated=True)
    assert set(g.dossiers["بارنامه"]) == {"BLA", "BLC"} and "M8" in set(g.expert["متریال"])
    # مبنای تاریخ: تخلیه BLB = 1405/05/01
    h = R.filtered_model(m, df, ex, REF, end="1405/05/10", basis="تخلیه")
    assert h.dossiers["بارنامه"].tolist() == ["BLB"]
    assert R.filtered_model(m, df, ex, REF) is m                              # بی‌بازه: همان مدل
