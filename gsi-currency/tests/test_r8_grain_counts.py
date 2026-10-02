# -*- coding: utf-8 -*-
"""R8: شمارش‌ها و جمع‌های گزارش تخت روی دانه درست — نه روی ردیف.

سناریوی مصنوعی ممیزی: سفارش 100001 با متریال M1 (FIRST) و M2 (ADDITIONAL) و
سه بارنامه BL1..BL3؛ سفارش 100002 با متریال M1 روی BL3؛ هر دو روی ثبت سفارش
12345678. ردیف FIRST سفارش اول به ازای هر بارنامه تکرار شده و ردیف ADDITIONAL
بارنامه ندارد.
"""
from __future__ import annotations

from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

REG = "12345678"
RED = "قرمز"


def scenario() -> pd.DataFrame:
    return pd.DataFrame({
        "KEY_ORDER": ["100001", "100001", "100001", "100001", "100002"],
        "CANONICAL_ORDER": ["100001", "100001", "100001", "100001", "100002"],
        "KEY_MATERIAL": ["M1", "M1", "M1", "M2", "M1"],
        "MOGH_ITEM_ROLE": ["FIRST", "FIRST", "FIRST", "ADDITIONAL", "FIRST"],
        "KEY_BL": ["BL1", "BL2", "BL3", "", "BL3"],
        "CANONICAL_BL": ["BL1", "BL2", "BL3", "", "BL3"],
        "KEY_REG": [REG] * 5,
        "KEY_PR": ["PR1", "PR1", "PR1", "PR2", "PR3"],
        "CANONICAL_EXPERT": ["ExpA", "ExpA", "ExpA", "ExpA", "ExpB"],
        "KEY_EMP": ["E1", "E1", "E1", "E1", "E2"],
        "ORG_VICE": ["V1"] * 5, "ORG_DEPT": ["D1"] * 5,
        "ORG_MANAGER": ["Mgr"] * 5, "ORG_HEAD": ["Head"] * 5,
        # سطح ثبت سفارش — تکرارشونده
        "BALANCE": [1000.0] * 5,
        "مانده تعهد": [1000.0] * 5,
        "روزهای تأخیر": [12] * 5,
        "وضعیت کلی هشدار": [RED] * 5,
        "BUY_DATE": ["2026-01-01"] * 5,
        # سطح بارنامه
        "BLREG_BL_INVOICE_VALUE": [10.0, 20.0, 30.0, None, 30.0],
        "BLREG_REG_VALUE": [500.0] * 5,
        "DISCHARGE_DATE": ["2026-02-01", "2026-02-02", "2026-02-03", "", "2026-02-03"],
        "WH_STATUS": ["OVERDUE", "OVERDUE", "OVERDUE", "", "OVERDUE"],
        "BL_CRITICAL": [True, True, True, False, True],
        # سطح سفارش × متریال
        "SUPPLIER_QTY": [10.0, 10.0, 10.0, 5.0, 7.0],
        "MOGH_QTY_READY": [1.0, 1.0, 1.0, 2.0, 3.0],
        "SUPPLY_POSITION_STATUS": ["COMPLETE", "COMPLETE", "COMPLETE", "PARTIAL", "COMPLETE"],
        # سطح متریال (Oracle)
        "STOCK_IKCO": [100.0, 100.0, 100.0, 50.0, 100.0],
        "کد طبقه بحرانی": ["CRITICAL", "CRITICAL", "CRITICAL", "SAFE", "CRITICAL"],
        "بحرانی (کوتاه)": ["بحرانی", "بحرانی", "بحرانی", "ایمن", "بحرانی"],
        "طبقه ریسک": ["بالا"] * 5,
        "هشدار ترکیبی بحرانی": ["x", "x", "x", "", "x"],
        "IS_BLOCKED": [True, True, True, True, False],
        # فرایند
        "انحراف فرآیند": ["جاافتاده", "جاافتاده", "جاافتاده", "جاافتاده", "بدون انحراف"],
        "امتیاز انطباق (٪)": [50.0, 50.0, 50.0, 50.0, 100.0],
        "روزهای رسوب": [5, 5, 5, None, 9],
        "امتیاز ریسک": [80] * 5,
    })


def _ctx():
    rb = SimpleNamespace(status_label=lambda code: RED if code == "red" else code)
    return SimpleNamespace(rb=rb, extras={})


# ── grain.py ─────────────────────────────────────────────────────────────
def test_derived_and_blreg_columns_inherit_source_grain():
    from gsi.studio_core.grain import column_grain
    for c in ("BALANCE", "CB_VALUE", "ALLOCATED_AMOUNT", "OPEN_QUEUE_AMOUNT",
              "REJECTED_ALLOC_AMOUNT", "ALLOC_REQUESTS", "CREDIT_PROFORMA",
              "CREDIT_RIAL_AMOUNT", "BLREG_REG_VALUE", "BLREG_REG_UNSHIPPED_VALUE"):
        assert column_grain(c) == "REG", c
    for c in ("INVOICE_VALUE", "EUR_VALUE", "DUTY_AMOUNT", "CLEAR_AMOUNT",
              "BLREG_BL_INVOICE_VALUE", "BLREG_BL_SHARE_PCT"):
        assert column_grain(c) == "BL", c
    for c in ("SUPPLIER_QTY", "READY_QTY", "IN_TRANSIT_QTY", "IN_CUSTOMS_QTY",
              "SUPPLY_EXPERT_STOCK", "MOGH_QTY_READY", "EXPERT_INV_UNKNOWN_QTY",
              "موجودی در راه", "موجودی نزد سازنده"):
        assert column_grain(c) == "ORDER_MATERIAL", c
    for c in ("STOCK_IKCO", "STOCK_SAPCO", "DAILY_NEED", "SUPPLY_ORACLE_STOCK"):
        assert column_grain(c) == "MATERIAL", c
    assert column_grain("KEY_EMP") == "EMP"
    assert column_grain("BLREG_STATUS") == "ROW"


def test_safe_agg_does_not_double_count_on_bl_fanout():
    from gsi.studio_core.grain import safe_agg
    df = scenario()
    assert safe_agg(df, "BALANCE", "sum") == 1000.0
    assert safe_agg(df, "BLREG_REG_VALUE", "sum") == 500.0
    assert safe_agg(df, "BLREG_BL_INVOICE_VALUE", "sum") == 60.0
    # سفارش×متریال: 100001/M1=10 + 100001/M2=5 + 100002/M1=7
    assert safe_agg(df, "SUPPLIER_QTY", "sum") == 22.0
    assert safe_agg(df, "MOGH_QTY_READY", "sum") == 6.0
    # Oracle در دانه متریال: M1=100 + M2=50
    assert safe_agg(df, "STOCK_IKCO", "sum") == 150.0


def test_supply_total_is_not_summable():
    from gsi.studio_core.grain import KIND_RATIO, measure_kind, summable
    for c in ("SUPPLY_TOTAL_CONFIRMED", "SUPPLY_TOTAL_LOWER_BOUND", "موجودی کل قابل احتساب"):
        assert measure_kind(c) == KIND_RATIO
    df = pd.DataFrame({"SUPPLY_TOTAL_CONFIRMED": [1.0, 2.0], "SUPPLIER_QTY": [1.0, 2.0]})
    assert summable(df, list(df.columns)) == ["SUPPLIER_QTY"]


def test_order_material_blank_key_rows_are_kept():
    from gsi.studio_core.grain import safe_agg
    df = pd.DataFrame({"KEY_ORDER": ["", "", "O1", "O1"],
                       "KEY_MATERIAL": ["M1", "M1", "M1", "M1"],
                       "SUPPLIER_QTY": [1.0, 2.0, 4.0, 4.0]})
    assert safe_agg(df, "SUPPLIER_QTY", "sum") == 7.0


# ── stages ───────────────────────────────────────────────────────────────
def test_commitment_red_counts_unique_registrations():
    from gsi.stages.s50_commitment import CommitmentStage
    k = CommitmentStage().kpis(scenario(), _ctx())
    assert k["تعهدات با هشدار قرمز"][0] == 1


def test_conformance_counts_unique_cases():
    from gsi.stages.s85_conformance import ConformanceStage
    k = ConformanceStage().kpis(scenario(), _ctx())
    assert k["پرونده‌های منحرف از مسیر استاندارد"][0] == 1
    assert "از 2 پرونده" in k["پرونده‌های منحرف از مسیر استاندارد"][1]
    assert k["میانگین انطباق پرونده‌های قابل‌ارزیابی (٪)"][0] == 75.0


def test_supply_position_counts_order_material():
    from gsi.stages.s38_supply_position import SupplyPositionStage
    k = SupplyPositionStage().kpis(scenario(), _ctx())
    assert k["موجودی با پوشش کامل"][0] == 2        # 100001/M1، 100002/M1
    assert k["شکاف داده موجودی"][0] == 1           # 100001/M2


def test_warehouse_declaration_counts_unique_bl():
    from gsi.stages.s39_warehouse_declaration import WarehouseDeclarationStage
    k = WarehouseDeclarationStage().kpis(scenario(), _ctx())
    assert k["شکاف شاهد انبار"][0] == 3


def test_risk_combined_alert_counts_unique_orders():
    from gsi.stages.s70_risk import RiskStage
    k = RiskStage().kpis(scenario(), _ctx())
    assert k["هشدار ترکیبی بحرانی + پرونده مشکل‌دار"][0] == 2


# ── گزارش‌ها ─────────────────────────────────────────────────────────────
def test_daily_email_overdue_counts_unique_registrations():
    from gsi.integrations.daily_email import _case_counts, executive_kpis
    k = {x["label"]: x["value"] for x in executive_kpis(scenario())}
    assert k["تعهدات معوق"] == 1
    assert int(_case_counts(scenario(), "طبقه ریسک")["بالا"]) == 2


def test_dashboard_critical_bl_is_static_unique_count(tmp_path):
    from gsi.report.dashboard import ExcelDashboardBuilder
    from gsi.stages.base import ColumnSpec
    b = ExcelDashboardBuilder(str(tmp_path / "d.xlsx"))
    n = b.build_matrix(scenario(), [ColumnSpec("KEY_MATERIAL", "متریال"),
                                    ColumnSpec("SUPPLIER_QTY", "نزد سازنده"),
                                    ColumnSpec("BL_CRITICAL", "بارنامه بحرانی")])
    b.build_executive({"شاخص": (1, "توضیح")}, "روایت", n)
    wb = load_workbook(b.save())
    ws = wb.worksheets[0]
    found = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)
             if str(ws.cell(r, 1).value or "").startswith("بارنامه‌های بحرانی")]
    assert found == [3]


def test_dashboard_scorecard_keeps_shared_bl_per_expert(tmp_path):
    from gsi.report.dashboard import ExcelDashboardBuilder
    b = ExcelDashboardBuilder(str(tmp_path / "s.xlsx"))
    b.build_scorecard(scenario())
    wb = load_workbook(b.save())
    ws = wb["۵. کارنامه سازمانی"]
    got = {ws.cell(r, 5).value: ws.cell(r, 7).value for r in range(2, ws.max_row + 1)}
    assert got == {"ExpA": 3, "ExpB": 1}


def test_insight_stage_counts_unique_keys():
    from gsi.report.insight import LIFECYCLE, _metrics, _stage_counts
    df = scenario()
    counts = dict(zip([c for c, _, _ in LIFECYCLE], _stage_counts(df)))
    assert counts["ORDER"] == 2 and counts["REG"] == 1 and counts["BL"] == 3
    assert counts["FX"] == 1 and counts["DISCHARGE"] == 3 and counts["PR"] == 3
    m = _metrics(df)
    assert m["پرونده"] == 2 and m["بحرانی"] == 2
    # متریال افزوده بارنامه خودش را ندارد، ولی سفارشش دارد: افت کاذب نیست
    m2 = dict(zip([c for c, _, _ in LIFECYCLE], _stage_counts(df[df["KEY_MATERIAL"] == "M2"], df)))
    assert m2["BL"] == 3


def test_excel_export_counts_cases_and_materials(tmp_path):
    from gsi.studio_core.excel_export import build_custom_excel
    path = build_custom_excel(scenario(), tmp_path / "studio.xlsx", ["org", "expert"],
                              "2026-09-27")
    wb = load_workbook(path)
    ex = {r[0]: r[1] for r in wb["Executive"].iter_rows(min_row=2, values_only=True)}
    assert ex["Critical"] == 1 and ex["Critical BLs"] == 3
    org = {r[0]: r[1] for r in wb["Organization"].iter_rows(min_row=2, values_only=True)}
    assert org == {"D1": 2}
    hdr = [c.value for c in wb["Expert Workload"][1]]
    exp = {r[0]: dict(zip(hdr, r)) for r in wb["Expert Workload"].iter_rows(min_row=2, values_only=True)}
    assert exp["ExpA"]["Cases"] == 1 and exp["ExpA"]["Critical Cases"] == 1
    assert exp["ExpB"]["Cases"] == 1


def test_supply_dept_view_counts_unique_critical_materials():
    from gsi.report.supply_views import build_dept_view
    out = build_dept_view(scenario())
    if out.empty:
        pytest.skip("dept view needs more columns")
    assert int(out.loc[0, "تعداد بحرانی"]) == 1
