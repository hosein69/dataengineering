# -*- coding: utf-8 -*-
"""R8 — رابطه ساختگی بارنامه↔متریال و ستون‌های بی‌برچسب در خروجی‌ها.

داده مصنوعی: سفارش 100001 با M1 (FIRST، ایمن) و M2 (ADDITIONAL، توقف خط) و دو
بارنامه BL1/BL2. بارنامه‌ها فقط روی ردیف FIRST نشسته‌اند (population._attach_abbasi)،
پس هیچ خروجی نباید بگوید بارنامه «متریال M1» را حمل می‌کند یا بارنامه ایمن است.
"""
from __future__ import annotations

from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.rulebook import get_rulebook
from gsi.stages.base import PipelineContext

ORDER = "100001"


def _population() -> pd.DataFrame:
    base = {"CANONICAL_ORDER": ORDER, "KEY_ORDER": ORDER, "KEY_REG": "R1", "STOCK_SAPCO": 0,
            "SUPPLIER_QTY": 0, "READY_QTY": 0, "IN_TRANSIT_QTY": 0, "IN_CUSTOMS_QTY": 0, "DAILY_NEED": 10,
            "INVOICE_VALUE": 100, "INVOICE_CURRENCY": "EUR"}
    return pd.DataFrame([
        {**base, "KEY_MATERIAL": "M1", "MOGH_ITEM_ROLE": "FIRST", "CANONICAL_BL": "BL1", "STOCK_IKCO": 100000},
        {**base, "KEY_MATERIAL": "M1", "MOGH_ITEM_ROLE": "FIRST", "CANONICAL_BL": "BL2", "STOCK_IKCO": 100000},
        {**base, "KEY_MATERIAL": "M2", "MOGH_ITEM_ROLE": "ADDITIONAL", "CANONICAL_BL": "", "STOCK_IKCO": 0},
        # سفارش دیگر، ایمن، روی بارنامه خودش
        {**base, "CANONICAL_ORDER": "100002", "KEY_ORDER": "100002", "KEY_REG": "R2", "KEY_MATERIAL": "M3",
         "MOGH_ITEM_ROLE": "FIRST", "CANONICAL_BL": "BL3", "STOCK_IKCO": 100000},
    ])


@pytest.fixture(scope="module")
def crit() -> pd.DataFrame:
    from gsi.stages.s40_criticality import CriticalityStage
    ctx = PipelineContext(rb=get_rulebook(), today=date(2026, 9, 29))
    return CriticalityStage().run(_population(), ctx)


def test_bl_is_critical_when_its_orders_second_material_is_stockout(crit):
    assert crit.loc[crit["KEY_MATERIAL"].eq("M2"), "کد طبقه بحرانی"].iloc[0] == "STOCKOUT"
    assert crit.loc[crit["KEY_MATERIAL"].eq("M1"), "کد طبقه بحرانی"].iloc[0] not in ("STOCKOUT", "CRITICAL")
    bl_rows = crit[crit["CANONICAL_BL"].isin(["BL1", "BL2"])]
    assert len(bl_rows) == 2 and bl_rows["BL_CRITICAL"].all()
    assert set(bl_rows["BL_CRITICAL_LEVEL"]) == {"STOCKOUT"}
    for _, r in bl_rows.iterrows():
        assert "M2" in r["BL_CRITICAL_MATERIALS"] and "M1" not in r["BL_CRITICAL_MATERIALS"]
        assert ORDER in r["BL_CRITICAL_REASON"] and "M2" in r["BL_CRITICAL_REASON"]
    # سفارش روی همه ردیف‌هایش (از جمله ADDITIONAL) بحرانی است
    assert crit.loc[crit["CANONICAL_ORDER"].eq(ORDER), "ORDER_CRITICAL"].all()
    assert not crit.loc[crit["CANONICAL_BL"].eq("BL3"), "BL_CRITICAL"].any()


def test_bl_labels_speak_about_orders_not_carried_materials():
    from gsi.stages.s40_criticality import CriticalityStage
    from gsi.studio_core.field_catalog import COMPUTED_LABELS
    titles = {c.key: c.title for c in CriticalityStage().columns()}
    assert "سفارش" in titles["BL_CRITICAL_MATERIALS"] and "سفارش" in COMPUTED_LABELS["BL_CRITICAL_MATERIALS"]
    assert titles["BL_CRITICAL_MATERIALS"] != "متریال بحرانی بارنامه"


def test_critical_board_uses_order_evidence(crit):
    from gsi.report import critical_board as CB
    mats = CB.critical_materials(crit).set_index("متریال")
    assert "بارنامه‌ها" not in mats.columns
    assert set(mats.loc["M2", CB.BL_OF_ORDERS].split("، ")) == {"BL1", "BL2"}
    bls = CB.critical_bls(crit).set_index("بارنامه")
    assert set(bls.index) == {"BL1", "BL2"}
    for b in ("BL1", "BL2"):
        assert "M1" not in bls.loc[b, CB.MATS_OF_BL_ORDERS]
        assert bls.loc[b, CB.MIN_RES_OF_BL_ORDERS] == 0          # مقاومت M2، نه M1 ردیف بارنامه


def test_link_materials_are_the_orders_materials(crit):
    from gsi.stages.s59_bl_registration_link import BLRegistrationLinkStage

    class _RB:
        def normalize_currency(self, v):
            t = "" if v is None else str(v).strip()
            return t.upper() if t.lower() not in ("", "nan", "none") else ""

    ctx = PipelineContext(rb=_RB(), today=date(2026, 9, 29))
    BLRegistrationLinkStage().run(crit.copy(), ctx)
    link = ctx.extras["bl_registration_link"].set_index("KEY_BL")
    for b in ("BL1", "BL2"):
        assert link.loc[b, "MATERIALS"].split("، ") == ["M1", "M2"]   # همه متریال‌های سفارش، نه «M1 بارنامه»


def _fx(mart: pd.DataFrame, link: pd.DataFrame):
    from gsi.report import fx_insight as X
    empty = pd.DataFrame()
    fx = X.FxData(lc=empty, link=link, recon=empty, ledger=empty, money=empty, queue=empty, control=empty,
                  decisions=empty, positions=empty, mart=mart)
    return X, fx


def test_fx_material_bls_are_order_bls(crit):
    X, fx = _fx(crit, pd.DataFrame(columns=["KEY_REG", "KEY_BL"]))
    (o,) = [o for o in X.orders(fx, "R1") if o["key"] == ORDER]
    assert sorted(o["bls"]) == ["BL1", "BL2"]
    m = {x["key"]: x for x in o["materials"]}
    assert sorted(m["M2"]["bls"]) == ["BL1", "BL2"] and m["M1"]["bls"] == m["M2"]["bls"]
    assert "بارنامه‌ها" not in X.materials_frame(fx, ["R1"]).columns


def test_stage_counts_are_unique_over_registrations():
    mart = pd.DataFrame([
        {"KEY_REG": "R1", "KEY_ORDER": "O1", "KEY_MATERIAL": "M9", "کد طبقه بحرانی": "STOCKOUT", "CANONICAL_BL": "BL1"},
        {"KEY_REG": "R2", "KEY_ORDER": "O2", "KEY_MATERIAL": "M9", "کد طبقه بحرانی": "STOCKOUT", "CANONICAL_BL": "BL1"},
    ])
    link = pd.DataFrame([{"KEY_REG": "R1", "KEY_BL": "BL1"}, {"KEY_REG": "R2", "KEY_BL": "BL1"}])
    X, fx = _fx(mart, link)
    fx.__dict__["reg_table"] = pd.DataFrame([
        {"KEY_REG": r, "STAGE_CODE": X.STAGES[0][0], "STAGE_DAYS": 1, "STAGE_STATUS": "", "GAP_COUNT": 0,
         "CRITICAL_LEVEL": "STOCKOUT", "CRITICAL_MATERIALS": 1, "ORDER_COUNT": 1, "BL_COUNT": 1} for r in ("R1", "R2")])
    row = X.stage_summary(fx).iloc[0]
    assert row["BLS"] == 1 and row["ORDERS"] == 2 and row["CRITICAL_MATERIALS"] == 1
    assert X.unique_counts(fx) == {"bls": 1, "orders": 2, "critical_materials": 1}


# ─────────────────────────── برچسب یکتا ───────────────────────────
def test_unique_labels_disambiguate_colliding_source_labels():
    from gsi.studio_core.field_catalog import dedupe_label_map, unique_labels
    cols = ["BL_STATUS", "COT_STATUS", "FX_STATUS", "IL_STATUS", "BL_GOODS_DESC", "SATA_GOODS_DESC",
            "CL_GOODS_DESC", "FX_GOODS_DESC", "KEY_ORDER"]
    labels = unique_labels(cols)
    assert set(labels) == set(cols) and len(set(labels.values())) == len(cols)
    assert all(v.startswith("وضعیت") for k, v in labels.items() if k.endswith("_STATUS"))
    assert all(not v.endswith(" ") for v in labels.values())
    d = dedupe_label_map({"A": "وضعیت", "B": "وضعیت", "C": "وضعیت · B"}, ["A", "B", "C", "D"])
    assert len(set(d.values())) == 4 and d["A"] == "وضعیت · A" and d["D"] == "D"


# ─────────────────────────── فایل کارشناس ───────────────────────────
def _wide(crit: pd.DataFrame) -> pd.DataFrame:
    extra = {"KEY_EMP": "E1", "CANONICAL_EXPERT": "کارشناس الف", "BL_STATUS": "x", "COT_STATUS": "y",
             **{f"JUNK_TECH_{i}": i for i in range(300)}}
    return pd.concat([crit, pd.DataFrame(extra, index=crit.index)], axis=1)


def _headers(out_dir) -> list:
    import glob
    import os
    (path,) = glob.glob(os.path.join(str(out_dir), "*.xlsx"))
    ws = load_workbook(path).active
    return [c.value for c in ws[3] if c.value is not None]


def test_expert_extract_default_is_the_official_matrix_in_persian(tmp_path, monkeypatch, crit):
    from gsi.i18n.columns import is_persian
    from gsi.report.extracts import write_expert_extracts
    monkeypatch.delenv("GSI_EXTRACT_ALL_COLUMNS", raising=False)
    df = _wide(crit)
    write_expert_extracts(df, str(tmp_path / "d"))
    heads = _headers(tmp_path / "d")
    assert len(heads) == len(set(heads))
    assert len(heads) < 60 < len(df.columns)
    assert all(is_persian(h) for h in heads), [h for h in heads if not is_persian(h)]
    assert "متریال بحرانی سفارش‌های این بارنامه" in heads and "کلید سفارش" in heads
    assert not any("JUNK" in h for h in heads)


def test_expert_extract_env_switch_restores_full_dump(tmp_path, monkeypatch, crit):
    from gsi.report.extracts import write_expert_extracts
    monkeypatch.setenv("GSI_EXTRACT_ALL_COLUMNS", "1")
    df = _wide(crit)
    write_expert_extracts(df, str(tmp_path / "f"))
    heads = _headers(tmp_path / "f")
    assert len(heads) == len(df.columns) and len(heads) == len(set(heads))
    assert "کلید سفارش" in heads and "متریال بحرانی سفارش‌های این بارنامه" in heads
    assert sum(h.startswith("وضعیت") for h in heads) >= 2


def test_snapshot_export_passes_unique_persian_labels(tmp_path, monkeypatch, crit):
    from gsi.warehouse import export_snapshot as ES
    from gsi.warehouse import service
    import gsi.studio_core.excel_export as EE
    seen = {}

    def fake_build(data, output_path, modules, ref, **kw):
        seen.update(kw)
        return str(output_path)

    df = _wide(crit)
    monkeypatch.setattr(service, "last_report", lambda ref=None: (df, df, {}, None))
    monkeypatch.setattr(EE, "build_custom_excel", fake_build)
    ES.export_published_snapshot(tmp_path / "x.xlsx", "2026-09-29")
    labels = seen["field_labels"]
    # R10: ستون موازی (CANONICAL_ORDER همان KEY_ORDER) و بی‌داده در Snapshot نمی‌آیند؛ برچسب‌ها یکتا می‌مانند.
    assert set(labels) == set(seen["selected_fields"]) and len(set(labels.values())) == len(labels)
    assert "CANONICAL_ORDER" not in labels and "KEY_ORDER" in labels
    monkeypatch.setenv("GSI_OUTPUT_ALL_COLUMNS", "1")
    ES.export_published_snapshot(tmp_path / "y.xlsx", "2026-09-29")
    assert set(seen["field_labels"]) == set(df.columns)
    labels = seen["field_labels"]
    assert labels["KEY_ORDER"] == "کلید سفارش" and labels["BL_STATUS"] != labels["COT_STATUS"]
