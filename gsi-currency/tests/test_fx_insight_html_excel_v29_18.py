# -*- coding: utf-8 -*-
"""29.18 — یک بینش در Studio، HTML ارسالی، گزارش‌ساز، ایمیل و Excel؛ فونت ایران‌سنس.

هر آزمون یک خطای مدیریتی یا نمایشی واقعی را قفل می‌کند:
  * ارزش PI سفارش و ارزش بارنامه روی ردیف‌های متریال تکرار نشود (سلسله‌مراتب و Excel)؛
  * ارزش هر مرحله فقط داخل یک ارز جمع شود و ارز نامعلوم ردیف جدا و بی‌جمع داشته باشد؛
  * حمل بیش از ارزش یک ثبت سفارش، مانده حمل‌نشده ثبت سفارش دیگر را کم نکند؛
  * شمار «کلیک روی مرحله» با خلاصه مرحله و Excel یکی باشد؛
  * تاریخ ورود به صف پس از تاریخ مرجع، انتظار منفی نسازد؛
  * HTML ارسالی بدون منبع بیرونی، بدون SVG (که Streamlit پاک می‌کند) و فقط با رنگ توکن‌ها باشد؛
  * بدنه ایمیل Outlook فقط جدول درون‌خطی داشته باشد؛
  * فونت نصب‌شده سیستم هرگز در فایل ارسالی جاسازی نشود.
"""
from __future__ import annotations

import base64
import io
import math
import re
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.design import fonts as F
from gsi.design import icons as I
from gsi.design import tokens as T
from gsi.i18n import columns as C
from gsi.report import critical_board as CB
from gsi.report import fx_excel as XE
from gsi.report import fx_html as H
from gsi.report import fx_insight as X
from gsi.stages.base import PipelineContext
from gsi.stages.s59_bl_registration_link import BLRegistrationLinkStage

TODAY = date(2026, 9, 27)
REF = "2026-09-27"
PERSIAN = re.compile(r"[؀-ۿ]")


# ═══════════════════════════ داده آزمون ═══════════════════════════
class _RB:
    """نرمال‌ساز ساده ارز برای آزمون واحد (نسخه واقعی در RuleBook است)."""
    ALIASES = {"یورو": "EUR", "یوان": "CNY", "درهم": "AED"}

    def normalize_currency(self, v):
        t = "" if v is None else str(v).strip()
        return self.ALIASES.get(t, t.upper() if t.lower() not in ("", "nan", "none") else "")


def _row(reg, bl, order, mat, value, ccy, level, resistance, pi=None, **kw):
    r = {"KEY_REG": reg, "CANONICAL_BL": bl, "KEY_ORDER": order, "KEY_MATERIAL": mat,
         "INVOICE_VALUE": value, "INVOICE_CURRENCY": ccy, "MATERIAL_DESC": "شرح " + mat,
         "MOGH_CURRENCY": ccy, "ORDER_CRITICAL_LEVEL": level,
         "کد طبقه بحرانی": level, "مقاومت (روز)": resistance}
    if pi is not None:
        r["MOGH_PI_VALUE_SUM"] = pi
    r.update(kw)
    return r


def _done(reg, steps):
    return [{"KEY_REG": reg, "STAGE_CODE": code, "STATUS": "DONE", "EVENT_DATE": dt, "DUE_DATE": "",
             "EVIDENCE": "شاهد"} for code, dt in steps.items()]


def _build():
    """چهار ثبت سفارش در سه مرحله و چهار ارز:

    R1 یورو ۱۰۰۰: BL1=۴۰۰ روی دو متریال (fan-out)، BL2=۱۰۰؛ سفارش O1 با PI ۹۰۰ روی سه ردیف؛
       ورود به صف بعد از تاریخ مرجع؛ مرحله جاری «ورود / کوتاژ» با گام‌های بی‌شاهد.
    R2 یورو ۱۰۰۰: BL3=۱۳۰۰ (حمل بیش از ارزش)؛ مرحله «ترخیص».
    R3 یوان ۵۰۰۰: BL4=۱۰۰۰؛ متریال بحرانی؛ هفت روز در صف؛ مرحله «تخصیص ارز».
    R4 ارزش و ارز ثبت سفارش نامعلوم؛ BL5=۳۰۰ درهم؛ توقف خط؛ مرحله «ترخیص».
    """
    rows = [
        _row("R1", "BL1", "O1", "M1", 400, "EUR", "SAFE", 90, pi=900),
        _row("R1", "BL1", "O1", "M2", 400, "EUR", "SAFE", 80, pi=900),
        _row("R1", "BL2", "O1", "M1", 100, "EUR", "SAFE", 90, pi=900, ALLOC_QUEUE_STATE="IN_QUEUE",
             ALLOC_QUEUE_ENTER_DATE="2026-10-02", NTSW_REQ_CURRENCY="EUR", OPEN_QUEUE_AMOUNT=50),
        _row("R2", "BL3", "O2", "M3", 1300, "EUR", "WATCH", 25, pi=1000),
        _row("R3", "BL4", "O3", "M4", 1000, "CNY", "CRITICAL", 4, pi=5000, ALLOC_QUEUE_STATE="IN_QUEUE",
             ALLOC_QUEUE_ENTER_DATE="2026-09-20", NTSW_REQ_CURRENCY="CNY", OPEN_QUEUE_AMOUNT=700),
        _row("R4", "BL5", "O4", "M5", 300, "AED", "STOCKOUT", 0),
    ]
    credit = pd.DataFrame([{"KEY_REG": r, "CRD_PROFORMA_VALUE": v, "CRD_CURRENCY": c}
                           for r, v, c in [("R1", 1000, "EUR"), ("R2", 1000, "EUR"), ("R3", 5000, "CNY")]])
    full = {"ORDER_REG": "2026-04-01", "ALLOCATION_QUEUE": "2026-04-05", "ALLOCATION": "2026-04-20",
            "FX_PURCHASE": "2026-05-01", "FUNDING": "2026-05-05", "SWIFT_CONVERSION": "2026-05-10",
            "SHIPMENT": "2026-06-01", "CUSTOMS": "2026-07-01"}
    tl = pd.DataFrame(
        _done("R1", {"ORDER_REG": "2026-05-01", "ALLOCATION_QUEUE": "2026-05-10", "ALLOCATION": "2026-06-01",
                     "SHIPMENT": "2026-08-01"})
        + _done("R2", full)
        + _done("R3", {"ORDER_REG": "2026-08-01", "ALLOCATION_QUEUE": "2026-09-20"})
        + _done("R4", dict(full, CUSTOMS="2026-06-01")))
    ctx = PipelineContext(rb=_RB(), today=TODAY)
    ctx.sources["credit"] = {"main": credit}
    ctx.extras["fx_stage_timeline"] = tl
    out = BLRegistrationLinkStage().run(pd.DataFrame(rows), ctx)
    money = pd.DataFrame([
        dict(KEY_REG="R1", CURRENCY="EUR", REGISTRATION_AMOUNT=1000, REQUESTED_AMOUNT=1000, ALLOCATED_AMOUNT=950,
             PURCHASED_AMOUNT=None, SUPPLIER_PAID_AMOUNT=None, COMMITMENT_INITIAL=1000, COMMITMENT_RELEASED=0,
             COMMITMENT_BALANCE=1000, EVIDENCE_GAPS="PURCHASED_AMOUNT | SUPPLIER_PAID_AMOUNT",
             RECON_STATUS="EVIDENCE_GAP"),
        dict(KEY_REG="R2", CURRENCY="EUR", REGISTRATION_AMOUNT=1000, REQUESTED_AMOUNT=1000, ALLOCATED_AMOUNT=1000,
             PURCHASED_AMOUNT=1000, SUPPLIER_PAID_AMOUNT=1000, COMMITMENT_INITIAL=1000, COMMITMENT_RELEASED=600,
             COMMITMENT_BALANCE=400, EVIDENCE_GAPS="", RECON_STATUS="RECONCILED"),
        dict(KEY_REG="R3", CURRENCY="CNY", REGISTRATION_AMOUNT=5000, REQUESTED_AMOUNT=5000, ALLOCATED_AMOUNT=None,
             PURCHASED_AMOUNT=None, SUPPLIER_PAID_AMOUNT=None, COMMITMENT_INITIAL=5000, COMMITMENT_RELEASED=None,
             COMMITMENT_BALANCE=None, EVIDENCE_GAPS="ALLOCATED_AMOUNT", RECON_STATUS="EVIDENCE_GAP"),
        dict(KEY_REG="R4", CURRENCY="", REGISTRATION_AMOUNT=None, REQUESTED_AMOUNT=300, ALLOCATED_AMOUNT=None,
             PURCHASED_AMOUNT=None, SUPPLIER_PAID_AMOUNT=None, COMMITMENT_INITIAL=None, COMMITMENT_RELEASED=None,
             COMMITMENT_BALANCE=None, EVIDENCE_GAPS="", RECON_STATUS="EVIDENCE_GAP"),
    ])
    extras = dict(ctx.extras)
    extras["fx_money_reconciliation"] = money
    return out, extras


@pytest.fixture(scope="module")
def built():
    return _build()


@pytest.fixture(scope="module")
def fx(built):
    out, extras = built
    return X.load(out, extras, REF)


@pytest.fixture(scope="module")
def report(built):
    out, extras = built
    return H.build_report(out, extras, REF, embed_fonts=False)


def _book(data: bytes):
    return load_workbook(io.BytesIO(data))


def _header_row(ws) -> int:
    teal = T.BRAND_TEAL.lstrip("#").upper()
    for r in range(1, 8):
        rgb = str(ws.cell(row=r, column=1).fill.fgColor.rgb or "")
        if rgb.upper().endswith(teal):
            return r
    raise AssertionError(f"سرستون فیروزه‌ای در شیت {ws.title} نیست")


def _table(ws) -> pd.DataFrame:
    head = _header_row(ws)
    rows = list(ws.iter_rows(min_row=head, values_only=True))
    return pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]], dtype=object)   # خانه خالی None می‌ماند


def _no_nan_text(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, float) and math.isnan(v)), ws.title
                assert not (isinstance(v, str) and v.strip() in ("nan", "NaN", "None", "NaT")), (ws.title, v)


# ═══════════════════════════ دانه‌ها و ارزها ═══════════════════════════
def test_current_stage_is_the_step_after_the_most_advanced_done_step(fx):
    stage = dict(zip(fx.reg_table["KEY_REG"], fx.reg_table["STAGE_CODE"]))
    assert stage == {"R1": "CUSTOMS", "R2": "CLEARANCE", "R3": "ALLOCATION", "R4": "CLEARANCE"}
    r1 = fx.reg_table.set_index("KEY_REG").loc["R1"]
    assert r1["GAP_COUNT"] >= 3                      # خرید ارز، تأمین وجه و سوئیفت بی‌شاهدند


def test_order_pi_is_counted_once_and_material_rows_carry_no_money(fx):
    rows = XE.hierarchy_rows(fx)
    o1 = [rec for _lvl, kind, rec in rows if kind == "order" and rec.get("سفارش") == "O1"]
    assert len(o1) == 1 and o1[0]["مبلغ"] == 900     # سه ردیف مارت، یک PI
    mats = [rec for _lvl, kind, rec in rows if kind == "mat"]
    assert mats and all(rec.get("مبلغ") is None for rec in mats)
    mf = X.materials_frame(fx)
    assert not [c for c in mf.columns if any(w in c for w in ("مبلغ", "ارزش"))]
    assert sorted(mf[mf["ثبت سفارش"].eq("R1")]["متریال"]) == ["M1", "M2"]


def test_bl_value_is_counted_once_despite_material_fan_out(fx):
    r1 = fx.reg_table.set_index("KEY_REG").loc["R1"]
    assert r1["SHIPPED_VALUE"] == 500                # BL1 روی دو متریال هنوز ۴۰۰ است
    bls = {b["key"]: b["value"] for b in X.bls(fx, "R1")}
    assert bls == {"BL1": 400, "BL2": 100}
    bl_rows = [rec for _lvl, kind, rec in XE.hierarchy_rows(fx) if kind == "bl" and rec["ثبت سفارش"] == "R1"]
    assert sorted(rec["مبلغ"] for rec in bl_rows) == [100, 400]


def test_stage_values_never_mix_currencies_and_unknown_stays_separate(fx):
    v = X.stage_values(fx)
    assert not v.duplicated(["STAGE_CODE", "CURRENCY"]).any()
    clear = v[v["STAGE_CODE"].eq("CLEARANCE")].set_index("CURRENCY")
    assert set(clear.index) == {"EUR", "نامشخص"}
    unknown = clear.loc["نامشخص"]
    assert pd.isna(unknown["REG_VALUE"]) and unknown["UNKNOWN_VALUE_REGS"] == 1
    assert pd.isna(unknown["SHIPPED_VALUE"])          # بارنامه درهم به هیچ ارزی اضافه نمی‌شود
    assert clear.loc["EUR", "REG_VALUE"] == 1000
    assert v[v["STAGE_CODE"].eq("ALLOCATION")].iloc[0]["CURRENCY"] == "CNY"


def test_overshipment_is_not_netted_against_other_registrations(fx):
    clear = X.stage_values(fx).set_index(["STAGE_CODE", "CURRENCY"]).loc[("CLEARANCE", "EUR")]
    assert clear["UNSHIPPED_VALUE"] == 0 and clear["OVERSHIPPED_REGS"] == 1
    # R1 را هم در «ترخیص» بگذار: مانده ۵۰۰ آن نباید با ۳۰۰- R2 به ۲۰۰ برسد
    moved = X.FxData(**{k: getattr(fx, k) for k in ("lc", "link", "recon", "ledger", "money", "queue", "control",
                                                    "decisions", "positions", "mart")}, ref_date=fx.ref_date)
    table = fx.reg_table.copy()
    table.loc[table["KEY_REG"].eq("R1"), "STAGE_CODE"] = "CLEARANCE"
    moved.__dict__["reg_table"] = table
    eur = X.stage_values(moved).set_index(["STAGE_CODE", "CURRENCY"]).loc[("CLEARANCE", "EUR")]
    assert eur["REGISTRATIONS"] == 2 and eur["REG_VALUE"] == 2000 and eur["SHIPPED_VALUE"] == 1800
    assert eur["UNSHIPPED_VALUE"] == 500 and eur["OVERSHIPPED_REGS"] == 1


def test_stage_click_counts_match_the_stage_summary(fx):
    summary = X.stage_summary(fx).set_index("STAGE_CODE")
    model = X.explorer(fx)
    assert [s["code"] for s in model] == [c for c, _ in X.STAGES]
    for s in model:
        assert s["count"] == summary.loc[s["code"], "REGISTRATIONS"]
        assert len(s["regs"]) + s["hidden"] == s["count"]
    capped = {s["code"]: s for s in X.explorer(fx, max_regs=1)}
    assert capped["CLEARANCE"]["count"] == 2 and len(capped["CLEARANCE"]["regs"]) == 1
    assert capped["CLEARANCE"]["hidden"] == 1
    node = next(r for s in model for r in s["regs"] if r["KEY_REG"] == "R1")
    assert [o["key"] for o in node["orders"]] == ["O1"]
    assert [m["key"] for m in node["orders"][0]["materials"]] == ["M1", "M2"]
    assert [b["key"] for b in node["bls"]] == ["BL1", "BL2"]
    crit = {s["code"]: s["count"] for s in X.explorer(fx, critical_only=True)}
    assert crit["ALLOCATION"] == 1 and crit["CLEARANCE"] == 1 and crit["CUSTOMS"] == 0


def test_future_queue_entry_is_an_unknown_wait_not_a_negative_one(fx):
    q = fx.queue.set_index("KEY_REG")
    assert pd.isna(q.loc["R1", "WAIT_DAYS"]) and X.s(q.loc["R1", "WAIT_NOTE"])
    assert q.loc["R3", "WAIT_DAYS"] == 7
    assert (pd.to_numeric(q["WAIT_DAYS"], errors="coerce").dropna() >= 0).all()
    assert "ورود به صف بعد از تاریخ مرجع" in H.queue_section(fx)


# ═══════════════════════════ Excel ═══════════════════════════
def test_stage_workbook_is_rtl_teal_and_the_hierarchy_folds_in_four_levels(fx):
    wb = _book(XE.stage_workbook(fx))
    assert wb.sheetnames == ["خلاصه مراحل", "ارزش به تفکیک ارز", "ثبت سفارش‌ها", "سلسله‌مراتب", "سفارش‌ها",
                             "متریال‌ها", "بارنامه‌ها"]
    for ws in wb.worksheets:
        assert ws.sheet_view.rightToLeft, ws.title
        _header_row(ws)
    _no_nan_text(wb)
    stages = _table(wb["خلاصه مراحل"]).set_index("مرحله")["ثبت سفارش"]
    assert stages[X.STAGE_FA["CLEARANCE"]] == 2 and stages.sum() == 4
    ws = wb["سلسله‌مراتب"]
    assert ws.sheet_properties.outlinePr.summaryBelow is False
    head = _header_row(ws)
    rows = XE.hierarchy_rows(fx)
    levels = [ws.row_dimensions[i].outline_level or 0 for i in range(head + 1, head + 1 + len(rows))]
    assert levels == [lvl for lvl, _k, _r in rows]
    assert set(levels) == {0, 1, 2, 3}
    values = _table(wb["ارزش به تفکیک ارز"])
    unknown = values[values["ارز"].eq("نامشخص")].iloc[0]
    assert unknown["ارزش ثبت سفارش"] is None          # نامعلوم خانه خالی است، نه صفر


def test_single_stage_workbook_holds_only_that_stage(fx):
    wb = _book(XE.stage_workbook(fx, stages=["CLEARANCE"]))
    regs = _table(wb["ثبت سفارش‌ها"])
    assert sorted(regs["ثبت سفارش"]) == ["R2", "R4"]
    assert set(_table(wb["سلسله‌مراتب"])["مرحله"].dropna()) == {X.STAGE_FA["CLEARANCE"]}


def test_registration_workbook_keeps_unknown_steps_empty(fx):
    wb = _book(XE.registration_workbook(fx, ["R1"]))
    assert wb.sheetnames == ["خلاصه", "جریان پول", "رویدادهای مبلغی", "بارنامه‌ها", "سفارش‌ها", "متریال‌ها",
                             "مراحل چرخه", "تصمیم‌های مالی", "صف تخصیص"]
    _no_nan_text(wb)
    flow = _table(wb["جریان پول"])
    assert flow["ثبت سفارش"].tolist() == ["R1"]
    r1 = flow.iloc[0]
    assert r1[X.MONEY_FA["PURCHASED_AMOUNT"]] is None and r1[X.MONEY_FA["ALLOCATED_AMOUNT"]] == 950
    assert r1[X.MONEY_FA["COMMITMENT_RELEASED"]] == 0  # صفرِ ثبت‌شده صفر می‌ماند
    assert sorted(_table(wb["بارنامه‌ها"])["بارنامه"]) == ["BL1", "BL2"]
    path = _table(wb["مراحل چرخه"])
    assert path["مرحله"].tolist() == [fa for _c, fa in X.LIFECYCLE]
    all_regs = _book(XE.registration_workbook(fx, None))
    assert "جمع به تفکیک ارز" in all_regs.sheetnames
    totals = _table(all_regs["جمع به تفکیک ارز"])
    assert not totals.duplicated(subset=[totals.columns[0]]).any()   # یک ردیف برای هر ارز


def test_critical_workbook_counts_only_critical_levels_by_stage(built, fx):
    out, _extras = built
    wb = _book(XE.critical_workbook(fx, out))
    assert wb.sheetnames == ["ماتریس مرحله × سطح", "بحرانی به تفکیک مرحله", "سلسله‌مراتب بحرانی",
                             "متریال‌های بحرانی", "بارنامه‌های بحرانی", "ثبت سفارش‌ها"]
    _no_nan_text(wb)
    by_stage = _table(wb["بحرانی به تفکیک مرحله"])
    assert sorted(by_stage["متریال"]) == ["M4", "M5"]           # SAFE و WATCH بیرون می‌مانند
    matrix = _table(wb["ماتریس مرحله × سطح"]).set_index("مرحله")
    assert matrix.loc[X.STAGE_FA["ALLOCATION"], CB.level_label("CRITICAL")] == 1
    assert matrix.loc[X.STAGE_FA["CLEARANCE"], CB.level_label("STOCKOUT")] == 1
    kinds = [k for _l, k, _r in XE.hierarchy_rows(fx, critical_only=True)]
    assert "bl" not in kinds and kinds.count("mat") == 2
    one = _book(XE.critical_workbook(fx, out, stages=["ALLOCATION"]))
    assert _table(one["ثبت سفارش‌ها"])["ثبت سفارش"].tolist() == ["R3"]


def test_english_workbooks_translate_labels_and_keep_data(built, fx):
    out, _extras = built
    books = [XE.stage_workbook(fx, lang=C.EN), XE.registration_workbook(fx, ["R1"], lang=C.EN),
             XE.registration_workbook(fx, None, lang=C.EN), XE.critical_workbook(fx, out, lang=C.EN)]
    data_text = {f"شرح M{i}" for i in range(1, 6)} | {"شاهد"}
    for data in books:
        wb = _book(data)
        assert all(not PERSIAN.search(n) for n in wb.sheetnames), wb.sheetnames
        assert not any(ws.sheet_view.rightToLeft for ws in wb.worksheets)   # انگلیسی چپ‌به‌راست
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and PERSIAN.search(v):
                        assert v in data_text, (ws.title, v)
    hier = _book(books[0])["Hierarchy"]
    assert hier.sheet_properties.outlinePr.summaryBelow is False
    assert any(isinstance(c.value, str) and c.value.startswith("Registrations: ")
               for row in hier.iter_rows() for c in row)
    lists = _table(_book(books[0])["Materials"])
    assert "BL1, BL2" in lists.iloc[:, 8].tolist()                # فهرست کدها با ویرگول لاتین


# ═══════════════════════════ HTML ═══════════════════════════
def _allowed_colours():
    allowed = {c.lower() for c in (T.SURFACE_PAGE, T.SURFACE_RAISED, T.SURFACE_SUNKEN, T.SURFACE_INVERSE,
                                   T.SURFACE_PAPER, T.SURFACE_PAPER_SOFT, T.PAPER_RULE, T.PENCIL, T.AQUA_MIST,
                                   T.LAPIS_WASH, T.BORDER, T.BORDER_STRONG, T.BORDER_FOCUS, T.TEXT,
                                   T.TEXT_SECONDARY, T.TEXT_MUTED, T.TEXT_ON_DARK, T.BRAND_NAVY, T.BRAND_TEAL,
                                   T.BRAND_GOLD, T.TEAL_INK, T.GOLD_INK, T.TEAL_WASH, T.GOLD_WASH, T.NAVY_WASH,
                                   T.TEXT_ON_BRAND)}
    allowed |= {v.lower() for s in T.STATUS_SCALE for v in (s.ink, s.fill, s.wash)}
    allowed |= {c.lower() for c in T.CATEGORICAL + T.SEQUENTIAL + T.TEAL_PALETTE}
    allowed |= {c.lower() for c in T.BRAND_BAND}  # R10: سربرگ برند با پالت موشن‌گرافیک
    return allowed


def test_standalone_report_is_offline_rtl_svg_free_and_token_only(report):
    assert report.startswith("<!DOCTYPE html>") and 'dir="rtl"' in report
    assert "<script" not in report and "<svg" not in report
    assert not re.search(r'(?:src|href)="(?:https?:)?//', report) and "@import" not in report
    used = {m.lower() for m in re.findall(r"#[0-9A-Fa-f]{6}\b", report)}
    assert not (used - _allowed_colours()), sorted(used - _allowed_colours())
    panes = re.split(r'<section class="gx-tpane gx-tp\d+">', report)[1:]
    assert len(panes) == len(H.TABS)
    assert panes[0].count('<details class="gx-reg') == 4           # هر ثبت سفارش یک بار، در مرحله خودش
    assert panes[5].count('<details class="gx-reg') == 2           # تب بحرانی: فقط R3 و R4
    assert report.count('type="radio" class="gx-r ') == 2 * len(X.STAGES)


def test_standalone_report_embeds_real_workbooks(report):
    links = dict(re.findall(r'download="([^"]+\.xlsx)" href="data:[^;]+;base64,([^"]+)"', report))
    for name in (f"GSI_FX_STAGES_{REF}.xlsx", f"GSI_FX_FINANCE_ALL_{REF}.xlsx", f"GSI_CRITICAL_BY_STAGE_{REF}.xlsx",
                 f"GSI_FX_REG_R1_{REF}.xlsx", f"GSI_FX_STAGE_CLEARANCE_{REF}.xlsx",
                 f"GSI_CRITICAL_STAGE_ALLOCATION_{REF}.xlsx"):
        assert name in links, name
    assert f"GSI_CRITICAL_STAGE_CUSTOMS_{REF}.xlsx" not in links  # مرحله بدون اقلام بحرانی فایل ندارد
    wb = _book(base64.b64decode(links[f"GSI_FX_REG_R1_{REF}.xlsx"]))
    assert _table(wb["جریان پول"])["ثبت سفارش"].tolist() == ["R1"]


def test_streamlit_safe_css_and_fragments(built):
    out, extras = built
    css = H.css()
    assert "<" not in css                                  # DOMPurify استایل دارای «<» را حذف می‌کند
    assert "<" not in CB.critical_css()
    assert ".gx-x{container-type:inline-size}" in css and "mi-cring" in css
    block = H.composer_block(out, extras, REF, key="t")
    assert "<svg" not in block and "<script" not in block
    assert block.count('<details class="gx-reg') == 4
    assert 'role="tablist"' not in block                   # قاعده سراسری Studio شکل ریل را نمی‌شکند


def test_english_report_translates_labels_but_keeps_codes(built):
    out, extras = built
    page = H.build_report(out, extras, REF, lang=C.EN, embed_fonts=False)
    assert 'lang="en"' in page and "Registrations per stage" in page
    body = re.sub(r"<style>.*?</style>", "", page, flags=re.S)
    body = re.sub(r'href="data:[^"]*"', "", body)
    left = {t.strip() for t in re.split(r"<[^>]+>", body) if PERSIAN.search(t)}
    assert left <= {f"شرح M{i}" for i in range(1, 6)}, sorted(left)
    assert C.phrase("BL1", C.EN) == "BL1" and C.phrase("EUR", C.EN) == "EUR"
    assert C.phrase("ترخیص — تخصیص ارز", C.EN) == f'{C.phrase("ترخیص", C.EN)} — {C.phrase("تخصیص ارز", C.EN)}'



def test_english_pages_read_left_to_right_and_persian_pages_right_to_left(built):
    from gsi.design.css import stylesheet
    out, extras = built
    en = H.build_report(out, extras, REF, lang=C.EN, embed_fonts=False, embed_excel=False)
    fa = H.build_report(out, extras, REF, embed_fonts=False, embed_excel=False)
    assert '<html lang="en" dir="ltr">' in en and '<html lang="fa" dir="rtl">' in fa
    # ریشه .gx و body خودشان direction:rtl دارند؛ بدون این قاعده‌ها dir="ltr" بی‌اثر می‌ماند
    css = H.css()
    assert '[dir="ltr"] .gx,.gx[dir="ltr"]{direction:ltr}' in css
    assert '[dir="ltr"] .gx-reg.is-crit' in css
    assert 'html[dir="ltr"] body{direction:ltr}' in stylesheet()
    assert '<html lang="en" dir="ltr">' in CB.build_critical_html(out, REF, lang=C.EN)
    assert '<html lang="fa" dir="rtl">' in CB.build_critical_html(out, REF)


def test_stepper_connector_joins_each_step_to_the_previous_one():
    css = I.minimal_css()
    rule = re.search(r"\.mi-step::before \{[^}]*\}", css).group(0)
    assert "inset-inline-end:50%" in rule and "width:100%" in rule   # در هر دو جهت به گام قبلی
    assert ".mi-step:first-child::before { display:none; }" in css


def test_registration_boxes_wrap_instead_of_pinning_empty_grid_columns(built):
    out, extras = built
    css = H.css()
    # grid با auto-fit و بخش تمام‌عرض، ستون‌های خالی را نگه می‌داشت و جدول بارنامه سرریز می‌شد
    assert ".gx-grid{display:flex;flex-wrap:wrap" in css and "grid-column:1/-1" not in css
    block = H.composer_block(out, extras, REF, key="t", embed_excel=False)
    assert block.count('<section class="gx-box is-bls">') == 4


def test_report_builder_block_uses_the_same_components(built):
    from gsi.studio_core import composer
    from gsi.studio_core.html_export import build_dynamic_html
    out, extras = built
    assert "fx_lifecycle" in composer.BLOCKS   # چیدمان‌های ذخیره‌شده؛ تب تازه از 29.20 همین بخش‌ها را جدا دارد
    assert all({"fx_kpi", "fx_stages", "fx_money"} <= set(blocks) for blocks in composer.DEFAULT_BLOCKS.values())
    page = build_dynamic_html(out, REF, tabs=[{"id": "fx", "title": "FX", "blocks": ["fx_lifecycle"]}],
                              process_extras=extras)
    assert "داده چرخه ارز برای این گزارش در دسترس نیست" not in page
    assert page.count('<details class="gx-reg') == 4
    assert page.count(".gx-x{container-type:inline-size}") == 1   # CSS بلوک یک بار در <style> صفحه


def test_email_section_is_outlook_safe(fx):
    mail = H.email_section(fx, attachment=f"GSI_FX_LIFECYCLE_{REF}.html")
    for bad in ("class=", "<style", "<input", "<svg", "<details", "<script"):
        assert bad not in mail, bad
    assert X.STAGE_FA["CLEARANCE"] in mail and "EUR" in mail and "CNY" in mail
    assert f"GSI_FX_LIFECYCLE_{REF}.html" in mail
    assert H.email_section(X.load(None, {}, REF)) == ""


def test_studio_email_attachment_stays_inside_the_report_scope(built, tmp_path, monkeypatch):
    from gsi.integrations import daily_email as DE
    out, extras = built
    scope = out[out["KEY_REG"].eq("R1")]
    body, path = DE._fx_lifecycle_part(TODAY, scope, extras, tmp_path, scoped=True)
    page = path.read_text(encoding="utf-8")
    assert 'class="gx-key">R1<' in page
    assert not [r for r in ("R2", "R3", "R4") if f'class="gx-key">{r}<' in page]
    assert X.STAGE_FA["CUSTOMS"] in body and X.STAGE_FA["CLEARANCE"] not in body
    monkeypatch.setenv("GSI_EMAIL_FX", "0")
    assert DE._fx_lifecycle_part(TODAY, scope, extras, tmp_path, scoped=True) == ("", None)


def test_write_report_names_the_file_by_date_and_language(built, tmp_path):
    out, extras = built
    fa = H.write_report(out, extras, REF, tmp_path, embed_excel=False, embed_fonts=False)
    en = H.write_report(out, extras, REF, tmp_path, lang=C.EN, embed_excel=False, embed_fonts=False)
    assert fa.name == f"GSI_FX_LIFECYCLE_{REF}.html" and en.name == f"GSI_FX_LIFECYCLE_{REF}_EN.html"
    assert "data:application/vnd.openxmlformats" not in fa.read_text(encoding="utf-8")


# ═══════════════════════════ آیکن CSS ═══════════════════════════
def test_css_icon_mode_is_scoped_and_never_leaks():
    assert not I.css_mode() and "<svg" in I.icon("file", 16)
    with I.css_icons():
        assert I.css_mode()
        frag = I.icon("file", 16)
        assert "<svg" not in frag and "mi-ic-file" in frag
        assert "mi-cring" in I.ring(40) and "<svg" not in I.ring(40)
    assert not I.css_mode() and "<svg" in I.icon("file", 16)
    with pytest.raises(RuntimeError):
        with I.css_icons():
            raise RuntimeError("boom")
    assert not I.css_mode()


def test_critical_fragment_in_css_mode_has_no_svg_but_the_standalone_board_keeps_it(built):
    out, _extras = built
    with I.css_icons():
        frag = CB.critical_fragment(out)
    assert frag and "<svg" not in frag
    assert "<svg" in CB.build_critical_html(out, REF)


# ═══════════════════════════ فونت ═══════════════════════════
@pytest.fixture
def font_env(monkeypatch, tmp_path):
    """فقط پوشه‌های همین آزمون دیده شوند؛ کش پیش و پس از آزمون پاک می‌شود."""
    for var in ("GSI_FONT_DIR", "GSI_FONT_PATH", "GSI_EMBED_FONTS", "GSI_FONT_SCAN_SYSTEM", "GSI_FONT_FANUM"):
        monkeypatch.delenv(var, raising=False)
    web, system = tmp_path / "web", tmp_path / "system"
    web.mkdir()
    system.mkdir()
    monkeypatch.setattr(F, "package_dirs", lambda: [])
    monkeypatch.setattr(F, "system_dirs", lambda: [system])
    F.refresh()
    yield web, system
    F.refresh()


def test_explicit_font_files_are_embedded(font_env, monkeypatch):
    web, _system = font_env
    (web / "IRANSansWeb.woff2").write_bytes(b"wOF2-regular")
    (web / "IRANSansWeb_Bold.woff2").write_bytes(b"wOF2-bold")
    monkeypatch.setenv("GSI_FONT_DIR", str(web))
    F.refresh()
    css = F.html_font_css()
    assert css.count("@font-face") == 2 and "font/woff2" in css and "font-weight:700" in css
    assert f"font-family:'{F.FAMILY}'" in css
    st = F.status()
    assert st["ok"] and st["embedded"] and st["embedded_in_html"]
    monkeypatch.setenv("GSI_EMBED_FONTS", "0")
    assert F.html_font_css() == "" and F.font_face_css()   # Studio جاسازی می‌کند، فایل ارسالی نه


def test_installed_system_font_is_never_embedded_but_charts_use_it(font_env):
    _web, system = font_env
    ttf = system / "IRANSans.ttf"
    ttf.write_bytes(b"\x00\x01\x00\x00ttf")
    F.refresh()
    assert F.html_font_css() == "" and F.font_face_css() == ""
    assert F.chart_font_path() == ttf
    st = F.status()
    assert st["ok"] and not st["embedded"] and not st["embedded_in_html"]


def test_font_discovery_is_cached_until_refresh(font_env, monkeypatch):
    web, _system = font_env
    monkeypatch.setenv("GSI_FONT_DIR", str(web))
    (web / "IRANSansWeb.woff2").write_bytes(b"regular")
    F.refresh()
    assert sorted(F.cached_discover()["iransans"]) == [400]
    (web / "IRANSansWeb_Bold.woff2").write_bytes(b"bold")
    assert sorted(F.cached_discover()["iransans"]) == [400]
    F.refresh()
    assert sorted(F.cached_discover()["iransans"]) == [400, 700]


def test_font_file_names_are_classified_by_weight_variant_and_digits(tmp_path):
    from pathlib import Path
    f = F.classify(Path("IRANSansWeb(FaNum)_Bold.woff2"))
    assert (f.weight, f.variant, f.fanum, f.fmt) == (700, "web", True, "woff2")
    f = F.classify(Path("IRANSans_Medium.ttf"))
    assert (f.weight, f.variant, f.fanum, f.fmt) == (500, "desktop", False, "truetype")
    assert F.classify(Path("IRANSansX-UltraLight.woff")).weight == 200
    f = F.classify(Path("RaviFaNum-Regular.woff"))
    assert (f.family, f.weight, f.fanum, f.fmt) == ("ravi", 400, True, "woff")
    assert (F.classify(Path("Ravi-SemiBold.ttf")).weight, F.classify(Path("Ravi-ExtraBlack.woff2")).weight) == (600, 950)
    assert F.classify(Path("Ravi-Bold.woff2")).family == "ravi"
    assert F.classify(Path("RAVIE.TTF")) is None          # فونت تزئینی ویندوز، نه راوی
    assert F.classify(Path("Vazirmatn-Regular.ttf")) is None
    assert F.classify(Path("IRANSans.txt")) is None


def test_font_stack_puts_ravi_first_then_iransans():
    for name in ("Ravi", "IRANSansWeb", "IRANSans", "Ravi FaNum", "IRANSansWeb(FaNum)", "IRANSans(FaNum)"):
        assert f"'{name}'" in T.FONT_STACK, name
    stack = T.FONT_STACK
    assert stack.startswith("'Ravi',") and T.FONT_STACK_NUM.startswith("'Ravi',")
    assert stack.index("'IRANSansWeb'") < stack.index("'Ravi FaNum'") < stack.index("Tahoma")


def test_ravi_comes_first_and_iransans_fills_the_gaps(font_env, monkeypatch):
    web, system = font_env
    # همان نام‌هایی که مالک فرستاد؛ راوی معمولی فقط نسخه ارقام فارسی دارد
    for name in ("RaviFaNum-Regular.woff", "Ravi-SemiBold.woff2", "Ravi-SemiBold.ttf", "Ravi-Bold.woff2",
                 "Ravi-ExtraBlack.woff2", "IRANSans-web.woff"):
        (web / name).write_bytes(name.encode())
    (system / "IRANSans.ttf").write_bytes(b"\x00\x01\x00\x00ttf")
    monkeypatch.setenv("GSI_FONT_DIR", str(web))
    F.refresh()
    css = F.html_font_css()
    faces = re.findall(r"font-family:'([^']+)';font-style:normal;font-weight:(\d+)", css)
    assert faces == [("Ravi", "400"), ("Ravi", "600"), ("Ravi", "700"), ("IRANSansWeb", "400")]
    assert "font-weight:950" not in css                                 # ExtraBlack در طراحی نیست
    # ارقام لاتین متن معمولی از ایران‌سنس می‌آید، نه از فایل FaNum راوی
    assert css.count("unicode-range") == 1 and f"unicode-range:{F.NO_ASCII_DIGITS}" in css.split("\n")[0]
    st = F.status()
    assert st["family"] == "Ravi" and st["families"] == ["Ravi", "IRANSansWeb"]
    assert st["short"].startswith("راوی و ایران‌سنس")
    assert F.chart_font_path() == web / "Ravi-SemiBold.ttf"         # راوی پیش از ایران‌سنس نصب‌شده
    monkeypatch.setenv("GSI_FONT_FANUM", "1")
    F.refresh()
    assert "unicode-range" not in F.html_font_css()


def test_iransans_alone_still_works_when_ravi_is_absent(font_env, monkeypatch):
    web, _system = font_env
    (web / "IRANSans-web.woff").write_bytes(b"woff")
    monkeypatch.setenv("GSI_FONT_DIR", str(web))
    F.refresh()
    css = F.html_font_css()
    assert "font-family:'Ravi'" not in css and css.count("font-family:'IRANSansWeb'") == 1
    assert F.status()["family"] == "IRANSansWeb"
