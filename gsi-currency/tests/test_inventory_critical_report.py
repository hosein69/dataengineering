# -*- coding: utf-8 -*-
"""کنترل موجودی و اقلام بحرانی — فقط داده مصنوعی.

سفارش O1: M1 (FIRST، ایمن، روی BL1 و BL2 — ردیف FIRST به ازای هر بارنامه تکرار شده)
و M2 (ADDITIONAL، بحرانی، بی‌بارنامه). سفارش O2: M3 (FIRST، روی BL1) و M4 (ADDITIONAL)،
هر دو ایمن. سفارش O3: M5 (FIRST، روی BL3، بدون هیچ داده موجودی، سطح نامشخص).
"""
from __future__ import annotations

import io
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from gsi.report import inventory_critical_report as R

NAN = float("nan")


def _row(order, mat, role, bl, code, res, ikco, sapco, need, transit, customs=0.0, conflict="", **kw):
    return {"KEY_ORDER": order, "CANONICAL_ORDER": order, "KEY_MATERIAL": mat, "MOGH_ITEM_ROLE": role,
            "CANONICAL_BL": bl, "KEY_BL": bl, "KEY_REG": "R-" + order, "کد طبقه بحرانی": code,
            "مقاومت (روز)": res, "STOCK_IKCO": ikco, "STOCK_SAPCO": sapco, "DAILY_NEED": need,
            "MOGH_QTY_AT_SUPPLIER": 0.0 if not math.isnan(transit) else NAN,
            "MOGH_QTY_READY": 0.0 if not math.isnan(transit) else NAN,
            "MOGH_QTY_IN_TRANSIT": transit, "MOGH_QTY_IN_CUSTOMS": customs if not math.isnan(transit) else NAN,
            "MOGH_QTY_STATE_UNKNOWN": 0.0 if not math.isnan(transit) else NAN,
            "MOGH_INVENTORY_CONFLICT": conflict, "CANONICAL_EXPERT": "کارشناس " + order,
            "ORG_DEPT": "اداره " + order, "STATUS_WHERE": "گمرک و ترخیص", "NEXT_ACTION_TITLE": "پیگیری " + order,
            "DISCHARGE_DATE": "1405/06/01" if bl else "", "IS_FULL_CLEARED": False, "IS_PARTIAL_CLEARED": False,
            "MATERIAL_DESC": "شرح " + mat, **kw}


def _frame() -> pd.DataFrame:
    return pd.DataFrame([
        _row("O1", "M1", "FIRST", "BL1", "SAFE", 90.0, 900.0, 0.0, 10.0, 5.0),
        _row("O1", "M1", "FIRST", "BL2", "SAFE", 90.0, 900.0, 0.0, 10.0, 5.0),   # پخش روی بارنامه دوم
        _row("O1", "M2", "ADDITIONAL", "", "CRITICAL", 4.0, 40.0, 0.0, 10.0, 7.0, conflict="Part 9: در راه | در گمرک"),
        _row("O2", "M3", "FIRST", "BL1", "SAFE", 60.0, 600.0, 0.0, 10.0, 1.0),
        _row("O2", "M4", "ADDITIONAL", "", "SAFE", 70.0, 700.0, 0.0, 10.0, 2.0),
        _row("O3", "M5", "FIRST", "BL3", "", NAN, NAN, NAN, NAN, NAN),
    ])


@pytest.fixture(scope="module")
def model() -> R.InventoryModel:
    return R.build_model(_frame(), None, "1405/06/10")


def test_one_row_per_key_and_unique_kpis(model):
    assert model.materials["متریال"].is_unique and len(model.materials) == 5
    assert model.orders["سفارش"].is_unique and len(model.orders) == 3
    assert model.bls["بارنامه"].is_unique and set(model.bls["بارنامه"]) == {"BL1", "BL2", "BL3"}
    k = model.kpis
    assert k["متریال در گزارش"] == 5            # نه ۶ ردیف
    assert k["سفارش با متریال بحرانی"] == 1
    assert k["بارنامه با سفارش بحرانی"] == 2      # BL1 و BL2؛ نه شمار ردیف‌ها
    o1 = model.orders.set_index("سفارش").loc["O1"]
    assert o1["تعداد متریال"] == 2 and o1["کد سطح"] == "CRITICAL"


def test_bl_critical_if_any_order_critical(model):
    bl = model.bls.set_index("بارنامه")
    # BL1 دو سفارش دارد؛ O1 بحرانی است (از ردیف ADDITIONAL بی‌بارنامه M2)
    assert bl.loc["BL1", "بارنامه بحرانی"] == "بله" and bl.loc["BL1", "کد سطح"] == "CRITICAL"
    assert bl.loc["BL2", "کد سطح"] == "CRITICAL"
    assert bl.loc["BL1", R.BL_MIN_RES] == 4.0
    assert set(bl.loc["BL1", R.BL_ORDERS].split("، ")) == {"O1", "O2"}
    assert set(bl.loc["BL1", R.BL_MATS].split("، ")) == {"M1", "M2", "M3", "M4"}
    assert bl.loc["BL3", "بارنامه بحرانی"] == "خیر"
    assert bl.loc["BL1", "روز از تخلیه"] == 9
    assert R.BL_MATS == "متریال‌های سفارش‌های این بارنامه"
    assert R.MAT_BLS == "بارنامه‌های سفارش‌های این متریال"


def test_quantities_use_order_material_grain(model):
    m = model.materials.set_index("متریال")
    assert m.loc["M1", "در راه"] == 5.0           # ردیف تکراری روی BL2 دو برابرش نمی‌کند
    assert m.loc["M2", "در راه"] == 7.0
    # یک متریال در دو سفارش: جمع روی جفت‌های یکتا
    df = pd.concat([_frame(), pd.DataFrame([_row("O4", "M1", "FIRST", "BL4", "SAFE", 90.0, 900.0, 0.0, 10.0, 3.0)])],
                   ignore_index=True)
    m2 = R.build_model(df, None, "1405/06/10").materials.set_index("متریال")
    assert m2.loc["M1", "در راه"] == 8.0
    assert m2.loc["M1", "تعداد سفارش"] == 2
    assert m2.loc["M1", "پوشش داده کارشناس"] == "2/2"


def test_material_without_inventory_is_gap_not_zero(model):
    m5 = model.materials.set_index("متریال").loc["M5"]
    assert m5["وضعیت داده موجودی"] == R.INV_NONE
    for c in ("موجودی ایران‌خودرو", "موجودی ساپکو", "در راه", "در گمرک", "نزد سازنده"):
        assert m5[c] is None or pd.isna(m5[c])
    gaps = model.checks[model.checks["بررسی"].eq(R.CHK_NO_DATA)]
    assert gaps["متریال"].tolist() == ["M5"]
    assert model.kpis[R.CHK_NO_DATA] == 1


def test_conflict_and_followup(model):
    conf = model.checks[model.checks["بررسی"].eq(R.CHK_CONFLICT)]
    assert conf[["متریال", "سفارش"]].values.tolist() == [["M2", "O1"]]
    assert model.kpis["سفارش×متریال با تعارض"] == 1
    fu = model.followup
    # R10: بحرانی‌ها اول؛ ردیف با وضعیت نامشخص (M5) بعد از آن‌ها با «ممکن است بحرانی» می‌آید
    assert fu["متریال"].tolist() == ["M2", "M5"] and fu.iloc[0]["کارشناس مسئول"] == "کارشناس O1"
    assert fu.iloc[1]["کد سطح"] == "UNKNOWN" and fu.iloc[1]["اقدام پیشنهادی"] == R.MAYBE_CRIT
    assert fu.iloc[0]["اولویت"] == 1
    stock = model.checks[model.checks["بررسی"].eq(R.CHK_STOCK_CRIT)]
    assert stock["متریال"].tolist() == ["M2"]


def test_followup_sorted_by_resistance():
    df = _frame()
    df.loc[df["KEY_MATERIAL"].eq("M3"), ["کد طبقه بحرانی", "مقاومت (روز)"]] = ["STOCKOUT", 0.0]
    fu = R.build_model(df, None, "1405/06/10").followup
    assert fu["متریال"].tolist() == ["M3", "M2", "M5"]      # R10: نامشخص (M5) بعد از بحرانی‌ها
    assert fu["اولویت"].tolist() == [1, 2, 3]


@pytest.mark.parametrize("lang,dir_", [("fa", 'dir="rtl"'), ("en", 'dir="ltr"')])
def test_html_builds(model, lang, dir_):
    page = R.build_html(model, lang=lang, embed_fonts=False)
    assert page.startswith("<!DOCTYPE html>") and dir_ in page
    assert "<script src" not in page and "<link" not in page and "https://" not in page   # بدون CDN
    assert "M5" in page and "BL3" in page
    if lang == "fa":
        assert R.BL_MATS in page and R.MAT_BLS in page
    else:
        assert "Materials of this B/L&#x27;s orders" in page or "Materials of this B/L's orders" in page


def test_excel_sheets(model):
    wb = load_workbook(io.BytesIO(R.build_excel(model)))
    assert wb.sheetnames == list(R.SHEETS)
    assert wb["متریال‌ها"].sheet_view.rightToLeft
    heads = [c.value for c in wb["متریال‌ها"][4]]
    assert "در راه" in heads and "کد سطح" not in heads
    en = load_workbook(io.BytesIO(R.build_excel(model, "en")))
    assert len(en.sheetnames) == len(R.SHEETS) and not en[en.sheetnames[1]].sheet_view.rightToLeft


@pytest.mark.parametrize("df", [None, pd.DataFrame()])
def test_empty_gives_explicit_no_data(df):
    m = R.build_model(df, None, "1405/06/10")
    assert m.empty and m.kpis["متریال در گزارش"] == 0
    assert R.NO_DATA in R.build_html(m, embed_fonts=False)
    assert R.NO_DATA in R.kpis_section(m)
    wb = load_workbook(io.BytesIO(R.build_excel(m)))
    assert wb.sheetnames == list(R.SHEETS)
    assert R.NO_DATA in [c.value for row in wb["خلاصه"].iter_rows() for c in row]


def test_quantities_from_expert_ledger_when_mart_has_none():
    df = _frame().drop(columns=[c for c in _frame().columns if c.startswith("MOGH_QTY_")])
    ledger = pd.DataFrame([{"KEY_ORDER": "O1", "KEY_MATERIAL": "M1", "MOGH_QTY_IN_TRANSIT": 5.0},
                           {"KEY_ORDER": "O9", "KEY_MATERIAL": "M1", "MOGH_QTY_IN_TRANSIT": 50.0}])  # خارج از داده
    m = R.build_model(df, {"expert_material_positions": ledger}, "1405/06/10")
    assert m.quantity_source == "expert_material_positions"
    mats = m.materials.set_index("متریال")
    assert mats.loc["M1", "در راه"] == 5.0
    assert pd.isna(mats.loc["M3", "در راه"])       # بی‌داده، نه صفر
