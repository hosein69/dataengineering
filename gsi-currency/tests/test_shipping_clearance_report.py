# -*- coding: utf-8 -*-
"""گزارش جداگانه حمل و ترخیص: شمار یکتای بارنامه، بحرانی از سفارش، مدت فقط با دو تاریخ،
HTML دوزبانه، Excel و پیام صریح برای قاب خالی. فقط داده ساختگی کوچک."""
import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from gsi.report import shipping_clearance_report as R

REF = "1405/06/10"


def _frame() -> pd.DataFrame:
    rows = []
    # بارنامه A: دو سفارش × دو متریال = چهار ردیف؛ یک متریال سفارش O2 بحرانی است
    for order, mats in (("O1", ("M1", "M2")), ("O2", ("M3", "M4"))):
        for m in mats:
            rows.append({"CANONICAL_BL": "BLA", "KEY_BL": "BLA", "CANONICAL_ORDER": order, "KEY_ORDER": order,
                         "KEY_MATERIAL": m, "KEY_REG": "R1", "TRANSPORT_MODE": "SEA",
                         "ARRIVAL_DATE": "1405/05/28", "DISCHARGE_DATE": "1405/06/01", "DO_DATE": "1405/06/03",
                         "CL_COTAGE_DATE": "1405/06/05", "کد طبقه بحرانی": "CRITICAL" if m == "M3" else "SAFE",
                         "CANONICAL_EXPERT": "کارشناس الف" if order == "O1" else "کارشناس ب",
                         "ORG_DEPT": "مدیریت خرید خارجی", "IS_FULL_CLEARED": False, "IS_PARTIAL_CLEARED": False})
    # بارنامه B: یک سفارش ایمن، ترخیص کامل؛ تاریخ کوتاژ ندارد
    rows.append({"CANONICAL_BL": "BLB", "CANONICAL_ORDER": "O3", "KEY_MATERIAL": "M5", "TRANSPORT_MODE": "AIR",
                 "DISCHARGE_DATE": "1405/05/01", "FULL_CLEAR_DATE": "1405/05/20", "کد طبقه بحرانی": "SAFE",
                 "IS_FULL_CLEARED": True})
    # بارنامه C: بدون شاهد ورود/تخلیه ← در راه
    rows.append({"CANONICAL_BL": "BLC", "CANONICAL_ORDER": "O4", "KEY_MATERIAL": "M6", "TRANSPORT_MODE": "SEA",
                 "کد طبقه بحرانی": "WATCH"})
    # ردیف متریال افزوده سفارش O4 بدون بارنامه، بحرانی ← بارنامه C هم بحرانی است
    rows.append({"CANONICAL_BL": "", "CANONICAL_ORDER": "O4", "KEY_MATERIAL": "M7", "کد طبقه بحرانی": "STOCKOUT"})
    # سفارشی که هیچ بارنامه‌ای ندارد
    rows.append({"CANONICAL_BL": "", "CANONICAL_ORDER": "O5", "KEY_MATERIAL": "M8", "کد طبقه بحرانی": "SAFE"})
    return pd.DataFrame(rows)


class ShippingClearanceReportTests(unittest.TestCase):
    def setUp(self):
        self.m = R.build_model(_frame(), None, REF)

    def _row(self, bl):
        d = self.m.dossiers
        return d.loc[d["بارنامه"].eq(bl)].iloc[0]

    def test_unique_bl_counts_do_not_double(self):
        self.assertEqual(self.m.total, 3)
        self.assertEqual(len(self.m.dossiers), 3)
        self.assertEqual(self.m.counts[R.ST_ARRIVED], 1)          # BLA با چهار ردیف یک بار
        self.assertEqual(self.m.counts[R.ST_FULL], 1)
        self.assertEqual(self.m.counts[R.ST_TRANSIT], 1)
        self.assertEqual(sum(self.m.counts[s] for s in R.STATUSES), self.m.total)
        self.assertEqual(int(self.m.ages["بارنامه"].sum()), self.m.counts["open"])
        self.assertEqual(self.m.counts["orders_without_bl"], 1)
        cov = self.m.coverage.set_index("شاهد")
        self.assertEqual(int(cov.loc["تخلیه", "بارنامه"]), 3)

    def test_bl_is_critical_when_any_order_is_critical(self):
        a = self._row("BLA")
        self.assertEqual(a["بحرانی"], R.YES)
        self.assertEqual(a["بدترین سطح بحرانی سفارش‌ها"], "بحرانی")
        self.assertEqual(a[R.CRIT_MATS_OF_BL_ORDERS], "M3")
        self.assertEqual(self._row("BLB")["بحرانی"], R.NO)
        # متریال بحرانی روی ردیف بی‌بارنامه همان سفارش هم حساب است
        self.assertEqual(self._row("BLC")["بحرانی"], R.YES)
        self.assertEqual(self.m.counts["critical_open"], 2)

    def test_durations_only_when_both_dates_exist(self):
        d = self.m.durations.set_index(["مرحله", "روش حمل"])
        allm = d.loc[("تخلیه تا کوتاژ", "همه روش‌ها")]
        self.assertEqual(allm["دارای هر دو تاریخ"], 1)             # فقط BLA
        self.assertEqual(allm["میانه (روز)"], 4)
        self.assertEqual(allm["بدون تاریخ پایان"], 2)
        full = d.loc[("تخلیه تا ترخیص کامل", "همه روش‌ها")]
        self.assertEqual((full["دارای هر دو تاریخ"], full["بیشینه (روز)"]), (1, 19))
        arr = d.loc[("ورود تا دریافت ترخیصیه", "همه روش‌ها")]
        self.assertEqual((arr["دارای هر دو تاریخ"], arr["بدون تاریخ شروع"]), (1, 2))
        # روش حملی که هیچ جفت تاریخ ندارد عدد نمی‌گیرد، صفر هم نمی‌گیرد
        self.assertTrue(pd.isna(d.loc[("تخلیه تا کوتاژ", R._mode_label("AIR")), "میانه (روز)"]))

    def test_days_waiting_and_follow_up(self):
        self.assertEqual(int(self._row("BLA")[R.WAIT]), 9)
        self.assertTrue(pd.isna(self._row("BLB")[R.WAIT]))           # ترخیص کامل
        fu = self.m.follow_up
        self.assertEqual(fu["بارنامه"].tolist(), ["BLA", "BLC"])      # بی‌تاریخ تخلیه آخر
        self.assertIn("کارشناس الف", fu.iloc[0]["کارشناس سفارش‌ها"])
        self.assertIn("کارشناس ب", fu.iloc[0]["کارشناس سفارش‌ها"])

    def test_filter_frame_by_text_and_date(self):
        d = self.m.dossiers
        self.assertEqual(R.filter_frame(d, text="bla")["بارنامه"].tolist(), ["BLA"])
        got = R.filter_frame(d, start="۱۴۰۵/۰۵/۱۵", include_undated=False)["بارنامه"].tolist()
        self.assertEqual(got, ["BLA"])
        self.assertEqual(len(R.filter_frame(d, start="1405/05/15")), 1)   # R10: بی‌تاریخ به‌طور پیش‌فرض بیرون
        self.assertEqual(len(R.filter_frame(d, start="1405/05/15", include_undated=True)), 2)

    def test_html_fa_and_en(self):
        fa = R.build_html(self.m, lang="fa", embed_fonts=False)
        self.assertIn('dir="rtl"', fa)
        self.assertIn("BLA", fa)
        self.assertIn(R.CRIT_MATS_OF_BL_ORDERS, fa)
        self.assertNotIn("<script src", fa)
        self.assertNotIn("<link", fa)
        en = R.build_html(self.m, lang="en", embed_fonts=False, embed_excel=False)
        self.assertIn('dir="ltr"', en)
        self.assertIn("Arrived, not cleared", en)
        for frag in (R.kpi_section, R.stages_section, R.dossier_section, R.followup_section, R.gaps_section):
            self.assertTrue(frag(self.m, "fa").strip())

    def test_excel_has_sheets(self):
        wb = load_workbook(io.BytesIO(R.build_excel(self.m, "fa")))
        self.assertEqual(wb.sheetnames, [fa for fa, _ in R.SHEETS])
        self.assertTrue(wb["پرونده بارنامه"].sheet_view.rightToLeft)
        wb_en = load_workbook(io.BytesIO(R.build_excel(self.m, "en")))
        self.assertEqual(wb_en.sheetnames, [en for _, en in R.SHEETS])
        self.assertFalse(wb_en["BL dossiers"].sheet_view.rightToLeft)

    def test_empty_frame_says_no_data(self):
        for frame in (pd.DataFrame(), None, pd.DataFrame([{"CANONICAL_BL": "", "CANONICAL_ORDER": "O1"}])):
            m = R.build_model(frame, None, REF)
            self.assertTrue(m.empty)
            self.assertIn(R.NOTE_EMPTY, m.notes)
            page = R.build_html(m, lang="fa", embed_fonts=False)
            self.assertIn(R.NOTE_EMPTY, page)
            self.assertNotIn('class="mi-kpi-val"', page)                    # هیچ کارت صفری
            self.assertIn(R.NOTE_EMPTY, R.kpi_section(m, "fa"))
            load_workbook(io.BytesIO(R.build_excel(m, "fa")))


if __name__ == "__main__":
    unittest.main()
