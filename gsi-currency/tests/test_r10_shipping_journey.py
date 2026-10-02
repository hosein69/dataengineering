# -*- coding: utf-8 -*-
"""R10 (مالک، ۱۴۰۵/۰۷/۰۸): مسیر شواهد بارنامه در گزارش حمل. قاعده «اگر شک داری انجام نده»: عبور ضمنی
فقط قطعی، تاریخ‌های متفاوت متعارض، ترخیص فقط با پرچم، رسید SAP فقط با نسبت دادن قطعی. داده ساختگی."""
import io

import pandas as pd
from openpyxl import load_workbook

from gsi.report import shipping_clearance_report as R
from gsi.report import shipping_journey as J

REF = "1405/06/10"


#: ستون مشتق مارت (s20) ← ستون خام منبع، مثل مارت منتشرشده
_DERIVED = {"ARRIVAL_DATE": "CL_ARRIVAL_DATE", "DISCHARGE_DATE": "BL_DISCHARGE_DATE", "DO_DATE": "BL_DO_DATE",
            "FULL_CLEAR_DATE": "CL_CLEAR_DATE"}


def _r(bl, order, mat, mode="SEA", **kw):
    row = {"CANONICAL_BL": bl, "KEY_BL": bl, "CANONICAL_ORDER": order, "KEY_ORDER": order, "KEY_MATERIAL": mat,
           "TRANSPORT_MODE": mode, "کد طبقه بحرانی": "SAFE", "MOGH_BL_COUNT": 1, **kw}
    row.update({d: kw[s] for d, s in _DERIVED.items() if s in kw})
    return row


def _frame():
    return pd.DataFrame([
        _r("BL1", "O1", "M1", BL_BL_DATE="1405/04/01", CL_ARRIVAL_DATE="1405/05/01", BL_DISCHARGE_DATE="1405/05/03",
           BL_DO_DATE="1405/05/06", CL_COTAGE_DATE="1405/05/10", COT_COTAGE_DATE="1405/05/12"),
        _r("BL2", "O2", "M2", IS_FULL_CLEARED=True, CL_CLEAR_DATE="1405/05/20",
           GR_DLV_DELIVERED_QTY=10.0, GR_DLV_LAST_DATE="2026-08-20"),
        _r("BL3", "O3", "M3", mode="AIR", CL_ARRIVAL_DATE="1405/05/20", CL_CLEAR_DATE="1405/05/25"),
        _r("BL4", "O4", "M4", BL_DISCHARGE_DATE="1405/03/01", CL_ARRIVAL_DATE="1405/03/05"),
        _r("BL5", "O5", "M5", IS_FULL_CLEARED=True, MOGH_BL_COUNT=2, GR_DLV_DELIVERED_QTY=4.0,
           GR_DLV_LAST_DATE="2026-08-25"),
    ])


def _extras():
    return {"expert_shipments": pd.DataFrame([
        {"KEY_ORDER": "O3", "KEY_MATERIAL": "M3", "PART_STATE": "CLEARED", "BL": "BL3"},
        {"KEY_ORDER": "O4", "KEY_MATERIAL": "M4", "PART_STATE": "IN_TRANSIT", "BL": "BL4"},
    ])}


def _rows():
    m = R.build_model(_frame(), _extras(), REF)
    return m, m.journey.rows.set_index("بارنامه")


def test_states_follow_only_certain_evidence():
    m, rows = _rows()
    c1 = rows.loc["BL1", "_CELLS"]
    assert c1["COTAGE"]["state"] == J.CONFLICT_S                        # دو منبع، دو تاریخ؛ هیچ‌کدام برگزیده نمی‌شود
    assert {s for s, _ in c1["COTAGE"]["src"]} == {"cl", "cot"}
    assert rows.loc["BL1", "_CUR"] == "COTAGE" and rows.loc["BL1", "_NEXT"] == "CLEAR"
    c2 = rows.loc["BL2", "_CELLS"]
    assert [c2[c]["state"] for c in ("SHIP", "ARRIVE", "DISCHARGE", "COTAGE")] == [J.IMPLIED_S] * 4
    assert c2["DO"]["state"] == J.NONE_S                                # ترخیصیه از ترخیص استنتاج نمی‌شود
    assert c2["RECEIPT"]["state"] == J.DATED_S and rows.loc["BL2", "_CUR"] == "RECEIPT"
    c3 = rows.loc["BL3", "_CELLS"]
    assert c3["CLEAR"]["state"] == J.CONFLICT_S                          # تاریخ و ادعای کارشناس بدون پرچم
    assert rows.loc["BL3", "_CUR"] == "ARRIVE" and rows.loc["BL3", "_NEXT"] == "COTAGE"   # هوایی: تخلیه/ترخیصیه رد
    assert rows.loc["BL4", "_REV"] == {"ARRIVE", "DISCHARGE"}            # تخلیه پیش از ورود
    assert rows.loc["BL5", "_CELLS"]["RECEIPT"]["state"] == J.UNATTR_S   # سفارش دو بارنامه دارد
    k = m.journey.kpis
    assert k["bl"] == 5 and k["reversal"] == 1 and k["expert_ahead"] == 1 and k["expert_behind"] == 1
    assert k["conflict"] == 2


def test_process_discovery_and_shapley():
    m, _ = _rows()
    flow = m.journey.flow
    e = flow[(flow["_A"] == "SHIP") & (flow["_B"] == "ARRIVE")].iloc[0]
    assert e["شمار"] == 1 and e["میانه روز"] == 31
    assert ((flow["_A"] == "DISCHARGE") & (flow["_B"] == "ARRIVE")).any()   # وارونگی دیده می‌شود
    src = m.journey.sources.set_index("_K")
    assert abs(src["سهم شپلی"].sum() - 100) < 0.5
    assert src.loc["cot", "خانه با تنها یک منبع"] == 0 and src.loc["bl", "خانه با تنها یک منبع"] > 0


def test_journey_renders_without_svg_and_exports():
    m, _ = _rows()
    frag = R.journey_section(m, "fa")
    assert "<svg" not in frag and 'data-key="journey"' in frag and 'data-cur="RECEIPT"' in frag
    page = R.build_html(m, lang="en", embed_fonts=False, embed_excel=False)
    assert "Evidence route" in page and "Conflicting evidence" in page
    wb = load_workbook(io.BytesIO(R.build_excel(m, "fa")))
    ws = wb["مسیر شواهد"]
    assert ws.max_row >= 6
    f = R.filtered_model(m, _frame(), _extras(), REF, start="1405/05/01")
    assert set(f.journey.rows["بارنامه"]) == {"BL1", "BL3"}
