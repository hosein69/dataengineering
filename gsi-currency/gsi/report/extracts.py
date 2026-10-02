# -*- coding: utf-8 -*-
"""فایل اکسل مجزا برای هر کارشناس + گزارش ممیزی تعارضات (§۱۰).

نسخه ۲۰.۱ هیچ‌کدام را نداشت (پوشه expert_extracts اصلاً ساخته نمی‌شد).
گروه‌بندی بر اساس **کد پرسنلی** انجام می‌شود، نه نام — چون نام‌ها تکراری و
پرنویز هستند.
"""
from __future__ import annotations

import os
import re
from typing import List

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..core.excel_text import keep_text
from ..dataio.logging_setup import log
from .palette import LuxuryPalette as P, NUM_FORMAT_CURRENCY

_SAFE = re.compile(r"[^\w\u0600-\u06FF\-]+")


def _safe_name(s: str) -> str:
    return _SAFE.sub("_", str(s)).strip("_")[:60] or "unknown"


#: R8: کلیدهایی که کنار ستون‌های ماتریس رسمی همیشه در پرونده کارشناس می‌آیند
EXTRACT_KEY_COLUMNS = ("KEY_EMP", "CANONICAL_EXPERT", "KEY_REG", "CANONICAL_ORDER", "KEY_ORDER",
                       "MOGH_ITEM_ROLE", "KEY_MATERIAL", "CANONICAL_BL", "KEY_BL")


def _extract_all_columns() -> bool:
    """R8: ``GSI_EXTRACT_ALL_COLUMNS=1`` همان خروجی کامل قدیمی (همه ستون‌ها) را برمی‌گرداند."""
    return str(os.environ.get("GSI_EXTRACT_ALL_COLUMNS", "")).strip().lower() in ("1", "true", "yes", "on")


def _matrix_specs() -> list:
    """R8: همان ColumnSpecهای شیت ماتریس رسمی (``collect_columns`` همه مرحله‌ها)."""
    try:
        from ..stages import collect_columns, discover
        return collect_columns(discover())
    except Exception as ex:           # گزارش کارشناس نباید به‌خاطر کشف مرحله‌ها بیفتد
        log.warning(f"⚠️ ستون‌های ماتریس رسمی برای فایل کارشناس خوانده نشد: {ex}")
        return []


def extract_columns(df: pd.DataFrame, full: bool | None = None):
    """R8: (ستون‌ها، {ستون: سرستون فارسی یکتا}) برای فایل کارشناس.

    پیش‌فرض: همان ستون‌ها و عنوان‌های فارسی شیت ماتریس رسمی به‌اضافه کلیدها.
    با ``GSI_EXTRACT_ALL_COLUMNS=1``: همه ستون‌ها (جز RAW_) مثل قبل، ولی با
    برچسب فارسی یکتای کاتالوگ هر جا که هست. هیچ ستونی از df حذف نمی‌شود؛
    گزارش‌ساز همچنان همه ستون‌ها را برای انتخاب دارد."""
    from ..studio_core.field_catalog import build_catalog, dedupe_label_map, unique_labels
    full = _extract_all_columns() if full is None else full
    specs = _matrix_specs()
    titles = {sp.key: sp.title for sp in specs if sp.key in df.columns}
    if full:
        cols = [c for c in df.columns if not str(c).startswith(("RAW_", "MOGH_RAW_"))]
    else:
        keys = [c for c in EXTRACT_KEY_COLUMNS if c in df.columns]
        cols = list(dict.fromkeys(keys + list(titles)))
        if not cols:        # بی‌مرحله/بی‌کلید: به‌جای فایل خالی، همان خروجی کامل
            cols = [c for c in df.columns if not str(c).startswith(("RAW_", "MOGH_RAW_"))]
    # R10: ستون موازی یا بی‌داده در پرونده کارشناس نمی‌آید (کلیدها همیشه می‌مانند)؛
    # GSI_EXTRACT_ALL_COLUMNS=1 همچنان همه ستون‌ها را می‌دهد.
    from ..studio_core.column_tidy import tidy_columns, tidy_enabled
    if tidy_enabled() and not full:
        cols, _ = tidy_columns(df, cols, protect=EXTRACT_KEY_COLUMNS)
    try:
        cat = unique_labels(build_catalog(df[cols]))
    except Exception:
        cat = {}
    labels = {c: titles.get(c) or cat.get(c) or str(c) for c in cols}
    return cols, dedupe_label_map(labels, cols)


def write_expert_extracts(df: pd.DataFrame, out_dir: str,
                          columns: List[str] | None = None) -> List[str]:
    if df.empty:
        log.warning("⚠️ داده‌ای برای استخراج کارشناسان وجود ندارد.")
        return []
    os.makedirs(out_dir, exist_ok=True)
    # R8: پیش‌فرض = ستون‌ها و سرستون‌های فارسی ماتریس رسمی + کلیدها (نه ۶۰۰+ ستون فنی)
    if columns:
        from ..studio_core.field_catalog import build_catalog, dedupe_label_map, unique_labels
        cols = [c for c in columns if c in df.columns]
        heads = dedupe_label_map(unique_labels(build_catalog(df[cols])), cols)
    else:
        cols, heads = extract_columns(df)
    paths: List[str] = []

    key = "KEY_EMP" if "KEY_EMP" in df.columns else "CANONICAL_EXPERT"
    for emp, g in df.groupby(key, dropna=False):
        expert = str(g["CANONICAL_EXPERT"].iloc[0]) if "CANONICAL_EXPERT" in g else ""
        fname = f"{_safe_name(emp) or 'no_code'}__{_safe_name(expert)}.xlsx"
        path = os.path.join(out_dir, fname)

        wb = Workbook()
        ws = wb.active
        ws.title = "پرونده کارشناس"
        ws.sheet_view.rightToLeft = True

        ws.cell(row=1, column=1, value=f"پرونده کارشناس: {expert} (کد {emp})").font = P.font_title(13)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(cols), 8))

        for i, h in enumerate(cols, start=1):
            c = ws.cell(row=3, column=i, value=heads.get(h, h))
            c.font = P.font_header(1)
            c.fill = P.fill_header()
            c.alignment = P.align("center", wrap=True)
            ws.column_dimensions[get_column_letter(i)].width = 24
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}3"

        for r, (_, row) in enumerate(g.iterrows(), start=4):
            crit = "بحرانی" in str(row.get("وضعیت هوشمند", ""))
            for i, col in enumerate(cols, start=1):
                v = row.get(col, "")
                cell = ws.cell(row=r, column=i, value=v if isinstance(v, (int, float, str)) else str(v))
                cell.font = P.font_body()
                cell.border = P.thin_border()
                cell.fill = P.fill(P.CRITICAL_FILL if crit else P.GREEN_L4)
                if col in ("مانده تعهد", "جریمه برآوردی"):
                    cell.number_format = NUM_FORMAT_CURRENCY
        keep_text(wb)                 # متن منبع «=…» فرمول نشود
        wb.save(path)
        paths.append(path)

    log.info(f"📁 {len(paths)} فایل اختصاصی کارشناس در {out_dir} ساخته شد.")
    return paths


def write_audit_report(audit_df: pd.DataFrame, path: str) -> str:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "ممیزی تعارضات"
    ws.sheet_view.rightToLeft = True

    if audit_df.empty:
        ws["A1"] = "هیچ تعارضی میان سورس‌ها شناسایی نشد."
        ws["A1"].font = P.font_title(13)
        keep_text(wb)                 # متن منبع «=…» فرمول نشود
        wb.save(path)
        return path

    headers = list(audit_df.columns)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = P.font_header(1)
        c.fill = P.fill_header()
        c.alignment = P.align("center", wrap=True)
        ws.column_dimensions[get_column_letter(i)].width = 30 if i > 2 else 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for r, (_, row) in enumerate(audit_df.iterrows(), start=2):
        critical = str(row.get("نوع اهمیت")) == "CRITICAL"
        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=i, value=str(row[h]))
            cell.font = P.font_body()
            cell.border = P.thin_border()
            cell.alignment = P.align("right", wrap=True)
            cell.fill = P.fill(P.CRITICAL_FILL if critical else P.GREEN_L4)
    keep_text(wb)                 # متن منبع «=…» فرمول نشود
    wb.save(path)
    log.info(f"📋 گزارش ممیزی تعارضات ذخیره شد: {path} ({len(audit_df)} مورد)")
    return path
