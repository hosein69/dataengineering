# -*- coding: utf-8 -*-
"""R10: ستون‌های موازی و بی‌داده در خروجی‌های «همه ستون‌ها».

مالک (۱۴۰۵-۰۷-۰۸): «در خروجی ستون‌های با نام مشابه زیاد تولید شده … نام‌های صحیح که داده دارند رو نگه
دار و نام‌های موازی یا بدون ریشه داده را حذف کن».

چرا این ستون‌ها هست: جدول تخت عمداً هر ستون خام سورس (با پیشوند سورس، مثل SATA_INVOICE_VALUE) را کنار
ستون معیاری که از آن ساخته شده (INVOICE_VALUE) نگه می‌دارد تا ریشه هر عدد در انبار و محاسبات معلوم بماند؛
چند موتور قدیمی هم همان عدد را با نام فارسی (مثل «موجودی ایران خودرو» کنار STOCK_IKCO) نوشته‌اند. داخل
خط لوله این‌ها لازم‌اند؛ در خروجی که همه ستون‌ها را می‌ریزد، تکرارند.

این ماژول فقط **انتخاب ستون خروجی** است؛ هیچ ستونی از داده یا انبار حذف نمی‌شود و ستونی که کاربر خودش در
Composer انتخاب کرده دست نمی‌خورد.

* **موازی:** ستونی که ریشه‌اش (نگاشت s10/s20 یا نام قدیمی ثبت‌شده در ALIASES) ستون معیار
  دیگری است و در هیچ ردیفی مقداری ندارد که ستون معیار همان ردیف نداشته باشد. اگر مقدار تازه‌ای دارد
  می‌ماند (دیگر موازی نیست).
* **بی‌داده:** ستونی که در این خروجی هیچ مقداری ندارد (خالی، نه صفر یا «خیر»).
"""
from __future__ import annotations

import re
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd

#: کلیدها همیشه می‌مانند.
PROTECTED = ("KEY_ORDER", "KEY_BL", "KEY_REG", "KEY_MATERIAL", "KEY_EMP", "KEY_PR", "KEY_REG_FILE",
             "MOGH_ITEM_ROLE")

#: نام‌های قدیمی (فارسی یا فنی) که همان مقدار یک ستون معیار را دارند (ستون معیار ← نام موازی).
#: هم‌نامی تنها ریشه نیست: CRD_CURRENCY (ارز اعتبار) و CURRENCY (ارز تعهد) دو معنا دارند.
ALIASES = {
    "STOCK_IKCO": ("موجودی ایران خودرو",),
    "STOCK_SAPCO": ("موجودی ساپکو",),
    "DAILY_NEED": ("نیاز روزانه",),
    "BALANCE": ("مانده تعهد",),
    "SUPPLY_TOTAL_CONFIRMED": ("موجودی کل قابل احتساب",),
    "SUPPLY_TOTAL_LOWER_BOUND": ("حداقل موجودی قابل اثبات",),
    "SUPPLY_POSITION_COVERAGE_PCT": ("پوشش اجزای موجودی (٪)",),
    "KEY_ORDER": ("CASE_KEY",),
    "KEY_REG": ("CANONICAL_REG",),
    "KEY_MATERIAL": ("CANONICAL_PART_NO",),
}

REASON_PARALLEL = "موازی"
REASON_EMPTY = "بی‌داده"


def _lineage() -> List[Tuple[str, str, str]]:
    """(ستون معیار، ستون موازی، ریشه) از نگاشت‌های خود خط لوله."""
    out: List[Tuple[str, str, str]] = []
    try:
        from ..stages.s20_derive import DERIVED
        for target, spec in DERIVED.items():
            for cand in (spec[0] if spec else []) or []:
                if cand != target:
                    out.append((target, cand, "s20_derive"))
    except Exception:
        pass
    try:
        from ..stages.s10_resolve import ResolveStage
        for target, (_l, _k, cands) in ResolveStage.MAP.items():
            for col, _src in cands:
                if col != target:
                    out.append((target, col, "s10_resolve"))
    except Exception:
        pass
    for canon, names in ALIASES.items():
        for n in names:
            out.append((canon, n, "نام قدیمی"))
    return out


def _norm(s: pd.Series) -> pd.Series:
    """مقدار قابل مقایسه؛ خالی = NA. عدد با عدد، متن با متن فشرده."""
    num = pd.to_numeric(s, errors="coerce")
    txt = s.astype(str).str.strip()
    empty = s.isna() | txt.isin(["", "nan", "NaN", "None", "NaT", "<NA>"])
    out = txt.str.replace(r"\s+", " ", regex=True)
    out = out.where(num.isna(), num.round(6).astype(str))
    return out.mask(empty)


def _has_data(s: pd.Series) -> bool:
    return bool(_norm(s).notna().any())


def _adds_nothing(par: pd.Series, canon: pd.Series) -> bool:
    """ستون par هیچ مقداری ندارد که canon در همان ردیف نداشته باشد."""
    p, c = _norm(par), _norm(canon)
    have = p.notna()
    return bool(have.any()) and bool((p[have] == c[have]).all())


def tidy_columns(df: pd.DataFrame, columns: Optional[Iterable[str]] = None,
                 protect: Iterable[str] = ()) -> Tuple[List[str], List[Dict[str, str]]]:
    """(ستون‌های ماندنی به همان ترتیب، [{ستون، علت، ستون معیار، ریشه}]).

    ``protect``: ستون‌هایی که به هر حال می‌مانند (کلیدها، انتخاب صریح کاربر)."""
    cols = [str(c) for c in (columns if columns is not None else df.columns) if str(c) in df.columns]
    keep = set(cols)
    guard = set(PROTECTED) | {str(c) for c in protect}
    dropped: Dict[str, Dict[str, str]] = {}
    series = {c: df[c] for c in cols if not isinstance(df[c], pd.DataFrame)}

    def drop(col, reason, canon="", root=""):
        if col in guard or col not in keep:
            return
        keep.discard(col)
        dropped[col] = {"ستون": col, "علت": reason, "ستون معیار": canon, "ریشه": root}

    for c in cols:
        if c in series and not _has_data(series[c]):
            drop(c, REASON_EMPTY, "", "هیچ ردیفی مقدار ندارد")

    pairs = list(_lineage())
    for canon, par, root in pairs:
        if canon not in keep or par not in keep or canon not in series or par not in series:
            continue
        if par in guard:
            # کلید می‌ماند؛ ستون معیاری که عیناً همان کلید است موازی کلید است.
            if _adds_nothing(series[canon], series[par]) and _adds_nothing(series[par], series[canon]):
                drop(canon, REASON_PARALLEL, par, root)
            continue
        if _adds_nothing(series[par], series[canon]):
            drop(par, REASON_PARALLEL, canon, root)
    # زنجیره (MOGH_MATERIAL ← CANONICAL_PART_NO ← KEY_MATERIAL): ستون معیار همیشه ستونی است که ماند.
    for d in dropped.values():
        seen = set()
        while d["ستون معیار"] in dropped and d["ستون معیار"] not in seen:
            seen.add(d["ستون معیار"])
            d["ستون معیار"] = dropped[d["ستون معیار"]]["ستون معیار"]
    return [c for c in cols if c in keep], list(dropped.values())


def dropped_frame(dropped: List[Dict[str, str]]) -> pd.DataFrame:
    cols = ["ستون", "علت", "ستون معیار", "ریشه"]
    return pd.DataFrame(dropped, columns=cols) if dropped else pd.DataFrame(columns=cols)


def tidy_enabled() -> bool:
    """``GSI_OUTPUT_ALL_COLUMNS=1`` همه ستون‌ها را مثل قبل در خروجی می‌گذارد."""
    import os
    return str(os.environ.get("GSI_OUTPUT_ALL_COLUMNS", "")).strip().lower() not in ("1", "true", "yes", "on")


DROPPED_SHEET = "ستون‌های کنار گذاشته"


def write_dropped_sheet(xlsx_path, dropped: List[Dict[str, str]]) -> None:
    """برگه‌ای که می‌گوید کدام ستون چرا در این خروجی نیامد؛ هیچ حذفی بی‌صدا نیست."""
    if not dropped:
        return
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(xlsx_path)
    if DROPPED_SHEET in wb.sheetnames:
        del wb[DROPPED_SHEET]
    ws = wb.create_sheet(DROPPED_SHEET)
    ws.sheet_view.rightToLeft = True
    ws.cell(1, 1, "این ستون‌ها در داده و انبار هستند و در Studio قابل انتخاب‌اند؛ فقط در این خروجی تکرار یا "
                  "خالی بودند. برای خروجی کامل: GSI_OUTPUT_ALL_COLUMNS=1").font = Font(bold=True)
    frame = dropped_frame(dropped)
    for j, h in enumerate(frame.columns, 1):
        ws.cell(3, j, h).font = Font(bold=True)
    for i, row in enumerate(frame.itertuples(index=False, name=None), 4):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, str(v))
    for col, w in zip("ABCD", (34, 10, 30, 26)):
        ws.column_dimensions[col].width = w
    wb.save(xlsx_path)
