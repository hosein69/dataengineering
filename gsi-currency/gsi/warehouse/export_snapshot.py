# -*- coding: utf-8 -*-
"""Explicit Excel export from the last published warehouse snapshot.

This module never reads operational source workbooks and never runs ETL. It is
safe to use while Streamlit is open because it only hydrates the published mart.
"""
from __future__ import annotations
from pathlib import Path
from datetime import date
import pandas as pd

from ..design import tokens as T


def export_published_snapshot(output_path=None, ref_date=None, max_rows=200000):
    from .service import last_report
    from .store import Warehouse
    from ..studio_core.excel_export import build_custom_excel
    exact = last_report(ref_date) if ref_date else None
    stored = exact or last_report(None)
    if stored is None:
        raise RuntimeError("No published GSI snapshot exists. Run refresh first.")
    df, main, extras, _official = stored
    data = main if main is not None and not main.empty else df
    published_ref = str(extras.get("published_reference_date") or ref_date or date.today().isoformat())
    if output_path is None:
        root = Warehouse(initialize=False).path.parent / "exports"
        output_path = root / f"{published_ref}_GSI_PUBLISHED_SNAPSHOT.xlsx"
    output_path = Path(output_path)
    process_keys = ("eventlog", "case_table", "bottlenecks", "variants",
                    "conformance_cases", "conformance_root_causes")
    process_tables = {}
    for key in process_keys:
        try:
            value = extras.get(key)
        except Exception:
            value = None
        if value is not None:
            process_tables[key] = value
    modules = ["criticality", "case_alerts", "resistance", "commitment",
               "org", "expert", "process", "table"]
    # R8: سرستون فارسی یکتا از کاتالوگ (قبلاً همه ستون‌ها با نام فنی خام نوشته می‌شد)
    from ..studio_core.field_catalog import build_catalog, dedupe_label_map, unique_labels
    fields = list(data.columns)
    # R10: ستون‌های موازی (همان مقدار ستون معیار) و بی‌داده در این خروجی نمی‌آیند؛ فهرستشان در برگه جداست.
    from ..studio_core.column_tidy import tidy_columns, tidy_enabled, write_dropped_sheet
    dropped = []
    if tidy_enabled():
        fields, dropped = tidy_columns(data, fields)
    try:
        labels = dedupe_label_map(unique_labels(build_catalog(data[fields])), fields)
    except Exception:
        labels = None
    built = build_custom_excel(
        data, output_path, modules, published_ref, max_rows=int(max_rows),
        selected_fields=fields, process_tables=process_tables,
        allow_official_overwrite=False, field_labels=labels)

    try:
        write_dropped_sheet(built, dropped)
    except Exception:
        pass
    # The published flat mart is ORDER/BL oriented.  Preserve the independently
    # published Order×Material expert ledger as a separate sheet so the explicit
    # snapshot export cannot silently lose sibling materials/descriptions.
    try:
        positions = extras.get("expert_material_positions")
    except Exception:
        positions = None
    if positions is not None:
        from ..report.supply_views import build_expert_material_evidence_view
        evidence = build_expert_material_evidence_view(positions)
        if not evidence.empty:
            from openpyxl import load_workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo
            wb = load_workbook(built)
            name = "Material Evidence"
            if name in wb.sheetnames:
                del wb[name]
            ws = wb.create_sheet(name)
            ws.sheet_view.rightToLeft = True
            for j, h in enumerate(evidence.columns, 1):
                c = ws.cell(1, j, h)
                c.font = Font(name="IRANSans Light", size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=T.BRAND_NAVY.lstrip("#").upper())
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, row in enumerate(evidence.itertuples(index=False, name=None), 2):
                for j, v in enumerate(row, 1):
                    ws.cell(i, j, None if pd.isna(v) else v).alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            end = len(evidence) + 1
            ref = f"A1:{get_column_letter(len(evidence.columns))}{end}"
            tab = Table(displayName="PublishedExpertMaterialEvidence", ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tab)
            for j in range(1, len(evidence.columns) + 1):
                sample = [str(ws.cell(r, j).value or "") for r in range(1, min(end, 40) + 1)]
                ws.column_dimensions[get_column_letter(j)].width = min(max(max(map(len, sample)) + 2, 12), 42)
            # فقط برگه تازه: فرمول‌های خود داشبورد (SUBTOTAL/COUNTIF) در برگه‌های دیگر فرمول می‌مانند
            from ..core.excel_text import keep_text
            keep_text(ws)
            wb.save(built)
    return built


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="Export Excel from published GSI warehouse snapshot; no ETL/source read")
    ap.add_argument("--output")
    ap.add_argument("--date", dest="ref_date")
    ap.add_argument("--max-rows", type=int, default=200000)
    ns = ap.parse_args(argv)
    out = export_published_snapshot(ns.output, ns.ref_date, ns.max_rows)
    print(out)
    return 0
