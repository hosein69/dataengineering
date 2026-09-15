# -*- coding: utf-8 -*-
"""Custom Excel export for AIBL Studio."""
from __future__ import annotations
import os
from datetime import date
from pathlib import Path
from typing import Iterable, List, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

FONT = os.environ.get("AIBL_FONT_NAME", "IRANSans")

#: نام پیش‌فرض خروجی سفارشی Studio — عمداً با گزارش رسمی روزانه فرق دارد.
#: تست رگرسیون همین ثابت را با نام گزارش رسمی می‌سنجد (نه با grep روی UI).
DEFAULT_CUSTOM_NAME = "AIBL Studio"


class OfficialReportOverwrite(Exception):
    """تلاش برای بازنویسی گزارش رسمی روزانه با یک خروجی فیلترشده.

    خروجی Studio همیشه *زیرمجموعه* داده است (فیلتر کاربر + چند شیت). گزارش
    رسمی خط لوله ۱۷ شیت کامل دارد و ایمیل مدیریتی به همان فایل لینک می‌دهد.
    اگر نام پیش‌فرض یکی شود، یک کلیک روی «ساخت Excel سفارشی» گزارش رسمی را
    بی‌صدا نابود می‌کند. این استثنا جلوی آن را می‌گیرد.
    """
COLORS = {
    "critical": "C0392B", "warning": "F39C12", "watch": "F1C40F",
    "good": "27AE60", "inactive": "95A5A6", "header": "406057",
    "header_fill": "406057", "white": "FFFFFF", "grid": "D8E5E1"
}


def _status_color(value: str) -> str:
    s = str(value).lower()
    if "توقف" in s or "بحرانی" in s or "critical" in s: return COLORS["critical"]
    if "در حال" in s or "warning" in s or "becoming" in s: return COLORS["warning"]
    if "نظر" in s or "watch" in s: return COLORS["watch"]
    if "ایمن" in s or "safe" in s: return COLORS["good"]
    return COLORS["inactive"] if "مصرف" in s or "unknown" in s else COLORS["good"]


def _style_sheet(ws):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    thin = Side(style="thin", color=COLORS["grid"])
    for cell in ws[1]:
        cell.font = Font(name=FONT, size=10, bold=True, color=COLORS["white"])
        cell.fill = PatternFill("solid", fgColor=COLORS["header_fill"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=9, color="111917")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max([len(str(c.value or "")) for c in col[:80]] + [8])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 32)
    ws.row_dimensions[1].height = 30


def _official_report_paths(ref_date: str) -> set:
    """مسیرهای گزارش رسمی روزانه که هرگز نباید با خروجی Studio بازنویسی شوند."""
    from ..config.settings import SETTINGS
    out = set()
    try:
        from datetime import date as _date
        day = _date.fromisoformat(str(ref_date)[:10])
    except Exception:
        day = None
    for p in ([SETTINGS.daily_report_path(day)] if day else []) + [SETTINGS.daily_report_path()]:
        out.add(os.path.normcase(os.path.abspath(p)))
    return out


def _add_email_charts_sheet(wb, df: pd.DataFrame, selected):
    """نمودارهای ایمیل را به‌صورت Chart بومی Excel هم می‌سازد.

    به‌جای تصویر PNG، برچسب‌ها داخل خود Excel نوشته می‌شوند؛ بنابراین RTL و
    فارسی را Excel با فونت سیستم کاربر رندر می‌کند و مشکل وارونه/جدا شدن حروف
    در نام‌گذاری نمودارهای ایمیل حذف می‌شود.
    """
    from .designs import EMAIL_CHARTS
    name = "نمودارها"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "نمودارهای ایمیل — نسخه قابل ویرایش در Excel"
    ws["A2"] = "عنوان‌ها و برچسب‌ها فارسی هستند؛ خود نمودارها Native Excel هستند، نه تصویر."
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"].font = Font(name=FONT, size=10, italic=True)
    row = 4

    def add_bar(title, labels, values, anchor):
        nonlocal row
        if not labels: return
        start=row
        ws.cell(row,1,"عنوان"); ws.cell(row,2,"مقدار")
        for c in (ws.cell(row,1),ws.cell(row,2)):
            c.font=Font(name=FONT,bold=True); c.alignment=Alignment(horizontal="center")
        for lab,val in zip(labels,values):
            row += 1; ws.cell(row,1,str(lab)); ws.cell(row,2,float(val) if pd.notna(val) else 0)
        chart=BarChart(); chart.type="bar"; chart.style=10; chart.title=title; chart.y_axis.title=""; chart.x_axis.title="مقدار"
        data=Reference(ws,min_col=2,min_row=start,max_row=row); cats=Reference(ws,min_col=1,min_row=start+1,max_row=row)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.height=7.2; chart.width=14; chart.legend=None
        chart.dataLabels=DataLabelList(); chart.dataLabels.showVal=True
        try:
            from ..report.charts import _apply_chart_font
            _apply_chart_font(chart)
        except Exception:
            pass
        ws.add_chart(chart,anchor); row += 3

    if "criticality" in selected and "کد طبقه بحرانی" in df.columns:
        order=["STOCKOUT","CRITICAL","BECOMING_CRITICAL","WATCH","SAFE","NO_CONSUMPTION","UNKNOWN"]
        labels={"STOCKOUT":"توقف خط","CRITICAL":"بحرانی","BECOMING_CRITICAL":"در آستانه","WATCH":"تحت نظر","SAFE":"ایمن","NO_CONSUMPTION":"بدون مصرف","UNKNOWN":"نامشخص"}
        present=[x for x in order if (df["کد طبقه بحرانی"].astype(str)==x).any()]
        add_bar(EMAIL_CHARTS["criticality"],[labels[x] for x in present],[(df["کد طبقه بحرانی"].astype(str)==x).sum() for x in present],"D4")

    if "low_resistance" in selected and {"KEY_MATERIAL","مقاومت (روز)"}.issubset(df.columns):
        x=df[["KEY_MATERIAL","مقاومت (روز)"]].copy(); x["مقاومت (روز)"]=pd.to_numeric(x["مقاومت (روز)"],errors="coerce"); x=x.dropna().drop_duplicates("KEY_MATERIAL").sort_values("مقاومت (روز)").head(10)
        add_bar(EMAIL_CHARTS["low_resistance"],x["KEY_MATERIAL"].astype(str).tolist(),x["مقاومت (روز)"].tolist(),"D20")

    if "risk_mix" in selected and "طبقه ریسک" in df.columns:
        vc=df["طبقه ریسک"].fillna("").astype(str).replace("","نامشخص").value_counts()
        add_bar(EMAIL_CHARTS["risk_mix"],vc.index.astype(str).tolist(),vc.values.tolist(),"D36")

    if "org_workload" in selected and "ORG_DEPT" in df.columns:
        vc=df["ORG_DEPT"].fillna("").astype(str).replace("","نامشخص").value_counts().head(12).sort_values()
        add_bar(EMAIL_CHARTS["org_workload"],vc.index.astype(str).tolist(),vc.values.tolist(),"D52")

    if "commitment" in selected and {"مانده تعهد","روزهای تأخیر"}.issubset(df.columns):
        from .grain import safe_agg
        bal=pd.to_numeric(df["مانده تعهد"],errors="coerce").fillna(0); overdue=pd.to_numeric(df["روزهای تأخیر"],errors="coerce").fillna(0)
        masks=[overdue>0,(overdue<=0)&(bal>0),bal<=0]; labs=["معوق","در مهلت","تسویه‌شده"]
        vals=[safe_agg(df.loc[m].copy(), "مانده تعهد", "sum") for m in masks]
        add_bar(EMAIL_CHARTS["commitment"],labs,vals,"D68")

    for col in ("A","B"):
        ws.column_dimensions[col].width = 34 if col=="A" else 18
    ws.freeze_panes="A4"


def build_custom_excel(df: pd.DataFrame, output_path: str | Path, modules: Iterable[str], ref_date: str, max_rows: int = 10000, selected_fields=None, allow_official_overwrite: bool = False, field_labels: Optional[dict] = None, process_tables: Optional[dict] = None, email_charts=None) -> str:
    path = Path(output_path)
    if not allow_official_overwrite and os.path.normcase(os.path.abspath(path)) in _official_report_paths(ref_date):
        raise OfficialReportOverwrite(
            f"نام انتخابی دقیقاً همان گزارش رسمی روزانه است:\n    {path}\n"
            f"این فایل ۱۷ شیت کامل خط لوله را دارد و ایمیل مدیریتی به آن لینک می‌دهد؛ "
            f"خروجی Studio فیلترشده است و جایگزین آن نمی‌شود.\n"
            f"یک نام دیگر بگذارید (مثلاً «AIBL Studio»).")
    path.parent.mkdir(parents=True, exist_ok=True)
    modules = list(modules)
    selected_fields = [c for c in list(selected_fields or []) if c in df.columns]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Always provide a compact executive sheet.
        k = []
        if "کد طبقه بحرانی" in df.columns:
            for code, label in [("STOCKOUT", "توقف خط"), ("CRITICAL", "بحرانی"), ("BECOMING_CRITICAL", "در آستانه"), ("WATCH", "تحت نظر"), ("SAFE", "ایمن")]:
                k.append((label, int((df["کد طبقه بحرانی"] == code).sum())))
        if "BL_CRITICAL" in df.columns and "CANONICAL_BL" in df.columns:
            k.append(("بارنامه‌های بحرانی", int(df.loc[df["BL_CRITICAL"].fillna(False).astype(bool), "CANONICAL_BL"].replace("", pd.NA).nunique())))
        if "ORDER_CRITICAL" in df.columns and "CANONICAL_ORDER" in df.columns:
            k.append(("سفارش‌های بحرانی", int(df.loc[df["ORDER_CRITICAL"].fillna(False).astype(bool), "CANONICAL_ORDER"].replace("", pd.NA).nunique())))
        pd.DataFrame(k, columns=["شاخص", "مقدار"]).to_excel(writer, sheet_name="خلاصه مدیریتی", index=False)
        if "criticality" in modules and "بحرانی (کوتاه)" in df.columns:
            cols = [c for c in ["KEY_MATERIAL", "بحرانی (کوتاه)", "مقاومت (روز)", "BL_CRITICAL", "BL_CRITICAL_MATERIALS", "BL_CRITICAL_REASON"] if c in df.columns]
            df[cols].head(max_rows).to_excel(writer, sheet_name="بحرانی بودن", index=False)
        if "case_alerts" in modules:
            cols = [c for c in ["CANONICAL_BL", "CANONICAL_ORDER", "BL_CRITICAL", "ORDER_CRITICAL", "BL_CRITICAL_MATERIALS", "ORDER_CRITICAL_MATERIALS", "BL_CRITICAL_REASON", "ORDER_CRITICAL_REASON"] if c in df.columns]
            df[cols].drop_duplicates().head(max_rows).to_excel(writer, sheet_name="علل پرونده", index=False)
        if "resistance" in modules and "مقاومت (روز)" in df.columns:
            cols = [c for c in ["KEY_MATERIAL", "مقاومت (روز)", "بحرانی (کوتاه)", "نیاز روزانه", "موجودی کل قابل احتساب"] if c in df.columns]
            x = df[cols].copy(); x["مقاومت (روز)"] = pd.to_numeric(x["مقاومت (روز)"], errors="coerce")
            x = x.sort_values("مقاومت (روز)").head(max_rows)
            x.to_excel(writer, sheet_name="کمترین مقاومت", index=False)
        if "commitment" in modules:
            cols = [c for c in ["CANONICAL_ORDER", "مانده تعهد", "روزهای تأخیر", "جریمه برآوردی", "طبقه ریسک"] if c in df.columns]
            df[cols].head(max_rows).to_excel(writer, sheet_name="تعهدات", index=False)
        if "org" in modules and "ORG_DEPT" in df.columns:
            x = df.groupby("ORG_DEPT", dropna=False).size().reset_index(name="تعداد پرونده").sort_values("تعداد پرونده", ascending=False)
            x.to_excel(writer, sheet_name="سازمان", index=False)
        if "expert" in modules and "CANONICAL_EXPERT" in df.columns:
            x = df.groupby("CANONICAL_EXPERT", dropna=False).agg(**{"تعداد پرونده": ("CANONICAL_EXPERT", "size")})
            if "بحرانی (کوتاه)" in df.columns:
                crit = df["بحرانی (کوتاه)"].isin(["بحرانی", "توقف خط"]).groupby(df["CANONICAL_EXPERT"]).sum()
                x["پرونده بحرانی"] = crit
            if "مقاومت (روز)" in df.columns:
                resistance = pd.to_numeric(df["مقاومت (روز)"], errors="coerce").groupby(df["CANONICAL_EXPERT"]).min()
                x["کمترین مقاومت"] = resistance
            x.reset_index().sort_values("تعداد پرونده", ascending=False).to_excel(writer, sheet_name="بار کارشناسان", index=False)
        if "process" in modules:
            # جدول‌های فرآیندکاوی که خط لوله می‌سازد — قبلاً فقط یک یادداشت بود.
            wrote_any = False
            for key, sheet in (("eventlog", "لاگ رویداد"), ("case_table", "پرونده‌ها"),
                               ("bottlenecks", "گلوگاه‌ها"), ("variants", "مسیرها"),
                               ("conformance_cases", "انطباق"),
                               ("conformance_root_causes", "علل ریشه‌ای")):
                t = (process_tables or {}).get(key)
                if t is not None and hasattr(t, "empty") and not t.empty:
                    t.head(max_rows).to_excel(writer, sheet_name=sheet, index=False)
                    wrote_any = True
            if not wrote_any:
                pd.DataFrame({"توضیح": [
                    "لاگ رویداد برای این اجرا ساخته نشده است.",
                    f"تاریخ مرجع: {ref_date}"]}).to_excel(
                        writer, sheet_name="فرآیند", index=False)
        if "table" in modules:
            default_cols = [
                "KEY_MATERIAL", "CANONICAL_ORDER", "CANONICAL_BL", "KEY_REG", "CANONICAL_EXPERT",
                "ORG_DEPT", "TRANSPORT_MODE", "بحرانی (کوتاه)", "مقاومت (روز)", "BL_CRITICAL",
                "BL_CRITICAL_MATERIALS", "BL_CRITICAL_REASON", "ORDER_CRITICAL",
                "ORDER_CRITICAL_MATERIALS", "ORDER_CRITICAL_REASON"
            ]
            cols = [c for c in (selected_fields or default_cols) if c in df.columns]
            if not cols: cols = list(df.columns[:20])
            live = df[cols].head(max_rows)
            if field_labels:
                # سرستون فارسی — نگاشت باید یکتا باشد وگرنه pandas ستون تکراری می‌سازد
                ren = {c: field_labels[c] for c in cols if c in field_labels}
                if len(set(ren.values())) == len(ren):
                    live = live.rename(columns=ren)
            live.to_excel(writer, sheet_name="داده زنده", index=False)
    wb = load_workbook(path)
    if email_charts:
        _add_email_charts_sheet(wb, df, list(email_charts))
    for ws in wb.worksheets:
        _style_sheet(ws)
        # Status-aware fills on common status columns.
        for col_idx, cell in enumerate(ws[1], start=1):
            if str(cell.value) in {"بحرانی (کوتاه)", "طبقه بحرانی"}:
                for row in range(2, ws.max_row + 1):
                    c = ws.cell(row=row, column=col_idx)
                    color = _status_color(c.value)
                    c.fill = PatternFill("solid", fgColor=color)
                    c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    wb.save(path)
    return str(path)
