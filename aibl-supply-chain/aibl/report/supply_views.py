# -*- coding: utf-8 -*-
"""سه نمای تصمیم‌ساز زنجیره تأمین: متریال، بارنامه و اداره.

اصل طراحی: **«الان کجاست؟ از کِی؟ دست کیست؟ معطل چه کسی است؟ اگر
نمی‌دانیم، دقیقاً چه داده‌ای کم است و از چه کسی باید خواست؟»**

هیچ علت یا کارشناسی از روی نبودِ داده حدس زده نمی‌شود.

## دو بازنویسی نسبت به نسخه ۲۶٫۹

**۱) موقعیت از رویداد می‌آید، نه از موجودی.** نسخه قبل موقعیت را از
ستون‌های موجودی حدس می‌زد (`موجودی در گمرک > 0 ⇒ گمرک`). آن حدس تاریخ
ندارد، پس به سؤال «از کِی؟» جواب نمی‌دهد. حالا ماژول ``resolve.part_status``
آخرین رویدادِ تاریخ‌دار را می‌یابد و «کجا/کِی/چه کسی» از همان می‌آید.
اگر آن ستون‌ها موجود نباشند (مثلاً فراخوانی مستقیم از Studio روی یک
زیرمجموعه)، همان‌جا محاسبه می‌شوند.

**۲) سرعت.** پیاده‌سازی قبلی برای هر ردیف یک ``pd.DataFrame([row])``
می‌ساخت و روی آن ``_first`` می‌زد. اندازه‌گیری واقعی روی همین ماشین:

    ۵۰۰ ردیف   → ۱۴٫۸ ثانیه
    ۲٬۰۰۰ ردیف → ۳۶٫۵ ثانیه

و این تابع سه بار در هر بازتولید Studio و یک بار در ``save()`` گزارش
رسمی صدا زده می‌شد. روی داده واقعی یعنی گزارش روزانه عملاً تمام نمی‌شد.
نسخه فعلی کاملاً برداری است (groupby + transform) و همان ۲٬۰۰۰ ردیف را
در کسری از ثانیه می‌سازد.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from ..resolve import part_status as ps
from ..resolve.expert_scope import OWNER_GAP, OWNER_NAME, SCOPES
from .palette import LuxuryPalette as P

RED = "C0392B"

_BLANKS = {"nan": "", "None": "", "NaT": "", "—": "", "-": ""}


def _s(df: pd.DataFrame, c: str) -> pd.Series:
    """ستون متنی تمیز؛ اگر نبود، ستون خالی هم‌طول."""
    if c not in df.columns:
        return pd.Series("", index=df.index, dtype=object)
    return df[c].fillna("").astype(str).str.strip().replace(_BLANKS)


def _num(df: pd.DataFrame, c: str) -> pd.Series:
    """ستون عددی؛ اگر نبود، ستون NaN هم‌طول.

    ``df.get(c)`` برای ستون غایب ``None`` می‌دهد و ``pd.to_numeric(None)``
    یک اسکالر برمی‌گرداند نه Series — و بعد ``.groupby`` روی آن می‌ترکد.
    """
    if c not in df.columns:
        return pd.Series(pd.NA, index=df.index, dtype="Float64")
    return pd.to_numeric(df[c], errors="coerce")


def _first_col(df: pd.DataFrame, cols, default: str = "نامشخص") -> pd.Series:
    """اولین ستون موجود که برای آن ردیف مقدار دارد — برداری."""
    out = pd.Series("", index=df.index, dtype=object)
    for c in cols:
        if c not in df.columns:
            continue
        out = out.where(out.ne(""), _s(df, c))
    return out.replace("", default)


def _group_mode(s: pd.Series, keys: pd.Series) -> pd.Series:
    """پرتکرارترین مقدار غیرخالی در هر گروه، پخش‌شده روی ردیف‌های گروه."""
    tmp = pd.DataFrame({"k": keys, "v": s})
    non = tmp[tmp["v"].ne("")]
    if non.empty:
        return pd.Series("", index=s.index, dtype=object)
    mode = non.groupby("k")["v"].agg(lambda x: x.value_counts().index[0])
    return keys.map(mode).fillna("")


def _ensure_status(df: pd.DataFrame) -> pd.DataFrame:
    """ستون‌های وضعیت را تضمین می‌کند؛ اگر نبودند، همین‌جا می‌سازد."""
    if ps.WHERE in df.columns:
        return df
    return ps.resolve(df.copy())


def _base(df: pd.DataFrame, group_col: str, kind: str) -> pd.DataFrame:
    if group_col not in df.columns:
        return pd.DataFrame()
    work = df[_s(df, group_col).ne("")].copy()
    if work.empty:
        return pd.DataFrame()
    work = _ensure_status(work)
    keys = _s(work, group_col)

    where = _s(work, ps.WHERE).replace("", "نامشخص")
    when = _s(work, ps.WHEN)
    age = _num(work, ps.AGE)

    title = (_first_col(work, ["MATERIAL_DESC", "CANONICAL_GOODS_DESC"], "")
             if kind == "متریال" else pd.Series("", index=work.index))
    title = title.where(title.ne(""), keys)

    out = pd.DataFrame({
        "کلید گروه": keys,
        "عنوان گروه": title,
        "متریال": _first_col(work, ["KEY_MATERIAL"]),
        "سفارش": _first_col(work, ["CANONICAL_ORDER"]),
        "ثبت سفارش": _first_col(work, ["KEY_REG"]),
        "بارنامه": _first_col(work, ["CANONICAL_BL"]),
        # ── کجا / کِی / چه کسی ──
        "موقعیت فعلی": where,
        "آخرین فعالیت": _s(work, ps.ACTIVITY).replace("", "—"),
        "تاریخ آخرین رویداد": when.replace("", "—"),
        "سن وضعیت (روز)": age,
        "کارشناس مسئول وضعیت": _s(work, ps.WHO).replace("", "—"),
        "حوزه مسئول وضعیت": _s(work, ps.WHO_SCOPE).replace("", "—"),
        "معطل حوزه": _s(work, ps.WAITING_SCOPE).replace("", "—"),
        "معطل کارشناس": _s(work, ps.WAITING_WHO).replace("", "—"),
        "فعالیت بعدی مورد انتظار": _s(work, ps.NEXT_ACT).replace("", "—"),
        # ── مالکیت ──
        "مالک قطعه (کارشناس خرید)": _s(work, OWNER_NAME).replace("", "شناسایی نشد"),
        "شکاف داده مالک": _s(work, OWNER_GAP),
        # ── چرایی و شواهد ──
        "چرایی / مبنای وضعیت": _reason(work),
        "شواهد تعیین وضعیت": _s(work, ps.BASIS).replace("", "شاهد تاریخ‌دار ثبت نشده است"),
        "داده مفقود برای تعیین وضعیت": _s(work, ps.MISSING),
        # ── زمینه ──
        "مرحله فعلی": _first_col(work, ["STAGE_FA"]),
        "مدیریت": _first_col(work, ["ORG_DEPT"]),
        "بحرانی": _first_col(work, ["بحرانی (کوتاه)", "کد طبقه بحرانی"]),
        "مقاومت (روز)": _num(work, "مقاومت (روز)"),
        "روزهای رسوب": _num(work, "روزهای رسوب"),
    })
    for scope in SCOPES:
        out[scope.fa] = _s(work, scope.key).replace("", "—")
    return out.sort_values(["کلید گروه"], kind="stable", ignore_index=True)


def _reason(work: pd.DataFrame) -> pd.Series:
    """علت ثبت‌شده مقدم است؛ سپس علت داده‌ای و قابل ممیزی."""
    r = _first_col(work, ["BL_CRITICAL_REASON", "ORDER_CRITICAL_REASON", "روایت",
                          "شرح هشدارها", "ALERTS", "MOGH_COMMERCIAL_NOTE",
                          "MOGH_LOGISTICS_NOTE"], "")
    waiting = _s(work, ps.WAITING_SCOPE)
    nxt = _s(work, ps.NEXT_ACT)
    derived = pd.Series("علت اختصاصی در داده‌های موجود ثبت نشده است.",
                        index=work.index, dtype=object)
    has_next = waiting.ne("") & nxt.ne("")
    derived = derived.mask(
        has_next, "منتظر «" + nxt + "» است؛ این کار بر عهده " + waiting + " است.")
    no_evidence = _s(work, ps.WHERE).isin(["", "نامشخص"])
    derived = derived.mask(
        no_evidence,
        "هیچ رویداد تاریخ‌داری ثبت نشده است؛ تا ثبت اولین تاریخ، موقعیت قابل تعیین نیست.")
    return r.where(r.ne(""), derived)


def build_material_view(df: pd.DataFrame) -> pd.DataFrame:
    return _base(df, "KEY_MATERIAL", "متریال")


def build_bl_view(df: pd.DataFrame) -> pd.DataFrame:
    return _base(df, "CANONICAL_BL", "بارنامه")


def build_dept_view(df: pd.DataFrame) -> pd.DataFrame:
    """نمای اداره — یک ردیف به ازای هر اداره/مدیریت."""
    key = "ORG_DEPT"
    if key not in df.columns:
        return pd.DataFrame()
    work = df[_s(df, key).ne("")].copy()
    if work.empty:
        return pd.DataFrame()
    work = _ensure_status(work)
    keys = _s(work, key)
    work = work.assign(_k=keys)

    where_mode = _group_mode(_s(work, ps.WHERE), keys)
    wait_mode = _group_mode(_s(work, ps.WAITING_SCOPE), keys)
    age = _num(work, ps.AGE)

    def nunique(col: str) -> pd.Series:
        if col not in work.columns:
            return pd.Series(0, index=work["_k"].unique())
        s = _s(work, col).replace("", pd.NA)
        return s.groupby(work["_k"]).nunique()

    crit = (_s(work, "کد طبقه بحرانی")
            .isin(["STOCKOUT", "CRITICAL", "BECOMING_CRITICAL"])
            .groupby(work["_k"]).sum())
    gap = _s(work, OWNER_GAP).ne("").groupby(work["_k"]).sum()
    ownerless = _s(work, OWNER_NAME).eq("").groupby(work["_k"]).sum()

    idx = pd.Index(sorted(work["_k"].unique()), name=None)
    out = pd.DataFrame({
        "کلید گروه": idx,
        "اداره / مدیریت": idx,
        "معاونت": _group_mode(_s(work, "ORG_VICE"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "مدیر": _group_mode(_s(work, "ORG_MANAGER"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "رئیس": _group_mode(_s(work, "ORG_HEAD"), keys).groupby(work["_k"]).first().reindex(idx).fillna(""),
        "موقعیت غالب پرونده‌ها": where_mode.groupby(work["_k"]).first().reindex(idx).fillna("نامشخص"),
        "معطل غالب حوزه": wait_mode.groupby(work["_k"]).first().reindex(idx).replace("", "—").fillna("—"),
        "میانگین سن وضعیت (روز)": age.groupby(work["_k"]).mean().reindex(idx).round(1),
        "بیشترین سن وضعیت (روز)": age.groupby(work["_k"]).max().reindex(idx),
        "تعداد متریال": nunique("KEY_MATERIAL").reindex(idx).fillna(0).astype(int),
        "تعداد بارنامه": nunique("CANONICAL_BL").reindex(idx).fillna(0).astype(int),
        "تعداد سفارش": nunique("CANONICAL_ORDER").reindex(idx).fillna(0).astype(int),
        "تعداد بحرانی": crit.reindex(idx).fillna(0).astype(int),
        "ردیف با شکاف داده مالک": gap.reindex(idx).fillna(0).astype(int),
        "ردیف بدون مالک شناسایی‌شده": ownerless.reindex(idx).fillna(0).astype(int),
        "کمترین مقاومت (روز)": _num(work, "مقاومت (روز)").groupby(work["_k"]).min().reindex(idx),
        "بیشترین رسوب (روز)": _num(work, "روزهای رسوب").groupby(work["_k"]).max().reindex(idx),
    }).reset_index(drop=True)
    out["چرایی بار کاری / وضعیت"] = [
        f"غالب پرونده‌ها در «{w}» است"
        + (f" و معطل {t} " if t not in ("—", "") else " ")
        + f"— {c} پرونده بحرانی، {g} ردیف با شکاف داده مالک."
        for w, t, c, g in zip(out["موقعیت غالب پرونده‌ها"], out["معطل غالب حوزه"],
                              out["تعداد بحرانی"], out["ردیف با شکاف داده مالک"])]
    return out


# ═══════════════════════════════════════════════════════════════════════════
#  نوشتن در Excel
# ═══════════════════════════════════════════════════════════════════════════
def style_grouped(ws, df: pd.DataFrame, key_col: str = "کلید گروه") -> None:
    if df.empty:
        return
    red = Side(style="thick", color=RED)
    vals = df[key_col].astype(str).tolist() if key_col in df.columns else []
    for r in range(2, len(df) + 2):
        first = r == 2 or (r - 2 < len(vals) and vals[r - 2] != vals[r - 3])
        last = r == len(df) + 1 or (r - 1 < len(vals) and vals[r - 1] != vals[r - 2])
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = P.font_body()
            cell.alignment = P.align("right", wrap=True)
            b = cell.border
            cell.border = Border(left=b.left, right=b.right,
                                 bottom=red if last else b.bottom,
                                 top=red if first else b.top)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for i in range(1, ws.max_column + 1):
        sample = [str(ws.cell(x, i).value or "") for x in range(1, min(ws.max_row, 25) + 1)]
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(14, max(map(len, sample)) + 2), 48)


def _write_table(ws, t: pd.DataFrame) -> None:
    for j, h in enumerate(t.columns, 1):
        c = ws.cell(1, j, h)
        c.font = P.font_header(1)
        c.fill = P.fill_header()
        c.alignment = P.align("center", wrap=True)
    for i, row in enumerate(t.itertuples(index=False, name=None), 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, None if pd.isna(v) else v)
    style_grouped(ws, t, "کلید گروه" if "کلید گروه" in t.columns else "__none__")


SHEETS = ("۱۴. متریال محور", "۱۵. بارنامه محور", "۱۶. اداره محور")


def write_supply_sheets(wb, df: pd.DataFrame) -> dict:
    specs = [(SHEETS[0], build_material_view(df)),
             (SHEETS[1], build_bl_view(df)),
             (SHEETS[2], build_dept_view(df))]
    for name, t in specs:
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = False
        if t.empty:
            ws["A1"] = "داده کافی برای این نما وجود ندارد."
            continue
        _write_table(ws, t)
    return {name: t for name, t in specs}
