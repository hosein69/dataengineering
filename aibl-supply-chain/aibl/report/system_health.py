# -*- coding: utf-8 -*-
"""شیت «۱۷. سلامت سیستم» — این عدد از کجا آمده و امروز چه چیزی قابل اتکاست.

هر گزارش عددی یک فرض پنهان دارد: «داده‌ای که این عدد از آن آمده سالم
بود». وقتی آن فرض جایی نوشته نشود، خواننده ناچار است فرضش کند — و
دقیقاً همان‌جا تصمیم غلط گرفته می‌شود.

این شیت پنج بخش دارد و به یک سؤال جواب می‌دهد: **اگر عددی در این گزارش
عجیب است، مشکل از داده است، سورس، ادغام، ترتیب زمانی، مرحله، یا قاعده
کسب‌وکار؟**
"""
from __future__ import annotations

from typing import List, Tuple

import pandas as pd
from openpyxl.utils import get_column_letter

from .. import health
from .palette import LuxuryPalette as P

SHEET = "۱۷. سلامت سیستم"

_VERDICT_NOTE = {
    health.OK: "همه سورس‌های الزامی سالم‌اند و هیچ مرحله‌ای شکست نخورده؛ "
               "اعداد این گزارش بر مبنای داده کامل ساخته شده‌اند.",
    health.DEGRADED: "گزارش ساخته شد ولی بخشی از داده ناقص است. "
                     "KPIهای وابسته به سورس‌های ناقص را با احتیاط بخوانید.",
    health.FAILED: "سورس الزامی یا مرحله‌ای شکست خورده است. "
                   "اعداد وابسته به آن بخش در این اجرا قابل اتکا نیستند.",
}


def _write_block(ws, row: int, title: str, table: pd.DataFrame,
                 note: str = "") -> int:
    c = ws.cell(row=row, column=1, value=title)
    c.font = P.font_title(12)
    row += 1
    if note:
        n = ws.cell(row=row, column=1, value=note)
        n.font = P.font_body()
        n.alignment = P.align("right", wrap=True)
        row += 1
    if table is None or table.empty:
        ws.cell(row=row, column=1, value="— موردی ثبت نشد —").font = P.font_body()
        return row + 2
    for j, h in enumerate(table.columns, 1):
        cell = ws.cell(row=row, column=j, value=str(h))
        cell.font = P.font_header(1)
        cell.fill = P.fill_header()
        cell.alignment = P.align("center", wrap=True)
    row += 1
    status_col = next((i for i, h in enumerate(table.columns, 1)
                       if str(h) in ("وضعیت", "شدت")), None)
    for _, r in table.iterrows():
        for j, v in enumerate(r, 1):
            cell = ws.cell(row=row, column=j,
                           value=None if pd.isna(v) else v)
            cell.font = P.font_body()
            cell.alignment = P.align("right", wrap=True)
        if status_col:
            st = str(r.iloc[status_col - 1])
            if st in (health.FAILED, health.SKIPPED, "خطا"):
                for j in range(1, len(table.columns) + 1):
                    ws.cell(row=row, column=j).fill = P.fill(P.CRITICAL_FILL)
            elif st in (health.DEGRADED, "هشدار"):
                for j in range(1, len(table.columns) + 1):
                    ws.cell(row=row, column=j).fill = P.fill(P.STATUS_WARNING_FILL)
        row += 1
    return row + 1


def build(wb, hp: "health.SystemHealth" = None) -> None:
    hp = hp or health.current()
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    verdict = hp.verdict()
    head = ws.cell(row=1, column=1, value=f"وضعیت کلی این اجرا: {verdict}")
    head.font = P.font_title(16)
    note = ws.cell(row=2, column=1, value=_VERDICT_NOTE.get(verdict, ""))
    note.font = P.font_body()
    note.alignment = P.align("right", wrap=True)

    row = 4
    blocking = hp.blocking()
    if blocking:
        b = ws.cell(row=row, column=1,
                    value="⛔ سورس الزامیِ در دسترس نبود: " + "، ".join(blocking)
                          + " — هر KPI وابسته به این سورس‌ها امروز بی‌اعتبار است.")
        b.font = P.font_title(12)
        b.fill = P.fill(P.CRITICAL_FILL)
        b.alignment = P.align("right", wrap=True)
        row += 2

    counts = pd.DataFrame([{"شاخص": k, "تعداد": v} for k, v in hp.counts().items()])
    row = _write_block(ws, row, "۰) خلاصه", counts)
    row = _write_block(ws, row, "۱) سلامت سورس‌ها", hp.sources_table(),
                       "«رد شد» یعنی فایل پیدا نشد؛ «ناقص» یعنی فایل بود ولی "
                       "شیت/ستون مورد انتظار نبود. این دو درمانِ متفاوت دارند.")
    row = _write_block(ws, row, "۲) سلامت مرحله‌های خط لوله", hp.stages_table(),
                       "ستون زمان می‌گوید کدام مرحله گلوگاه اجراست.")
    row = _write_block(ws, row, "۳) سلامت ادغام‌ها", hp.joins_table(),
                       "ادغامی که «رکورد منطبق» صفر دارد یعنی کلید دو طرف "
                       "هم‌جنس نیست؛ ستون‌های آن سورس خالی می‌مانند.")
    row = _write_block(ws, row, "۴) یافته‌ها", hp.findings_table())

    widths = [26, 30, 14, 12, 34, 12, 12, 14, 40, 46]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


def summary_rows(hp: "health.SystemHealth" = None) -> List[Tuple[str, str]]:
    """خلاصه کوتاه برای ایمیل و داشبورد."""
    hp = hp or health.current()
    c = hp.counts()
    rows = [("وضعیت داده امروز", hp.verdict())]
    if hp.blocking():
        rows.append(("سورس الزامیِ ناموجود", "، ".join(hp.blocking())))
    rows.append(("سورس سالم / ناقص / خراب",
                 f"{c['سورس سالم']} / {c['سورس ناقص']} / {c['سورس خراب/ردشده']}"))
    if c["خطا"] or c["هشدار"]:
        rows.append(("خطا / هشدار سیستمی", f"{c['خطا']} / {c['هشدار']}"))
    return rows
