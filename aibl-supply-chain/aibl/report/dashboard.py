# -*- coding: utf-8 -*-
"""سازنده داشبورد اکسل — با استفاده از تمام ظرفیت‌های اکسل (§۱۴).

امکاناتی که در نسخه ۲۰.۱ اصلاً وجود نداشت و اینجا اضافه شده است:
Freeze Panes، AutoFilter، Data Validation (لیست کشویی)، Conditional Formatting
(مقیاس رنگی + قانون متنی)، Formula Injection (SUBTOTAL/COUNTIF زنده)،
Number Formatting و گروه‌بندی سطر/ستون با outline_level.
"""
from __future__ import annotations

__contract__ = 3   # ← aibl/contracts.py

import os
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..dataio.logging_setup import log
from ..rulebook import get_rulebook
from .palette import (LuxuryPalette as P, NUM_FORMAT_CURRENCY, NUM_FORMAT_INT,
                      NUM_FORMAT_PCT)

SHEET_EXEC = "۱. خلاصه اجرایی"
SHEET_MATRIX = "۲. کالبدشکافی ۳ لایه‌ای ماتریسی"
SHEET_RESOLVE = "۳. تعیین تکلیف"
SHEET_COMMIT = "۴. رفع تعهد ارزی"
SHEET_SCORECARD = "۵. کارنامه سازمانی"
SHEET_MATH = "۶. پشتیبان ریاضی"


class ExcelDashboardBuilder:

    def __init__(self, output_path: str) -> None:
        self.output_path = output_path
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)  # FIX-B7
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    # ═══════════ ابزارهای مشترک ═══════════
    def _new_sheet(self, title: str) -> Any:
        ws = self.wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        return ws

    def _write_header(self, ws, headers: List[str], row: int = 1,
                      widths: Optional[List[int]] = None,
                      groups: Optional[Dict[int, int]] = None) -> None:
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.font = P.font_header(1)
            c.fill = P.fill_header()
            c.alignment = P.align("center", wrap=True)
            c.border = P.thin_border()
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = (widths[i - 1] if widths and i <= len(widths) else 20)
            if groups and i in groups:
                ws.column_dimensions[letter].outline_level = groups[i]
                ws.column_dimensions[letter].hidden = True
        ws.row_dimensions[row].height = 34
        ws.freeze_panes = ws.cell(row=row + 1, column=1)          # Freeze Panes
        ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"  # AutoFilter
        ws.sheet_properties.outlinePr.summaryRight = False

    @staticmethod
    def _style_row(ws, r: int, ncols: int, fill_color: Optional[str] = None) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = P.font_body()
            cell.border = P.thin_border()
            cell.alignment = P.align("right", wrap=False)
            if fill_color:
                cell.fill = P.fill(fill_color)

    # ═══════════ شیت ۱ — خلاصه اجرایی ═══════════
    def build_executive(self, kpis: Dict[str, Any], narrative: str,
                        matrix_rows: int) -> None:
        ws = self.wb.create_sheet(SHEET_EXEC, 0)
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 60

        ws["A1"] = "🏛️ داشبورد مدیریتی هوش لجستیک و حاکمیت داده — ایران خودرو"
        ws["A1"].font = P.font_title(16)
        ws.merge_cells("A1:C1")
        ws.row_dimensions[1].height = 30

        ws["A3"] = narrative
        ws["A3"].font = P.font_body()
        ws["A3"].alignment = P.align("right", wrap=True)
        ws.merge_cells("A3:C3")
        ws.row_dimensions[3].height = 90

        ws["A5"] = "شاخص"
        ws["B5"] = "مقدار"
        ws["C5"] = "توضیح"
        for col in "ABC":
            ws[f"{col}5"].font = P.font_header(1)
            ws[f"{col}5"].fill = P.fill_header()
            ws[f"{col}5"].alignment = P.align("center")

        r = 6
        for label, (value, note) in kpis.items():
            ws.cell(row=r, column=1, value=label).font = P.font_body(bold=True)
            v = ws.cell(row=r, column=2, value=value)
            v.font = P.font_kpi()
            v.alignment = P.align("center")
            if isinstance(value, (int, float)):
                v.number_format = NUM_FORMAT_INT if float(value).is_integer() else NUM_FORMAT_PCT
            ws.cell(row=r, column=3, value=note).font = P.font_body()
            for c in range(1, 4):
                ws.cell(row=r, column=c).border = P.thin_border()
                ws.cell(row=r, column=c).fill = P.fill_row_level(2)
            r += 1

        # Formula Injection — شمارنده‌های زنده از شیت ماتریس
        if matrix_rows > 0:
            rng = f"'{SHEET_MATRIX}'!$E$2:$E${matrix_rows + 1}"
            ws.cell(row=r + 1, column=1, value="بارنامه‌های بحرانی (فرمول زنده)").font = P.font_body(bold=True)
            ws.cell(row=r + 1, column=2, value=f'=COUNTIF({rng},"*بحرانی*")').font = P.font_kpi()
            ws.cell(row=r + 2, column=1, value="ردیف‌های قابل مشاهده پس از فیلتر").font = P.font_body(bold=True)
            ws.cell(row=r + 2, column=2,
                    value=f"=SUBTOTAL(103,'{SHEET_MATRIX}'!$B$2:$B${matrix_rows + 1})").font = P.font_kpi()

    # ═══════════ شیت ۲ — کالبدشکافی ماتریسی (ستون‌محور) ═══════════
    def build_matrix(self, df: pd.DataFrame, specs: list) -> int:
        """ستون‌ها را **مرحله‌ها** اعلام می‌کنند، نه این فایل.

        افزودن ستون جدید به گزارش = افزودن یک ColumnSpec در همان مرحله‌ای که
        ستون را می‌سازد. این متد هرگز برای قابلیت جدید تغییر نمی‌کند.
        """
        ws = self._new_sheet(SHEET_MATRIX)
        specs = [sp for sp in specs if sp.key in df.columns]
        if not specs:
            log.warning("⚠️ هیچ ستون گزارشی توسط مرحله‌ها اعلام نشد.")
            return 0

        headers = [sp.title for sp in specs]
        widths = [sp.width for sp in specs]
        groups = {i: sp.group for i, sp in enumerate(specs, start=1) if sp.group > 0}
        self._write_header(ws, headers, widths=widths, groups=groups)

        fmt_map = {"int": NUM_FORMAT_INT, "decimal": NUM_FORMAT_PCT,
                   "currency": NUM_FORMAT_CURRENCY}
        band_fills = {b["code"]: b.get("fill", "")
                      for b in (get_rulebook().get("criticality.bands", []) or [])}
        crit_col = next((i for i, sp in enumerate(specs, start=1)
                         if sp.key == "طبقه بحرانی"), None)

        r = 2
        for _, row in df.iterrows():
            critical = "بحرانی" in str(row.get("وضعیت هوشمند", ""))
            for i, sp in enumerate(specs, start=1):
                ws.cell(row=r, column=i, value=self._cell_value(row.get(sp.key, "")))
            self._style_row(ws, r, len(specs),
                            P.CRITICAL_FILL if critical else P.GREEN_L4)
            for i, sp in enumerate(specs, start=1):
                cell = ws.cell(row=r, column=i)
                if sp.fmt in fmt_map:
                    cell.number_format = fmt_map[sp.fmt]
                if sp.wrap:
                    cell.alignment = P.align("right", wrap=True)
            if crit_col:
                fill = band_fills.get(str(row.get("کد طبقه بحرانی", "")))
                if fill:
                    c = ws.cell(row=r, column=crit_col)
                    c.fill = P.fill(fill)
                    c.font = P.font_body(bold=True)
            r += 1

        last = r - 1
        if last >= 2:
            self._apply_rules(ws, specs, last)
            self._add_decision_column(ws, len(specs), last)
        log.info(f"📄 شیت «{SHEET_MATRIX}» ساخته شد — {last - 1} ردیف × {len(specs)} ستون.")
        return last - 1

    @staticmethod
    def _apply_rules(ws, specs: list, last: int) -> None:
        """قوانین رنگی که هر ستون در ColumnSpec خودش اعلام کرده است."""
        for i, sp in enumerate(specs, start=1):
            if not sp.color_rule:
                continue
            col = get_column_letter(i)
            rng = f"{col}2:{col}{last}"
            if sp.color_rule == "scale_low_bad":
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="num", start_value=0, start_color="F5B7B1",
                    mid_type="num", mid_value=20, mid_color="FDEBD0",
                    end_type="num", end_value=60, end_color="D5F5E3"))
            elif sp.color_rule == "scale_high_bad":
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="num", start_value=0, start_color="C6EFCE",
                    mid_type="num", mid_value=50, mid_color="FFEB9C",
                    end_type="num", end_value=100, end_color="FFC7CE"))
            elif sp.color_rule == "flag_nonempty":
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f"LEN(${col}2)>0"], fill=P.fill("F5B7B1"),
                    font=P.font_body(bold=True)))

    def _add_decision_column(self, ws, n_cols: int, last: int) -> None:
        dv_col = n_cols + 1
        ws.cell(row=1, column=dv_col, value="تصمیم کارشناس").font = P.font_header(1)
        ws.cell(row=1, column=dv_col).fill = P.fill_header()
        ws.column_dimensions[get_column_letter(dv_col)].width = 22
        dv = DataValidation(
            type="list",
            formula1='"در دست اقدام,ارجاع به گمرک,ارجاع به بانک,مختومه,نیازمند استعلام"',
            allow_blank=True, showDropDown=False)
        dv.error = "لطفاً یکی از گزینه‌های تعریف‌شده را انتخاب کنید."
        dv.errorTitle = "مقدار نامعتبر"
        ws.add_data_validation(dv)
        letter = get_column_letter(dv_col)
        dv.add(f"{letter}2:{letter}{last}")

    @staticmethod
    def _cell_value(v: Any) -> Any:
        if v is None:
            return ""
        if isinstance(v, (int, float, str)):
            return v
        return str(v)

    # ═══════════ شیت ۳ — تعیین تکلیف ═══════════
    def build_to_resolve(self, df: pd.DataFrame, main_df: Optional[pd.DataFrame] = None) -> None:
        ws = self._new_sheet(SHEET_RESOLVE)
        headers = ["شماره سفارش کانونی", "شماره بارنامه کانونی", "مسئولیت سازمانی",
                   "شرح کالا", "افراز کلید", "شرح علت عدم تعیین تکلیف",
                   "پیشنهاد عملیاتی سیستم"]
        self._write_header(ws, headers, widths=[20, 22, 45, 30, 26, 60, 50])

        r = 2
        for _, row in df.iterrows():
            bl = row.get("CANONICAL_BL") or "نامشخص"
            expert = row.get("CANONICAL_EXPERT") or "نامشخص"
            reason = (f"بارنامه {bl} در فایل مقاومت یافت نشد و همزمان فاقد کد ساتا و "
                      f"کوتاژ گمرکی است؛ هیچ سند مالی/گمرکی معتبری در شبکه سورس‌ها ندارد.")
            rec = (f"کارشناس {expert} موظف است فیزیک اسناد حمل را از شرکت حمل بین‌المللی "
                   f"استعلام و ردیف را در فایل مقاومت درج نماید.")
            vals = [row.get("CANONICAL_ORDER") or "نامشخص", bl,
                    row.get("ORG_CHAIN", ""), row.get("CANONICAL_GOODS_DESC", ""),
                    row.get("PARTITION_KEY", ""), reason, rec]
            for i, v in enumerate(vals, start=1):
                ws.cell(row=r, column=i, value=self._cell_value(v))
            self._style_row(ws, r, len(headers), P.AMBER_FILL)
            ws.cell(row=r, column=6).alignment = P.align("right", wrap=True)
            ws.cell(row=r, column=7).alignment = P.align("right", wrap=True)
            r += 1
        # ── سفارش‌های خارج از Commercial Expert Data ──
        # این بخش عمداً از «تعیین تکلیف» جداست: سفارش در جریان اصلی هست،
        # اما در سورس کارشناسان خرید بازرگانی ثبت نشده است. هیچ کارشناس خریدی
        # از روی این فقدان داده حدس زده نمی‌شود.
        source_for_missing = main_df if main_df is not None else df
        missing = source_for_missing[source_for_missing.get("ORDER_MISSING_COMMERCIAL_EXPERT", False).astype(bool)].copy() if "ORDER_MISSING_COMMERCIAL_EXPERT" in source_for_missing.columns else pd.DataFrame()
        if not missing.empty:
            r += 1
            title_row = r
            ws.cell(row=r, column=1, value="⚠️ سفارش‌های موجود در جریان اصلی ولی درج‌نشده در Commercial Expert Data")
            ws.cell(row=r, column=1).font = P.font_title(12)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
            r += 1
            miss_headers = ["شماره سفارش کانونی", "شماره بارنامه کانونی", "مسئولیت سازمانی",
                            "شرح کالا", "افراز کلید", "علت عدم درج در Commercial Expert Data",
                            "وضعیت انتساب کارشناس خرید"]
            for i, h in enumerate(miss_headers, 1):
                c = ws.cell(row=r, column=i, value=h); c.font=P.font_header(1); c.fill=P.fill(P.CRITICAL_FILL); c.alignment=P.align("center", wrap=True)
            miss_start = r
            r += 1
            for _, row in missing.drop_duplicates(subset=["CANONICAL_ORDER"], keep="first").iterrows():
                vals = [row.get("CANONICAL_ORDER", ""), row.get("CANONICAL_BL", ""),
                        row.get("ORG_CHAIN", ""), row.get("CANONICAL_GOODS_DESC", ""),
                        row.get("PARTITION_KEY", ""),
                        row.get("ORDER_MISSING_COMMERCIAL_REASON", "این سفارش در Commercial Expert Data درج نشده است."),
                        "کارشناس خرید قابل انتساب نیست — داده مبدأ وجود ندارد"]
                for i,v in enumerate(vals,1):
                    ws.cell(row=r,column=i,value=self._cell_value(v))
                self._style_row(ws,r,len(miss_headers),P.CRITICAL_FILL)
                ws.cell(row=r,column=6).alignment=P.align("right",wrap=True)
                ws.cell(row=r,column=7).alignment=P.align("right",wrap=True)
                r += 1
            ws.auto_filter.ref = f"A{miss_start}:{get_column_letter(len(miss_headers))}{r-1}"
        log.info(f"📄 شیت «{SHEET_RESOLVE}» ساخته شد — {r - 2} ردیف، سفارش خارج از Commercial Expert Data: {len(missing.drop_duplicates('CANONICAL_ORDER')) if not missing.empty else 0}.")

    # ═══════════ شیت ۴ — رفع تعهد ارزی ═══════════
    def build_commitment(self, df: pd.DataFrame) -> None:
        ws = self._new_sheet(SHEET_COMMIT)
        headers = ["شماره ثبت سفارش", "شماره بارنامه", "نوع پرونده", "ارز",
                   "تعهد اولیه", "مانده تعهد", "تاریخ ایجاد تعهد", "مهلت رفع تعهد",
                   "مهلت قانونی محاسبه‌شده", "روزهای تأخیر", "جریمه برآوردی",
                   "وضعیت رفع تعهد", "تخصیص ارز", "وضعیت کلی هشدار", "شرح هشدارها"]
        self._write_header(ws, headers,
                           widths=[18, 20, 14, 10, 18, 18, 16, 16, 20, 14, 18, 18, 14, 16, 70])

        cols = ["CANONICAL_REG", "CANONICAL_BL", "نوع پرونده", "NTSW_CURRENCY",
                "NTSW_INITIAL_COMMIT", "مانده تعهد", "NTSW_COMMIT_DATE", "NTSW_DEADLINE",
                "مهلت قانونی رفع تعهد", "روزهای تأخیر", "جریمه برآوردی",
                "NTSW_RELEASE_STATUS", "ALLOC_STATUS", "وضعیت کلی هشدار", "شرح هشدارها"]
        r = 2
        for _, row in df.iterrows():
            for i, key in enumerate(cols, start=1):
                ws.cell(row=r, column=i, value=self._cell_value(row.get(key, "")))
            status = str(row.get("وضعیت کلی هشدار", "سبز"))
            color = {"قرمز": P.CRITICAL_FILL, "زرد": "FCF3CF"}.get(status, P.GREEN_L3)
            self._style_row(ws, r, len(headers), color)
            for i in (5, 6, 11):
                ws.cell(row=r, column=i).number_format = NUM_FORMAT_CURRENCY
            ws.cell(row=r, column=15).alignment = P.align("right", wrap=True)
            r += 1

        last = r - 1
        if last >= 2:
            # جمع زنده مانده تعهد و جریمه (SUBTOTAL با فیلتر همگام است)
            ws.cell(row=last + 2, column=4, value="جمع کل:").font = P.font_body(bold=True)
            ws.cell(row=last + 2, column=6, value=f"=SUBTOTAL(109,F2:F{last})").number_format = NUM_FORMAT_CURRENCY
            ws.cell(row=last + 2, column=11, value=f"=SUBTOTAL(109,K2:K{last})").number_format = NUM_FORMAT_CURRENCY
            ws.conditional_formatting.add(
                f"J2:J{last}",
                CellIsRule(operator="greaterThan", formula=["0"], fill=P.fill("F5B7B1")))
        log.info(f"📄 شیت «{SHEET_COMMIT}» ساخته شد — {last - 1} ردیف.")

    # ═══════════ شیت ۵ — کارنامه سازمانی ═══════════
    def build_scorecard(self, df: pd.DataFrame) -> None:
        ws = self._new_sheet(SHEET_SCORECARD)
        headers = ["معاونت", "مدیریت", "مدیر", "رئیس", "کارشناس", "کد پرسنلی",
                   "تعداد بارنامه یکتا", "بحرانی", "میانگین روز رسوب",
                   "میانگین امتیاز ریسک", "جمع مانده تعهد", "جمع جریمه"]
        self._write_header(ws, headers, widths=[22, 24, 20, 20, 22, 14, 18, 12, 18, 18, 20, 20])

        # FIX-7 + FIX-8: تجمیع بر اساس بارنامه یکتا، اما وضعیت بحرانی از
        # «پرونده بحرانی» خوانده می‌شود؛ نه از اولین ردیف بارنامه.
        # یک BL می‌تواند چند متریال داشته باشد و متریال بحرانی ممکن است
        # در ردیفی غیر از اولین ردیف قرار گرفته باشد.
        uniq = df.drop_duplicates(subset=["CANONICAL_BL"], keep="first") if "CANONICAL_BL" in df else df.copy()
        if uniq.empty:
            return
        if {"CANONICAL_BL", "BL_CRITICAL"}.issubset(df.columns):
            crit_by_bl = (df.assign(_BL_CRITICAL=df["BL_CRITICAL"].astype(bool))
                            .groupby("CANONICAL_BL", dropna=False)["_BL_CRITICAL"].any()
                            .rename("_BL_CRITICAL_GROUP"))
            uniq = uniq.merge(crit_by_bl, left_on="CANONICAL_BL", right_index=True, how="left")
        else:
            base = uniq.get("وضعیت هوشمند", pd.Series(index=uniq.index, dtype=object)).astype(str)
            uniq["_BL_CRITICAL_GROUP"] = base.str.contains("بحرانی")

        grp = uniq.groupby(["ORG_VICE", "ORG_DEPT", "ORG_MANAGER", "ORG_HEAD",
                            "CANONICAL_EXPERT", "KEY_EMP"], dropna=False)
        r = 2
        for keys, g in grp:
            crit = int(g.get("_BL_CRITICAL_GROUP", pd.Series(False, index=g.index)).astype(bool).sum())
            levels = []
            if {"CANONICAL_BL", "BL_CRITICAL_LEVEL"}.issubset(df.columns):
                level_rows = (df[df["CANONICAL_BL"].isin(g["CANONICAL_BL"])]
                              [["CANONICAL_BL", "BL_CRITICAL_LEVEL"]]
                              .drop_duplicates("CANONICAL_BL"))
                levels = [str(x) for x in level_rows["BL_CRITICAL_LEVEL"].tolist()]
            if crit:
                row_fill, crit_color = P.STATUS_CRITICAL_FILL, P.STATUS_CRITICAL
            elif "BECOMING_CRITICAL" in levels:
                row_fill, crit_color = P.STATUS_WARNING_FILL, P.STATUS_WARNING
            elif "WATCH" in levels:
                row_fill, crit_color = P.STATUS_WATCH_FILL, P.STATUS_WATCH
            elif any(x in {"NO_CONSUMPTION", "UNKNOWN"} for x in levels):
                row_fill, crit_color = P.STATUS_INACTIVE_FILL, P.STATUS_INACTIVE
            else:
                row_fill, crit_color = P.STATUS_GOOD_FILL, P.STATUS_GOOD

            vals = list(keys) + [
                int(g["CANONICAL_BL"].nunique()), crit,
                round(pd.to_numeric(g["روزهای رسوب"], errors="coerce").mean() or 0, 1),
                round(pd.to_numeric(g.get("امتیاز ریسک", 0), errors="coerce").mean() or 0, 1),
                round(pd.to_numeric(g.get("مانده تعهد", 0), errors="coerce").sum() or 0, 2),
                round(pd.to_numeric(g.get("جریمه برآوردی", 0), errors="coerce").sum() or 0, 2),
            ]
            for i, v in enumerate(vals, start=1):
                ws.cell(row=r, column=i, value=self._cell_value(v))
            self._style_row(ws, r, len(headers), row_fill)
            crit_cell = ws.cell(row=r, column=8)
            crit_cell.font = P.font_body(bold=True)
            crit_cell.font = crit_cell.font.copy(color=crit_color)
            crit_cell.alignment = P.align("center")
            for i in (11, 12):
                ws.cell(row=r, column=i).number_format = NUM_FORMAT_CURRENCY
            r += 1
        last = r - 1
        if last >= 2:
            ws.conditional_formatting.add(
                f"J2:J{last}",
                ColorScaleRule(start_type="num", start_value=0, start_color="C6EFCE",
                               mid_type="num", mid_value=50, mid_color="FFEB9C",
                               end_type="num", end_value=100, end_color="FFC7CE"))
        log.info(f"📄 شیت «{SHEET_SCORECARD}» ساخته شد — {last - 1} ردیف.")

    # ═══════════ شیت ۶ — پشتیبان ریاضی ═══════════
    def build_math(self, layers: Dict[str, Dict[str, Any]]) -> None:
        ws = self._new_sheet(SHEET_MATH)
        for col, w in zip("ABC", (45, 55, 75)):
            ws.column_dimensions[col].width = w
        r = 1
        for algo, params in layers.items():
            t = ws.cell(row=r, column=1, value=algo)
            t.font = P.font_title(12)
            t.fill = P.fill_row_level(1)
            r += 1
            for name, val in params.items():
                ws.cell(row=r, column=2, value=name).font = P.font_body(bold=True)
                ws.cell(row=r, column=3, value=str(val)).font = P.font_body()
                for c in (2, 3):
                    ws.cell(row=r, column=c).border = P.thin_border()
                r += 1
            r += 1
        log.info(f"📄 شیت «{SHEET_MATH}» ساخته شد.")

    # ═══════════ ذخیره ═══════════
    def save(self) -> str:
        """ذخیره اتمیک و مقاوم در برابر فایل قفل‌شده توسط Excel/Outlook.

        ابتدا Workbook در فایل موقتِ همان پوشه نوشته می‌شود؛ سپس جایگزینی
        اتمیک انجام می‌گیرد. اگر فایل مقصد در Windows توسط Excel قفل باشد،
        نسخه زمان‌دار ساخته می‌شود. خطای دوم دیگر در سکوت گم نمی‌شود.
        """
        import tempfile
        from datetime import datetime
        folder = os.path.dirname(self.output_path) or "."
        base = os.path.basename(self.output_path)
        tmp_path = ""
        try:
            fd, tmp_path = tempfile.mkstemp(prefix=".aibl_", suffix=".xlsx", dir=folder)
            os.close(fd)
            self.wb.save(tmp_path)
            try:
                os.replace(tmp_path, self.output_path)
                tmp_path = ""
                log.info(f"🏆 داشبورد ذخیره شد: {self.output_path}")
                return self.output_path
            except PermissionError as exc:
                alt = os.path.join(folder, base.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx"))
                self.wb.save(alt)
                log.warning(f"⚠️ فایل مقصد قفل بود؛ نسخه جایگزین ساخته شد: {alt} | {exc}")
                return alt
        except PermissionError as exc:
            alt = os.path.join(folder, base.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx"))
            try:
                self.wb.save(alt)
                log.warning(f"⚠️ مسیر مقصد قابل جایگزینی نبود؛ نسخه جایگزین ساخته شد: {alt} | {exc}")
                return alt
            except Exception as exc2:
                raise RuntimeError(f"ذخیره Excel شکست خورد. مقصد={self.output_path} | خطای اصلی={exc!r} | خطای جایگزین={exc2!r}") from exc2
        except Exception as exc:
            raise RuntimeError(f"ذخیره Workbook شکست خورد: {type(exc).__name__}: {exc}") from exc
        finally:
            if tmp_path:
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass


# ═══════════════════════════════════════════════════════════════════════════
#  شیت‌های افزوده در نسخه ۲۲ (سورس مقاومت جدید + کتابخانه قوانین)
# ═══════════════════════════════════════════════════════════════════════════
SHEET_LINES = "۷. اقلام سفارش (سطح PR و PI)"
SHEET_RULES = "۸. کتابخانه قوانین"


def _build_order_lines(self, lines: "pd.DataFrame") -> None:
    """جدول سطح-قلم سورس مقاومت — drill-down زیر هر سفارش."""
    ws = self._new_sheet(SHEET_LINES)
    headers = [
        "مرجع سفارش", "ریشه سفارش", "شماره PR", "قلم PR", "متریال", "شرح کالا",
        "شماره فنی سازنده", "کد فروشنده", "شماره PI فروشنده", "ارز",
        "تعداد PI", "قیمت واحد", "ارزش قلم", "هزینه‌های افزوده",
        "تعداد سفارش", "تعداد پارت", "تعداد ترخیص‌شده",
        "روش حمل", "حمل‌کننده", "شماره بارنامه (معتبر)", "مقدار مشکوک BL",
        "علت رد BL", "وضعیت بازرگانی", "وضعیت لجستیک", "مرحله", "پیشرفت (٪)",
        "تعرفه پیشنهادی", "کلیدواژه تعرفه", "کد پرسنلی",
    ]
    widths = [16, 14, 14, 10, 14, 34, 20, 14, 18, 8,
              12, 14, 16, 16, 12, 12, 14, 12, 16, 20, 20,
              24, 30, 26, 14, 12, 14, 16, 14]
    groups = {i: 1 for i in range(11, 18)}
    groups.update({i: 2 for i in range(20, 30)})
    self._write_header(ws, headers, widths=widths, groups=groups)

    p = "MOGH_"
    cols = [p + c for c in [
        "ORDER_REF", "ORDER_BASE", "PR_NO", "PR_ITEM", "MATERIAL", "MATERIAL_DESC",
        "MFR_PART_NO", "VENDOR_CODE", "VENDOR_PI_NO", "CURRENCY",
        "PI_QTY", "PI_UNIT_PRICE", "PI_LINE_VALUE", "PI_ADDITIONAL",
        "QTY_IN_ORDER", "QTY_IN_PART", "CLEARED_QTY",
        "TRANSPORT_MODE_CODE", "CARRIER", "BL_NO", "BL_SUSPECT",
        "BL_REJECT_REASON", "COMMERCIAL_NOTE", "LOGISTICS_NOTE", "STAGE_FA",
        "PROGRESS", "HS_SUGGESTED", "HS_KEYWORD", "KEY_EMP"]]

    numeric = {11, 12, 13, 14, 15, 16, 17, 26}
    currency_cols = {12, 13, 14}
    r = 2
    for _, row in lines.iterrows():
        suspect = str(row.get(p + "BL_SUSPECT", "")).strip() != ""
        cancelled = bool(row.get(p + "EXCLUDED_FROM_KPI", False))
        for i, key in enumerate(cols, start=1):
            ws.cell(row=r, column=i, value=self._cell_value(row.get(key, "")))
        fill = (P.CRITICAL_FILL if cancelled else
                "FDEBD0" if suspect else P.GREEN_L4)
        self._style_row(ws, r, len(headers), fill)
        for i in numeric:
            ws.cell(row=r, column=i).number_format = (
                NUM_FORMAT_CURRENCY if i in currency_cols else NUM_FORMAT_INT)
        r += 1

    last = r - 1
    if last >= 2:
        ws.conditional_formatting.add(
            f"Z2:Z{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="FFC7CE",
                           mid_type="num", mid_value=50, mid_color="FFEB9C",
                           end_type="num", end_value=100, end_color="C6EFCE"))
        ws.conditional_formatting.add(
            f"U2:U{last}",
            FormulaRule(formula=['LEN($U2)>0'], fill=P.fill("F5B7B1")))
        ws.cell(row=last + 2, column=12, value="جمع ارزش:").font = P.font_body(bold=True)
        ws.cell(row=last + 2, column=13,
                value=f"=SUBTOTAL(109,M2:M{last})").number_format = NUM_FORMAT_CURRENCY
    log.info(f"📄 شیت «{SHEET_LINES}» ساخته شد — {last - 1} قلم.")


def _build_rulebook_sheet(self, rb) -> None:
    """شفافیت حاکمیتی: هر عددی که در محاسبات استفاده شد، با منبع و وضعیت."""
    ws = self._new_sheet(SHEET_RULES)
    headers = ["بسته", "کلید قاعده", "مقدار", "شرح", "وضعیت اعتبار", "منبع"]
    self._write_header(ws, headers, widths=[18, 34, 16, 52, 20, 34])

    rows = []
    for key, node in (rb.get("fx_governance.deadlines", {}) or {}).items():
        rows.append(("fx_governance", f"deadlines.{key}", node.get("days"),
                     node.get("fa", ""), node.get("status", "internal"),
                     node.get("source", "")))
    for tier in rb.get("fx_governance.penalties.delay_tiers", []) or []:
        rows.append(("fx_governance", "penalties.delay_tier",
                     f"{tier['monthly_rate'] * 100:.0f}٪ ماهانه", tier.get("fa", ""),
                     rb.get("fx_governance.penalties.status", "internal"), ""))
    for seg in ("production", "commercial"):
        for k, t in (rb.thresholds(seg) or {}).items():
            rows.append(("alarms", f"{seg}.{k}",
                         f"legal={t['legal']} / red={t['red']}",
                         f"{t.get('fa', k)} — بازه زرد {t.get('yellow')}",
                         "internal", ""))
    for k, v in (rb.risk_weights() or {}).items():
        rows.append(("alarms", f"risk.weights.{k}", v, "وزن مؤلفه ریسک", "internal", ""))
    rows.append(("customs", "survival_model.beta", rb.get("customs.survival_model.beta"),
                 "پارامتر شکل ویبول", "internal", ""))
    rows.append(("customs", "survival_model.eta_days", rb.get("customs.survival_model.eta_days"),
                 "پارامتر مقیاس ویبول (روز)", "internal", ""))
    rows.append(("customs", "demurrage.critical_days", rb.demurrage_critical_days(),
                 "آستانه رسوب شدید", "internal", ""))
    rows.append(("hs_codes", "current_edition", rb.get("hs_codes.current_edition"),
                 "ویرایش جاری نظام هماهنگ‌شده", "verified", "WCO"))
    rows.append(("hs_codes", "next_edition", rb.get("hs_codes.next_edition"),
                 f"لازم‌الاجرا از {rb.get('hs_codes.next_edition_effective')}", "verified", "WCO"))
    rows.append(("incoterms", "edition", rb.get("incoterms.edition"),
                 "ویرایش جاری ترم‌های حمل", "verified", "ICC"))

    r = 2
    for pack, key, value, desc, status, source in rows:
        for i, v in enumerate((pack, key, value, desc, status, source), start=1):
            ws.cell(row=r, column=i, value=self._cell_value(v))
        color = {"needs_verification": "FDEBD0", "verified": P.GREEN_L2}.get(status, P.GREEN_L4)
        self._style_row(ws, r, len(headers), color)
        ws.cell(row=r, column=4).alignment = P.align("right", wrap=True)
        r += 1

    pending = rb.needs_verification()
    if pending:
        ws.cell(row=r + 1, column=1,
                value=f"⚠️ {len(pending)} قاعده نیازمند تطبیق با آخرین بخشنامه:").font = P.font_title(11)
        r += 2
        for issue in pending:
            ws.cell(row=r, column=1, value=issue.pack).font = P.font_body()
            ws.cell(row=r, column=2, value=issue.path).font = P.font_body()
            c = ws.cell(row=r, column=4, value=issue.message)
            c.font = P.font_body()
            c.alignment = P.align("right", wrap=True)
            for i in range(1, 7):
                ws.cell(row=r, column=i).fill = P.fill("FDEBD0")
                ws.cell(row=r, column=i).border = P.thin_border()
            r += 1
    log.info(f"📄 شیت «{SHEET_RULES}» ساخته شد — {len(rows)} قاعده.")


ExcelDashboardBuilder.build_order_lines = _build_order_lines
ExcelDashboardBuilder.build_rulebook_sheet = _build_rulebook_sheet


SHEET_CRITICAL = "۹. قطعات بحرانی"


def _build_criticality(self, df: "pd.DataFrame") -> None:
    """شیت اختصاصی قطعات بحرانی — مرتب‌شده از توقف خط تا ایمن.

    فقط طبقاتی نمایش داده می‌شوند که نیاز به اقدام دارند
    (توقف خط / بحرانی / در حال بحرانی شدن)، به علاوه خلاصه آماری هر طبقه.
    """
    rb = get_rulebook()
    ws = self._new_sheet(SHEET_CRITICAL)

    bands = sorted(rb.get("criticality.bands", []) or [],
                   key=lambda b: b.get("sort", 99))
    fills = {b["code"]: b.get("fill", "") for b in bands}

    # ── خلاصه آماری بالای شیت ──
    ws["A1"] = "🔧 وضعیت مقاومت قطعات — مرتب‌شده بر اساس بحرانی بودن"
    ws["A1"].font = P.font_title(14)
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "طبقه"
    ws["B3"] = "تعداد قطعه یکتا"
    ws["C3"] = "تعداد ردیف"
    ws["D3"] = "اقدام پیشنهادی"
    for col in "ABCD":
        ws[f"{col}3"].font = P.font_header(1)
        ws[f"{col}3"].fill = P.fill_header()
        ws[f"{col}3"].alignment = P.align("center")
    for col, w in zip("ABCDEF", (34, 18, 14, 62, 20, 20)):
        ws.column_dimensions[col].width = w

    r = 4
    for b in bands:
        sub = df[df.get("کد طبقه بحرانی", pd.Series(dtype=str)) == b["code"]]
        n_parts = int(sub["KEY_MATERIAL"].replace("", pd.NA).nunique()) if "KEY_MATERIAL" in sub else 0
        ws.cell(row=r, column=1, value=b.get("fa", b["code"])).font = P.font_body(bold=True)
        ws.cell(row=r, column=2, value=n_parts).font = P.font_kpi()
        ws.cell(row=r, column=3, value=int(len(sub))).font = P.font_body()
        c = ws.cell(row=r, column=4, value=b.get("action", ""))
        c.font = P.font_body()
        c.alignment = P.align("right", wrap=True)
        for i in range(1, 5):
            ws.cell(row=r, column=i).border = P.thin_border()
            if b.get("fill"):
                ws.cell(row=r, column=i).fill = P.fill(b["fill"])
        r += 1

    # ── جدول تفصیلی قطعات نیازمند اقدام ──
    action_bands = ["STOCKOUT", "CRITICAL", "BECOMING_CRITICAL"]
    detail = df[df.get("کد طبقه بحرانی", pd.Series(dtype=str)).isin(action_bands)] \
        if "کد طبقه بحرانی" in df.columns else df.iloc[0:0]

    start = r + 2
    ws.cell(row=start - 1, column=1,
            value=f"قطعات نیازمند اقدام فوری ({len(detail)} ردیف)").font = P.font_title(12)

    # ⚠️ این ستون‌ها باید دقیقاً کلیدهایی باشند که CriticalityResult تولید
    # می‌کند. نسخه قبل «موجودی» و «مصرف روزانه» را می‌خواند که دیگر وجود
    # نداشتند، پس خالی چاپ می‌شد و کاربر «مقاومت ۴ روز با موجودی صفر»
    # می‌دید. tests/test_contracts_report.py حالا این تطابق را می‌سنجد.
    #
    # هر چهار جزء صورت کسر نمایش داده می‌شوند تا حساب مقاومت روی کاغذ
    # قابل بازبینی باشد: (ایران‌خودرو + ساپکو + در راه + در گمرک) ÷ نیاز روزانه
    headers = ["طبقه بحرانی", "کد متریال", "شرح کالا",
               "مقاومت (روز)", "مقاومت انبار (روز)",
               "موجودی ایران خودرو", "موجودی ساپکو", "در راه", "در گمرک",
               "موجودی کل", "نیاز روزانه",
               "انبار", "شماره سفارش", "شماره بارنامه",
               "مرحله سفارش", "پیشرفت (٪)", "روزهای رسوب", "طبقه ریسک",
               # ── دو ستون، عمداً جدا ────────────────────────────────────
               # یک ستون به اسم «کارشناس» یعنی خواننده باید حدس بزند این
               # نام مالک قطعه است یا کسی که اتفاقاً مرحله فعلی دستش است.
               # در گزارش بحرانی این حدس گران تمام می‌شود، چون بر مبنایش
               # به کسی تلفن می‌زنند. پس هر دو نوشته می‌شوند و هرکدام
               # اسم خودش را دارد.
               "مالک قطعه (کارشناس خرید)", "کارشناس مالک مرحله فعلی",
               "هشدار ترکیبی", "اقدام لازم"]
    widths = [26, 16, 30, 14, 16, 15, 13, 11, 11, 13, 12, 12, 16, 20, 20, 12, 14, 16, 24, 22, 46, 52]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=start, column=i, value=h)
        c.font = P.font_header(1)
        c.fill = P.fill_header()
        c.alignment = P.align("center", wrap=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[start].height = 32
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start}"

    cols = ["طبقه بحرانی", "KEY_MATERIAL", "CANONICAL_GOODS_DESC",
            "مقاومت (روز)", "مقاومت انبار (روز)",
            "موجودی ایران خودرو", "موجودی ساپکو", "موجودی در راه", "موجودی در گمرک",
            "موجودی کل قابل احتساب", "نیاز روزانه",
            "WAREHOUSE", "CANONICAL_ORDER", "CANONICAL_BL",
            "ORDER_STAGE_FA", "ORDER_PROGRESS", "روزهای رسوب", "طبقه ریسک",
            "PART_OWNER", "CANONICAL_EXPERT",
            "هشدار ترکیبی بحرانی", "اقدام هشدار ترکیبی"]

    rr = start + 1
    for _, row in detail.iterrows():
        for i, key in enumerate(cols, start=1):
            value = row.get(key, "")
            # مالک قطعه fail-safe است: اگر از سورس خرید قابل انتساب نباشد،
            # هیچ نامی از HR/ترخیص/اعتبارات جایش گذاشته نمی‌شود.
            if key == "PART_OWNER" and not str(value or "").strip():
                state = str(row.get("COMMERCIAL_COVERAGE_STATE", "") or "")
                value = ("— (سنجیده نشد: سورس خرید در دسترس نبود)"
                         if state and state != "measured"
                         else "— (در Commercial Expert Data درج نشده)")
            ws.cell(row=rr, column=i, value=self._cell_value(value))
        self._style_row(ws, rr, len(headers),
                        fills.get(str(row.get("کد طبقه بحرانی", "")), P.GREEN_L4))
        for i in (4, 5, 6, 7, 8, 9, 10, 11, 16, 17):
            ws.cell(row=rr, column=i).number_format = NUM_FORMAT_PCT
        for i in (20, 21):
            ws.cell(row=rr, column=i).alignment = P.align("right", wrap=True)
        rr += 1

    if rr > start + 1:
        ws.conditional_formatting.add(
            f"D{start + 1}:D{rr - 1}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F5B7B1",
                           mid_type="num", mid_value=10, mid_color="FDEBD0",
                           end_type="num", end_value=20, end_color="FCF3CF"))
        ws.cell(row=rr + 1, column=3, value="جمع قطعات نیازمند اقدام:").font = P.font_body(bold=True)
        ws.cell(row=rr + 1, column=4,
                value=f"=SUBTOTAL(103,B{start + 1}:B{rr - 1})").font = P.font_kpi()

    log.info(f"📄 شیت «{SHEET_CRITICAL}» ساخته شد — {len(detail)} قطعه نیازمند اقدام.")


ExcelDashboardBuilder.build_criticality = _build_criticality


SHEET_PROCESS = "۱۰. نقشه فرآیند و گلوگاه"


def _build_process(self, extras: dict, stage_map: "pd.DataFrame") -> None:
    """شیت فرآیندی — نگاه علّی به‌جای نگاه وضعیتی.

    سه بخش: نقشه مرحله‌ها (ورودی/خروجی هر گام)، گلوگاه‌ها (طولانی‌ترین فاصله
    میان دو فعالیت متوالی)، و مسیرهای فرآیند (variants) با سهم هرکدام.
    """
    ws = self._new_sheet(SHEET_PROCESS)
    for col, w in zip("ABCDEF", (34, 34, 18, 16, 16, 60)):
        ws.column_dimensions[col].width = w

    ws["A1"] = "🔄 تحلیل فرآیندی — کجا زمان از دست می‌رود"
    ws["A1"].font = P.font_title(14)
    ws.merge_cells("A1:F1")
    r = 3

    def section(title: str, nonlocal_r: int) -> int:
        c = ws.cell(row=nonlocal_r, column=1, value=title)
        c.font = P.font_title(12)
        c.fill = P.fill_row_level(1)
        return nonlocal_r + 1

    def table(df: "pd.DataFrame", start: int, highlight_first: bool = False) -> int:
        if df is None or df.empty:
            ws.cell(row=start, column=1, value="داده‌ای موجود نیست.").font = P.font_body()
            return start + 2
        for i, h in enumerate(df.columns, start=1):
            c = ws.cell(row=start, column=i, value=str(h))
            c.font = P.font_header(1)
            c.fill = P.fill_header()
            c.alignment = P.align("center", wrap=True)
        rr = start + 1
        for _, row in df.iterrows():
            for i, h in enumerate(df.columns, start=1):
                cell = ws.cell(row=rr, column=i, value=self._cell_value(row[h]))
                cell.font = P.font_body()
                cell.border = P.thin_border()
                cell.alignment = P.align("right", wrap=True)
                cell.fill = P.fill(
                    P.CRITICAL_FILL if (highlight_first and rr == start + 1) else P.GREEN_L4)
            rr += 1
        return rr + 2

    # ── ۱) گلوگاه‌ها ──
    r = section("۱) گلوگاه‌های فرآیند — میانگین فاصله میان دو فعالیت متوالی", r)
    bn = extras.get("bottlenecks")
    r = table(bn, r, highlight_first=True)

    # ── ۲) مسیرهای فرآیند ──
    r = section("۲) مسیرهای طی‌شده (Variants) — هرچه بیشتر، فرآیند بی‌انضباط‌تر", r)
    var = extras.get("variants")
    if var is not None and not var.empty:
        var = var.head(20)
    r = table(var, r)

    # ── ۳) نقشه مرحله‌ها ──
    r = section("۳) نقشه علّی مرحله‌ها — ورودی و خروجی هر گام", r)
    r = table(stage_map, r)

    ev = extras.get("eventlog")
    if ev is not None and not ev.empty:
        ws.cell(row=r, column=1,
                value=f"جدول فعالیت: {len(ev)} رویداد — "
                      f"فایل AIBL_EventLog.csv برای بارگذاری در Celonis ذخیره شد "
                      f"(_CASE_KEY / ACTIVITY_EN / EVENTTIME / _SORTING).").font = P.font_body(bold=True)
    log.info(f"📄 شیت «{SHEET_PROCESS}» ساخته شد.")


ExcelDashboardBuilder.build_process = _build_process


# ═══════════════════════════════════════════════════════════════════════════
#  شیت ۱۱ — نمودارهای تحلیلی (پیاده‌سازی در report/charts.py)
# ═══════════════════════════════════════════════════════════════════════════
from .charts import _build_charts as _charts_impl  # noqa: E402

ExcelDashboardBuilder.build_charts = _charts_impl

from .insight import _build_insight as _insight_impl  # noqa: E402
from .insight import _build_material as _material_impl  # noqa: E402
from .supply_views import SHEETS as SUPPLY_SHEETS  # noqa: E402
from .supply_views import write_supply_sheets  # noqa: E402

ExcelDashboardBuilder.build_insight = _insight_impl
ExcelDashboardBuilder.build_material = _material_impl


def _build_supply_views(self, df) -> None:
    """شیت‌های ۱۴ تا ۱۶ — «کجا / کِی / دست کیست».

    نسخه ۲۶٫۹ این کار را با وصله زدن به ``save()`` انجام می‌داد: یک کپی
    کامل از دیتافریم روی builder نگه می‌داشت و هنگام ذخیره نماها را
    می‌ساخت. دو اشکال داشت — کپی کامل داده در حافظه، و مهم‌تر اینکه
    **هر خطایی در این سه نما، ذخیره کل گزارش رسمی را از بین می‌برد**.
    حالا مرحله‌ای صریح در خط لوله است و شکستش گزارش را زمین نمی‌زند.
    """
    try:
        write_supply_sheets(self.wb, df)
        log.info(f"📄 شیت‌های «{'» و «'.join(SUPPLY_SHEETS)}» ساخته شد.")
    except Exception as ex:      # noqa: BLE001 — گزارش رسمی نباید قربانی شود
        log.warning(f"⚠️ نماهای تأمین ساخته نشد ({type(ex).__name__}: {ex}) — "
                    f"بقیه گزارش دست‌نخورده ذخیره می‌شود.")


ExcelDashboardBuilder.build_supply_views = _build_supply_views


def _build_system_health(self) -> None:
    """شیت ۱۷ — مبنای اعتماد به بقیه شیت‌ها.

    مثل نماهای تأمین، شکستِ این شیت نباید گزارش رسمی را زمین بزند.
    """
    try:
        from .. import health as _h
        from .system_health import SHEET as _SH
        from .system_health import build as _build
        _build(self.wb)
        # دفتر کنار گزارش هم نوشته می‌شود: ایمیل و داشبورد Streamlit در
        # فرآیند دیگری اجرا می‌شوند و دفترِ درون‌حافظه‌ای را نمی‌بینند.
        _h.save(os.path.dirname(self.output_path) or ".")
        log.info(f"📄 شیت «{_SH}» ساخته شد.")
    except Exception as ex:      # noqa: BLE001
        log.warning(f"⚠️ شیت سلامت سیستم ساخته نشد ({type(ex).__name__}: {ex}).")


ExcelDashboardBuilder.build_system_health = _build_system_health
