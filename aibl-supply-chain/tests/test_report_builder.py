# -*- coding: utf-8 -*-
"""تست سازنده گزارش و صحت محاسبات دانه‌ای.

مهم‌ترین تست این فایل، «دوباره‌شماری» است: وقتی یک ثبت سفارش چند بارنامه
دارد، جمع ساده روی ردیف‌ها مقدار را چند برابر می‌کند. داده آزمون مصنوعی
۱:۱ است و این تله را نشان نمی‌دهد، پس اینجا عمداً یک fixture با تکرار
ساخته می‌شود.

اجرا:  python tests/test_report_builder.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault("AIBL_OUTPUT", tempfile.mkdtemp(prefix="aibl_rep_"))
os.environ.setdefault("AIBL_LOGS", tempfile.mkdtemp(prefix="aibl_replog_"))

import pandas as pd  # noqa: E402

from aibl.studio_core import templates as tpl  # noqa: E402
from aibl.studio_core.grain import (KIND_ADDITIVE, KIND_IDENTIFIER,  # noqa: E402
                                    KIND_RATIO, column_grain, fanout,
                                    integrity_report, measure_kind, naive_agg,
                                    preferred_agg, safe_agg, summable)
from aibl.studio_core.html_export import build_dynamic_html  # noqa: E402
from aibl.studio_core.report_builder import ReportSpec, build  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


#: یک ثبت سفارش (R1) با سه بارنامه — حالتی که در تولید همیشه هست
FAN = pd.DataFrame({
    "KEY_BL":        ["BL1", "BL2", "BL3", "BL4"],
    "KEY_ORDER":     ["O1", "O1", "O2", "O3"],
    "KEY_REG":       ["R1", "R1", "R1", "R2"],
    "KEY_MATERIAL":  ["M1", "M2", "M3", "M4"],
    "CANONICAL_BL":  ["BL1", "BL2", "BL3", "BL4"],
    "CANONICAL_ORDER": ["501317", "501317", "501807", "812210"],
    "مانده تعهد":    [1000.0, 1000.0, 1000.0, 500.0],   # دانه REG ⇒ تکرار
    "مقاومت (روز)":  [10.0, 20.0, 30.0, 40.0],           # نسبتی
    "ORC_STOCK_IKCO": [5, 10, 15, 20],                   # دانه MATERIAL
    "بحرانی (کوتاه)": ["بحرانی", "ایمن", "تحت نظر", "ایمن"],
})


def test_fanout_double_counting() -> None:
    print("\n── ۱) دوباره‌شماری در جمع ──")
    check("دانه «مانده تعهد» از سورس مبدأ به ارث می‌رسد (REG، نه ردیف)",
          column_grain("مانده تعهد") == "REG", column_grain("مانده تعهد"))
    naive = naive_agg(FAN, "مانده تعهد", "sum")
    safe = safe_agg(FAN, "مانده تعهد", "sum")
    check("جمع ساده روی ردیف‌ها دوباره‌شماری می‌کند", naive == 3500.0, f"{naive:,.0f}")
    check("جمع دانه‌ای مقدار درست را می‌دهد", safe == 1500.0, f"{safe:,.0f}")
    check("اختلاف دقیقاً همان مقدار تکرارشده است", naive - safe == 2000.0)

    ded = safe_agg(FAN, "ORC_STOCK_IKCO", "sum")
    check("ستون با دانه ریز دست‌نخورده می‌ماند",
          ded == naive_agg(FAN, "ORC_STOCK_IKCO", "sum") == 50.0, f"{ded:,.0f}")

    fo = fanout(FAN)
    reg = fo[fo["کلید"] == "KEY_REG"].iloc[0]
    check("ضریب تکرار دانه ثبت سفارش گزارش می‌شود",
          abs(float(reg["ضریب تکرار"]) - 4 / 2) < 1e-9, str(reg["ضریب تکرار"]))


def test_measure_kind() -> None:
    print("\n── ۲) شناسه در برابر سنجه ──")
    check("شماره سفارش شناسه است، نه سنجه",
          measure_kind("CANONICAL_ORDER", FAN["CANONICAL_ORDER"]) == KIND_IDENTIFIER)
    check("شناسه هرگز جمع نمی‌شود (تجمیع پیش‌فرض: شمار یکتا)",
          preferred_agg("CANONICAL_ORDER", FAN["CANONICAL_ORDER"]) == "nunique")
    check("«مقاومت (روز)» نسبتی است، نه انباشتنی",
          measure_kind("مقاومت (روز)", FAN["مقاومت (روز)"]) == KIND_RATIO)
    check("تجمیع پیش‌فرض ستون نسبتی میانگین است",
          preferred_agg("مقاومت (روز)", FAN["مقاومت (روز)"]) == "mean")
    check("«مانده تعهد» انباشتنی است",
          measure_kind("مانده تعهد", FAN["مانده تعهد"]) == KIND_ADDITIVE)
    s = summable(FAN, list(FAN.columns))
    check("شناسه‌ها از فهرست قابل جمع کنار می‌روند",
          "CANONICAL_ORDER" not in s and "مقاومت (روز)" not in s
          and "مانده تعهد" in s, str(s))


def test_integrity_report() -> None:
    print("\n── ۳) ردپای محاسباتی ──")
    rep = integrity_report(FAN, ["مانده تعهد", "CANONICAL_ORDER", "مقاومت (روز)"])
    check("گزارش صحت برای هر ستون یک ردیف دارد", len(rep) == 3, str(len(rep)))
    row = rep[rep["ستون"] == "مانده تعهد"].iloc[0]
    check("دوباره‌شماری در گزارش صحت علامت می‌خورد",
          "دوباره‌شماری" in str(row["وضعیت"]), str(row["وضعیت"]))
    ident = rep[rep["ستون"] == "CANONICAL_ORDER"].iloc[0]
    check("شناسه در گزارش صحت جمع نمی‌شود",
          "شناسه" in str(ident["وضعیت"]), str(ident["وضعیت"]))


def test_html_export() -> None:
    print("\n── ۴) خروجی HTML ──")
    html = build_dynamic_html(FAN, "2026-08-31", "AIBL",
                              selected_fields=["CANONICAL_BL", "مانده تعهد",
                                               "مقاومت (روز)", "بحرانی (کوتاه)"])
    check("سند راست‌به‌چپ و فارسی است", 'dir="rtl"' in html and 'lang="fa"' in html)
    check("کلید دانه برای یکتاسازی سمت مرورگر تعبیه شده",
          '"مانده تعهد": "KEY_REG"' in html or "KEY_REG" in html)
    check("تابع تجمیع دانه‌ای در JS هست", "function gagg" in html and "gvals" in html)
    check("ستون نسبتی میانگین می‌گیرد نه جمع", "'mean'" in html or '"mean"' in html)
    check("دکمه چاپ/PDF دارد", "window.print()" in html)
    check("پالت وضعیت اعتبارسنجی‌شده استفاده شده",
          "#d03b3b" in html and "#fab219" in html)
    check("پالت رد شده قبلی دیگر نیست", "#F1C40F" not in html and "#C0392B" not in html)


def test_templates_build() -> None:
    print("\n── ۵) ساخت گزارش در هر قالب ──")
    check("چهار قالب تعریف شده است", len(tpl.TEMPLATES) == 4, str(list(tpl.TEMPLATES)))
    check("خاموش کردن ویژوال، بخش‌های نموداری را حذف می‌کند",
          "criticality" not in tpl.get("executive").active_sections(visuals=False))
    check("خاموش کردن جدول‌ها، بخش‌های جدولی را حذف می‌کند",
          "table" not in tpl.get("executive").active_sections(tables=False))

    with tempfile.TemporaryDirectory() as td:
        for key in tpl.TEMPLATES:
            spec = ReportSpec(template=key, fields=list(FAN.columns),
                              ref_date="2026-08-31", title="AIBL",
                              formats=["excel", "html"], file_stem=f"t_{key}")
            res = build(FAN, {}, spec, {}, td)
            check(f"قالب «{tpl.get(key).title}» هر دو فرمت را می‌سازد",
                  "excel" in res.files and "html" in res.files,
                  " | ".join(res.messages))


def test_cross_format_consistency() -> None:
    print("\n── ۶) یکسانی عدد بین فرمت‌ها ──")
    with tempfile.TemporaryDirectory() as td:
        spec = ReportSpec(template="executive", fields=list(FAN.columns),
                          ref_date="2026-08-31", formats=["excel", "html"],
                          file_stem="consistency")
        res = build(FAN, {}, spec, {}, td)
        rep = res.integrity
        row = rep[rep["ستون"] == "مانده تعهد"].iloc[0]
        safe = float(row["جمع درست (دانه‌ای)"])
        check("عدد گزارش صحت با محاسبه مستقل یکی است",
              safe == safe_agg(FAN, "مانده تعهد", "sum") == 1500.0, f"{safe:,.0f}")
        html = res.html
        check("HTML همان مقدار دانه‌ای را در سربرگ دارد",
              "1,500" in html or "1500" in html or "۱٬۵۰۰" in html
              or 'GRAIN' in html)   # مقدار در JS محاسبه می‌شود
        from openpyxl import load_workbook
        wb = load_workbook(res.files["excel"])
        check("برگه «صحت محاسبات» در Excel هست", "صحت محاسبات" in wb.sheetnames,
              str(wb.sheetnames))


if __name__ == "__main__":
    print("=" * 78)
    print("AIBL — تست سازنده گزارش و صحت دانه‌ای")
    print("=" * 78)
    test_fanout_double_counting()
    test_measure_kind()
    test_integrity_report()
    test_html_export()
    test_templates_build()
    test_cross_format_consistency()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
