# -*- coding: utf-8 -*-
"""تست نقاط کور سیستمی: fail-closed خواندن، سلامت سورس/مرحله/ادغام،
ناسازگاری زمانی، نسب متریال، و حالت سوم پوشش.

اجرا:  python tests/test_system_health.py
"""
from __future__ import annotations

import ast
import os
import subprocess
import sys
import tempfile
import warnings

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from aibl import health  # noqa: E402
from aibl.dataio.reader import read_sheet  # noqa: E402
from aibl.report import system_health as sh  # noqa: E402
from aibl.resolve import commercial_coverage as cc  # noqa: E402
from aibl.resolve import part_status as ps  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


# ═══════════ ۱) خواندن fail-closed ═══════════
def test_reader_fail_closed() -> None:
    print("\n── ۱) شیتِ نبوده هرگز با شیت دیگر جایگزین نمی‌شود ──")
    health.reset()
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "sample.xlsx")
        with pd.ExcelWriter(path, engine="openpyxl") as w:
            pd.DataFrame({"A": [1, 2]}).to_excel(w, index=False, sheet_name="Sheet1")
            pd.DataFrame({"B": [3]}).to_excel(w, index=False, sheet_name="Other")

        got = read_sheet(path, "Other", "t")
        check("شیت موجود خوانده می‌شود", got is not None and "B" in got.columns)

        got = read_sheet(path, "  other ", "t")
        check("تطبیق بدون حساسیت به حروف/فاصله کار می‌کند",
              got is not None and "B" in got.columns)

        # ── قلب موضوع ──
        got = read_sheet(path, "Expert Data", "t")
        check("شیت ناموجود ⇒ None، نه fallback به شیت اول",
              got is None, "None" if got is None else f"برگشت {list(got.columns)}")
        rec = health.current().sources.get("t")
        check("شکاف اسکیما در سلامت سیستم ثبت می‌شود",
              bool(rec and rec.schema_gaps), str(rec.schema_gaps if rec else None))
        check("وضعیت سورس «ناقص» می‌شود، نه سالم",
              bool(rec and rec.status == health.DEGRADED),
              rec.status if rec else "—")

        # sheet=None یعنی «شیت اول، هرچه باشد» — این عمدی است و باید بماند
        got = read_sheet(path, None, "t")
        check("sheet=None همچنان شیت اول را می‌خواند (رفتار عمدی)",
              got is not None and "A" in got.columns)


# ═══════════ ۲) دفتر سلامت ═══════════
def test_health_registry() -> None:
    print("\n── ۲) دفتر سلامت سیستم ──")
    h = health.reset()
    h.source("a", title="سورس الف", required=True, status=health.OK, rows=10, frames=1)
    h.source("b", title="سورس ب", required=False, status=health.SKIPPED)
    check("اجرای با سورس اختیاریِ ردشده «ناقص» است، نه خراب",
          h.verdict() == health.DEGRADED, h.verdict())
    check("سورس اختیاری مسدودکننده نیست", h.blocking() == [], str(h.blocking()))

    h.source("c", title="سورس ج", required=True, status=health.FAILED)
    check("سورس الزامیِ خراب، مسدودکننده است", h.blocking() == ["c"], str(h.blocking()))
    check("وضعیت کلی «خراب» می‌شود", h.verdict() == health.FAILED, h.verdict())

    h = health.reset()
    h.source("a", required=True, status=health.OK)
    check("همه سالم ⇒ وضعیت سالم", h.verdict() == health.OK, h.verdict())

    raised = False
    try:
        with h.timed_stage("boom", "تست", 10, rows_in=5):
            raise ValueError("خرابی عمدی")
    except ValueError:
        raised = True
    check("استثنای مرحله بلعیده نمی‌شود", raised)
    check("مرحله شکست‌خورده ثبت می‌شود",
          any(s.name == "boom" and s.status == health.FAILED for s in h.stages))
    check("خطای مرحله وضعیت کلی را خراب می‌کند", h.verdict() == health.FAILED)

    for name, t in (("سورس", h.sources_table()), ("مرحله", h.stages_table()),
                    ("ادغام", h.joins_table()), ("یافته", h.findings_table())):
        check(f"جدول {name} ساخته می‌شود", isinstance(t, pd.DataFrame))


# ═══════════ ۳) شیت ۱۷ ═══════════
def test_health_sheet() -> None:
    print("\n── ۳) شیت «۱۷. سلامت سیستم» ──")
    h = health.reset()
    h.source("moghavemat", title="سورس خرید", required=False, status=health.DEGRADED)
    h.schema_gap("moghavemat", "شیت «Expert Data» نبود")
    h.join(health.JoinHealth(label="x", key="KEY_ORDER", rows_before=10,
                             rows_after=10, matched=0, status=health.DEGRADED))
    wb = Workbook()
    wb.remove(wb.active)
    sh.build(wb, h)
    check("شیت ساخته می‌شود", sh.SHEET in wb.sheetnames, str(wb.sheetnames))
    ws = wb[sh.SHEET]
    body = " ".join(str(ws.cell(r, c).value or "")
                    for r in range(1, ws.max_row + 1) for c in range(1, 11))
    for probe in ("وضعیت کلی این اجرا", "سلامت سورس", "سلامت مرحله",
                  "سلامت ادغام", "یافته", "Expert Data"):
        check(f"بخش «{probe}» در شیت هست", probe in body)
    check("خلاصه کوتاه برای ایمیل ساخته می‌شود", len(sh.summary_rows(h)) >= 2)

    h2 = health.reset()
    h2.source("a", required=True, status=health.OK)
    wb2 = Workbook()
    sh.build(wb2, h2)
    check("اجرای سالم هم شیت می‌سازد (بدون استثنا)", sh.SHEET in wb2.sheetnames)


# ═══════════ ۴) ناسازگاری زمانی ═══════════
def test_timeline_anomaly() -> None:
    print("\n── ۴) ترتیب زمانی معکوس و تاریخ برنامه‌ای ──")
    health.reset()
    d = pd.DataFrame({
        "COT_DATE":        ["1405/06/10", "1405/01/01", "1405/01/01"],
        "FULL_CLEAR_DATE": ["1405/06/03", "1405/02/02", ""],
        "BL_DATE":         ["1404/12/01", "1404/12/01", "1410/01/01"],
    })
    out = ps.resolve(d, today=pd.Timestamp("2026-09-08"))
    check("ترتیب معکوس شناسایی می‌شود",
          "پیش از" in str(out[ps.ANOMALY].iloc[0]), out[ps.ANOMALY].iloc[0])
    check("ردیف سالم علامت نمی‌خورد",
          out[ps.ANOMALY].iloc[1] == "", out[ps.ANOMALY].iloc[1])
    check("تاریخ آینده به‌عنوان «برنامه‌ای» ثبت می‌شود",
          out[ps.PLANNED].iloc[2] != "", out[ps.PLANNED].iloc[2])
    check("تاریخ آینده وضعیت فعلی را تعیین نمی‌کند",
          "1410" not in str(out[ps.WHEN].iloc[2]), str(out[ps.WHEN].iloc[2]))
    msgs = [f.message for f in health.current().findings]
    check("ناسازگاری زمانی وارد سلامت سیستم می‌شود",
          any("ترتیب زمانی" in m for m in msgs), str(msgs))
    check("تاریخ برنامه‌ای هم گزارش می‌شود",
          any("برنامه‌ای" in m for m in msgs), str(msgs))
    check("ستون‌های تازه در فهرست خروجی اعلام شده‌اند",
          ps.ANOMALY in ps.OUTPUT_COLUMNS and ps.PLANNED in ps.OUTPUT_COLUMNS)


# ═══════════ ۵) نسب متریال ═══════════
def test_material_lineage() -> None:
    print("\n── ۵) سفارش چندمتریاله پنهان نمی‌ماند ──")
    health.reset()
    from aibl.adapters.moghavemat import MoghavematAdapter as M
    got = M._uniq_values(pd.Series(["a", "", None, "a", " b "]))
    check("تابع یکتاسازی مقادیر خالی را دور می‌ریزد", got == ["a", "b"], str(got))
    check("خروجی مرتب است (نسب پایدار، نه وابسته به ترتیب ردیف)",
          M._uniq_values(pd.Series(["z", "a"])) == ["a", "z"])


# ═══════════ ۶) حالت سوم پوشش ═══════════
def test_coverage_states() -> None:
    print("\n── ۶) سه حالت پوشش سورس خرید ──")
    d = pd.DataFrame({"CANONICAL_ORDER": ["O1", "O2"]})

    health.reset()
    out, n = cc.annotate(d, pd.DataFrame({"KEY_ORDER": ["O1"]}))
    check("سنجیده‌شده: حالت measured", out[cc.STATE].iloc[0] == cc.MEASURED)
    check("سنجیده‌شده: شمارش درست", n == 1, str(n))

    health.reset()
    out, n = cc.annotate(d, None)
    check("فایل نبود ⇒ source_unavailable",
          out[cc.STATE].iloc[0] == cc.UNAVAILABLE, out[cc.STATE].iloc[0])
    check("فایل نبود ⇒ هیچ سفارشی پرچم نمی‌گیرد", n == 0 and not out[cc.FLAG].any())
    check("KPI «سنجیده نشد» می‌دهد", cc.kpi(out, n)[0] == "سنجیده نشد")

    health.reset()
    health.current().schema_gap("moghavemat", "شیت «Expert Data» نبود")
    out, n = cc.annotate(d, None)
    check("ساختار عوض شد ⇒ source_schema_gap",
          out[cc.STATE].iloc[0] == cc.SCHEMA_GAP, out[cc.STATE].iloc[0])
    check("پیام KPI بین «فایل نبود» و «ساختار عوض شد» فرق می‌گذارد",
          "ساختار" in cc.kpi(out, n)[1], cc.kpi(out, n)[1])
    check("هیچ‌کدام از دو حالتِ نسنجیده، measured نیست", cc.measured(out) is False)


# ═══════════ ۷) بهداشت اجرا ═══════════
def test_runtime_hygiene() -> None:
    print("\n── ۷) بهداشت اجرا ──")
    bad = []
    for base, dirs, files in os.walk(os.path.join(ROOT, "aibl")):
        dirs[:] = [x for x in dirs if x != "__pycache__"]
        for f in files:
            if not f.endswith(".py"):
                continue
            full = os.path.join(base, f)
            with warnings.catch_warnings(record=True) as w:
                warnings.simplefilter("always")
                compile(open(full, encoding="utf-8").read(), full, "exec")
                for x in w:
                    if issubclass(x.category, (DeprecationWarning, SyntaxWarning)):
                        bad.append(f"{os.path.relpath(full, ROOT)}:{x.lineno} {x.message}")
    # «\o» در رشته غیرخام، در پایتون‌های بعدی SyntaxError می‌شود.
    check("هیچ escape نامعتبری در سورس نیست", not bad, str(bad[:3]))

    out = subprocess.run([sys.executable, "-W", "error::RuntimeWarning",
                          "-m", "aibl.factsheet"], cwd=ROOT,
                         capture_output=True, text=True, timeout=120)
    check("«python -m aibl.factsheet» بدون RuntimeWarning اجرا می‌شود",
          "RuntimeWarning" not in (out.stderr or ""), (out.stderr or "")[:120])

    src = open(os.path.join(ROOT, "run_all_tests.py"), encoding="utf-8").read()
    check("هر مجموعه تست سقف زمانی دارد",
          "timeout=" in src and "TimeoutExpired" in src)
    check("خروجی TIMEOUT صریح گزارش می‌شود", "TIMEOUT" in src)

    tree = ast.parse(open(os.path.join(ROOT, "aibl", "__init__.py"),
                          encoding="utf-8").read())
    lazy = any(isinstance(n, ast.FunctionDef) and n.name == "__getattr__"
               for n in tree.body)
    check("نسخه پکیج تنبل خوانده می‌شود", lazy)


if __name__ == "__main__":
    print("=" * 78)
    print("AIBL — نقاط کور سیستمی و فرآیندی")
    print("=" * 78)
    test_reader_fail_closed()
    test_health_registry()
    test_health_sheet()
    test_timeline_anomaly()
    test_material_lineage()
    test_coverage_states()
    test_runtime_hygiene()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
