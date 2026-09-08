# -*- coding: utf-8 -*-
"""تست‌های کتابخانه قوانین، سورس مقاومت جدید و تضمین ماژولاریتی.

اجرا:  python tests/test_rules_and_moghavemat.py
"""
from __future__ import annotations

import os
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

# ── بارگذاری مطمئن make_synthetic بدون اتکا به sys.path ──
# روی ویندوز و در اجرای subprocess، پوشه tests همیشه در sys.path نیست.
_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)


def _load_make_synthetic():
    import importlib.util
    path = os.path.join(_TESTS_DIR, "make_synthetic.py")
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"فایل make_synthetic.py کنار تست‌ها نیست: {path}\n"
            f"پوشه tests/ را کامل کپی کنید.")
    spec = importlib.util.spec_from_file_location("aibl_make_synthetic", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_ms = _load_make_synthetic()

build = _ms.build

_TMP = tempfile.mkdtemp(prefix="aibl_rules_")
DIRS = build(_TMP)
os.environ.update({
    "AIBL_FOREIGN": DIRS["foreign"], "AIBL_BLS": DIRS["bls"],
    "AIBL_CLEARANCE": DIRS["clearance"], "AIBL_HR": DIRS["hr"],
    "AIBL_ESMAEILI": DIRS["esmaeili"],
    "AIBL_GS_COMBINE": DIRS["gs_combine"], "AIBL_MOHAMADI": DIRS["mohamadi"],
    "AIBL_OUTPUT": DIRS["output"], "AIBL_LOGS": DIRS["logs"],
    "AIBL_TODAY": "2026-08-31",
})

import logging  # noqa: E402
import shutil  # noqa: E402
from datetime import date  # noqa: E402

import yaml  # noqa: E402

from aibl.adapters import discover  # noqa: E402
from aibl.config.sources import MERGE_ORDER, SOURCES, get_source  # noqa: E402
from aibl.core.text import clean_order_ref, order_ref_base  # noqa: E402
from aibl.rulebook import RuleBook, get_rulebook  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


# ═══════════ ۱) کتابخانه قوانین ═══════════
def test_rulebook() -> None:
    print("\n── ۱) کتابخانه قوانین (YAML) ──")
    rb = get_rulebook(reload=True)

    check("هر ۹ بسته قانونی بارگذاری شد", len(rb.packs) == 9, ", ".join(rb.packs))
    errors = [i for i in rb.validate() if i.level == "error"]
    check("اعتبارسنجی ساختاری بدون خطا", not errors,
          "؛ ".join(f"{e.path}: {e.message}" for e in errors) or "۰ خطا")
    check("مجموع وزن‌های ریسک دقیقاً ۱٫۰ است",
          abs(sum(rb.risk_weights().values()) - 1.0) < 1e-9,
          f"{sum(rb.risk_weights().values())}")
    check("قواعد نیازمند تطبیق علامت‌گذاری شده‌اند", len(rb.needs_verification()) > 0,
          f"{len(rb.needs_verification())} قاعده")

    check("مهلت رفع تعهد تولیدی/بازرگانی از YAML",
          rb.release_deadline_days("production") == 540
          and rb.release_deadline_days("commercial") == 360)
    check("نرمال‌سازی ارز از currencies.yaml",
          rb.normalize_currency("یورو") == "EUR"
          and rb.normalize_currency("درهم امارات") == "AED"
          and rb.normalize_currency("Yuan") == "CNY")
    check("ین ژاپن بدون اعشار تعریف شده", rb.currency_minor_units("JPY") == 0)

    cip = rb.incoterm("CIP")
    cif = rb.incoterm("CIF")
    check("Incoterms 2020: CIP بیمه ICC(A) و CIF بیمه ICC(C)",
          "(A)" in cip["insurance_minimum"] and "(C)" in cif["insurance_minimum"],
          f"CIP → {cip['insurance_minimum'][:30]}")
    check("نام قدیمی DAT به DPU نگاشت می‌شود", rb.incoterm("DAT")["code"] == "DPU")
    check("EXW ترخیص صادراتی بر عهده خریدار است",
          rb.incoterm("EXW")["export_clearance_by"] == "buyer")

    check("ویرایش جاری HS و ویرایش بعدی ثبت شده",
          rb.get("hs_codes.current_edition") == "HS 2022"
          and rb.get("hs_codes.next_edition_effective") == "2028-01-01",
          f"{rb.get('hs_codes.current_edition')} → {rb.get('hs_codes.next_edition')}")
    code, kw = rb.infer_hs("تیغچه الماسه Sandvik")
    check("استنتاج تعرفه از شرح کالا", code == "8209", f"«{kw}» → {code}")
    code2, _ = rb.infer_hs("مته الماسه Guhring")
    check("استنتاج تعرفه برای مته", code2 == "8207", code2)

    check("سال مالی از YAML خوانده می‌شود",
          rb.fiscal_year_start() == date(2026, 3, 21), str(rb.fiscal_year_start()))
    check("آستانه رسوب شدید از YAML", rb.demurrage_critical_days() == 45)
    check("تشخیص سگمنت از کلیدواژه",
          rb.detect_segment("واحد تولیدی") == "production"
          and rb.detect_segment("بازرگانی") == "commercial")


# ═══════════ ۲) اعتبارسنجی بارنامه ═══════════
def test_bl_validation() -> None:
    print("\n── ۲) اعتبارسنجی بارنامه (قرنطینه مقادیر جعلی) ──")
    rb = get_rulebook()
    cases = {
        "MSCU1234567": True,   # پیشوند SCAC معتبر
        "COSU9998887": True,
        "541339": False,       # شماره فنی Walter
        "2036866948": False,   # شماره فنی Bosch
        "603111/1": False,     # مرجع داخلی
        "KBL5001350": True,    # الگوی استاندارد
        "": False,
        "AB12": False,
    }
    ok = True
    details = []
    for value, expected in cases.items():
        got, reason = rb.validate_bl(value)
        if got != expected:
            ok = False
            details.append(f"{value!r}: انتظار {expected} ولی {got} ({reason})")
    check("تفکیک بارنامه واقعی از شماره فنی/پروفرما", ok,
          "؛ ".join(details) or "همه ۸ حالت درست")


# ═══════════ ۳) واژگان وضعیت ═══════════
def test_status_lexicon() -> None:
    print("\n── ۳) واژگان وضعیت فارسی ──")
    rb = get_rulebook()

    a = rb.parse_status_note("ترخیص درصدی انجام شد//در انتظار خرید ارز")
    check("تفکیک دو بخش «//»",
          a["COMMERCIAL_NOTE"] == "ترخیص درصدی انجام شد"
          and a["LOGISTICS_NOTE"] == "در انتظار خرید ارز")
    check("«در انتظار خرید ارز» ⇒ پرونده بلوکه", a["BLOCKING"] is True)
    check("«ترخیص درصدی» ⇒ راهنمای نوع ترخیص", a["CLEARANCE_HINT"] == "PARTIAL",
          a["CLEARANCE_HINT"])
    check("هشدار متنی تولید شد", bool(a["ALERTS"]), a["ALERTS"][:50])

    b = rb.parse_status_note("ابطال شد//")
    check("«ابطال شد» ⇒ خارج از KPI", b["EXCLUDED_FROM_KPI"] is True and b["TERMINAL"] is True)

    c = rb.parse_status_note("تمام//")
    check("«تمام» ⇒ مرحله پایانی با پیشرفت ۱۰۰", c["PROGRESS"] == 100 and c["TERMINAL"],
          f"progress={c['PROGRESS']}")

    d = rb.parse_status_note("در مرحله مذاکره بازرگانی//")
    check("«مذاکره بازرگانی» ⇒ مرحله NEGOTIATION", d["STAGE"] == "NEGOTIATION",
          f"{d['STAGE']} / {d['PROGRESS']}٪")

    e = rb.parse_status_note("آماده حمل کالا//در انتظار ثبت سفارش")
    check("تضاد وضعیت شناسایی می‌شود (آماده حمل ولی بدون ثبت سفارش)",
          e["BLOCKING"] and e["PROGRESS"] >= 50, f"progress={e['PROGRESS']} blocking={e['BLOCKING']}")


# ═══════════ ۴) مرجع سفارش ═══════════
def test_order_ref() -> None:
    print("\n── ۴) نرمال‌سازی مرجع سفارش ──")
    check("پسوند قلم حذف می‌شود", clean_order_ref("603128A-\nItem 1") == "603128A",
          clean_order_ref("603128A-\nItem 1"))
    check("«Items 2 & 3» حذف می‌شود",
          clean_order_ref("603130B-Items 2 & 3 ") == "603130B")
    check("ریشه عددی سفارش استخراج می‌شود",
          order_ref_base("603128A") == "603128" and order_ref_base("602164B") == "602164")
    check("پسوند '.0' اکسل حذف می‌شود", clean_order_ref("10200001.0") == "10200001")


# ═══════════ ۵) سورس مقاومت جدید ═══════════
def test_moghavemat():
    print("\n── ۵) سورس مقاومت (۳۵ ستون، سطح قلم) ──")
    logging.disable(logging.INFO)
    from aibl.pipeline import Pipeline

    p = Pipeline()
    res = p.run(build_report=True)
    lines = res.mogh_lines
    df = res.df

    check("۳۵ ستون خوانده و استاندارد شد", not lines.empty, f"{len(lines)} قلم")
    check("قلم‌ها در سطح سفارش تجمیع شدند (بدون تکثیر سطر)", len(df) == 6,
          f"{len(lines)} قلم → {len(df)} ردیف نهایی")

    susp = lines[lines["MOGH_BL_SUSPECT"].astype(str).str.strip() != ""]
    check("مقادیر جعلی ستون BL No. قرنطینه شدند", len(susp) == 4,
          "، ".join(susp["MOGH_BL_SUSPECT"].astype(str)))
    valid = lines[lines["MOGH_BL_NO"].astype(str).str.strip() != ""]
    check("فقط بارنامه‌های معتبر وارد کلید شدند", len(valid) == 1,
          "، ".join(valid["MOGH_BL_NO"].astype(str)))

    # مرجع سفارش چندخطی 603128A همراه با شکست خط باید نرمال شود
    check("مرجع سفارش چندخطی نرمال شد",
          any(str(x).startswith("603128A") for x in lines["MOGH_ORDER_REF"])
          or "603128A" in set(lines["KEY_ORDER"]),
          str(sorted(set(lines["KEY_ORDER"]))[:6]))

    check("درصد پیشرفت سفارش از واژگان وضعیت ساخته شد",
          df["ORDER_PROGRESS"].astype(float).max() > 0,
          f"بیشینه {df['ORDER_PROGRESS'].astype(float).max():.0f}٪")
    check("پرونده‌های بلوکه شناسایی شدند", int(df["IS_BLOCKED"].sum()) >= 1,
          f"{int(df['IS_BLOCKED'].sum())} ردیف")
    check("ستون ارز خالی به «NAN» تبدیل نشد",
          set(lines["MOGH_CURRENCY"]) <= {"EUR", "USD", "AED", "CNY", ""},
          str(sorted(set(lines["MOGH_CURRENCY"]))))
    check("تعرفه پیشنهادی برای اقلام تولید شد",
          lines["MOGH_HS_SUGGESTED"].astype(str).str.strip().ne("").sum() >= 1,
          str(sorted(set(lines["MOGH_HS_SUGGESTED"]))))
    check("کد پرسنلی از فرمت '..._GS' استخراج شد",
          "10201069" in set(lines["MOGH_KEY_EMP"]),
          str(sorted(set(lines["MOGH_KEY_EMP"]))))
    return res


# ═══════════ ۶) تضمین ماژولاریتی ═══════════
def test_modularity() -> None:
    print("\n── ۶) ماژولاریتی: حذف و افزودن سورس بدون تغییر کد ──")
    registry = discover()
    check("adapterها خودکار کشف می‌شوند", len(registry) == 13,
          f"{len(registry)} adapter")
    check("هر سورس فعال YAML یک adapter دارد",
          set(SOURCES) == set(registry),
          f"بدون adapter: {set(SOURCES) - set(registry)}")
    check("سورس بدون فایل (missmohammadi) با یک enabled:false غیرفعال شد",
          "missmohammadi" not in SOURCES,
          "تأییدشده در HEADERS_MAP: هیچ فایلی با آن الگو وجود ندارد")
    check("کلید اتصال مقاومت به «سفارش» تغییر کرد",
          get_source("moghavemat").join_on == "ORDER")
    check("ترتیب ادغام از YAML خوانده می‌شود", len(MERGE_ORDER) == 10,
          " → ".join(MERGE_ORDER[:4]) + " ...")

    # حذف یک سورس فقط با ویرایش YAML
    tmp_yaml = os.path.join(_TMP, "sources_test.yaml")
    src_yaml = os.path.join(ROOT, "aibl", "config", "sources.yaml")
    with open(src_yaml, encoding="utf-8") as f:
        data = yaml.safe_load(f)
    data["sources"]["credit"]["enabled"] = False
    data["merge_order"] = [m for m in data["merge_order"] if m != "credit"]
    with open(tmp_yaml, "w", encoding="utf-8") as f:
        yaml.safe_dump(data, f, allow_unicode=True)

    os.environ["AIBL_SOURCES_YAML"] = tmp_yaml
    import aibl.config.sources as cfg
    cfg.reload_sources()
    check("غیرفعال کردن یک سورس فقط با ویرایش YAML ممکن است",
          "credit" not in cfg.SOURCES and "credit" not in cfg.MERGE_ORDER,
          f"{len(cfg.SOURCES)} سورس فعال ماند")
    os.environ["AIBL_SOURCES_YAML"] = src_yaml
    cfg.reload_sources()
    check("بازگردانی رجیستری بدون تغییر کد", "credit" in cfg.SOURCES)

    # تغییر یک قاعده فقط با ویرایش YAML
    ext = os.path.join(_TMP, "rules_ext")
    shutil.copytree(os.path.join(ROOT, "aibl", "rules"), ext, dirs_exist_ok=True)
    path = os.path.join(ext, "fx_governance.yaml")
    with open(path, encoding="utf-8") as f:
        rules = yaml.safe_load(f)
    rules["deadlines"]["release_production"]["days"] = 480
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(rules, f, allow_unicode=True)

    rb_ext = RuleBook(rules_dir=ext)
    check("تغییر مهلت قانونی فقط با ویرایش YAML اعمال می‌شود",
          rb_ext.release_deadline_days("production") == 480,
          f"۵۴۰ → {rb_ext.release_deadline_days('production')} روز")

    # نسخه‌بندی زمانی
    rules["deadlines"]["release_commercial"] = {
        "versions": [
            {"effective_from": "2024-03-21", "value": 360},
            {"effective_from": "2026-03-21", "value": 300},
        ]}
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(rules, f, allow_unicode=True)
    old = RuleBook(rules_dir=ext, as_of=date(2025, 6, 1))
    new = RuleBook(rules_dir=ext, as_of=date(2026, 8, 31))
    check("نسخه‌بندی زمانی قواعد کار می‌کند",
          old.get("fx_governance.deadlines.release_commercial") == 360
          and new.get("fx_governance.deadlines.release_commercial") == 300,
          "۱۴۰۴ → ۳۶۰ روز | ۱۴۰۵ → ۳۰۰ روز")


def test_excel(res) -> None:
    print("\n── ۷) شیت‌های جدید اکسل ──")
    from openpyxl import load_workbook
    wb = load_workbook(res.dashboard_path)
    check("شیت «اقلام سفارش» ساخته شد",
          "۷. اقلام سفارش (سطح PR و PI)" in wb.sheetnames)
    check("شیت «کتابخانه قوانین» ساخته شد", "۸. کتابخانه قوانین" in wb.sheetnames)
    ws = wb["۸. کتابخانه قوانین"]
    values = [c.value for row in ws.iter_rows(max_col=3) for c in row if c.value]
    check("مهلت ۵۴۰ روز در شیت قوانین مستند شد", 540 in values)
    check("ویرایش Incoterms در شیت قوانین آمده", "Incoterms 2020" in values)
    ws2 = wb["۷. اقلام سفارش (سطح PR و PI)"]
    check("ستون قرنطینه BL در شیت اقلام وجود دارد",
          any(c.value == "مقدار مشکوک BL" for c in ws2[1]))


# ═══════════ ۷) کانفیگ کلیدها — تغییر کلید و کلید مرکب ═══════════
def test_key_registry() -> None:
    print("\n── ۷) رجیستری کلیدها (config/keys.yaml) ──")
    import pandas as _pd
    from aibl.config.keys import KeyRegistry, get_keys

    kr = get_keys(reload=True)
    check("رجیستری بدون خطای ساختاری بارگذاری شد",
          all("اگر ستون خام است" in i for i in kr.validate()),
          "؛ ".join(kr.validate()) or "بدون هشدار")
    check("شش کلید ساده و سه کلید مرکب تعریف شده‌اند",
          len(kr.simple_keys()) == 6 and len(kr.composite_keys()) == 3,
          f"ساده {kr.simple_keys()} | مرکب {kr.composite_keys()}")

    # الگوی اعتبارسنجی: عدد ۹ رقمی نباید کلید ثبت سفارش شود
    reg = kr.get("REG")
    check("کد ثبت سفارش ۸ رقمی معتبر و ۹ رقمی نامعتبر است",
          reg.is_valid("97687754") and not reg.is_valid("664823825"),
          "همان تله‌ای که کل لایه رفع تعهد را قطع کرده بود")
    check("ستون «شماره پرونده» در فهرست ممنوعه است",
          "IL_FILE_NO" in reg.forbidden_columns)

    # ساخت کلید ساده از چند منبع، به ترتیب اولویت
    df = _pd.DataFrame({"SATA_KEY_REG": ["97687754", "", ""],
                        "FX_KEY_REG": ["", "83680876", ""],
                        "IL_KEY_REG": ["", "", "89212481"]})
    df = kr.build(df, "REG")
    check("کلید از چند سورس به ترتیب اولویت پر می‌شود",
          list(df["KEY_REG"]) == ["97687754", "83680876", "89212481"],
          str(list(df["KEY_REG"])))

    # کلید مرکب دو ستونی
    df2 = _pd.DataFrame({"KEY_BL": ["HDM1511WXRQ9407", "HDM1511WXRQ9407", ""],
                         "KEY_ORDER": ["502805", "501807", ""]})
    df2 = kr.build(df2, "BL_ORDER")
    check("کلید مرکب «بارنامه + سفارش» ساخته شد",
          df2["KEY_BL_ORDER"].iloc[0] == "HDM1511WXRQ9407|502805",
          df2["KEY_BL_ORDER"].iloc[0])
    check("یک بارنامه با دو سفارش، دو کلید مرکب متمایز می‌گیرد",
          df2["KEY_BL_ORDER"].iloc[0] != df2["KEY_BL_ORDER"].iloc[1])
    check("ردیف تهی، کلید مرکب تهی می‌گیرد",
          df2["KEY_BL_ORDER"].iloc[2] == "")

    # ادغام روی کلید مرکب، بدون تکثیر سطر
    from aibl.dataio.merge import safe_merge
    left = _pd.DataFrame({"KEY_BL": ["B1", "B1"], "KEY_ORDER": ["O1", "O2"],
                          "X": [1, 2]})
    right = _pd.DataFrame({"KEY_BL": ["B1", "B1"], "KEY_ORDER": ["O1", "O2"],
                           "Y": ["a", "b"]})
    merged = safe_merge(left, right, ["KEY_BL", "KEY_ORDER"], "تست مرکب")
    check("ادغام با کلید مرکب سطر تکثیر نمی‌کند و درست می‌چسبد",
          len(merged) == 2 and list(merged["Y"]) == ["a", "b"],
          f"{len(merged)} ردیف | Y={list(merged.get('Y', []))}")

    # تغییر کلید فقط با ویرایش YAML
    import os as _os
    import shutil as _sh
    import yaml as _yaml
    ext = _os.path.join(_TMP, "keys_ext")
    _os.makedirs(ext, exist_ok=True)
    src = _os.path.join(ROOT, "aibl", "config", "keys.yaml")
    dst = _os.path.join(ext, "keys.yaml")
    _sh.copy(src, dst)
    with open(dst, encoding="utf-8") as f:
        cfg = _yaml.safe_load(f)
    cfg["keys"]["REG"]["sources"] = [
        {"source": "ilappend", "frame": "main", "column": "IL_KEY_REG"}]
    with open(dst, "w", encoding="utf-8") as f:
        _yaml.safe_dump(cfg, f, allow_unicode=True)
    kr2 = KeyRegistry(path=dst)
    check("تغییر منبع کلید فقط با ویرایش YAML اعمال می‌شود",
          kr2.get("REG").sources[0]["column"] == "IL_KEY_REG"
          and kr.get("REG").sources[0]["column"] == "SATA_KEY_REG",
          "رجیستری اصلی دست‌نخورده ماند")



if __name__ == "__main__":
    print("=" * 78)
    print("AIBL V22 — تست کتابخانه قوانین، سورس مقاومت جدید و ماژولاریتی")
    print("=" * 78)
    test_rulebook()
    test_bl_validation()
    test_status_lexicon()
    test_order_ref()
    result = test_moghavemat()
    test_modularity()
    test_excel(result)
    test_key_registry()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)

