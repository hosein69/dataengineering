# -*- coding: utf-8 -*-
"""تست حوزه مسئولیت، مالکیت قطعه، وضعیت «کجا/کِی/چه کسی» و نماهای تأمین.

اجرا:  python tests/test_supply_views.py
"""
from __future__ import annotations

import os
import sys
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from aibl.report.supply_views import (SHEETS, build_bl_view,  # noqa: E402
                                      build_dept_view, build_material_view,
                                      write_supply_sheets)
from aibl.resolve import part_status as ps  # noqa: E402
from aibl.resolve.commercial_coverage import (FLAG, MEASURED,  # noqa: E402
                                              STATE, UNAVAILABLE, annotate,
                                              measured)
from aibl.resolve.expert_scope import (OWNER_GAP, OWNER_NAME,  # noqa: E402
                                       OWNER_SOURCE, SCOPES, coverage,
                                       resolve_owner, resolve_scopes)

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


def sample() -> pd.DataFrame:
    """داده نمونه با تاریخ **شمسی** — همان چیزی که سورس واقعی می‌دهد."""
    return pd.DataFrame({
        "KEY_MATERIAL": ["M1", "M1", "M2"],
        "CANONICAL_ORDER": ["O1", "O1", "O2"],
        "CANONICAL_BL": ["BL1", "BL1", "BL2"],
        "KEY_REG": ["R1", "R1", "R2"],
        "STAGE_FA": ["حمل", "حمل", "ترخیص"],
        "ORG_DEPT": ["خرید", "خرید", "لجستیک"],
        # نقش‌های ریز — ورودی حوزه‌ها
        "EXPERT_BUYER": ["الف", "الف", ""],
        "EXPERT_CLEARANCE": ["ج", "ج", "ج"],
        "EXPERT_CREDIT": ["ب", "ب", "ب"],
        "MOGH_KEY_EMP": ["", "", "9001"],
        # تاریخ رویدادها (شمسی)
        "BL_DATE": ["1405/01/10", "1405/01/10", ""],
        "COT_DATE": ["", "", "1405/02/20"],
        "MOGH_KEY_PR": ["PR1", "PR1", ""],
        "MOGH_PO_SENT_DATE": ["1404/12/01", "1404/12/01", ""],
        "MOGH_VENDOR_CODE": ["V1", "V1", ""],
        "MOGH_PI_VALUE_SUM": [100, 100, 0],
        "MOGH_BL_NO": ["BL1", "BL1", ""],
        "MOGH_CLEARED_QTY_SUM": [5, 5, 0],
        "مقاومت (روز)": [5, 5, 8],
        "روزهای رسوب": [2, 2, 12],
    })


class _Mapper:
    """جایگزین سبک DynamicOrgMapper برای تست کد پرسنلی → نام."""

    def map(self, name="", code=""):
        return {"expert": {"9001": "دلاور"}.get(str(code).strip(), "نامشخص")}


# ═══════════ ۱) سه حوزه مسئولیت ═══════════
def test_scopes() -> None:
    print("\n── ۱) سه حوزه مسئولیت و مالک قطعه ──")
    d = resolve_scopes(sample(), _Mapper())

    check("سه ستون حوزه ساخته می‌شود",
          [s.key for s in SCOPES] == ["EXPERT_PURCHASING", "EXPERT_COMMERCIAL",
                                      "EXPERT_LOGISTICS"],
          str([s.key for s in SCOPES]))
    check("کارشناس ترخیص هرگز در ستون کارشناس خرید نمی‌نشیند",
          "ج" not in d["EXPERT_PURCHASING"].tolist(),
          str(d["EXPERT_PURCHASING"].tolist()))
    check("حوزه لجستیک فقط کارشناس ترخیص را دارد",
          d["EXPERT_LOGISTICS"].tolist() == ["ج", "ج", "ج"])
    check("حوزه بازرگانی از نقش اعتبارات پر می‌شود",
          d["EXPERT_COMMERCIAL"].tolist() == ["ب", "ب", "ب"])
    check("وقتی نام مستقیم نیست، کد پرسنلی از HR به نام تبدیل می‌شود",
          d["EXPERT_PURCHASING"].iloc[2] == "دلاور", d["EXPERT_PURCHASING"].iloc[2])

    d = resolve_owner(d, _Mapper())
    check("مالک قطعه همان کارشناس خرید است",
          d[OWNER_NAME].tolist() == ["الف", "الف", "دلاور"], str(d[OWNER_NAME].tolist()))
    check("منبع شناسایی مالک ثبت می‌شود",
          "کد پرسنلی" in d[OWNER_SOURCE].iloc[2], d[OWNER_SOURCE].iloc[2])

    # مهم‌ترین قاعده کسب‌وکار: داده ناقص مالک حذف نمی‌شود
    check("ردیف با داده ناقصِ مالک حذف نمی‌شود", len(d) == 3, str(len(d)))
    check("شکاف داده مالک برای ردیف ناقص ثبت می‌شود",
          d[OWNER_GAP].iloc[2] != "", d[OWNER_GAP].iloc[2])
    check("ردیف کامل شکاف ندارد", d[OWNER_GAP].iloc[0] == "", d[OWNER_GAP].iloc[0])

    cov = coverage(d)
    check("گزارش پوشش، مالک بودن حوزه خرید را اعلام می‌کند",
          cov.loc[cov["حوزه"] == "کارشناس خرید", "مالک قطعه"].iloc[0] == "بله")
    check("فقط یک حوزه مالک است", sum(1 for s in SCOPES if s.owner) == 1)


# ═══════════ ۲) وضعیت: کجا / کِی / چه کسی ═══════════
def test_status() -> None:
    print("\n── ۲) وضعیت قطعه: کجا / کِی / چه کسی ──")
    d = resolve_owner(resolve_scopes(sample(), _Mapper()), _Mapper())
    out = ps.resolve(d, today=pd.Timestamp("2026-09-08"))

    check("موقعیت از رویداد تاریخ‌دار می‌آید، نه از حدس موجودی",
          out[ps.WHERE].tolist() == ["حمل", "حمل", "گمرک و ترخیص"],
          str(out[ps.WHERE].tolist()))
    check("آخرین فعالیت نام‌گذاری می‌شود",
          out[ps.ACTIVITY].iloc[2] == "ثبت کوتاژ گمرکی", out[ps.ACTIVITY].iloc[2])

    # ── همان باگی که در اجرای واقعی دیده شد ──
    age = pd.to_numeric(out[ps.AGE], errors="coerce")
    check("تاریخ شمسی درست تبدیل می‌شود (سن وضعیت منطقی است، نه ۲۲۷٬۰۰۰ روز)",
          age.notna().all() and age.between(0, 3650).all(),
          str(age.tolist()))
    check("تاریخ به میلادی نرمال می‌شود",
          str(out[ps.WHEN].iloc[0])[:4] == "2026", str(out[ps.WHEN].iloc[0]))

    check("مسئول وضعیت، کارشناس همان حوزه است",
          out[ps.WHO_SCOPE].tolist() ==
          ["کارشناس حمل و لجستیک"] * 3, str(out[ps.WHO_SCOPE].tolist()))
    check("نام مسئول از ستون همان حوزه برداشته می‌شود",
          out[ps.WHO].tolist() == ["ج", "ج", "ج"], str(out[ps.WHO].tolist()))

    # «بعدی» فقط از میان فعالیت‌هایی انتخاب می‌شود که ستون تاریخشان در داده
    # هست. درباره مرحله‌ای که اصلاً قابل مشاهده نیست ادعایی نمی‌کنیم.
    check("فعالیت بعدی مورد انتظار تعیین می‌شود",
          out[ps.NEXT_ACT].iloc[0] == "ثبت کوتاژ گمرکی", out[ps.NEXT_ACT].iloc[0])
    check("چرخه تمام‌شده «کامل است» گزارش می‌شود",
          out[ps.NEXT_ACT].iloc[2] == "چرخه کامل است", out[ps.NEXT_ACT].iloc[2])
    check("«معطل چه حوزه‌ای» پاسخ دارد",
          out[ps.WAITING_SCOPE].iloc[0] == "کارشناس حمل و لجستیک",
          out[ps.WAITING_SCOPE].iloc[0])
    check("تاریخ‌های ثبت‌نشده فهرست می‌شوند", out[ps.MISSING].iloc[0] != "")

    # ردیف بدون هیچ تاریخ
    blank = ps.resolve(pd.DataFrame({"KEY_MATERIAL": ["X"], "BL_DATE": [""]}))
    check("ردیف بدون رویداد «نامشخص» می‌شود، نه یک حدس",
          blank[ps.WHERE].iloc[0] == "نامشخص", blank[ps.WHERE].iloc[0])
    check("مبنای نامشخص بودن صریح ثبت می‌شود",
          "ثبت نشده" in blank[ps.BASIS].iloc[0], blank[ps.BASIS].iloc[0])

    # تاریخ آینده شاهد وضعیت فعلی نیست
    fut = ps.resolve(pd.DataFrame({"BL_DATE": ["1500/01/01"]}),
                     today=pd.Timestamp("2026-09-08"))
    check("تاریخ آینده به‌عنوان وضعیت فعلی پذیرفته نمی‌شود",
          fut[ps.WHERE].iloc[0] == "نامشخص", fut[ps.WHERE].iloc[0])

    check("جدول زمانی با جدول فعالیت لاگ رویداد یکی است", _timeline_matches())


def _timeline_matches() -> bool:
    """TIMELINE و ACTIVITIES نباید از هم جدا بیفتند."""
    from aibl.stages.s80_eventlog import ACTIVITIES
    mine = [(c, o, st) for c, _fa, o, st in ps.TIMELINE]
    theirs = [(c, o, st) for c, _en, _fa, o, st in ACTIVITIES]
    return mine == theirs


# ═══════════ ۳) پوشش سورس خرید ═══════════
def test_commercial_coverage() -> None:
    print("\n── ۳) پوشش Commercial Expert Data ──")
    d = sample()
    lines = pd.DataFrame({"KEY_ORDER": ["O1"]})
    out, n = annotate(d, lines)
    check("سفارش خارج از سورس شمرده می‌شود", n == 1, str(n))
    check("سفارش موجود در سورس پرچم نمی‌گیرد", not bool(out[FLAG].iloc[0]))
    check("سفارش غایب پرچم می‌گیرد", bool(out[FLAG].iloc[2]))
    check("وضعیت سنجش «measured» است", out[STATE].iloc[0] == MEASURED)
    check("measured() درست گزارش می‌دهد", measured(out) is True)

    # ── باگ نسخه ۲۶٫۹: سورس نبود ⇒ همه پرچم می‌گرفتند ──
    for empty in (None, pd.DataFrame()):
        out2, n2 = annotate(d, empty)
        check(f"سورس در دسترس نباشد ⇒ هیچ سفارشی به‌غلط پرچم نمی‌گیرد ({type(empty).__name__})",
              n2 == 0 and not out2[FLAG].any(), f"n={n2}")
    out2, _ = annotate(d, None)
    check("حالت «سنجیده نشد» از «صفر» تفکیک می‌شود",
          out2[STATE].iloc[0] == UNAVAILABLE and measured(out2) is False)
    check("علت سنجیده‌نشدن برای کاربر توضیح داده می‌شود",
          "سنجیده نشده" in out2["ORDER_MISSING_COMMERCIAL_REASON"].iloc[0])


# ═══════════ ۴) سه نمای تأمین ═══════════
def test_views() -> None:
    print("\n── ۴) نماهای متریال / بارنامه / اداره ──")
    d = ps.resolve(resolve_owner(resolve_scopes(sample(), _Mapper()), _Mapper()))
    m, b, o = build_material_view(d), build_bl_view(d), build_dept_view(d)

    check("نمای متریال هر ردیف را نگه می‌دارد",
          len(m) == 3 and m["کلید گروه"].tolist() == ["M1", "M1", "M2"])
    check("نمای بارنامه هر ردیف را نگه می‌دارد",
          len(b) == 3 and b["کلید گروه"].tolist() == ["BL1", "BL1", "BL2"])
    check("نمای اداره یک ردیف به ازای هر اداره دارد", len(o) == 2, str(len(o)))

    for col in ("موقعیت فعلی", "تاریخ آخرین رویداد", "سن وضعیت (روز)",
                "کارشناس مسئول وضعیت", "معطل حوزه",
                "مالک قطعه (کارشناس خرید)", "شکاف داده مالک"):
        check(f"ستون «{col}» در نما هست", col in m.columns)
    for s in SCOPES:
        check(f"ستون حوزه «{s.fa}» در نما هست", s.fa in m.columns)
    check("نمای اداره چرایی بار کاری دارد", "چرایی بار کاری / وضعیت" in o.columns)
    check("نمای اداره شکاف مالک را می‌شمارد", "ردیف با شکاف داده مالک" in o.columns)

    wb = Workbook()
    wb.remove(wb.active)
    write_supply_sheets(wb, d)
    check("سه شیت با نام درست ساخته می‌شود", wb.sheetnames == list(SHEETS),
          str(wb.sheetnames))

    check("دیتافریم خالی خطا نمی‌دهد", build_material_view(pd.DataFrame()).empty)
    check("ستون گروه‌بندی نباشد ⇒ نمای خالی، نه استثنا",
          build_bl_view(pd.DataFrame({"X": [1]})).empty)


# ═══════════ ۵) کارایی ═══════════
def test_performance() -> None:
    print("\n── ۵) کارایی (رگرسیون نسخه ۲۶٫۹) ──")
    n = 5000
    rng = np.random.default_rng(0)
    d = pd.DataFrame({
        "KEY_MATERIAL": [f"M{i % 800}" for i in range(n)],
        "CANONICAL_BL": [f"BL{i % 900}" for i in range(n)],
        "CANONICAL_ORDER": [f"O{i % 1200}" for i in range(n)],
        "ORG_DEPT": rng.choice(["خرید", "لجستیک"], n),
        "BL_DATE": rng.choice(["1404/10/01", "1405/01/12", ""], n),
        "COT_DATE": rng.choice(["1405/02/20", ""], n),
        "مقاومت (روز)": rng.integers(0, 60, n),
    })
    t = time.time()
    build_material_view(d); build_bl_view(d); build_dept_view(d)
    el = time.time() - t
    # پیاده‌سازی ردیف‌به‌ردیف نسخه ۲۶٫۹ برای ۲٬۰۰۰ ردیف ۳۶ ثانیه می‌گرفت.
    check(f"سه نما روی {n:,} ردیف زیر ۱۰ ثانیه ساخته می‌شود", el < 10,
          f"{el:.2f}s")


# ═══════════ ۶) پوشش خودِ مجموعه تست ═══════════
def test_suite_registration() -> None:
    print("\n── ۶) هیچ فایل تستی از قلم نیفتد ──")
    import re
    body = open(os.path.join(ROOT, "run_all_tests.py"), encoding="utf-8").read()
    listed = set(re.findall(r'"(tests/test_\w+\.py)"', body))
    on_disk = {f"tests/{f}" for f in os.listdir(os.path.join(ROOT, "tests"))
               if f.startswith("test_") and f.endswith(".py")}
    missing = sorted(on_disk - listed)
    # نسخه ۲۶٫۹ همین فایل را به run_all_tests اضافه نکرده بود و تست‌هایش
    # هرگز اجرا نمی‌شد، در حالی که گزارش «همه سبز» می‌داد.
    check("همه فایل‌های تست در run_all_tests ثبت شده‌اند", not missing,
          f"ثبت‌نشده: {missing}" if missing else f"{len(on_disk)} فایل")


if __name__ == "__main__":
    print("=" * 78)
    print("AIBL — حوزه مسئولیت، مالکیت قطعه و نماهای تأمین")
    print("=" * 78)
    test_scopes()
    test_status()
    test_commercial_coverage()
    test_views()
    test_performance()
    test_suite_registration()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
