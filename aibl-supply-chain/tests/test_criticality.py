# -*- coding: utf-8 -*-
"""تست‌های لایه مقاومت قطعه و بحرانی بودن.

این لایه در نسخه ۲۰.۱ و در بازنویسی اول **کاملاً غایب** بود: سورس Oracle فقط
وضعیت بارنامه را می‌داد و موجودی/مصرف هرگز خوانده نمی‌شد.

اجرا:  python tests/test_criticality.py
"""
from __future__ import annotations

import os
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import tempfile  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)


def _load_make_synthetic():
    import importlib.util
    path = os.path.join(_TESTS_DIR, "make_synthetic.py")
    spec = importlib.util.spec_from_file_location("aibl_make_synthetic", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_ms = _load_make_synthetic()
build = _ms.build

_TMP = tempfile.mkdtemp(prefix="aibl_crit_")
DIRS = build(_TMP)
os.environ.update({
    "AIBL_FOREIGN": DIRS["foreign"], "AIBL_BLS": DIRS["bls"],
    "AIBL_CLEARANCE": DIRS["clearance"], "AIBL_HR": DIRS["hr"],
    "AIBL_ESMAEILI": DIRS["esmaeili"],
    "AIBL_GS_COMBINE": DIRS["gs_combine"], "AIBL_MOHAMADI": DIRS["mohamadi"],
    "AIBL_OUTPUT": DIRS["output"], "AIBL_LOGS": DIRS["logs"],
    "AIBL_TODAY": "2026-08-31",
})

import logging  # noqa: E402
import shutil  # noqa: E402

import pandas as pd  # noqa: E402
import yaml  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from aibl.engines.criticality import CriticalityEngine  # noqa: E402
from aibl.rulebook import RuleBook, get_rulebook  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


# ═══════════ ۱) فرمول و طبقه‌بندی ═══════════
def test_engine() -> None:
    print("\n── ۱) فرمول مقاومت و طبقه‌بندی ──")
    eng = CriticalityEngine(get_rulebook(reload=True))

    # (ایران‌خودرو, ساپکو, در راه, در گمرک, نیاز روزانه, مقاومت, طبقه)
    cases = [
        (80, 20, 0, 0, 20, 5.0, "CRITICAL"),              # زیر ۱۰ روز
        (150, 30, 0, 0, 20, 9.0, "CRITICAL"),             # مرز پایین
        (150, 50, 0, 0, 20, 10.0, "BECOMING_CRITICAL"),   # دقیقاً ۱۰
        (200, 100, 0, 0, 20, 15.0, "BECOMING_CRITICAL"),
        (300, 99, 0, 0, 20, 19.95, "BECOMING_CRITICAL"),
        (300, 100, 0, 0, 20, 20.0, "WATCH"),              # دقیقاً ۲۰
        (600, 400, 0, 0, 25, 40.0, "WATCH"),
        (5000, 0, 0, 0, 10, 500.0, "SAFE"),
        (0, 0, 0, 0, 5, 0.0, "STOCKOUT"),
        # اجزای «در راه» و «در گمرک» هم در صورت کسر می‌آیند
        (100, 0, 200, 100, 20, 20.0, "WATCH"),
    ]
    ok, bad = True, []
    for ikco, sapco, transit, customs, need, exp_days, exp_band in cases:
        stock, daily = ikco + sapco, need
        r = eng.evaluate({"STOCK_IKCO": ikco, "STOCK_SAPCO": sapco,
                          "IN_TRANSIT_QTY": transit, "IN_CUSTOMS_QTY": customs,
                          "DAILY_NEED": need})
        got = -1.0 if r.resistance_days is None else r.resistance_days  # 0.0 falsy است
        if abs(got - exp_days) > 0.06 or r.band != exp_band:
            ok = False
            bad.append(f"{stock}+{transit}+{customs}/{daily} → {got} {r.band} "
                       f"(انتظار {exp_days} {exp_band})")
    check("مرزهای ۱۰ و ۲۰ روز دقیقاً طبق تعریف کاربر", ok,
          "؛ ".join(bad) or f"{len(cases)} حالت درست")

    r = eng.evaluate({"STOCK_IKCO": 100, "STOCK_SAPCO": 0, "DAILY_NEED": 0})
    check("نیاز روزانه صفر ⇒ «بدون مصرف» نه «بحرانی»",
          r.band == "NO_CONSUMPTION" and r.resistance_days is None, r.band_short)

    r = eng.evaluate({"STOCK_IKCO": "", "STOCK_SAPCO": "", "DAILY_NEED": ""})
    check("داده غایب ⇒ «نامشخص» نه «توقف خط»",
          r.band == "UNKNOWN", r.band_short)

    r = eng.evaluate({"STOCK_IKCO": 0, "STOCK_SAPCO": 0, "IN_TRANSIT_QTY": 200,
                      "DAILY_NEED": 5})
    check("مقاومت انبار جدا از مقاومت کل گزارش می‌شود",
          r.resistance_warehouse == 0.0 and r.resistance_days == 40.0,
          f"انبار {r.resistance_warehouse} | کل {r.resistance_days}")


    check("ترتیب مرتب‌سازی طبقات درست است",
          [eng.evaluate({"STOCK_IKCO": s, "STOCK_SAPCO": 0,
                         "DAILY_NEED": 10}).sort_rank
           for s in (0, 50, 150, 300, 1000)] == [0, 1, 2, 3, 4],
          "توقف خط → بحرانی → در حال بحرانی → تحت نظر → ایمن")


# ═══════════ ۲) هشدارهای ترکیبی ═══════════
def test_combined_alerts() -> None:
    print("\n── ۲) هشدار ترکیبی «قطعه بحرانی × پرونده مشکل‌دار» ──")
    eng = CriticalityEngine(get_rulebook())

    a = eng.combined_alerts({"کد طبقه بحرانی": "CRITICAL", "IS_BLOCKED": True,
                             "روزهای رسوب": 0, "IS_IN_MOGHAVEMAT": True})
    check("قطعه بحرانی + پرونده بلوکه ⇒ هشدار CRITICAL",
          any(x["severity"] == "CRITICAL" for x in a),
          a[0]["fa"] if a else "هیچ")

    b = eng.combined_alerts({"کد طبقه بحرانی": "CRITICAL", "IS_BLOCKED": False,
                             "روزهای رسوب": 30, "IS_IN_MOGHAVEMAT": True})
    check("قطعه بحرانی + رسوب گمرکی ⇒ هشدار", bool(b), b[0]["fa"] if b else "هیچ")

    c = eng.combined_alerts({"کد طبقه بحرانی": "CRITICAL", "IS_BLOCKED": False,
                             "روزهای رسوب": 0, "IS_IN_MOGHAVEMAT": False})
    check("قطعه بحرانی بدون سفارش باز ⇒ هشدار PR اضطراری",
          any("سفارش باز" in x["fa"] for x in c), c[0]["fa"] if c else "هیچ")

    d = eng.combined_alerts({"کد طبقه بحرانی": "SAFE", "IS_BLOCKED": True,
                             "روزهای رسوب": 100, "IS_IN_MOGHAVEMAT": True})
    check("قطعه ایمن هشدار ترکیبی نمی‌گیرد", not d)


# ═══════════ ۳) خط لوله کامل ═══════════
def test_pipeline():
    print("\n── ۳) اتصال سورس Oracle → مقاومت در خط لوله ──")
    logging.disable(logging.INFO)
    from aibl.pipeline import Pipeline

    res = Pipeline().run(build_report=True)
    df = res.df

    ikco = pd.to_numeric(df["موجودی ایران خودرو"], errors="coerce").sum()
    sapco = pd.to_numeric(df["موجودی ساپکو"], errors="coerce").sum()
    check("موجودی ایران‌خودرو و ساپکو هر دو از Oracle آمدند",
          ikco > 0 and sapco > 0, f"IKCO={ikco:,.0f} ساپکو={sapco:,.0f}")
    check("نیاز روزانه از Oracle خوانده شد",
          pd.to_numeric(df["نیاز روزانه"], errors="coerce").sum() > 0)
    linked = int(df["KEY_MATERIAL"].astype(str).str.strip().ne("").sum())
    check("کلید متریال از زنجیره بارنامه ← سفارش ← مقاومت ← Oracle ساخته شد",
          linked >= 5, f"{linked} از {len(df)} ردیف "
                       f"(ردیف بدون سفارش در مقاومت، متریال ندارد)")

    bands = list(df["کد طبقه بحرانی"])
    check("هر شش طبقه در داده آزمون تولید شد",
          len(set(bands)) >= 6, str(sorted(set(df["بحرانی (کوتاه)"]))))

    # مرتب‌سازی
    ranks = list(pd.to_numeric(df["CRITICALITY_SORT"], errors="coerce"))
    check("گزارش بر اساس بحرانی بودن مرتب شده است",
          ranks == sorted(ranks), str(ranks))
    check("بحرانی‌ترین قطعه در ردیف اول است",
          df.iloc[0]["کد طبقه بحرانی"] == "STOCKOUT",
          f"{df.iloc[0]['بحرانی (کوتاه)']} — متریال {df.iloc[0]['KEY_MATERIAL']}")

    # صحت عددی
    row = df[df["KEY_MATERIAL"] == "IK013581CR"].iloc[0]
    check("(۴۰۰ ایران‌خودرو + ۱۰۰ ساپکو) ÷ ۱۰۰ = ۵ روز ⇒ بحرانی",
          abs(float(row["مقاومت (روز)"]) - 5.0) < 0.01
          and row["کد طبقه بحرانی"] == "CRITICAL",
          f"{row['مقاومت (روز)']} روز → {row['بحرانی (کوتاه)']}")
    row2 = df[df["KEY_MATERIAL"] == "9656956180"].iloc[0]
    check("(۱۲۰۰ + ۳۰۰) ÷ ۱۰۰ = ۱۵ روز ⇒ در حال بحرانی شدن",
          abs(float(row2["مقاومت (روز)"]) - 15.0) < 0.01
          and row2["کد طبقه بحرانی"] == "BECOMING_CRITICAL")

    # موجودی چند انبار
    check("قطعه با موجودی صفر ⇒ توقف خط",
          df[df["KEY_MATERIAL"] == "K914564758A"].iloc[0]["کد طبقه بحرانی"] == "STOCKOUT")

    # اثر بر ریسک
    crit = df[df["کد طبقه بحرانی"].isin(["STOCKOUT", "CRITICAL"])]["امتیاز ریسک"].mean()
    safe = df[df["کد طبقه بحرانی"] == "SAFE"]["امتیاز ریسک"].mean()
    check("قطعه بحرانی امتیاز ریسک را بالا می‌برد", crit > safe,
          f"بحرانی {crit:.1f} در برابر ایمن {safe:.1f}")

    # نقش دوم Oracle حفظ شده
    check("گروه‌بندی قطعه از Oracle وارد گزارش شد",
          "PART_GROUP" in df.columns
          and df["PART_GROUP"].astype(str).str.strip().ne("").any(),
          str(sorted({str(x) for x in df.get("PART_GROUP", [])})[:3]))

    n_alert = df["هشدار ترکیبی بحرانی"].astype(str).str.strip().ne("").sum()
    check("هشدارهای ترکیبی در خط لوله تولید شدند", n_alert >= 1, f"{n_alert} ردیف")
    return res


# ═══════════ ۴) شیت اکسل ═══════════
def test_excel(res) -> None:
    print("\n── ۴) شیت «قطعات بحرانی» ──")
    wb = load_workbook(res.dashboard_path)
    check("شیت ۹ ساخته شد", "۹. قطعات بحرانی" in wb.sheetnames, str(wb.sheetnames))

    ws = wb["۹. قطعات بحرانی"]
    labels = [ws.cell(row=r, column=1).value for r in range(4, 11)]
    check("خلاصه هر هفت طبقه در بالای شیت آمده",
          sum(1 for x in labels if x) == 7, str([x for x in labels if x][:3]))
    check("جدول تفصیلی قطعات نیازمند اقدام دارد",
          any(c.value == "مقاومت (روز)" for row in ws.iter_rows(max_row=20) for c in row))

    m = wb["۲. کالبدشکافی ۳ لایه‌ای ماتریسی"]
    first_headers = [c.value for c in m[1][:5]]
    check("ستون «طبقه بحرانی» اولین ستون شیت ماتریس است",
          first_headers[0] == "طبقه بحرانی", str(first_headers))
    check("ستون مقاومت بلافاصله بعد از آن است", first_headers[1] == "مقاومت (روز)")

    exec_ws = wb["۱. خلاصه اجرایی"]
    kpi_labels = [exec_ws.cell(row=r, column=1).value for r in range(6, 24)]
    check("KPI قطعات بحرانی در خلاصه اجرایی آمده",
          any(x and "مقاومت زیر ۱۰ روز" in str(x) for x in kpi_labels),
          str([x for x in kpi_labels if x and "بحرانی" in str(x)][:2]))


# ═══════════ ۵) پیکربندی‌پذیری آستانه‌ها ═══════════
def test_configurable() -> None:
    print("\n── ۵) تغییر آستانه‌ها فقط با YAML ──")
    ext = os.path.join(_TMP, "rules_ext")
    shutil.copytree(os.path.join(ROOT, "aibl", "rules"), ext, dirs_exist_ok=True)
    path = os.path.join(ext, "criticality.yaml")
    with open(path, encoding="utf-8") as f:
        rules = yaml.safe_load(f)
    for b in rules["bands"]:
        if b["code"] == "CRITICAL":
            b["max_days"] = 15          # به‌جای ۱۰
        if b["code"] == "BECOMING_CRITICAL":
            b["max_days"] = 30          # به‌جای ۲۰
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(rules, f, allow_unicode=True)

    eng = CriticalityEngine(RuleBook(rules_dir=ext))
    r = eng.evaluate({"STOCK_IKCO": 240, "STOCK_SAPCO": 0, "DAILY_NEED": 20})
    check("با تغییر آستانه به ۱۵ روز، ۱۲ روز «بحرانی» می‌شود",
          r.band == "CRITICAL", f"{r.resistance_days} روز → {r.band}")
    r2 = eng.evaluate({"STOCK_IKCO": 500, "STOCK_SAPCO": 0, "DAILY_NEED": 20})
    check("۲۵ روز با آستانه جدید «در حال بحرانی شدن» است",
          r2.band == "BECOMING_CRITICAL", f"{r2.resistance_days} روز → {r2.band}")

    eng_default = CriticalityEngine(get_rulebook())
    r3 = eng_default.evaluate({"STOCK_IKCO": 240, "STOCK_SAPCO": 0, "DAILY_NEED": 20})
    check("قوانین اصلی دست‌نخورده ماند (۱۲ روز = در حال بحرانی شدن)",
          r3.band == "BECOMING_CRITICAL", r3.band_short)



def test_group_criticality() -> None:
    print("\n── ۶) سرایت بحرانی بودن از متریال به بارنامه/سفارش ──")
    from aibl.stages.s40_criticality import CriticalityStage
    from aibl.stages.base import PipelineContext
    import pandas as pd
    ctx = PipelineContext(rb=get_rulebook(), today=__import__('datetime').date(2026,8,31))
    df = pd.DataFrame({
        "KEY_MATERIAL": ["M-RED", "M-GREEN", "M-SAFE"],
        "CANONICAL_BL": ["BL-1", "BL-1", "BL-2"],
        "CANONICAL_ORDER": ["ORD-1", "ORD-1", "ORD-2"],
        "STOCK_IKCO": [0, 1000, 1000], "STOCK_SAPCO": [0,0,0],
        "IN_TRANSIT_QTY": [0,0,0], "IN_CUSTOMS_QTY": [0,0,0],
        "DAILY_NEED": [10,10,100],
    })
    out = CriticalityStage().run(df, ctx)
    r1 = out.iloc[0]
    check("یک متریال بحرانی کل بارنامه را بحرانی می‌کند", bool(r1["BL_CRITICAL"]))
    check("یک متریال بحرانی کل سفارش را بحرانی می‌کند", bool(r1["ORDER_CRITICAL"]))
    check("علت بحرانی بودن تا سطح متریال ثبت شده", "M-RED" in r1["BL_CRITICAL_MATERIALS"] and "M-RED" in r1["BL_CRITICAL_REASON"])
    check("بارنامه بدون متریال بحرانی بحرانی نمی‌شود", not bool(out.iloc[2]["BL_CRITICAL"]))

if __name__ == "__main__":
    print("=" * 78)
    print("AIBL V23 — تست مقاومت قطعه و بحرانی بودن")
    print("=" * 78)
    test_engine()
    test_combined_alerts()
    result = test_pipeline()
    test_excel(result)
    test_configurable()
    test_group_criticality()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
