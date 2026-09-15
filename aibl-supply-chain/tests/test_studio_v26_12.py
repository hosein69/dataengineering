# -*- coding: utf-8 -*-
import os, sys, tempfile
from pathlib import Path
import pandas as pd
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
os.environ.setdefault("AIBL_DESIGNS", tempfile.mkdtemp(prefix="aibl_designs_"))
from aibl.rulebook import get_rulebook
from aibl.studio_core.filters import FilterState, apply_filters, filter_options
from aibl.studio_core.html_export import build_dynamic_html
from aibl.studio_core.report_builder import ReportSpec, build
from aibl.studio_core.designs import ReportDesign, save_design, load_design, EMAIL_CHARTS

def test_transport():
    rb=get_rulebook(reload=True)
    assert rb.transport_mode("Land")=="ROAD"
    assert rb.transport_mode_fa("ROAD").startswith("زمینی")
    df=pd.DataFrame({"TRANSPORT_MODE":["SEA","AIR","ROAD","SEA"]})
    opts=filter_options(df)["transport"]
    assert "دریایی" in opts and "هوایی" in opts and any(x.startswith("زمینی") for x in opts)
    assert len(apply_filters(df,FilterState(transport=["دریایی"])))==2
    assert len(apply_filters(df,FilterState(transport=["هوایی"])))==1
    assert len(apply_filters(df,FilterState(transport=["زمینی (جاده‌ای)"])))==1

def test_tabs_and_design():
    df=pd.DataFrame({"KEY_MATERIAL":["M1","M2"],"TRANSPORT_MODE":["SEA","AIR"],
                     "CANONICAL_ORDER":["O1","O2"],"مانده تعهد":[10,20]})
    labels={"KEY_MATERIAL":"کد متریال","TRANSPORT_MODE":"روش حمل","CANONICAL_ORDER":"شماره سفارش"}
    with tempfile.TemporaryDirectory() as td:
        spec=ReportSpec(fields=list(df.columns),ref_date="2026-09-13",formats=["html","excel"],
                        file_stem="tabs",tabs=[{"title":"حمل","fields":["TRANSPORT_MODE","KEY_MATERIAL"]},
                                               {"title":"سفارش","fields":["CANONICAL_ORDER"]}])
        r=build(df,{},spec,labels,td)
        h=Path(r.files["html"]).read_text(encoding="utf8")
        assert h.count('role="tab"')==2 and h.count('role="tabpanel"')==2
        from openpyxl import load_workbook
        wb=load_workbook(r.files["excel"])
        assert "حمل" in wb.sheetnames and "سفارش" in wb.sheetnames
    d=ReportDesign(name="reuse",template="executive",fields=["A"],
                   tabs=[{"title":"T","fields":["A"]}],email_charts=["criticality"])
    save_design(d); assert load_design("reuse").email_charts==["criticality"]

if __name__=="__main__":
    test_transport(); test_tabs_and_design(); print("OK")
