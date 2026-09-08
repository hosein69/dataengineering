# -*- coding: utf-8 -*-
"""تست معماری: «هر تغییر، یک فایل» و تشخیص اختلاف نسخه.

این مجموعه دقیقاً همان قولی را می‌سنجد که طراحی داده است:
    ۱. افزودن یک قابلیت = یک فایل جدید در stages/ ، بدون لمس pipeline یا dashboard
    ۲. حذف یک قابلیت = حذف همان فایل، بدون خطا
    ۳. اختلاف نسخه فایل‌ها پیش از اجرا و با پیام صریح گرفته می‌شود
       (علت خطای واقعی: 'SourceSpec' object has no attribute 'frame_map')

اجرا:  python tests/test_architecture.py
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PKG = os.path.join(ROOT, "aibl")
sys.path.insert(0, ROOT)

_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)


def _load_make_synthetic():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "aibl_make_synthetic", os.path.join(_TESTS_DIR, "make_synthetic.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


build = _load_make_synthetic().build

_TMP = tempfile.mkdtemp(prefix="aibl_arch_")
DIRS = build(_TMP)
ENV = {
    "AIBL_FOREIGN": DIRS["foreign"], "AIBL_BLS": DIRS["bls"],
    "AIBL_CLEARANCE": DIRS["clearance"], "AIBL_HR": DIRS["hr"],
    "AIBL_ESMAEILI": DIRS["esmaeili"],
    "AIBL_GS_COMBINE": DIRS["gs_combine"], "AIBL_MOHAMADI": DIRS["mohamadi"],
    "AIBL_OUTPUT": DIRS["output"], "AIBL_LOGS": DIRS["logs"],
    "AIBL_TODAY": "2026-08-31",
}
os.environ.update(ENV)

import logging  # noqa: E402

import pandas as pd  # noqa: E402

from aibl.stages import (collect_columns, discover, validate_graph)  # noqa: E402
from aibl.stages.base import REGISTRY, StageContractError  # noqa: E402
from aibl.version import check_contracts  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


BASE_COLS = ["KEY_BL", "KEY_ORDER", "KEY_REG", "KEY_EMP", "KEY_MATERIAL",
             "MOGH_PRESENT", "COT_DATE", "NTSW_ALLOC_DATE", "ORC_STATUS"]


# ═══════════ ۱) قرارداد مرحله‌ها ═══════════
def test_stage_contract() -> None:
    print("\n── ۱) قرارداد مرحله‌ها ──")
    stages = discover()
    # ⚠️ عمداً تعداد دقیق را بررسی نمی‌کنیم: افزودن یک مرحله جدید نباید
    # تست را بشکند. آنچه باید ثابت بماند، *ثابت‌های ترتیبی* است.
    names = [s.name for s in stages]
    check("مرحله‌ها خودکار کشف می‌شوند", len(stages) >= 9,
          f"{len(stages)} مرحله: " + " → ".join(names))
    order_ok = all(names.index(a) < names.index(b) for a, b in [
        ("resolve", "derive"), ("derive", "criticality"),
        ("criticality", "risk"), ("narrate", "sort"),
        ("eventlog", "sort")] if a in names and b in names)
    check("ثابت‌های ترتیبی زنجیره علّی رعایت شده", order_ok,
          "resolve→derive→criticality→risk و eventlog→sort")
    check("ترتیب اجرا صعودی و بدون تکرار است",
          [s.order for s in stages] == sorted(set(s.order for s in stages)),
          str([s.order for s in stages]))
    check("هر مرحله عنوان فارسی دارد", all(s.title for s in stages))
    check("هر مرحله ورودی/خروجی خود را اعلام کرده",
          all(isinstance(s.requires, list) and isinstance(s.provides, list)
              for s in stages))

    warns = validate_graph(stages, BASE_COLS)
    check("زنجیره علّی کامل است (هر ورودی توسط مرحله قبلی تولید می‌شود)",
          not warns, "؛ ".join(warns) or "بدون هشدار")

    cols = collect_columns(stages)
    check("ستون‌های گزارش را مرحله‌ها اعلام می‌کنند، نه dashboard",
          len(cols) >= 30, f"{len(cols)} ستون از {len(stages)} مرحله")
    check("ستون‌های بحرانی در ابتدای ترتیب نمایش‌اند",
          [c.title for c in cols[:2]] == ["طبقه بحرانی", "مقاومت (روز)"],
          str([c.title for c in cols[:4]]))


# ═══════════ ۲) گارد ترتیب — باگ B1 دیگر ممکن نیست ═══════════
def test_order_guard() -> None:
    print("\n── ۲) گارد ترتیب علّی ──")
    from aibl.stages.base import Stage

    class BadStage(Stage):
        name = "bad_test_stage"
        title = "مرحله‌ای که ورودی ناموجود می‌خواهد"
        order = 5                       # قبل از derive
        requires = ["COTAGE_NO"]        # که derive در ترتیب ۲۰ می‌سازد
        provides = []

        def run(self, df, ctx):
            return df

    stages = sorted(discover() + [BadStage()], key=lambda s: s.order)
    try:
        validate_graph(stages, BASE_COLS)
        check("مرحله با ترتیب اشتباه رد می‌شود", False, "خطا صادر نشد!")
    except StageContractError as ex:
        msg = str(ex)
        check("مرحله با ترتیب اشتباه پیش از خواندن داده رد می‌شود", True)
        check("پیام خطا مرحله تولیدکننده را نام می‌برد", "derive" in msg,
              msg[:110])

    check("راوی نمی‌تواند پیش از ساخت ستون‌های شاهد اجرا شود (باگ B1)",
          all(c in discover()[1].provides
              for c in ("COTAGE_NO", "SATA_NO", "DISCHARGE_DATE"))
          and {"COTAGE_NO", "SATA_NO", "DISCHARGE_DATE"} <=
          set(next(s for s in discover() if s.name == "narrate").requires))


# ═══════════ ۳) افزودن قابلیت = یک فایل ═══════════
def test_add_remove_feature() -> None:
    print("\n── ۳) افزودن و حذف قابلیت با یک فایل ──")
    new_file = os.path.join(PKG, "stages", "s95_demo_feature.py")
    content = '''# -*- coding: utf-8 -*-
"""مرحله آزمایشی — اثبات اینکه یک قابلیت فقط یک فایل است."""
from __future__ import annotations

from typing import Dict, List

import pandas as pd

from .base import ColumnSpec, GROUP_MAIN, PipelineContext, Stage, register


@register
class DemoStage(Stage):
    name = "demo_feature"
    title = "شاخص آزمایشی فشار تأمین"
    order = 95
    requires = ["مقاومت (روز)", "امتیاز ریسک"]
    provides = ["شاخص فشار تأمین"]

    def run(self, df: pd.DataFrame, ctx: PipelineContext) -> pd.DataFrame:
        res = pd.to_numeric(df["مقاومت (روز)"], errors="coerce").fillna(999)
        risk = pd.to_numeric(df["امتیاز ریسک"], errors="coerce").fillna(0)
        df["شاخص فشار تأمین"] = (risk / (res + 1)).round(2)
        return df

    def columns(self) -> List[ColumnSpec]:
        return [ColumnSpec("شاخص فشار تأمین", "شاخص فشار تأمین", 16,
                           GROUP_MAIN, fmt="decimal", order=6)]

    def kpis(self, df: pd.DataFrame, ctx: PipelineContext) -> Dict[str, tuple]:
        s = pd.to_numeric(df.get("شاخص فشار تأمین"), errors="coerce")
        return {"بیشینه فشار تأمین": (round(float(s.max() or 0), 2), "ریسک ÷ مقاومت")}
'''
    with open(new_file, "w", encoding="utf-8") as f:
        f.write(content)

    try:
        code = (
            "import sys; sys.path.insert(0, r'%s')\n"
            "import logging; logging.disable(logging.INFO)\n"
            "from aibl.pipeline import Pipeline\n"
            "r = Pipeline().run(build_report=True)\n"
            "print('COL_OK', 'شاخص فشار تأمین' in r.df.columns)\n"
            "from openpyxl import load_workbook\n"
            "ws = load_workbook(r.dashboard_path)['۲. کالبدشکافی ۳ لایه‌ای ماتریسی']\n"
            "print('SHEET_OK', any(c.value == 'شاخص فشار تأمین' for c in ws[1]))\n"
            "ex = load_workbook(r.dashboard_path)['۱. خلاصه اجرایی']\n"
            "vals = [ex.cell(row=i, column=1).value for i in range(6, 30)]\n"
            "print('KPI_OK', any(v and 'فشار تأمین' in str(v) for v in vals))\n"
        ) % ROOT
        r = subprocess.run([sys.executable, "-c", code], cwd=ROOT,
                           capture_output=True, text=True,
                           env=dict(os.environ, PYTHONIOENCODING="utf-8"),
                           encoding="utf-8", errors="replace")
        out = r.stdout
        check("قابلیت جدید فقط با افزودن یک فایل فعال شد",
              "COL_OK True" in out, out.strip()[-120:] or r.stderr[-120:])
        check("ستونش خودکار در شیت ماتریس آمد (بدون لمس dashboard.py)",
              "SHEET_OK True" in out)
        check("KPI اش خودکار در خلاصه اجرایی آمد (بدون لمس pipeline.py)",
              "KPI_OK True" in out)
    finally:
        os.remove(new_file)

    # ⚠️ مبنا باید در یک پروسه تازه گرفته شود، نه از discover() این پروسه:
    # رجیستری مرحله‌ها کش می‌شود و هنوز فایلِ حذف‌شده را در خود دارد،
    # پس baseline یکی بیشتر از واقعیت درمی‌آمد و تست کاذب قرمز می‌شد.
    code0 = ("import sys; sys.path.insert(0, r'%s')\n"
             "from aibl.stages import discover\n"
             "print('BASE', len(discover()))\n") % ROOT
    r0 = subprocess.run([sys.executable, "-c", code0], cwd=ROOT,
                        capture_output=True, text=True,
                        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
                        encoding="utf-8", errors="replace")
    baseline = int(r0.stdout.split("BASE")[1].split()[0]) if "BASE" in r0.stdout else -1
    code2 = ("import sys; sys.path.insert(0, r'%s')\n"
             "from aibl.stages import discover\n"
             "print('STAGES', len(discover()))\n") % ROOT
    r2 = subprocess.run([sys.executable, "-c", code2], cwd=ROOT,
                        capture_output=True, text=True,
                        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
                        encoding="utf-8", errors="replace")
    check("حذف فایل، قابلیت را بدون هیچ خطایی برداشت",
          f"STAGES {baseline}" in r2.stdout,
          r2.stdout.strip() + f" (مبنا {baseline})")


# ═══════════ ۴) تشخیص اختلاف نسخه فایل‌ها ═══════════
def test_version_skew() -> None:
    print("\n── ۴) تشخیص اختلاف نسخه (علت خطای frame_map) ──")
    check("قراردادها در نسخه سالم بدون مشکل‌اند", not check_contracts(),
          f"{len(check_contracts())} مشکل")

    # شبیه‌سازی دقیق سناریوی واقعی: sources.py قدیمی، بقیه جدید
    sandbox = os.path.join(_TMP, "skew")
    shutil.copytree(PKG, os.path.join(sandbox, "aibl"),
                    ignore=shutil.ignore_patterns("__pycache__"))
    stale = os.path.join(sandbox, "aibl", "config", "sources.py")
    with open(stale, encoding="utf-8") as f:
        src = f.read()
    # نسخه قرارداد را هرچه باشد یک واحد پایین می‌آورد (وابسته به عدد ثابت نیست)
    import re as _re
    src = _re.sub(r"^__contract__ = (\d+)",
                  lambda m: f"__contract__ = {int(m.group(1)) - 1}",
                  src, count=1, flags=_re.M)
    src = src.replace("    def frame_map(self)", "    def _frame_map_removed(self)")
    with open(stale, "w", encoding="utf-8") as f:
        f.write(src)

    env = dict(os.environ, PYTHONPATH=sandbox, PYTHONIOENCODING="utf-8")
    d = subprocess.run([sys.executable, "-m", "aibl.doctor"], cwd=sandbox,
                       capture_output=True, text=True, env=env,
                       encoding="utf-8", errors="replace")
    check("Doctor فایل قدیمی را تشخیص می‌دهد", d.returncode == 1)
    check("پیام دقیقاً نام همان یک فایل را می‌گوید",
          "config/sources.py" in d.stdout,
          next((l.strip() for l in d.stdout.splitlines()
                if "sources.py" in l), "")[:120])
    check("پیام می‌گوید نیازی به کپی مجدد کل پکیج نیست",
          "نیازی به کپی مجدد کل پکیج نیست" in d.stdout)

    # خط لوله هم باید پیش از خواندن داده متوقف شود، نه با AttributeError
    code = ("import sys; sys.path.insert(0, r'%s')\n"
            "from aibl.pipeline import Pipeline\n"
            "Pipeline().run(build_report=False)\n") % sandbox
    p = subprocess.run([sys.executable, "-c", code], cwd=sandbox,
                       capture_output=True, text=True, env=env,
                       encoding="utf-8", errors="replace")
    combined = p.stdout + p.stderr
    check("خط لوله با پیام صریح متوقف می‌شود، نه AttributeError مبهم",
          "AttributeError" not in combined and "نسخه فایل" in combined,
          [l for l in combined.splitlines() if "نسخه فایل" in l][:1])

    # مانیفست هم باید همان فایل را نشان دهد
    check("مانیفست فایل تغییرکرده را نام می‌برد",
          "config/sources.py" in d.stdout)


# ═══════════ ۵) لاگ رویداد استاندارد Celonis ═══════════
def test_eventlog():
    print("\n── ۵) لاگ رویداد (استاندارد Celonis) ──")
    logging.disable(logging.INFO)
    from aibl.pipeline import Pipeline

    res = Pipeline().run(build_report=True)
    ev = res.extras.get("eventlog")
    cases = res.extras.get("case_table")

    check("جدول فعالیت ساخته شد", ev is not None and not ev.empty,
          f"{len(ev)} رویداد" if ev is not None else "خالی")
    required = ["_CASE_KEY", "ACTIVITY_EN", "EVENTTIME", "_SORTING"]
    check("هر چهار ستون استاندارد Celonis موجود است",
          all(c in ev.columns for c in required), str(required))
    check("ستون _SORTING پر است (رویدادهای هم‌تاریخ ترتیب قطعی دارند)",
          ev["_SORTING"].notna().all() and ev["_SORTING"].nunique() > 1,
          f"{ev['_SORTING'].nunique()} مقدار ترتیب متمایز")

    ok = True
    for case, g in ev.groupby("_CASE_KEY"):
        seq = list(zip(g["EVENTTIME"], g["_SORTING"]))
        if seq != sorted(seq):
            ok = False
    check("رویدادهای هر پرونده بر اساس (زمان، ترتیب) مرتب‌اند", ok)

    check("جدول پرونده با throughput ساخته شد",
          cases is not None and "THROUGHPUT_DAYS" in cases.columns,
          f"میانگین {cases['THROUGHPUT_DAYS'].mean():.0f} روز" if cases is not None else "")
    check("طول چرخه = فاصله نخستین تا آخرین رویداد",
          bool(((cases["LAST_EVENT"] - cases["FIRST_EVENT"]).dt.days
                == cases["THROUGHPUT_DAYS"]).all()))

    bn = res.extras.get("bottlenecks")
    check("گلوگاه‌ها محاسبه و نزولی مرتب شدند",
          bn is not None and not bn.empty
          and list(bn["میانگین روز"]) == sorted(bn["میانگین روز"], reverse=True),
          f"{bn.iloc[0]['از فعالیت']} → {bn.iloc[0]['به فعالیت']}: "
          f"{bn.iloc[0]['میانگین روز']:.0f} روز" if bn is not None and len(bn) else "")

    var = res.extras.get("variants")
    check("مسیرهای فرآیند (variants) با سهم درصدی تولید شد",
          var is not None and "سهم (٪)" in var.columns,
          f"{len(var)} مسیر متمایز" if var is not None else "")

    csv_path = os.path.join(DIRS["output"], "AIBL_EventLog.csv")
    check("فایل CSV آماده بارگذاری در Celonis ذخیره شد",
          os.path.exists(csv_path), csv_path)
    if os.path.exists(csv_path):
        head = pd.read_csv(csv_path, nrows=1)
        check("سرستون CSV مطابق استاندارد Celonis است",
              all(c in head.columns for c in required), str(list(head.columns)[:5]))

    from openpyxl import load_workbook
    wb = load_workbook(res.dashboard_path)
    check("شیت «نقشه فرآیند و گلوگاه» ساخته شد",
          "۱۰. نقشه فرآیند و گلوگاه" in wb.sheetnames, str(wb.sheetnames[-2:]))
    # ⚠️ عدد از شناسنامه می‌آید، نه هاردکد: افزودن شیت جدید نباید تست را
    # بشکند، ولی ناهماهنگی مستندات با کد باید فوراً دیده شود.
    from aibl.factsheet import DASHBOARD_SHEETS
    check(f"داشبورد {DASHBOARD_SHEETS} شیتی کامل است",
          len(wb.sheetnames) == DASHBOARD_SHEETS, f"{len(wb.sheetnames)} شیت")
    return res


# ═══════════ ۶) لاغر ماندن ارکستراتور ═══════════
def test_thin_orchestrator() -> None:
    print("\n── ۶) ارکستراتور باید لاغر بماند ──")
    with open(os.path.join(PKG, "pipeline.py"), encoding="utf-8") as f:
        src = f.read()
    lines = [l for l in src.splitlines()
             if l.strip() and not l.strip().startswith("#")]
    check("pipeline.py زیر ۳۰۰ خط مؤثر است", len(lines) < 300, f"{len(lines)} خط")

    business_terms = ["ویبول", "بیزین", "۵۴۰", "540", "مقاومت =",
                      "زیر ۱۰ روز", "yellow", "STOCKOUT"]
    leaked = [t for t in business_terms if t in src]
    check("هیچ عدد یا قاعده کسب‌وکاری در ارکستراتور نیست", not leaked,
          f"نشتی: {leaked}" if leaked else "پاک")

    with open(os.path.join(PKG, "report", "dashboard.py"), encoding="utf-8") as f:
        dash = f.read()
    check("dashboard.py ستون‌های ماتریس را هاردکد نمی‌کند",
          "def build_matrix(self, df: pd.DataFrame, specs: list)" in dash)


if __name__ == "__main__":
    print("=" * 78)
    print("AIBL V23 — تست معماری، اختلاف نسخه و لاگ رویداد")
    print("=" * 78)
    test_stage_contract()
    test_order_guard()
    test_add_remove_feature()
    test_version_skew()
    test_eventlog()
    test_thin_orchestrator()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
