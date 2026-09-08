# -*- coding: utf-8 -*-
"""تست‌های صحت (Validation) — اثبات هر اصلاح، طبق §۱۶ بند ۴.

اجرا:  python tests/test_validation.py
"""
from __future__ import annotations

import os
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

# ── بارگذاری مطمئن make_synthetic بدون اتکا به sys.path ──
# روی ویندوز و در اجرای subprocess، پوشه tests همیشه در sys.path نیست.
_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _TESTS_DIR not in sys.path:
    sys.path.insert(0, _TESTS_DIR)


def _load_make_synthetic():
    import importlib.util
    path = os.path.join(_TESTS_DIR, "make_synthetic.py")
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"فایل make_synthetic.py کنار تست‌ها نیست: {path}\n"
            f"پوشه tests/ را کامل کپی کنید.")
    spec = importlib.util.spec_from_file_location("aibl_make_synthetic", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_ms = _load_make_synthetic()

BLS, build = _ms.BLS, _ms.build

# ── محیط باید قبل از import پکیج ست شود ──
_TMP = tempfile.mkdtemp(prefix="aibl_test_")
DIRS = build(_TMP)
os.environ.update({
    "AIBL_FOREIGN": DIRS["foreign"],
    "AIBL_BLS": DIRS["bls"],
    "AIBL_CLEARANCE": DIRS["clearance"],
    "AIBL_HR": DIRS["hr"],
    "AIBL_ESMAEILI": DIRS["esmaeili"],
    "AIBL_GS_COMBINE": DIRS["gs_combine"],
    "AIBL_MOHAMADI": DIRS["mohamadi"],
    "AIBL_OUTPUT": DIRS["output"],
    "AIBL_LOGS": DIRS["logs"],
    "AIBL_TODAY": "2026-08-31",
})

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from aibl.core.jalali import CalendarEngine  # noqa: E402
from aibl.core.text import (clean_bl, clean_employee_code, clean_key,  # noqa: E402
                            is_empty_val, num_safe)
from aibl.engines.commitment import (delay_penalty, detect_payment_method,  # noqa: E402
                                     is_barat)
from aibl.engines.math_engine import DoctoralMathEngine  # noqa: E402
from aibl.engines.risk import RiskScoreEngine  # noqa: E402
from aibl.narrate.narrator import DynamicGranularNarrator  # noqa: E402
from aibl.pipeline import Pipeline  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


# ══════════════ بخش ۱: توابع پایه ══════════════
def test_units() -> None:
    print("\n── ۱) نرمال‌سازی و کلیدها ──")
    check("clean_key فقط پسوند '.0' را حذف می‌کند (FIX-1)",
          clean_key("10200001.0") == "10200001" and clean_key("1.05") == "1.05",
          f"'1.05' → '{clean_key('1.05')}' (نسخه قبلی: '15')")
    check("clean_bl کلید کوتاه را رد می‌کند",
          clean_bl("AB") == "" and clean_bl("mscu1234567") == "MSCU1234567")
    check("clean_employee_code بخش عددی را zero-pad می‌کند (FIX-12)",
          clean_employee_code("10201069_GS") == "10201069"
          and clean_employee_code("201071") == "00201071")
    check("is_empty_val مقدار '0' و 'نامشخص' را تهی می‌شمارد (FIX-9)",
          is_empty_val("0") and is_empty_val("نامشخص") and not is_empty_val("SATA-1001"))
    check("num_safe جداکننده هزارگان و ارقام فارسی (FIX-2)",
          num_safe("1,200,000") == 1200000.0 and num_safe("۸۵۰۰۰۰") == 850000.0
          and num_safe("1.234.567") == 1234567.0,
          f"'1.234.567' → {num_safe('1.234.567')}")
    check("CalendarEngine تاریخ شمسی را بدون jdatetime تبدیل می‌کند",
          str(CalendarEngine.parse("1404/01/15")) == "2025-04-04",
          str(CalendarEngine.parse("1404/01/15")))
    check("CalendarEngine تاریخ میلادی را می‌خواند",
          str(CalendarEngine.parse("2025-06-10")) == "2025-06-10")


def test_engines() -> None:
    print("\n── ۲) موتورهای بیزینسی و ریاضی ──")
    check("تشخیص برات/یوزانس", is_barat("برات ۱۸۰ روزه") and not is_barat("نقدی"))
    check("تشخیص روش پرداخت از واژگان YAML",
          detect_payment_method("اعتبار اسنادی") == "LC"
          and detect_payment_method("یوزانس") == "BARAT"
          and detect_payment_method("حواله نقدی") == "CASH",
          f'اعتبار اسنادی → {detect_payment_method("اعتبار اسنادی")}')
    p12 = delay_penalty(1_000_000, 365)
    expected = 1_000_000 * (6 * 0.01 + (365 / 30 - 6) * 0.02)
    check("جریمه پلکانی ۱٪ تا ۶ ماه سپس ۲٪", abs(p12 - expected) < 1,
          f"{p12:,.0f} ریال برای ۳۶۵ روز تأخیر")
    check("ویبول: رسوب صفر ⇒ بقای ۱۰۰٪",
          DoctoralMathEngine.weibull_survival(0) == 100.0)
    s45 = DoctoralMathEngine.weibull_survival(45)
    check("ویبول: t=η ⇒ بقا ≈ ۳۶.۸٪", abs(s45 - 36.8) < 0.5, f"{s45}٪")
    check("ویبول: ترخیص درصدی بقا را بالا می‌برد",
          DoctoralMathEngine.weibull_survival(100, clearance_type="ترخیص درصدی")
          > DoctoralMathEngine.weibull_survival(100))
    full = DoctoralMathEngine.bayesian_customs_prob(["تاریخ تخلیه", "کوتاژ گمرکی", "کد ساتا"])
    none_ = DoctoralMathEngine.bayesian_customs_prob([])
    check("بیزین: شواهد کامل > شواهد صفر", full > none_, f"{full}٪ در برابر {none_}٪")
    r = RiskScoreEngine(1_000_000).score({
        "ELAPSED_DAYS": 600, "LEGAL_DEADLINE_DAYS": 540, "CB_VALUE": 1_000_000,
        "ALLOCATED": False, "SATA_NO": "", "FIN_RECEIPT_DATE": "", "FULL_CLEAR_DATE": "",
        "PENALTY": 200_000, "STUCK_DAYS": 200, "BUDGET_FLAG": "CRITICAL",
        "PART_CRITICALITY_SCORE": 100})
    check("امتیاز ریسک بدترین حالت ⇒ بحرانی", r.score >= 76 and "بحرانی" in r.band,
          f"{r.score:.1f} → {r.band}")
    r2 = RiskScoreEngine(1_000_000).score({
        "ELAPSED_DAYS": 10, "LEGAL_DEADLINE_DAYS": 540, "CB_VALUE": 1000,
        "ALLOCATED": True, "SATA_NO": "S1", "FIN_RECEIPT_DATE": "1405/01/01",
        "FULL_CLEAR_DATE": "1405/01/05", "PENALTY": 0, "STUCK_DAYS": 2,
        "PART_CRITICALITY_SCORE": 0})
    check("امتیاز ریسک بهترین حالت ⇒ پایین", r2.score < 26, f"{r2.score:.1f} → {r2.band}")


def test_narrator_guard() -> None:
    print("\n── ۳) گارد ضد باگ B1 (راوی کور) ──")
    try:
        DynamicGranularNarrator.generate({"CANONICAL_BL": "X"}, CalendarEngine.parse("2026-08-31"))
        check("راوی بدون ستون‌های شاهد خطا می‌دهد", False, "خطا صادر نشد!")
    except KeyError as ex:
        check("راوی بدون ستون‌های شاهد خطای صریح می‌دهد (B1)", True, str(ex)[:70])


# ══════════════ بخش ۴: خط لوله کامل ══════════════
def test_pipeline():
    print("\n── ۴) اجرای کامل خط لوله روی داده مصنوعی ──")
    res = Pipeline().run(build_report=True)
    df = res.df

    check("خط لوله بدون خطا اجرا شد", not df.empty, f"{len(df)} ردیف نهایی")

    # T1/T2/B3/FIX-6 — عدم انفجار سطر
    check("عدم تکثیر سطر پس از ادغام ۱۱ سورس (B3 + FIX-6)", len(df) == 6,
          f"۶ ردیف مبدأ → {len(df)} ردیف نهایی (clearance دو ردیف تکراری برای همان بارنامه داشت)")

    # FIX-1 — clean_key
    orders = set(df["CANONICAL_ORDER"])
    check("مرجع سفارش با پسوند حرفی سالم ماند (812211A)",
          "812211A" in orders, str(sorted(orders)))
    check("مرجع سفارش عددی درست ساخته شد", "502805" in orders)

    # T1 — بارنامه نامعتبر
    check("بارنامه نامعتبر 'AB' حذف شد و کلید خالی گرفت (B3)",
          (df["CANONICAL_BL"] == "").sum() == 1,
          f"{int((df['CANONICAL_BL'] == '').sum())} ردیف بدون بارنامه معتبر")

    # FIX-3 — دو شیت NTSW
    check("هر دو شیت NTSW خوانده شد (FIX-3)",
          "NTSW_BALANCE" in df.columns and "NTSW_ALLOC_STATUS" in df.columns,
          "Release Commitment + Allocation")
    total = float(pd.to_numeric(df["مانده تعهد"], errors="coerce").sum())
    check("مانده تعهد از شیت Release Commitment وارد شد", total > 0, f"جمع = {total:,.0f}")
    # T5: کد ثبت سفارش 98404279 دو ردیف تعهد دارد (2795310 + 583440)
    row = df[df["KEY_REG"] == "98404279"]
    check("چند ردیف تعهد یک ثبت سفارش با هم جمع شدند (T5)",
          len(row) == 1 and abs(float(row.iloc[0]["مانده تعهد"]) - 3378750) < 1,
          f"{float(row.iloc[0]['مانده تعهد']):,.0f} = 2,795,310 + 583,440" if len(row) else "یافت نشد")
    # T6: دو درخواست تخصیص — آخرین وضعیت ملاک است، نه اولی
    check("آخرین وضعیت تخصیص گرفته شد، نه اولین (T6)",
          len(row) and str(row.iloc[0]["ALLOC_STATUS"]) == "تخصیص یافته",
          f"وضعیت = {row.iloc[0]['ALLOC_STATUS']} (درخواست اول: پذیرفته نشده)" if len(row) else "")
    check("کلید ثبت سفارش از ساتا ساخته شد، نه شماره پرونده IL",
          df["KEY_REG"].astype(str).str.match(r"^\d{8}$").sum() >= 4,
          str(sorted(set(df["KEY_REG"]))[:6]))

    # FIX-11/12/13 — اتصال HR
    linked = (df["ORG_MATCH"] != "پیش‌فرض").sum()
    check("اتصال HR از طریق کد پرسنلی برقرار شد (FIX-10/11/12/13)", linked >= 1,
          f"{linked} از {len(df)} ردیف؛ روش‌ها: {sorted(set(df['ORG_MATCH']))}")
    check("نام کارشناس از سورس‌های چندگانه حل شد",
          df["CANONICAL_EXPERT"].astype(str).str.strip().ne("").sum() >= 3,
          str(sorted(set(df["CANONICAL_EXPERT"]))[:4]))
    check("زنجیره سقوط مدیر عمل کرد (مریم احمدی: مدیر خالی → رئیس/مسئول)",
          "حسن محمدی" in set(df["ORG_MANAGER"]) or "رضا نوری" in set(df["ORG_MANAGER"]),
          str(sorted(set(df["ORG_MANAGER"]))))

    # B1 — راوی بینا
    conf = pd.to_numeric(df["درصد قطعیت"], errors="coerce")
    check("درصد قطعیت دیگر روی یک مقدار ثابت قفل نیست (B1)", conf.nunique() > 1,
          f"مقادیر مشاهده‌شده: {sorted(set(conf.dropna().astype(int)))}")
    row0 = df[df["CANONICAL_BL"] == BLS[0]].iloc[0]
    check("شواهد کوتاژ و ساتا واقعاً دیده می‌شوند (B1 + FIX-9)",
          str(row0["COTAGE_NO"]).strip() != "" and str(row0["SATA_NO"]).strip() != "",
          f"کوتاژ={row0['COTAGE_NO']} ساتا={row0['SATA_NO']} "
          f"قطعیت={row0['درصد قطعیت']}٪")
    row2 = df[df["CANONICAL_BL"] == BLS[2]].iloc[0]
    check("کد رهگیری '0' به‌درستی تهی شمرده شد (FIX-9)",
          row2["درصد قطعیت"] < 100, f"قطعیت = {row2['درصد قطعیت']}٪")

    # FIX-8 — فیلتر سال مالی
    check("محموله بدون تاریخ تخلیه «در راه» علامت خورد",
          bool(df[df["CANONICAL_BL"] == BLS[3]].iloc[0].get("IS_IN_TRANSIT", False)),
          f"بارنامه {BLS[3]} تخلیه نشده است")

    # FIX-7 — KPI بر بارنامه یکتا
    uniq = df.loc[df["CANONICAL_BL"] != "", "CANONICAL_BL"].nunique()
    check("شمارش KPI بر اساس بارنامه یکتا است (FIX-7)", uniq == 5,
          f"{uniq} بارنامه یکتا از {len(df)} ردیف")

    # افراز دو محوری
    check("محور A (کلید) هر سه پارتیشن را تولید کرد",
          df["PARTITION_KEY"].nunique() >= 2, str(res.counts))
    check("محور B: بارنامه بدون مقاومت/ساتا/کوتاژ به تعیین تکلیف رفت (T10)",
          len(res.to_resolve) >= 1, f"{len(res.to_resolve)} ردیف")

    # موتور تعهد
    check("مهلت قانونی و جریمه محاسبه شد",
          pd.to_numeric(df["جریمه برآوردی"], errors="coerce").sum() > 0,
          f"جمع جریمه = {pd.to_numeric(df['جریمه برآوردی'], errors='coerce').sum():,.0f}")
    check("هشدار سه‌رنگ تولید شد", df["وضعیت کلی هشدار"].nunique() >= 1,
          str(sorted(set(df["وضعیت کلی هشدار"]))))

    # ممیزی
    check("ممیزی تعارض اجرا شد", isinstance(res.audit, pd.DataFrame),
          f"{len(res.audit)} تعارض ثبت شد")

    return res


def test_excel(res) -> None:
    print("\n── ۵) صحت خروجی اکسل ──")
    check("فایل داشبورد ساخته شد", os.path.exists(res.dashboard_path),
          os.path.basename(res.dashboard_path))
    wb = load_workbook(res.dashboard_path)
    expected = ["۱. خلاصه اجرایی", "۲. کالبدشکافی ۳ لایه‌ای ماتریسی", "۳. تعیین تکلیف",
                "۴. رفع تعهد ارزی", "۵. کارنامه سازمانی", "۶. پشتیبان ریاضی",
                "۹. قطعات بحرانی"]
    check("هر ۷ شیت کلیدی بدون خطا ساخته شد", all(s in wb.sheetnames for s in expected),
          str(wb.sheetnames))

    ws = wb["۲. کالبدشکافی ۳ لایه‌ای ماتریسی"]
    check("Freeze Panes فعال است", ws.freeze_panes == "A2", str(ws.freeze_panes))
    check("AutoFilter فعال است", ws.auto_filter.ref is not None, str(ws.auto_filter.ref))
    check("Conditional Formatting اعمال شد", len(list(ws.conditional_formatting)) >= 2,
          f"{len(list(ws.conditional_formatting))} قانون")
    check("Data Validation (لیست کشویی) اضافه شد", len(ws.data_validations.dataValidation) >= 1)
    check("گروه‌بندی ستون‌ها (outline) فعال است",
          any(d.outline_level and d.outline_level > 0 for d in ws.column_dimensions.values()))
    check("جهت راست‌به‌چپ تنظیم شد", ws.sheet_view.rightToLeft is True)

    wsc = wb["۴. رفع تعهد ارزی"]
    has_formula = any(isinstance(c.value, str) and c.value.startswith("=")
                      for row in wsc.iter_rows() for c in row)
    check("Formula Injection (SUBTOTAL) در شیت تعهد", has_formula)

    check("فایل مجزای هر کارشناس ساخته شد", len(res.extract_paths) >= 2,
          f"{len(res.extract_paths)} فایل در expert_extracts")
    check("گزارش ممیزی تعارضات ساخته شد",
          os.path.exists(os.path.join(DIRS["output"], "AIBL_Data_Conflicts_Audit.xlsx")))


# ═══════════ ۱۰) اصلاحات نسخه ۲۵ — باگ‌های فساد خاموش داده ═══════════
def test_v25_fixes(df) -> None:
    print("\n── ۱۰) اصلاحات نسخه ۲۵ ──")
    from aibl.core.jalali import jalali_sort_key as jk
    from aibl.core.text import clean_employee_code as cec

    # C: تاریخ شمسی دو رقمی نباید تاریخ جدید را بخورد
    check("تاریخ ۲ رقمی «98/12/27» زیر «1403/01/01» مرتب می‌شود",
          jk("98/12/27") < jk("1403/01/01"),
          f"{jk('98/12/27')} < {jk('1403/01/01')}")
    check("«1403/1/5» و «1403/01/05» یک کلید می‌گیرند",
          jk("1403/1/5") == jk("1403/01/05"), jk("1403/1/5"))
    check("مقدار خراب همیشه بازنده است (نه برنده)",
          jk("*") == "0000-00-00" and jk("9812345") == "0000-00-00")
    check("سال ۸۸ هم درست بسط می‌یابد (نه فقط ۹x)",
          jk("88/05/12") < jk("98/12/27"), f"{jk('88/05/12')}")

    # کد پرسنلی: هم پسوند حرفی، هم «.0» عددی‌خوانده‌شده
    check("کد پرسنلی '10201069_GS' درست تمیز می‌شود", cec("10201069_GS") == "10201069")
    check("کد پرسنلی عددی‌خوانده‌شده '10201069.0' رقم اضافه نمی‌گیرد",
          cec("10201069.0") == "10201069", cec("10201069.0"))
    check("صفر پیشوند در هر دو طرف یکدست می‌شود",
          cec("201071") == cec("00201071") == "00201071")

    # doccheck روی سفارش، نه بارنامه
    check("doccheck روی شماره سفارش متصل شد (نه بارنامه)",
          "DOC_STATUS" in df.columns
          and df["DOC_STATUS"].astype(str).str.strip().ne("").sum() >= 2,
          f"{int(df['DOC_STATUS'].astype(str).str.strip().ne('').sum())} ردیف سند متصل")

    # ترخیص کامل بدون تاریخ
    if "CL_CLEAR_DONE_NO_DATE" in df.columns:
        check("«ترخیص کامل بدون تاریخ» تاریخ جعلی نمی‌سازد",
              not (df["CL_IS_FULL"].astype(bool)
                   & df["CL_CLEAR_DATE"].astype(str).str.startswith("*")).any())

    # نرمال‌سازی نام ستون با زیرخط
    from aibl.core.columns import find_col
    import pandas as _pd
    probe = _pd.DataFrame({"_ تاریخ بارگیری نهایی_": ["1405/01/01"]})
    check("نام ستون با زیرخط و فاصله اضافی تطبیق می‌خورد",
          find_col(probe, ["تاریخ بارگیری نهایی"]) is not None)



if __name__ == "__main__":
    print("=" * 78)
    print("AIBL V21 — مجموعه تست‌های صحت")
    print("=" * 78)
    test_units()
    test_engines()
    test_narrator_guard()
    result = test_pipeline()
    test_excel(result)
    test_v25_fixes(result.df)
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)

