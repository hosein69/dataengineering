# -*- coding: utf-8 -*-
"""تست منطق داشبورد — بدون نیاز به نصب Streamlit.

Streamlit در زمان import کد سطح-ماژول را اجرا می‌کند، پس ``dashboard.py``
به‌تنهایی قابل تست خودکار نیست. تمام منطق در ``app/ui_kit.py`` است و اینجا
واقعاً اجرا و سنجیده می‌شود؛ ``dashboard.py`` فقط از نظر نحوی و وابستگی
بررسی می‌شود.

اجرا:  python tests/test_dashboard.py
"""
from __future__ import annotations

import ast
import logging
import os
import re
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import pandas as pd  # noqa: E402

from app.ui_kit import (BAND_COLORS, band_count, card_html, detail_columns,  # noqa: E402
                        export_html, kpis, num)

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


def sample() -> pd.DataFrame:
    """نمونه‌ای که همان تناقض گزارش‌شده کاربر را هم دارد."""
    return pd.DataFrame({
        "KEY_MATERIAL": ["IK001", "IK002", "IK003", "IK004", "IK004"],
        "CANONICAL_ORDER": ["502805", "501807", "501317", "812210", "812210"],
        "CANONICAL_BL": ["B1", "B2", "", "B4", "B4"],
        "KEY_REG": ["98404279", "79668513", "", "97687754", "97687754"],
        "کد طبقه بحرانی": ["STOCKOUT", "CRITICAL", "BECOMING_CRITICAL",
                            "SAFE", "SAFE"],
        "بحرانی (کوتاه)": ["توقف خط", "بحرانی", "در حال بحرانی شدن",
                            "ایمن", "ایمن"],
        "طبقه بحرانی": ["🔴 توقف خط", "🔴 بحرانی", "🟠 در حال بحرانی شدن",
                         "🟢 ایمن", "🟢 ایمن"],
        # ← همان حالت کاربر: مقاومت ۴ روز با موجودی ایران‌خودرو صفر
        "مقاومت (روز)": [0.0, 4.0, 15.0, 500.0, 500.0],
        "موجودی ایران خودرو": [0.0, 0.0, 1200.0, 5000.0, 5000.0],
        "موجودی ساپکو": [0.0, 440.0, 300.0, 0.0, 0.0],
        "موجودی کل قابل احتساب": [0.0, 440.0, 1500.0, 5000.0, 5000.0],
        "نیاز روزانه": [411.0, 110.0, 100.0, 10.0, 10.0],
        "روزهای رسوب": [200.0, 30.0, 5.0, 0.0, 0.0],
        "روزهای تأخیر": [120.0, 0.0, 0.0, 0.0, 0.0],
        "مانده تعهد": [3378750.0, 0.0, 1000000.0, 0.0, 0.0],
        "امتیاز ریسک": [88.0, 61.0, 43.0, 12.0, 12.0],
        "طبقه ریسک": ["🔴 بحرانی", "🟠 بالا", "🟡 متوسط", "🟢 پایین", "🟢 پایین"],
        "CANONICAL_EXPERT": ["علی رضایی", "مریم احمدی", "سارا کریمی",
                              "علی رضایی", "علی رضایی"],
        "ORG_DEPT": ["لجستیک", "خرید خارجی", "لجستیک", "لجستیک", "لجستیک"],
    })


def test_metrics() -> None:
    print("\n── ۱) محاسبه شاخص‌ها ──")
    df = sample()
    check("شمارش طبقه بحرانی بر اساس متریال یکتاست، نه ردیف",
          band_count(df, "SAFE") == 1,
          "دو ردیف IK004 → یک متریال")
    check("توقف خط و بحرانی درست شمرده می‌شوند",
          band_count(df, "STOCKOUT") == 1 and band_count(df, "CRITICAL") == 1)
    check("طبقه ناموجود صفر برمی‌گرداند", band_count(df, "UNKNOWN") == 0)
    check("num روی ستون غایب خطا نمی‌دهد",
          num(df, "ستون ناموجود").isna().all() and len(num(df, "ستون ناموجود")) == len(df))

    k = kpis(df)
    check("کارت‌های مدیریتی کامل تولید می‌شوند", len(k) == 11, str(len(k)))
    by = {c["label"]: c for c in k}
    check("کارت «توقف خط» بحرانی علامت می‌خورد", by["توقف خط"]["critical"] is True)
    check("کمترین مقاومت درست است", by["کمترین مقاومت"]["value"] == "0.0 روز",
          str(by["کمترین مقاومت"]["value"]))
    check("جمع مانده تعهد درست است",
          abs(by["جمع مانده تعهد"]["value"] - 4378750.0) < 1,
          f"{by['جمع مانده تعهد']['value']:,.0f}")
    check("تعهد معوق شمرده شد و بحرانی علامت خورد",
          by["تعهدات معوق"]["value"] == 1 and by["تعهدات معوق"]["critical"])

    empty = kpis(pd.DataFrame())
    check("روی داده خالی خطا نمی‌دهد و «—» می‌گذارد",
          len(empty) == 11 and empty[3]["value"] == "—")


def test_cards() -> None:
    print("\n── ۲) کارت‌ها ──")
    h = card_html("قطعات بحرانی", 17, "زیر ۱۰ روز", critical=True)
    check("کلاس بحرانی اعمال می‌شود", 'class="kpi crit"' in h)
    check("عدد با جداکننده هزارگان قالب می‌گیرد",
          "1,234" in card_html("x", 1234), card_html("x", 1234)[:60])
    check("مقدار متنی هم پذیرفته می‌شود", "—" in card_html("x", "—"))
    bad = card_html("<script>alert(1)</script>", 5)
    check("ورودی متنی escape می‌شود (بدون تزریق HTML)",
          "<script>" not in bad and "&lt;script&gt;" in bad)


def test_export_html() -> None:
    print("\n── ۳) خروجی HTML ──")
    df = sample()
    out = export_html(df, "2026-09-06")

    check("سند کامل HTML است", out.startswith("<!DOCTYPE html>") and
          out.rstrip().endswith("</html>"))
    check("راست‌به‌چپ و فارسی تنظیم شده", 'dir="rtl"' in out and 'lang="fa"' in out)
    check("CSS درون‌خط است (بدون فایل بیرونی)",
          "<style>" in out and 'rel="stylesheet"' not in out)
    check("JS درون‌خط است", "<script>" in out and "window.print()" in out)
    check("دکمه ذخیره PDF دارد", "ذخیره به PDF" in out)
    check("در چاپ، دکمه پنهان می‌شود", "@media print" in out and ".btn{display:none}" in out.replace(" ", ""))
    check("پالت آکوا/سبز/سفید/خاکستری استفاده شده",
          all(c in out for c in ("#0F6E6E", "#D9EDE7", "#FFFFFF", "#5A6B6B")))
    check("جدول داده در خروجی هست", 'class="dataframe tbl"' in out or 'class="tbl"' in out)
    check("همه ردیف‌ها منتقل شده‌اند", out.count("<tr>") >= len(df))
    check("تاریخ مرجع در سند آمده", "2026-09-06" in out)

    big = export_html(pd.concat([df] * 200, ignore_index=True), "2026-09-06")
    check("خروجی برای داده بزرگ به ۵۰۰ ردیف محدود می‌شود",
          big.count("<tr>") <= 520, f"{big.count('<tr>')} سطر")

    check("روی داده خالی هم سند معتبر می‌سازد",
          export_html(pd.DataFrame({"a": []}), "x").startswith("<!DOCTYPE"))


def test_columns() -> None:
    print("\n── ۴) ستون‌های جدول ──")
    df = sample()
    cols = detail_columns(df)
    check("فقط ستون‌های موجود انتخاب می‌شوند",
          all(c in df.columns for c in cols) and len(cols) >= 10, str(len(cols)))
    check("ستون‌های اجزای مقاومت در جدول هستند",
          "موجودی کل قابل احتساب" in cols and "نیاز روزانه" in cols,
          "تا تناقض «مقاومت با موجودی صفر» دوباره رخ ندهد")
    check("روی دیتافریم بی‌ربط خالی برمی‌گرداند",
          detail_columns(pd.DataFrame({"z": [1]})) == [])


def test_dashboard_module() -> None:
    print("\n── ۵) سلامت فایل داشبورد (بدون اجرای Streamlit) ──")
    path = os.path.join(ROOT, "app", "dashboard.py")
    src = open(path, encoding="utf-8").read()
    try:
        tree = ast.parse(src)
        ok = True
    except SyntaxError as ex:
        ok, tree = False, None
        print(f"    {ex}")
    check("dashboard.py از نظر نحوی سالم است", ok)
    if not ok:
        return

    check("منطق از ui_kit وارد می‌شود (بدون کد تکراری)",
          "from app.ui_kit import" in src)
    check("پس‌زمینه متحرک تعریف شده",
          "@keyframes drift1" in src and "@keyframes drift2" in src)
    check("نبض فقط روی کارت بحرانی است", "@keyframes pulse" in src
          and ".kpi.crit::after" in src)
    check("هر سه خروجی وجود دارد (اکسل کامل، داده فیلترشده، HTML)",
          src.count("download_button") >= 3)
    check("st.set_page_config پیش از هر فراخوانی دیگر st است",
          src.index("st.set_page_config") < src.index("st.markdown"))
    check("راه‌انداز پورت آزاد پیدا می‌کند",
          "def free_port" in open(os.path.join(ROOT, "app", "run_dashboard.py"),
                                  encoding="utf-8").read())

    # نبود streamlit نباید تست را بشکند — فقط گزارش شود
    try:
        import streamlit  # noqa: F401
        print("    ℹ️ Streamlit نصب است؛ اجرای زنده هم ممکن است.")
    except ImportError:
        print("    ⚠️ Streamlit در این محیط نصب نیست — فقط تحلیل ایستا انجام شد. "
              "برای اجرا:  pip install streamlit plotly")


def test_scorecard_group_criticality() -> None:
    print("\n── ۶) کارنامه سازمانی و رنگ وضعیت ──")
    from tempfile import TemporaryDirectory
    from openpyxl import load_workbook
    from aibl.report.dashboard import ExcelDashboardBuilder

    df = pd.DataFrame({
        "CANONICAL_BL": ["BL1", "BL1", "BL2"],
        "KEY_MATERIAL": ["M_SAFE", "M_CRIT", "M_WATCH"],
        "CANONICAL_ORDER": ["O1", "O1", "O2"],
        "ORG_VICE": ["V1", "V1", "V2"], "ORG_DEPT": ["D1", "D1", "D2"],
        "ORG_MANAGER": ["Mgr", "Mgr", "Mgr2"], "ORG_HEAD": ["Head", "Head", "Head2"],
        "CANONICAL_EXPERT": ["Exp", "Exp", "Exp2"], "KEY_EMP": ["E1", "E1", "E2"],
        "BL_CRITICAL": [False, True, False],
        "BL_CRITICAL_LEVEL": ["SAFE", "CRITICAL", "WATCH"],
        "کد طبقه بحرانی": ["SAFE", "CRITICAL", "WATCH"],
        "بحرانی (کوتاه)": ["ایمن", "بحرانی", "تحت نظر"],
        "مقاومت (روز)": [50, 4, 30], "روزهای رسوب": [1, 20, 2],
        "امتیاز ریسک": [10, 90, 30], "مانده تعهد": [0, 1000, 200], "جریمه برآوردی": [0, 20, 0],
    })
    with TemporaryDirectory() as td:
        out = f"{td}/scorecard.xlsx"
        b = ExcelDashboardBuilder(out); b.build_scorecard(df); b.save()
        wb = load_workbook(out)
        ws = wb["۵. کارنامه سازمانی"]
        vals = list(ws.iter_rows(min_row=2, max_row=2, min_col=7, max_col=8, values_only=True))[0]
        check("کارنامه بارنامه بحرانی را از کل BL حساب می‌کند، نه اولین ردیف", vals == (1, 1), str(vals))
        check("فونت کارنامه IRANSans Light است", ws["A2"].font.name == "IRANSans Light", ws["A2"].font.name)
        check("عدد بحرانی با رنگ قرمز نمایش داده می‌شود", ws["H2"].font.color.rgb in {"00B3261E", "FFB3261E"}, str(ws["H2"].font.color.rgb))



# ═══════════════════════════════════════════════════════════════════════════
# ۷) رگرسیون‌های نسخه ۲۶٫۲٫۳ — هرکدام یک باگ واقعی تولید را قفل می‌کند
# ═══════════════════════════════════════════════════════════════════════════
def test_regressions_v26_2_3() -> None:
    print("\n── ۷) رگرسیون‌های تثبیت‌شده ──")

    # R1 ── داشبورد سفیدِ خالی رندر می‌شد.
    # CSS داشبورد «position» را روی .stApp بازنویسی می‌کرد. Streamlit خودش
    # آن را absolute/inset:0 می‌گذارد و #stAppViewContainer ارتفاعش را از
    # همان می‌گیرد؛ با relative ارتفاع صفر می‌شد و چون overflow آن hidden
    # است هیچ‌چیز نقاشی نمی‌شد — عناصر در DOM بودند ولی صفحه سفید بود.
    dash = os.path.join(ROOT, "app", "dashboard.py")
    src = open(dash, encoding="utf-8").read()
    in_stapp, overrides = False, []
    for raw in src.splitlines():
        line = raw.strip()
        if line.startswith(".stApp {"):
            in_stapp = True
            continue
        if in_stapp:
            if line.startswith("}"):
                in_stapp = False
            elif line.startswith("position:"):
                overrides.append(line)
    check("CSS داشبورد position را روی .stApp بازنویسی نمی‌کند",
          not overrides,
          "بازنویسی یعنی رندر سفید" if overrides else "پاک")

    # R2 ── یک کلیک روی «ساخت Excel سفارشی» گزارش رسمی ۱۳ شیتی را می‌بلعید.
    # نام پیش‌فرض حالا یک ثابت است، نه رشته‌ای داخل UI — پس به‌جای grep
    # روی فایل رابط، خود ثابت سنجیده می‌شود (مقاوم در برابر تغییر ویجت).
    from aibl.config.settings import SETTINGS
    from aibl.studio_core.excel_export import DEFAULT_CUSTOM_NAME
    official_stem = SETTINGS.SYSTEMMATIC_MATERIAL_BASENAME.rsplit(".", 1)[0]
    check("نام پیش‌فرض Excel سفارشی با گزارش رسمی یکی نیست",
          DEFAULT_CUSTOM_NAME.strip() != official_stem, DEFAULT_CUSTOM_NAME)

    # R3 ── نگهبان نوشتن روی گزارش رسمی، در خودِ writer (نه فقط در UI)
    from aibl.studio_core.excel_export import build_custom_excel, OfficialReportOverwrite
    from datetime import date as _date
    ref = "2026-08-31"
    blocked = False
    try:
        build_custom_excel(pd.DataFrame({"a": [1]}),
                           SETTINGS.daily_report_path(_date.fromisoformat(ref)),
                           ["kpi"], ref)
    except OfficialReportOverwrite:
        blocked = True
    except Exception:
        pass
    check("writer نوشتن روی گزارش رسمی روزانه را رد می‌کند", blocked)

    # R4 ── «python -m aibl email» روی نصب تازه با ModuleNotFoundError می‌افتاد.
    req = open(os.path.join(ROOT, "requirements.txt"), encoding="utf-8").read().lower()
    missing = [x for x in ("matplotlib", "yaml", "numpy") if x not in req]
    check("وابستگی‌های واقعی در requirements.txt اعلام شده‌اند",
          not missing, f"غایب: {missing}" if missing else "matplotlib/PyYAML/numpy")

    # R6 ── کاتالوگ فیلد ۹۴٪ داده را پنهان می‌کرد.
    # نسخه قبل یک دیکشنری دستی ۲۷تایی بود و از ۳۶۷ ستون خط لوله فقط ۲۲ تا
    # را نشان می‌داد؛ بقیه در UI وجود نداشتند و قابل گزارش نبودند. حالا
    # کاتالوگ از خود adapterها مشتق می‌شود و باید *هر* ستون را پوشش دهد.
    from aibl.studio_core.field_catalog import build_catalog, unique_labels
    probe = pd.DataFrame({
        "KEY_MATERIAL": ["M1"], "BL_GOODS_DESC": ["x"], "SATA_GOODS_DESC": ["y"],
        "NTSW_BALANCE": [1], "ORC_STOCK_IKCO": [5], "بحرانی (کوتاه)": ["بحرانی"],
        "یک_ستون_کاملاً_ناشناخته": ["z"],
    })
    specs = build_catalog(probe)
    covered = {sp.column for sp in specs}
    check("کاتالوگ هیچ ستونی را جا نمی‌اندازد",
          covered == set(probe.columns),
          f"{len(covered)}/{len(probe.columns)} — جامانده: {set(probe.columns)-covered}")
    check("ستون ناشناخته هم قابل گزارش می‌ماند (حذف نمی‌شود)",
          "یک_ستون_کاملاً_ناشناخته" in covered)
    check("برچسب از هدر واقعی سورس مشتق می‌شود",
          any(sp.column == "BL_GOODS_DESC" and sp.label == "شرح کالا" for sp in specs))

    # R7 ── چند سورس برچسب یکسان دارند («شرح کالا» در بارنامه و ساتا).
    # rename مستقیم ستون تکراری می‌ساخت و جدول با ValueError می‌ترکید.
    names = list(unique_labels(specs).values())
    check("نام‌های نمایشی یکتا هستند (جدول با ستون تکراری نمی‌ترکد)",
          len(set(names)) == len(names), f"{len(set(names))}/{len(names)}")
    try:
        probe.rename(columns=unique_labels(specs))[list(names)]
        renamed_ok = True
    except Exception:
        renamed_ok = False
    check("rename با نگاشت کاتالوگ روی دیتافریم واقعی کار می‌کند", renamed_ok)

    # R8 ── فهرست گیرندگان داده شخصی است و نباید در سورس باشد.
    # نسخه‌های قبل ۲۶ نشانی واقعی کارکنان را داخل daily_email.py داشتند.
    leaked = []
    for base, dirs, files in os.walk(ROOT):
        dirs[:] = [d for d in dirs
                   if d not in {"__pycache__", ".git"} and not d.startswith("D:")]
        for f in files:
            if not f.endswith((".py", ".yaml", ".yml", ".md", ".txt")):
                continue
            if f == "recipients.example.yaml":
                continue
            full = os.path.join(base, f)
            try:
                body = open(full, encoding="utf-8").read()
            except Exception:
                continue
            for m in re.finditer(r"[\w.+-]+@[\w-]+\.[\w.]+", body):
                addr = m.group(0)
                if not addr.lower().endswith((".invalid", ".example", "@example.com")):
                    leaked.append(f"{os.path.relpath(full, ROOT)}: {addr}")
    check("هیچ نشانی ایمیل واقعی در سورس نیست",
          not leaked, f"نشت: {leaked[:4]}" if leaked else "پاک")

    from aibl.integrations.daily_email import _recipients, NoRecipients
    _saved = os.environ.pop("AIBL_EMAIL_TO", None)
    try:
        check("بدون پیکربندی، فهرست گیرندگان خالی است (نه یک فهرست جاسازی‌شده)",
              _recipients() == [])
    finally:
        if _saved is not None:
            os.environ["AIBL_EMAIL_TO"] = _saved
    os.environ["AIBL_EMAIL_TO"] = "x@a.invalid; y@b.invalid"
    check("فهرست گیرندگان از متغیر محیطی خوانده می‌شود",
          _recipients() == ["x@a.invalid", "y@b.invalid"], str(_recipients()))
    os.environ.pop("AIBL_EMAIL_TO", None)

    # R9 ── جدول متقاطع: شمارش با بُعد ستون، pivot_table را با
    # «Grouper not 1-dimensional» می‌ترکاند چون index و values یکی می‌شدند.
    from app.analytics import _pivot
    pv = pd.DataFrame({
        "dept": ["A", "A", "B", "B", "B"],
        "band": ["بحرانی", "ایمن", "بحرانی", "بحرانی", "ایمن"],
        "val":  [10, 20, 30, 40, 50],
        "bl":   ["X", "X", "Y", "Z", "Z"],
    })
    combos = [
        ("dept", "band", "count", None),
        ("dept", None, "count", None),
        ("dept", "band", "sum", "val"),
        ("band", None, "mean", "val"),
        ("dept", None, "nunique", "bl"),
    ]
    broke = []
    for r_, c_, how_, m_ in combos:
        try:
            _pivot(pv, r_, c_, how_, m_)
        except Exception as ex:
            broke.append(f"{(r_, c_, how_, m_)}: {type(ex).__name__}")
    check("جدول متقاطع در هر ترکیب بُعد/سنجه کار می‌کند",
          not broke, f"شکست: {broke}" if broke else f"{len(combos)} ترکیب")
    check("شمارش با بُعد ستون درست جمع می‌زند",
          int(_pivot(pv, "dept", "band", "count", None).values.sum()) == len(pv))
    check("جمع سنجه با بُعد ستون درست است",
          float(_pivot(pv, "dept", "band", "sum", "val").values.sum()) == float(pv["val"].sum()))

    # R10 ── نام کارشناس ترخیص جای کارشناس خرید می‌نشست.
    # «CANONICAL_EXPERT» با «اولین مقدار غیرتهی» از پنج سورس پر می‌شد و
    # چون ORC_BUYER فقط ۱۷٪ پر است، برای بیشتر ردیف‌ها به CL_EXPERT
    # (کارشناس ترخیص) می‌افتاد و عملکرد ترخیص به پای خرید نوشته می‌شد.
    from aibl.resolve.expert_roles import (ROLES, coverage as role_coverage,
                                           current_owner, resolve_roles)
    probe = pd.DataFrame({
        "ORC_BUYER": ["اباذر بالی", "", ""],
        "CL_EXPERT": ["یعقوب طایفه", "عقیل بقاپور", "امیر آقامحمدی"],
        "SATA_CREDIT_EXPERT": ["حمید یحیائی", "مرتضی یزدی", ""],
        "DOC_EXPERT": ["نیما صابری", "", ""],
    })
    res = resolve_roles(probe.copy())
    check("هر نقش کارشناسی ستون مستقل دارد",
          all(r.key in res.columns for r in ROLES), f"{len(ROLES)} نقش")
    check("کارشناس ترخیص هرگز در ستون کارشناس خرید نمی‌نشیند",
          list(res["EXPERT_BUYER"]) == ["اباذر بالی", "", ""],
          str(list(res["EXPERT_BUYER"])))
    check("ستون ترخیص فقط کارشناس ترخیص را دارد",
          list(res["EXPERT_CLEARANCE"]) == ["یعقوب طایفه", "عقیل بقاپور",
                                            "امیر آقامحمدی"])
    check("ستون اعتبارات با ترخیص آلوده نمی‌شود",
          list(res["EXPERT_CREDIT"]) == ["حمید یحیائی", "مرتضی یزدی", ""])

    name, role = current_owner(res)
    check("مالک مرحله فعلی همراه با نام نقشش گزارش می‌شود",
          bool(str(role.iloc[0]).strip()) and bool(str(name.iloc[0]).strip()),
          f"{name.iloc[0]} / {role.iloc[0]}")
    check("نقش مالک با ستونی که نامش از آن آمده هم‌خوان است",
          str(role.iloc[1]) == "کارشناس ترخیص"
          and str(name.iloc[1]) == "عقیل بقاپور",
          f"{name.iloc[1]} / {role.iloc[1]}")

    cov = role_coverage(res)
    check("پوشش هر نقش جداگانه گزارش می‌شود",
          len(cov) == len(ROLES) and "پرشدگی (٪)" in cov.columns)
    empty_roles = cov[cov["پرشدگی (٪)"] == 0]["نقش"].tolist()
    check("نقش بدون داده صریحاً صفر گزارش می‌شود (نه پر از نقش دیگر)",
          "کارشناس ثبت سفارش" in empty_roles, str(empty_roles))

    # R5 ── settings.py با f-string تودرتوی هم‌نقل‌قول فقط روی پایتون ۳٫۱۲+
    # کامپایل می‌شد؛ روی ۳٫۹–۳٫۱۱ کل پکیج SyntaxError می‌داد.
    bad = []
    for base, dirs, files in os.walk(os.path.join(ROOT, "aibl")):
        dirs[:] = [d for d in dirs if d != "__pycache__"]
        for f in files:
            if not f.endswith(".py"):
                continue
            full = os.path.join(base, f)
            try:
                ast.parse(open(full, encoding="utf-8").read())
            except SyntaxError as ex:
                bad.append(f"{os.path.relpath(full, ROOT)}:{ex.lineno}")
    check("هر ماژول پکیج روی همین نسخه پایتون کامپایل می‌شود",
          not bad, f"خراب: {bad}" if bad else f"python {sys.version_info.major}.{sys.version_info.minor}")


def test_email_from_hr() -> None:
    """R11 ── گیرندگان ایمیل از سورس HR.

    فهرست گیرنده داده شخصی است: نه در سورس، نه در فایل جانبیِ کهنه.
    بهترین منبع، ستون Email همان فایل پرسنلی است که همیشه به‌روز است و
    پرسنل غیرفعال خودکار از آن حذف می‌شوند.
    """
    print("\n── ۸) گیرنده ایمیل از سورس HR ──")
    from aibl.integrations import daily_email as de

    ppl = pd.DataFrame([
        # فعال، مدیر، نشانی سالم → باید بیاید
        dict(HR_EMAIL="a.manager@x.invalid", HR_STATUS="فعال",
             HR_POST="مدیر مواد اولیه", HR_DEPT="مواد اولیه", HR_OFFICE="خرید"),
        # فعال، رئیس → باید بیاید
        dict(HR_EMAIL="b.head@x.invalid", HR_STATUS="فعال",
             HR_POST="رئیس اداره ترخیص", HR_DEPT="قطعات", HR_OFFICE="ترخیص"),
        # غیرفعال با پست مدیریتی → نباید بیاید
        dict(HR_EMAIL="c.left@x.invalid", HR_STATUS="غیرفعال",
             HR_POST="مدیر", HR_DEPT="مواد اولیه", HR_OFFICE="خرید"),
        # کارشناس فعال → با پیش‌فرضِ مدیریتی نباید بیاید
        dict(HR_EMAIL="d.expert@x.invalid", HR_STATUS="فعال",
             HR_POST="کارشناس خرید خارجی", HR_DEPT="مواد اولیه", HR_OFFICE="خرید"),
        # نشانی ناقص → هرگز
        dict(HR_EMAIL="not-an-email", HR_STATUS="فعال",
             HR_POST="مدیر", HR_DEPT="قطعات", HR_OFFICE="خرید"),
        # نشانی خالی → هرگز
        dict(HR_EMAIL="", HR_STATUS="فعال",
             HR_POST="مدیر", HR_DEPT="قطعات", HR_OFFICE="خرید"),
    ])
    keys = ["AIBL_EMAIL_TO", "AIBL_RECIPIENTS_FILE", "AIBL_EMAIL_FROM_HR",
            "AIBL_EMAIL_HR_POSTS", "AIBL_EMAIL_HR_MANAGEMENTS",
            "AIBL_EMAIL_HR_OFFICES", "AIBL_EMAIL_HR_MAX"]
    saved = {k: os.environ.pop(k, None) for k in keys}
    try:
        got = de.hr_recipients(ppl)
        check("پیش‌فرض فقط سطوح مدیریتیِ فعال است",
              got == ["a.manager@x.invalid", "b.head@x.invalid"], str(got))
        check("پرسنل غیرفعال حذف می‌شود", "c.left@x.invalid" not in got)
        check("نشانی نامعتبر حذف می‌شود",
              not any("not-an-email" in g for g in got))
        check("ردیف بدون نشانی حذف می‌شود", "" not in got)

        os.environ["AIBL_EMAIL_HR_POSTS"] = "کارشناس"
        got = de.hr_recipients(ppl)
        check("فیلتر شرح پست کار می‌کند", got == ["d.expert@x.invalid"], str(got))

        os.environ["AIBL_EMAIL_HR_MANAGEMENTS"] = "قطعات"
        check("فیلترها با هم AND می‌شوند", de.hr_recipients(ppl) == [])
        os.environ.pop("AIBL_EMAIL_HR_MANAGEMENTS")

        os.environ["AIBL_EMAIL_HR_POSTS"] = "مدیر,رئیس,کارشناس"
        os.environ["AIBL_EMAIL_HR_MAX"] = "2"
        check("سقف تعداد رعایت می‌شود", len(de.hr_recipients(ppl)) == 2)
        os.environ.pop("AIBL_EMAIL_HR_MAX")
        os.environ.pop("AIBL_EMAIL_HR_POSTS")

        check("جدول خالی ⇒ فهرست خالی", de.hr_recipients(pd.DataFrame()) == [])
        check("سورس بدون ستون Email ⇒ فهرست خالی",
              de.hr_recipients(pd.DataFrame({"HR_POST": ["مدیر"]})) == [])

        # ── زنجیره حل ──
        check("بدون AIBL_EMAIL_FROM_HR، سورس HR خوانده نمی‌شود",
              de._recipients() == [])
        os.environ["AIBL_EMAIL_FROM_HR"] = "1"
        _real = de.hr_recipients
        de.hr_recipients = lambda frame=None: ["from.hr@x.invalid"]
        try:
            check("با فعال‌سازی، زنجیره به سورس HR می‌رسد",
                  de._recipients() == ["from.hr@x.invalid"])
            os.environ["AIBL_EMAIL_TO"] = "override@x.invalid"
            check("متغیر محیطی صریح بر سورس HR اولویت دارد",
                  de._recipients() == ["override@x.invalid"])
            os.environ.pop("AIBL_EMAIL_TO")
        finally:
            de.hr_recipients = _real

        # ── نشانی‌ها فقط شمرده می‌شوند، هرگز لاگ نمی‌شوند ──
        rec: list = []
        _h = _LogCatcher(rec)
        de.log.addHandler(_h)
        try:
            de.hr_recipients(ppl)
        finally:
            de.log.removeHandler(_h)
        body = " ".join(rec)
        check("هیچ نشانی‌ای لاگ نمی‌شود", "@x.invalid" not in body, body[:90])
        check("تعداد گیرنده لاگ می‌شود", any("گیرنده" in m for m in rec), body[:90])
    finally:
        for k in keys:
            os.environ.pop(k, None)
            if saved.get(k) is not None:
                os.environ[k] = saved[k]


class _LogCatcher(logging.Handler):
    def __init__(self, sink: list) -> None:
        super().__init__()
        self.sink = sink

    def emit(self, record: logging.LogRecord) -> None:
        self.sink.append(record.getMessage())



if __name__ == "__main__":
    print("=" * 78)
    print("AIBL — تست منطق داشبورد")
    print("=" * 78)
    test_metrics()
    test_cards()
    test_export_html()
    test_columns()
    test_dashboard_module()
    test_scorecard_group_criticality()
    test_regressions_v26_2_3()
    test_email_from_hr()
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
