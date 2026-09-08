# -*- coding: utf-8 -*-
"""تست قرارداد گزارش — هر کلیدی که شیت‌ها می‌خوانند باید واقعاً تولید شود.

## باگی که این فایل می‌بندد

کاربر در گزارش واقعی دید: **«مقاومت ۴ روز» کنار «موجودی صفر»**.
محاسبه درست بود (۴۴۰ ساپکو ÷ ۱۱۰ نیاز روزانه = ۴)، ولی شیت «قطعات بحرانی»
ستون قدیمی ``"موجودی"`` را می‌خواند که پس از تغییر فرمول دیگر تولید نمی‌شد.
``row.get("موجودی", "")`` بی‌صدا رشته خالی برمی‌گرداند — نه خطا، نه هشدار.

این خطرناک‌ترین نوع خطاست: عدد **درست** کنار عدد **غایب** می‌نشیند و
گزارش، متناقض ولی موجه به نظر می‌رسد. یک کاربر می‌تواند بر اساس آن تصمیم
توقف خط بگیرد.

pandas و openpyxl هیچ‌کدام این را نمی‌گیرند، چون از نظر آن‌ها هیچ اتفاقی
نیفتاده. پس باید صریح تست شود.

اجرا:  python tests/test_contracts_report.py
"""
from __future__ import annotations

import os
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import tempfile  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_TESTS = os.path.dirname(os.path.abspath(__file__))
if _TESTS not in sys.path:
    sys.path.insert(0, _TESTS)


def _load_ms():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "aibl_make_synthetic", os.path.join(_TESTS, "make_synthetic.py"))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


DIRS = _load_ms().build(tempfile.mkdtemp(prefix="aibl_contract_"))
os.environ.update({
    "AIBL_FOREIGN": DIRS["foreign"], "AIBL_BLS": DIRS["bls"],
    "AIBL_CLEARANCE": DIRS["clearance"], "AIBL_HR": DIRS["hr"],
    "AIBL_ESMAEILI": DIRS["esmaeili"], "AIBL_GS_COMBINE": DIRS["gs_combine"],
    "AIBL_MOHAMADI": DIRS["mohamadi"], "AIBL_OUTPUT": DIRS["output"],
    "AIBL_LOGS": DIRS["logs"], "AIBL_TODAY": "2026-08-31",
})

import logging  # noqa: E402
import re  # noqa: E402

import pandas as pd  # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  → {detail}" if detail else ""))


# ═══════════ ۱) بازتولید دقیق باگ کاربر ═══════════
def test_reproduce_user_bug() -> None:
    print("\n── ۱) بازتولید «مقاومت ۴ روز با موجودی صفر» ──")
    from aibl.engines.criticality import CriticalityEngine

    r = CriticalityEngine().evaluate({
        "STOCK_IKCO": 0, "STOCK_SAPCO": 440,
        "IN_TRANSIT_QTY": 0, "IN_CUSTOMS_QTY": 0, "DAILY_NEED": 110})
    d = r.as_dict()

    check("محاسبه درست است: ۴۴۰ ساپکو ÷ ۱۱۰ = ۴ روز",
          r.resistance_days == 4.0 and r.band == "CRITICAL",
          f"{r.resistance_days} روز → {r.band_short}")
    check("موجودی ایران‌خودرو واقعاً صفر است (تناقض ظاهری بود نه محاسباتی)",
          d["موجودی ایران خودرو"] == 0.0 and d["موجودی ساپکو"] == 440.0)
    check("کلید قدیمی «موجودی» دیگر وجود ندارد — منشأ باگ",
          "موجودی" not in d,
          "شیت باید «موجودی ایران خودرو» و «موجودی ساپکو» را بخواند")
    check("موجودی کل قابل احتساب گزارش می‌شود تا حساب قابل بازبینی باشد",
          d["موجودی کل قابل احتساب"] == 440.0)


# ═══════════ ۲) قرارداد: هر کلید خوانده‌شده باید تولید شود ═══════════
def test_report_keys_exist() -> None:
    print("\n── ۲) قرارداد کلیدهای گزارش ──")
    logging.disable(logging.WARNING)
    from aibl.pipeline import Pipeline

    res = Pipeline().run(build_report=True)
    produced = set(res.df.columns)

    # کلیدهایی که شیت‌ها از ردیف می‌خوانند، از خود سورس استخراج می‌شوند
    src = open(os.path.join(ROOT, "aibl", "report", "dashboard.py"),
               encoding="utf-8").read()
    src += open(os.path.join(ROOT, "aibl", "report", "charts.py"),
                encoding="utf-8").read()

    read_keys = set()
    for m in re.finditer(r'row\.get\(\s*"([^"]+)"', src):
        read_keys.add(m.group(1))
    for m in re.finditer(r'cols\s*=\s*\[(.*?)\]', src, re.S):
        read_keys |= set(re.findall(r'"([^"]+)"', m.group(1)))

    # کلیدهای فنی که عمداً ممکن است نباشند (وابسته به سورس اختیاری)
    optional = {"MOGH_BL_SUSPECT", "MOGH_BL_REJECT_REASON", "SAFETY_STOCK",
                "WAREHOUSE", "MOGH_HS_KEYWORD"}
    missing = sorted(k for k in read_keys
                     if k not in produced and k not in optional
                     and not k.startswith("MOGH_") and " " in k or
                     (k not in produced and k not in optional and k.isupper() is False
                      and " " in k))
    missing = sorted({k for k in read_keys
                      if " " in k and k not in produced and k not in optional})

    check("هیچ شیتی کلید ناموجود نمی‌خواند", not missing,
          f"غایب: {missing}" if missing else f"{len(read_keys)} کلید بررسی شد")

    # ستون‌های حیاتی مقاومت واقعاً در خروجی هستند
    need = ["مقاومت (روز)", "مقاومت انبار (روز)", "موجودی ایران خودرو",
            "موجودی ساپکو", "موجودی در راه", "موجودی در گمرک",
            "موجودی کل قابل احتساب", "نیاز روزانه"]
    absent = [c for c in need if c not in produced]
    check("هر هشت ستون اجزای مقاومت در خروجی موجودند", not absent, str(absent))
    return res


# ═══════════ ۳) سازگاری عددی: صورت کسر با مقاومت بخواند ═══════════
def test_arithmetic_visible(res) -> None:
    print("\n── ۳) حساب مقاومت روی کاغذ قابل بازبینی است ──")
    df = res.df
    ok, bad = True, []
    for _, r in df.iterrows():
        need = pd.to_numeric(r.get("نیاز روزانه"), errors="coerce")
        total = pd.to_numeric(r.get("موجودی کل قابل احتساب"), errors="coerce")
        days = pd.to_numeric(r.get("مقاومت (روز)"), errors="coerce")
        if pd.isna(days) or pd.isna(need) or need <= 0:
            continue
        expect = round(min(total / need, 999), 1)
        if abs(expect - days) > 0.15:
            ok = False
            bad.append(f"{r.get('KEY_MATERIAL')}: {total}/{need}={expect} ≠ {days}")
    check("در هر ردیف، موجودی کل ÷ نیاز روزانه با مقاومت گزارش‌شده می‌خواند",
          ok, "؛ ".join(bad) or "همه ردیف‌ها سازگار")

    # اجزا باید با کل جمع بزنند
    ok2, bad2 = True, []
    for _, r in df.iterrows():
        parts = sum(pd.to_numeric(r.get(c), errors="coerce") or 0 for c in
                    ("موجودی ایران خودرو", "موجودی ساپکو",
                     "موجودی در راه", "موجودی در گمرک"))
        total = pd.to_numeric(r.get("موجودی کل قابل احتساب"), errors="coerce") or 0
        if abs(parts - total) > 0.01:
            ok2 = False
            bad2.append(f"{r.get('KEY_MATERIAL')}: {parts} ≠ {total}")
    check("جمع چهار جزء با «موجودی کل» برابر است", ok2,
          "؛ ".join(bad2) or "همه ردیف‌ها سازگار")


# ═══════════ ۴) شیت قطعات بحرانی خالی چاپ نمی‌کند ═══════════
def test_sheet_not_blank(res) -> None:
    print("\n── ۴) شیت «قطعات بحرانی» عدد خالی چاپ نمی‌کند ──")
    from openpyxl import load_workbook
    wb = load_workbook(res.dashboard_path)
    ws = wb["۹. قطعات بحرانی"]

    header_row = None
    for r in range(1, 30):
        if ws.cell(row=r, column=1).value == "طبقه بحرانی":
            header_row = r
            break
    check("سطر هدر جدول تفصیلی پیدا شد", header_row is not None, str(header_row))
    if header_row is None:
        return

    heads = [ws.cell(row=header_row, column=c).value
             for c in range(1, ws.max_column + 1)]
    for want in ("موجودی ایران خودرو", "موجودی ساپکو", "نیاز روزانه", "موجودی کل"):
        check(f"ستون «{want}» در شیت هست", want in heads)

    # اولین ردیف داده: اگر مقاومت عدد دارد، اجزا هم باید عدد داشته باشند
    i_res = heads.index("مقاومت (روز)") + 1
    i_tot = heads.index("موجودی کل") + 1
    i_need = heads.index("نیاز روزانه") + 1
    contradictions = []
    for r in range(header_row + 1, min(header_row + 40, ws.max_row + 1)):
        res_v = ws.cell(row=r, column=i_res).value
        tot_v = ws.cell(row=r, column=i_tot).value
        need_v = ws.cell(row=r, column=i_need).value
        if res_v in (None, ""):
            continue
        if isinstance(res_v, (int, float)) and res_v > 0:
            if tot_v in (None, "", 0) or need_v in (None, "", 0):
                contradictions.append(f"سطر {r}: مقاومت {res_v} ولی کل={tot_v} نیاز={need_v}")
    check("هیچ ردیفی «مقاومت مثبت با موجودی/نیاز خالی» ندارد",
          not contradictions, "؛ ".join(contradictions[:3]) or "بدون تناقض")


if __name__ == "__main__":
    print("=" * 78)
    print("AIBL — تست قرارداد گزارش (باگ «مقاومت ۴ روز با موجودی صفر»)")
    print("=" * 78)
    test_reproduce_user_bug()
    result = test_report_keys_exist()
    test_arithmetic_visible(result)
    test_sheet_not_blank(result)
    print("\n" + "=" * 78)
    print(f"نتیجه: {len(PASS)} موفق | {len(FAIL)} ناموفق")
    if FAIL:
        print("ناموفق: " + " ، ".join(FAIL))
    print("=" * 78)
    sys.exit(1 if FAIL else 0)
