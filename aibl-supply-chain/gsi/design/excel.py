# -*- coding: utf-8 -*-
"""GSI Design System — پل Excel.

همان توکن‌هایی که HTML و ایمیل از آن می‌خوانند، اینجا به سبک‌های openpyxl
ترجمه می‌شوند. هدف یک چیز است: **مدیری که Excel را باز می‌کند و مدیری که
HTML را باز می‌کند، یک گزارش ببینند، نه دو گزارش.**

پیش از این، «تحت نظر» در HTML ``#fab219`` بود و در Excel ``F1C40F``؛ هیچ‌کدام
غلط نبودند ولی گزارش یک صدا نداشت.

## تفاوت‌های اجباری با وب

Excel یک محیط چاپ‌محور است و چند محدودیت واقعی دارد که طراحی باید بپذیرد:

* **بدون شفافیت در fill شرطی** — پس پس‌زمینه‌های ملایم (``wash``) رنگ
  مات‌اند، نه ``rgba``.
* **ارتفاع سطر ثابت** — پس مقیاس تایپ Excel فشرده‌تر از وب است.
* **رنگ سلول بدون مرز، در چاپ سیاه‌وسفید محو می‌شود** — پس هر سلول رنگی
  مرز ``ink`` هم می‌گیرد، همان قاعده WCAG 1.4.11 که در وب اعمال شد.
"""
from __future__ import annotations

__contract__ = 1

from typing import Dict, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import tokens as T

#: Excel رنگ را بدون «#» و با حروف بزرگ می‌خواهد.
X = T.excel

FONT_NAME = T.FONT_EXCEL

# ── مقیاس تایپ Excel (نقطه، نه پیکسل) ─────────────────────────────────────
SIZE_TITLE = 15
SIZE_H2 = 12
SIZE_H3 = 11
SIZE_BODY = 10
SIZE_SMALL = 9
SIZE_KPI = 18

# ── قالب عدد ──────────────────────────────────────────────────────────────
NUM_INT = "#,##0"
NUM_CURRENCY = "#,##0.00"
NUM_PCT = "0.0"
NUM_DATE = "yyyy-mm-dd"


def font(*, size: int = SIZE_BODY, bold: bool = False,
         color: str = T.TEXT, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=X(color))


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=X(color))


def border(color: str = T.BORDER, style: str = "thin") -> Border:
    side = Side(style=style, color=X(color))
    return Border(left=side, right=side, top=side, bottom=side)


def align(h: str = "right", wrap: bool = False, v: str = "center") -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


# ═══════════════════════════════════════════════════════════════════════════
# سبک‌های نام‌دار — معادل کامپوننت‌های وب
# ═══════════════════════════════════════════════════════════════════════════
def title_style() -> Dict[str, object]:
    """تیتر صفحه — معادل ``.appbar``."""
    return {"font": font(size=SIZE_TITLE, bold=True, color=T.TEXT_ON_DARK),
            "fill": fill(T.BRAND_NAVY), "alignment": align("right")}


def header_style(level: int = 1) -> Dict[str, object]:
    """سر ستون جدول — معادل ``th``."""
    size = {1: SIZE_H3, 2: SIZE_BODY}.get(level, SIZE_SMALL)
    return {"font": font(size=size, bold=True, color=T.TEXT_ON_DARK),
            "fill": fill(T.BRAND_NAVY if level == 1 else T.TEXT_SECONDARY),
            "alignment": align("center", wrap=True),
            "border": border(T.BORDER_STRONG)}


def body_style(*, bold: bool = False, muted: bool = False) -> Dict[str, object]:
    return {"font": font(bold=bold, color=T.TEXT_MUTED if muted else T.TEXT),
            "alignment": align("right")}


def kpi_style(tone: str = "") -> Dict[str, object]:
    """کارت سنجه — عدد درشت با رنگ وضعیت."""
    ink = T.STATUS[tone].ink if tone in T.STATUS else T.BRAND_NAVY
    return {"font": font(size=SIZE_KPI, bold=True, color=ink),
            "alignment": align("center")}


def status_style(tone: str) -> Dict[str, object]:
    """سلول وضعیت — پس‌زمینه ملایم، متن تیره، و **مرز هم‌خانواده**.

    مرز اختیاری نیست: بدون آن، سلول «تحت نظر» در چاپ سیاه‌وسفید از سلول
    «ایمن» تفکیک نمی‌شود.
    """
    s = T.STATUS.get(tone)
    if not s:
        return body_style()
    return {"font": font(bold=True, color=s.ink), "fill": fill(s.wash),
            "border": border(s.ink), "alignment": align("center")}


def zebra_fill(row_index: int) -> Optional[PatternFill]:
    """نوار یک‌درمیان — همان ``tbody tr:nth-child(even)`` وب."""
    return fill("#fbfcfd") if row_index % 2 == 0 else None


def apply(cell, style: Dict[str, object]) -> None:
    """یک سبک نام‌دار را روی سلول می‌نشاند."""
    for key, value in style.items():
        setattr(cell, key, value)


#: رنگ سری نمودارهای بومی Excel — همان پالت دسته‌ای وب.
SERIES_COLORS = tuple(X(c) for c in T.CATEGORICAL)
#: رنگ سطح هر وضعیت، برای data point نمودار.
STATUS_FILLS: Dict[str, str] = {s.label: X(s.fill) for s in T.STATUS_SCALE}
STATUS_INKS: Dict[str, str] = {s.label: X(s.ink) for s in T.STATUS_SCALE}

# ── طیف قالب‌بندی شرطی (ColorScaleRule) ─────────────────────────────────────
#
# اکسل برای «کم = بد، زیاد = خوب» یک طیف سه‌نقطه‌ای می‌خواهد. قبلاً این سه
# نقطه هگزهای دستی بودند (``F5B7B1`` / ``FDEBD0`` / ``D5F5E3``) که هیچ نسبتی
# با بقیه گزارش نداشتند — کاربر یک صورتی می‌دید که در هیچ نشان و هیچ نموداری
# تکرار نشده بود.
#
# حالا هر سه نقطه همان ``wash`` وضعیت متناظرند. یعنی رنگی که در سلول
# «بحرانی» می‌بینید، دقیقاً همان رنگی است که در طیف هم معنی «بد» می‌دهد.
#
# این‌ها **پس‌زمینه**‌اند و متن روی‌شان :data:`TEXT` است؛ کم‌ترین کنتراست
# حاصل حدود ۱۴ است، یعنی خیلی بالاتر از کف ۴٫۵.
SCALE_BAD = X(T.STATUS["critical"].wash)
SCALE_MID = X(T.STATUS["warning"].wash)
SCALE_GOOD = X(T.STATUS["good"].wash)

#: ردیفی که «مشکوک» است ولی هنوز رد نشده — همان wash «تحت نظر».
SUSPECT_FILL = SCALE_MID

#: ردیفی که منتظر تطبیق رسمی است (قاعده needs_verification).
PENDING_FILL = SCALE_MID



def legacy_palette_map() -> Dict[str, str]:
    """نگاشت نام‌های قدیمی ``LuxuryPalette`` به توکن‌های جدید.

    برای مهاجرت بی‌شکست: کدی که هنوز ``P.STATUS_CRITICAL`` می‌خواند، رنگ
    درست و سنجیده‌شده می‌گیرد، نه مقدار قدیمی رد‌شده از کنتراست.
    """
    return {
        "STATUS_CRITICAL": X(T.STATUS["critical"].ink),
        "STATUS_WARNING": X(T.STATUS["serious"].ink),
        "STATUS_WATCH": X(T.STATUS["warning"].ink),
        "STATUS_GOOD": X(T.STATUS["good"].ink),
        "STATUS_INACTIVE": X(T.STATUS["neutral"].ink),
        "STATUS_CRITICAL_FILL": X(T.STATUS["critical"].wash),
        "STATUS_WARNING_FILL": X(T.STATUS["serious"].wash),
        "STATUS_WATCH_FILL": X(T.STATUS["warning"].wash),
        "STATUS_GOOD_FILL": X(T.STATUS["good"].wash),
        "STATUS_INACTIVE_FILL": X(T.STATUS["neutral"].wash),
        "TEXT_DARK": X(T.TEXT),
        "TEXT_MUTED": X(T.TEXT_SECONDARY),
        "TEXT_LIGHT": X(T.TEXT_ON_DARK),
        "BORDER_LIGHT": X(T.BORDER),
        "AMBER_HEADER": X(T.BRAND_NAVY),
        "AMBER_FILL": X(T.TEAL_WASH),
        "CRITICAL_FILL": X(T.STATUS["critical"].wash),
        "GREEN_L1": X(T.SEQUENTIAL[1]),
        "GREEN_L2": X(T.SEQUENTIAL[0]),
        "GREEN_L3": X(T.TEAL_WASH),
        "GREEN_L4": X(T.SURFACE_SUNKEN),
    }
