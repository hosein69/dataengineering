# -*- coding: utf-8 -*-
"""خروجی اکسل — راست‌به‌چپ، با همان هویت رنگی، برای بایگانی و توزیع سازمانی.

اکسل و HTML دو **مقصرند نه دو منبع**: هر دو از همان ``kpis``/``nodes``/
``edges`` ساخته می‌شوند که به Streamlit هم داده شده. تفاوت فقط در قالب
خروجی است، نه در محاسبه — تا عددی که مدیر در HTML می‌بیند با عددی که در
اکسل باز می‌کند هرگز واگرا نشود.
"""
from __future__ import annotations

__contract__ = 1

import io
from typing import Mapping, Sequence

from .. import persian as fa
from .. import tokens as T


def build_excel_report(*, kpis: Sequence[Mapping], nodes: Sequence[Mapping],
                        edges: Sequence[Mapping]) -> bytes:
    """گزارش اکسل را می‌سازد و بایت‌های فایل ``.xlsx`` را برمی‌گرداند."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor=T.INK.lstrip("#"))
    header_font = Font(color="FFFFFF", bold=True, name="Tahoma", size=11)
    body_font = Font(name="Tahoma", size=10.5)
    right = Alignment(horizontal="right", vertical="center")
    center = Alignment(horizontal="center", vertical="center")

    first_sheet_used = False

    def _sheet(name: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
        nonlocal first_sheet_used
        ws = wb.active if not first_sheet_used else wb.create_sheet()
        first_sheet_used = True
        ws.title = name
        ws.sheet_view.rightToLeft = True
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = body_font
                cell.alignment = right if c == 1 else center
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 22
        ws.freeze_panes = "A2"

    _sheet("KPI", ["شاخص", "مقدار", "تغییر نسبت به دوره قبل", "وضعیت"],
           [(k.get("label", ""), k.get("value", ""), k.get("delta", ""),
             T.status_of(k.get("status")).label) for k in kpis])

    labels = {n.get("id"): n.get("label", n.get("id")) for n in nodes}
    edge_rows = [(labels.get(e.get("source"), e.get("source")),
                  labels.get(e.get("target"), e.get("target")),
                  int(e.get("count", 0)), "بحرانی" if e.get("critical") else "عادی")
                 for e in sorted(edges, key=lambda e: -int(e.get("count", 0)))]
    _sheet("مسیرهای فرآیند", ["از مرحله", "به مرحله", "تعداد کیس", "طبقه"], edge_rows)

    node_rows = [(n.get("label", n.get("id")), n.get("count", ""), T.status_of(n.get("status")).label)
                 for n in nodes]
    _sheet("مراحل فرآیند", ["مرحله", "تعداد کیس فعلی", "وضعیت"], node_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
