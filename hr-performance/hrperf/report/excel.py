# -*- coding: utf-8 -*-
"""خروجی Excel چندبرگه با قالب‌بندی و رنگ رده عملکرد."""
from __future__ import annotations

__contract__ = 1

from pathlib import Path
from typing import Dict, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .theme import BANDS, band_of

FONT = "IRANSans Light"
HEADER_FILL = "0D366B"


def _style(ws) -> None:
    thin = Side(style="thin", color="E3E3DD")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=9, color="111917")
            c.alignment = Alignment(vertical="center")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max([len(str(c.value or "")) for c in col[:60]] + [9])
        ws.column_dimensions[letter].width = min(max(width + 2, 11), 34)
    ws.row_dimensions[1].height = 28


def _colour_performance(ws) -> None:
    """رنگ رده روی ستون «عملکرد» — همراه آیکن، نه فقط رنگ."""
    idx = None
    for i, cell in enumerate(ws[1], start=1):
        if str(cell.value).strip() == "عملکرد":
            idx = i
            break
    if idx is None:
        return
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=idx)
        try:
            v = float(c.value)
        except (TypeError, ValueError):
            continue
        color, icon, _label = band_of(v)
        c.value = f"{icon} {v:.1f}"
        c.font = Font(name=FONT, size=9, bold=True, color=color.replace("#", ""))


def build_excel(path: str | Path, sheets: Dict[str, pd.DataFrame]) -> str:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wrote = False
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        for name, df in sheets.items():
            if df is None or getattr(df, "empty", True):
                continue
            df.to_excel(w, sheet_name=str(name)[:31], index=False)
            wrote = True
        if not wrote:
            pd.DataFrame({"note": ["داده‌ای برای گزارش نبود."]}).to_excel(
                w, sheet_name="خالی", index=False)
    wb = load_workbook(p)
    for ws in wb.worksheets:
        _style(ws)
        _colour_performance(ws)
    wb.save(p)
    return str(p)
